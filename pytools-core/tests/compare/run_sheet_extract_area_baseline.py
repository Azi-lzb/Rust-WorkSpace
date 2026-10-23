"""Generate a synthetic XLSX fixture and run the read-only PyCode baseline."""

from __future__ import annotations

import hashlib
import json
import logging
import sys
import tempfile
from copy import copy
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook


WORKSPACE = Path(__file__).resolve().parents[3]
FIXTURE_DIR = WORKSPACE / "pytools-core" / "tests" / "fixtures" / "sheet_extract_area"
REPORT_PATH = WORKSPACE / "pytools-core" / "probe" / "reports" / "sheet_extract_area_python_baseline.json"
CASES_PATH = FIXTURE_DIR / "cases.json"
SOURCE_PROJECT = Path(r"C:\Users\AZI\Desktop\VibeCode\4.pytools-main")
SOURCE_PYCODE = SOURCE_PROJECT / "PyCode"
SOURCE_FUNCTION = SOURCE_PYCODE / "core" / "feature_1_6_extract_sheets.py"
CONFIG_FIXTURE = FIXTURE_DIR / "extract_config.xlsx"

sys.path.insert(0, str(SOURCE_PYCODE))
from core.feature_1_6_extract_sheets import run_extract_sheets  # noqa: E402


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def normalized_cell(value):
    if value is None:
        return None
    if isinstance(value, bool):
        return {"type": "boolean", "value": value}
    if isinstance(value, datetime):
        return {"type": "datetime", "value": value.isoformat(timespec="milliseconds")}
    if isinstance(value, date):
        return {"type": "date", "value": value.isoformat()}
    if isinstance(value, int):
        return {"type": "integer", "value": value}
    if isinstance(value, float):
        return {"type": "number", "value": value}
    if isinstance(value, str):
        return {"type": "string", "value": value}
    raise TypeError(f"unsupported baseline cell {value!r} ({type(value).__name__})")


def create_input(path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "月度资金报表"
    rows = [
        ["指标", "本月", None, "备注"],
        ["存款", 10, None, True],
        [None, None, None, "保留此行"],
        ["贷款", 20, None, False],
        ["制造业", 30, None, 0.125],
    ]
    for row_index, row in enumerate(rows, start=1):
        for column_index, value in enumerate(row, start=1):
            if value is not None:
                worksheet.cell(row=row_index, column=column_index, value=value)
    worksheet["E1"] = datetime(2024, 12, 31, 9, 30)
    font = copy(worksheet["A1"].font)
    font.bold = True
    worksheet["A1"].font = font
    worksheet.column_dimensions["B"].width = 18

    excluded = workbook.create_sheet("月度资金报表_旧版_归档")
    excluded["A1"] = "must be excluded"
    workbook.create_sheet("年度汇总")
    workbook.save(path)


def create_config(tasks: list[dict]) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "工作表提取"
    headers = [
        "是否启用",
        "工作簿关键字",
        "工作表关键字",
        "排除关键字",
        "提取区域",
        "多区域拼接方式",
        "输出文件夹",
        "文件名",
        "备注",
    ]
    worksheet.append(headers)
    for task in tasks:
        worksheet.append(
            [
                "是",
                task["workbook_keywords"],
                task["sheet_keywords"],
                task["exclude_keywords"],
                task["area_spec"],
                task["stitch_mode"],
                task["output_folder"],
                task["output_name"],
                task["remark"],
            ]
        )
    workbook.save(CONFIG_FIXTURE)


def main() -> None:
    cases = json.loads(CASES_PATH.read_text(encoding="utf-8"))
    FIXTURE_DIR.mkdir(parents=True, exist_ok=True)
    inputs = [FIXTURE_DIR / name for name in cases["input_files"]]
    create_input(inputs[0])
    create_input(inputs[1])
    inputs[2].write_bytes(b"not an XLSX package; expected open failure")
    create_config(cases["tasks"])

    source_hash_before = sha256(SOURCE_FUNCTION)
    input_hashes_before = {path.name: sha256(path) for path in inputs}
    tasks = cases["tasks"]

    with tempfile.TemporaryDirectory(prefix="pytools-sheet-extract-baseline-") as temporary:
        temp_root = Path(temporary)
        output_dir = temp_root / "output"
        log_dir = temp_root / "logs"
        stats = run_extract_sheets(
            inputs,
            tasks,
            output_dir,
            log_dir,
            append_existing=False,
            default_output_name="unused-default.xlsx",
        )
        logging.shutdown()
        output_path = output_dir / "区域提取" / cases["expected"]["output_file"]
        if not output_path.exists():
            raise AssertionError(f"Python baseline did not create {output_path}")
        workbook = load_workbook(output_path, data_only=True, read_only=True)
        worksheets = []
        for worksheet in workbook.worksheets:
            values = [
                [normalized_cell(value) for value in row]
                for row in worksheet.iter_rows(
                    min_row=1,
                    max_row=worksheet.max_row,
                    min_col=1,
                    max_col=worksheet.max_column,
                    values_only=True,
                )
            ]
            worksheets.append(
                {
                    "name": worksheet.title,
                    "max_row": worksheet.max_row,
                    "max_column": worksheet.max_column,
                    "values": values,
                }
            )
        workbook.close()

    expected = cases["expected"]
    for key in ("files", "extracted_sheets", "output_files", "failed_files"):
        if stats[key] != expected[key]:
            raise AssertionError(f"Python baseline {key}={stats[key]}, expected {expected[key]}")
    if [Path(path).name for path in stats["paths"]] != [expected["output_file"]]:
        raise AssertionError(f"unexpected Python output paths: {stats['paths']!r}")
    if sha256(SOURCE_FUNCTION) != source_hash_before:
        raise AssertionError("read-only Python source changed while generating the baseline")
    input_hashes_after = {path.name: sha256(path) for path in inputs}
    if input_hashes_after != input_hashes_before:
        raise AssertionError("Python baseline changed an input fixture")

    REPORT_PATH.parent.mkdir(parents=True, exist_ok=True)
    report = {
        "schema_version": 1,
        "status": "passed",
        "suite": "pytools.sheet_extract_area_python_baseline",
        "baseline": "PyCode/core/feature_1_6_extract_sheets.py::run_extract_sheets",
        "python_source_sha256": source_hash_before,
        "cases_sha256": sha256(CASES_PATH),
        "config_sha256": sha256(CONFIG_FIXTURE),
        "input_sha256": input_hashes_before,
        "stats": {key: stats[key] for key in ("files", "extracted_sheets", "output_files", "failed_files")},
        "output_files": [expected["output_file"]],
        "worksheets": worksheets,
        "input_unchanged": True,
        "source_unchanged": True,
        "limitations": [
            "synthetic XLSX fixture only",
            "area extraction and new workbook output only",
            "full-sheet style copying and append-existing are outside this Rust candidate",
            "formula text is not copied; the source also loads data_only values, but no formula fixture was included",
            "dates, durations, and error cells are explicitly rejected by this candidate",
        ],
    }
    REPORT_PATH.write_text(
        json.dumps(report, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    print(json.dumps({"status": "passed", "report": str(REPORT_PATH), "stats": report["stats"]}, ensure_ascii=False))


if __name__ == "__main__":
    main()
