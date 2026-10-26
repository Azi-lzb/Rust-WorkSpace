from __future__ import annotations

import argparse
import importlib
import json
import logging
import subprocess
import sys
import tempfile
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook

from compare_sheet_extract_area_durations import create_rust_config, project_workbook


def has_vba_project(path: Path) -> bool:
    with ZipFile(path) as archive:
        return "xl/vbaProject.bin" in archive.namelist()


def projection_counts(projection: dict[str, object]) -> dict[str, int]:
    sheets = projection["sheets"]
    assert isinstance(sheets, list)
    return {
        "sheets": len(sheets),
        "non_empty_cells": sum(len(sheet["cells"]) for sheet in sheets),
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--pytools-cli", type=Path, required=True)
    parser.add_argument("--python-project", type=Path, required=True)
    parser.add_argument("--temporary-root", type=Path, required=True)
    parser.add_argument("--report", type=Path, required=True)
    args = parser.parse_args()

    sys.dont_write_bytecode = True
    probe_root = Path(__file__).resolve().parent
    source_python = args.python_project.resolve() / "PyCode"
    sys.path.insert(0, str(source_python))
    legacy = importlib.import_module("core.feature_1_6_extract_sheets")
    source = probe_root.parent / "tests" / "fixtures" / "vba_config" / "VbaV2.4.2.xlsm"
    source_before = source.read_bytes()
    source_has_vba = has_vba_project(source)
    if not source_has_vba:
        raise AssertionError("the tracked XLSM fixture has no VBA project")

    source_book = load_workbook(source, data_only=True, read_only=True)
    try:
        source_sheet_count = len(source_book.worksheets)
    finally:
        source_book.close()

    args.temporary_root.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(
        prefix="pytools-real-xlsm-area-", dir=args.temporary_root
    ) as temporary:
        root = Path(temporary)
        python_output = root / "python-output"
        python_logs = root / "python-logs"
        rust_output = root / "rust-output"
        config_path = root / "config.xlsx"
        output_name = "real-xlsm-area.xlsx"
        output_folder = "xlsm-area"
        area_spec = "A1:XFD1048576"
        task = {
            "row_no": 2,
            "workbook_keywords": "",
            "sheet_keywords": "",
            "exclude_keywords": "",
            "area_spec": area_spec,
            "stitch_mode": "",
            "output_folder": output_folder,
            "output_name": output_name,
            "remark": "whole used range of real macro-enabled workbook fixture",
        }

        try:
            python_stats = legacy.run_extract_sheets(
                [source],
                [task],
                python_output,
                python_logs,
                append_existing=False,
                default_output_name="unused.xlsx",
            )
            create_rust_config(config_path, rust_output, 1)
            config_book = load_workbook(config_path)
            try:
                extraction = config_book["工作表提取"]
                extraction["E2"] = area_spec
                extraction["G2"] = output_folder
                extraction["H2"] = output_name
                extraction["I2"] = task["remark"]
                config_book.save(config_path)
            finally:
                config_book.close()

            rust_process = subprocess.run(
                [
                    str(args.pytools_cli.resolve()),
                    "extract-sheets-area",
                    "--config",
                    str(config_path),
                    "--output-dir",
                    str(rust_output),
                    str(source),
                ],
                capture_output=True,
                text=True,
                check=False,
            )
            if rust_process.returncode != 0:
                raise RuntimeError(
                    f"Rust CLI failed ({rust_process.returncode}): {rust_process.stderr}"
                )
            rust_stats = json.loads(rust_process.stdout)
            python_path = python_output / output_folder / output_name
            rust_path = rust_output / output_folder / output_name
            if not python_path.is_file() or not rust_path.is_file():
                raise AssertionError("one or both extraction outputs are missing")

            python_projection = project_workbook(python_path)
            rust_projection = project_workbook(rust_path)
            exact_match = python_projection == rust_projection
            if not exact_match:
                raise AssertionError(
                    "the full workbook value/type/coordinate/number-format projections differ"
                )

            python_vba_in_output = has_vba_project(python_path)
            rust_vba_in_output = has_vba_project(rust_path)
            source_unchanged = source.read_bytes() == source_before
            if not source_unchanged:
                raise AssertionError("the source XLSM changed during the read-only comparison")
            if python_vba_in_output or rust_vba_in_output:
                raise AssertionError("area extraction unexpectedly retained a VBA project")
            if python_stats["failed_files"] != 0 or rust_stats["failed_files"] != 0:
                raise AssertionError("an extraction side reported failed input files")

            report = {
                "schema_version": 1,
                "status": "passed",
                "feature": "PT-7-6 XLSM area extraction",
                "sample_kind": (
                    "real macro-enabled workbook fixture from the related VBA tool suite; "
                    "not an original PyTools report workbook"
                ),
                "source_extension": source.suffix.lower(),
                "source_sheet_count": source_sheet_count,
                "source_vba_project_present": source_has_vba,
                "area_spec": area_spec,
                "python_stats": {
                    key: value for key, value in python_stats.items() if key != "paths"
                },
                "rust_stats": {
                    key: value for key, value in rust_stats.items() if key != "paths"
                },
                "output_projection_exact_match": exact_match,
                "output_projection_counts": projection_counts(python_projection),
                "python_output_vba_project_present": python_vba_in_output,
                "rust_output_vba_project_present": rust_vba_in_output,
                "source_unchanged": source_unchanged,
                "scope_note": (
                    "This confirms full used-range value extraction parity on one real macro-enabled "
                    "fixture, not acceptance on an original PyTools XLSM report or VBA preservation. "
                    "The output contract is a new XLSX without VBA; other real templates and platform "
                    "delivery remain unverified."
                ),
            }
            args.report.parent.mkdir(parents=True, exist_ok=True)
            args.report.write_text(
                json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
            )
            print(json.dumps(report, ensure_ascii=False, indent=2))
        finally:
            logging.shutdown()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
