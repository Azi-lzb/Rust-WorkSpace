from __future__ import annotations

import argparse
import importlib
import json
import logging
import subprocess
import sys
import tempfile
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.utils.datetime import MAC_EPOCH

from compare_sheet_extract_area_durations import create_rust_config, project_workbook


_CONTENT_TYPES_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
_MACRO_ENABLED_WORKBOOK_TYPE = "application/vnd.ms-excel.sheet.macroEnabled.main+xml"


def create_macro_enabled_xlsm(source_xlsx: Path, output_xlsm: Path) -> None:
    with ZipFile(source_xlsx) as source_archive:
        parts = [
            (item, source_archive.read(item.filename))
            for item in source_archive.infolist()
        ]
    content_types = next(
        data for item, data in parts if item.filename == "[Content_Types].xml"
    )
    root = ET.fromstring(content_types)
    workbook_override = next(
        (
            node
            for node in root.findall(f"{{{_CONTENT_TYPES_NS}}}Override")
            if node.get("PartName") == "/xl/workbook.xml"
        ),
        None,
    )
    if workbook_override is None:
        raise ValueError("synthetic XLSX has no workbook content-type override")
    workbook_override.set("ContentType", _MACRO_ENABLED_WORKBOOK_TYPE)
    content_types = ET.tostring(root, encoding="utf-8", xml_declaration=True)

    with ZipFile(output_xlsm, "w", compression=ZIP_DEFLATED) as output_archive:
        for item, data in parts:
            output_archive.writestr(
                item,
                content_types if item.filename == "[Content_Types].xml" else data,
            )

def create_source(path: Path, date_1904: bool) -> bytes:
    workbook = Workbook()
    if date_1904:
        workbook.epoch = MAC_EPOCH
    worksheet = workbook.active
    worksheet.column_dimensions["B"].number_format = "h:mm:ss"
    worksheet["B1"] = 0.5
    cases = [
        (0.0, "h:mm:ss"),
        (0.5 / 86_400_000, "h:mm:ss"),
        (1.5 / 86_400_000, "h:mm:ss"),
        (86_399_999.5 / 86_400_000, "h:mm:ss"),
        (0.5, "yyyy-mm-dd"),
        (0.5, "h:mm"),
        (0.5, "[h]:mm:ss"),
    ]
    for row, (value, number_format) in enumerate(cases, start=1):
        worksheet.cell(row=row, column=1, value=value).number_format = number_format
    if path.suffix.lower() == ".xlsm":
        xlsx_path = path.with_suffix(".xlsx")
        workbook.save(xlsx_path)
        workbook.close()
        create_macro_enabled_xlsm(xlsx_path, path)
        xlsx_path.unlink()
    else:
        workbook.save(path)
        workbook.close()
    return path.read_bytes()


def macro_enabled_without_vba_project(path: Path) -> bool:
    with ZipFile(path) as archive:
        content_types = archive.read("[Content_Types].xml")
        parts = set(archive.namelist())
    return (
        b"macroEnabled.main+xml" in content_types
        and "xl/vbaProject.bin" not in parts
    )

def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--pytools-cli", type=Path, required=True)
    parser.add_argument("--python-project", type=Path, required=True)
    parser.add_argument("--report", type=Path)
    args = parser.parse_args()

    sys.dont_write_bytecode = True
    probe_root = Path(__file__).resolve().parent
    report_path = args.report or (
        probe_root / "reports" / "sheet_extract_area_time_only_parity_20261026.json"
    )
    source_python = args.python_project.resolve() / "PyCode"
    sys.path.insert(0, str(source_python))
    legacy = importlib.import_module("core.feature_1_6_extract_sheets")

    with tempfile.TemporaryDirectory(prefix="pytools-time-only-blocker-") as temporary:
        root = Path(temporary)
        sources = [
            root / "time-only-1900.xlsx",
            root / "time-only-1900-macro.xlsm",
            root / "time-only-1904.xlsx",
            root / "time-only-1904-macro.xlsm",
        ]
        source_snapshots = {
            source: create_source(source, date_1904="1904" in source.stem)
            for source in sources
        }
        xlsm_sources = [source for source in sources if source.suffix.lower() == ".xlsm"]
        xlsm_packages_valid = bool(xlsm_sources) and all(
            macro_enabled_without_vba_project(source) for source in xlsm_sources
        )

        python_output = root / "python-output"
        python_logs = root / "python-logs"
        rust_output = root / "rust-output"
        config_path = root / "config.xlsx"
        task = {
            "row_no": 2,
            "workbook_keywords": "",
            "sheet_keywords": "",
            "exclude_keywords": "",
            "area_spec": "A1:B7",
            "stitch_mode": "",
            "output_folder": "时长提取",
            "output_name": "duration-output.xlsx",
            "remark": "1900/1904 epoch time-only parity",
        }

        try:
            python_result = legacy.run_extract_sheets(
                sources,
                [task],
                python_output,
                python_logs,
                append_existing=False,
                default_output_name="unused.xlsx",
            )
            create_rust_config(config_path, rust_output, 7)
            config_workbook = load_workbook(config_path)
            config_workbook.worksheets[1]["E2"] = "A1:B7"
            config_workbook.save(config_path)
            config_workbook.close()
            rust_process = subprocess.run(
                [
                    str(args.pytools_cli.resolve()),
                    "extract-sheets-area",
                    "--config",
                    str(config_path),
                    "--output-dir",
                    str(rust_output),
                    *(str(source) for source in sources),
                ],
                capture_output=True,
                text=True,
                check=False,
            )

            python_path = python_output / "时长提取" / "duration-output.xlsx"
            rust_path = rust_output / "时长提取" / "duration-output.xlsx"
            python_projection = project_workbook(python_path)
            sources_unchanged = all(
                source.read_bytes() == source_snapshots[source] for source in sources
            )
            if rust_process.returncode == 0 and rust_path.is_file():
                rust_projection = project_workbook(rust_path)
                rust_stats = json.loads(rust_process.stdout)
            else:
                rust_projection = None
                rust_stats = None
            status = (
                "passed"
                if python_result["failed_files"] == 0
                and rust_process.returncode == 0
                and rust_stats is not None
                and rust_stats["failed_files"] == 0
                and python_projection == rust_projection
                and sources_unchanged
                and xlsm_packages_valid
                else "failed"
            )

            report = {
                "schema_version": 1,
                "status": status,
                "feature": "PT-7-6 area extraction time-only value classification",
                "sample_kind": (
                    "synthetic XLSX and macro-enabled XLSM files (without a VBA project) in "
                    "1900/1904 epochs with rounding, date-only, time, elapsed-duration and column styles"
                ),
                "source_python": "PyCode/core/feature_1_6_extract_sheets.py::run_extract_sheets",
                "rust_entry": "pytools extract-sheets-area",
                "python_stats": {
                    key: value for key, value in python_result.items() if key != "paths"
                },
                "rust_returncode": rust_process.returncode,
                "rust_stats": (
                    {key: value for key, value in rust_stats.items() if key != "paths"}
                    if rust_stats is not None
                    else None
                ),
                "python_projection": python_projection,
                "rust_error": rust_process.stderr.strip(),
                "rust_output_created": rust_path.is_file(),
                "rust_projection": rust_projection,
                "sources_unchanged": sources_unchanged,
                "xlsm_packages_macro_enabled_without_vba": xlsm_packages_valid,
                "scope_note": (
                    "Synthetic macro-enabled XLSM inputs have no VBA project. Real business time-only cells, "
                    "VBA-bearing XLSM parity, row default styles, unknown built-in formats, other input formats, "
                    "and target-platform delivery remain unverified. Unsupported metadata must fail closed."
                ),
            }
            report_path.parent.mkdir(parents=True, exist_ok=True)
            report_path.write_text(
                json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
            )
            print(json.dumps(report, ensure_ascii=False, indent=2))
            return 0 if status == "passed" else 1
        finally:
            logging.shutdown()


if __name__ == "__main__":
    raise SystemExit(main())
