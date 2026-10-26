from __future__ import annotations

import json
import math
import platform
import shutil
import subprocess
import tempfile
from datetime import date
from datetime import datetime
from pathlib import Path

import openpyxl
import xlrd
from openpyxl import Workbook, load_workbook

REPO = Path(__file__).resolve().parents[2]
FIXTURE = REPO / "pytools-core" / "tests" / "fixtures" / "timeline" / "wide_summary_legacy.xls"
RUST_CLI = REPO / "target" / "debug" / "pytools.exe"
REPORT = (
    REPO
    / "pytools-core"
    / "probe"
    / "reports"
    / "native_excel_value_conversion_parity_20261026.json"
)


def make_config(path: Path, root: Path) -> None:
    workbook = Workbook()
    global_sheet = workbook.active
    global_sheet.title = "全局配置"
    global_sheet.append(["配置项", "值"])
    for key, value in (
        ("输入目录", str(root)),
        ("输出目录", str(root / "output")),
        ("日志目录", str(root / "logs")),
        ("excel目的格式", "xlsx"),
        ("word目的格式", "docx"),
    ):
        global_sheet.append([key, value])
    workbook.save(path)


def compare_values(source_path: Path, output_path: Path) -> dict[str, object]:
    source = xlrd.open_workbook(source_path, on_demand=True)
    output = load_workbook(output_path, data_only=False, read_only=False)
    try:
        source_names = source.sheet_names()
        if source_names != output.sheetnames:
            raise AssertionError(
                f"worksheet names/order differ: Python={source_names}, Rust={output.sheetnames}"
            )

        nonempty_cells = 0
        numeric_cells = 0
        text_cells = 0
        for source_sheet, output_sheet in zip(source.sheets(), output.worksheets):
            expected: dict[str, tuple[int, object]] = {}
            for row_index in range(source_sheet.nrows):
                for column_index in range(source_sheet.ncols):
                    cell = source_sheet.cell(row_index, column_index)
                    if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                        continue
                    coordinate = f"{xlrd.colname(column_index)}{row_index + 1}"
                    expected[coordinate] = (cell.ctype, cell.value)

            actual = {
                cell.coordinate: cell
                for row in output_sheet.iter_rows()
                for cell in row
                if cell.value is not None
            }
            if set(expected) != set(actual):
                raise AssertionError(
                    f"non-empty cell coordinates differ on {source_sheet.name}: "
                    f"Python={sorted(expected)}, Rust={sorted(actual)}"
                )

            for coordinate, (source_type, source_value) in expected.items():
                target_cell = actual[coordinate]
                if source_type == xlrd.XL_CELL_TEXT:
                    if not isinstance(target_cell.value, str) or target_cell.value != source_value:
                        raise AssertionError(
                            f"text differs at {source_sheet.name}!{coordinate}: "
                            f"Python={source_value!r}, Rust={target_cell.value!r}"
                        )
                    text_cells += 1
                elif source_type == xlrd.XL_CELL_NUMBER:
                    if isinstance(target_cell.value, bool) or not isinstance(
                        target_cell.value, (int, float)
                    ):
                        raise AssertionError(
                            f"numeric type differs at {source_sheet.name}!{coordinate}: "
                            f"Rust={type(target_cell.value).__name__}"
                        )
                    if not math.isclose(
                        float(source_value), float(target_cell.value), rel_tol=0.0, abs_tol=0.0
                    ):
                        raise AssertionError(
                            f"numeric value differs at {source_sheet.name}!{coordinate}: "
                            f"Python={source_value!r}, Rust={target_cell.value!r}"
                        )
                    numeric_cells += 1
                else:
                    raise AssertionError(
                        f"fixture requires an explicit comparison for xlrd cell type "
                        f"{source_type} at {source_sheet.name}!{coordinate}"
                    )
                nonempty_cells += 1

        return {
            "scope": "Python xlrd source values compared with the Rust-generated XLSX reopened by openpyxl; this does not compare the legacy PyCode fidelity-conversion result.",
            "worksheet_names_and_order_equal": True,
            "nonempty_cell_coordinates_equal": True,
            "text_values_equal": True,
            "numeric_values_equal": True,
            "worksheets": len(source_names),
            "nonempty_cells": nonempty_cells,
            "text_cells": text_cells,
            "numeric_cells": numeric_cells,
            "numeric_type_projection": "XLS integer storage is read by xlrd as NUMBER and emitted by Rust as an XLSX numeric cell; values compare exactly.",
        }
    finally:
        source.release_resources()
        output.close()


def main() -> int:
    for required in (FIXTURE, RUST_CLI):
        if not required.is_file():
            raise FileNotFoundError(required)

    source_before = FIXTURE.read_bytes()
    with tempfile.TemporaryDirectory(prefix="pytools-native-xls-parity-") as temp_name:
        output_dir = Path(temp_name)
        config = output_dir / "config.xlsx"
        make_config(config, output_dir)
        result = subprocess.run(
            [
                str(RUST_CLI),
                "convert-excel",
                "--config",
                str(config),
                "--to",
                "xlsx",
                "--engine",
                "native",
                "--output-dir",
                str(output_dir),
                "--log-dir",
                str(output_dir / "logs"),
                "--log-view",
                "developer",
                str(FIXTURE),
            ],
            check=False,
            capture_output=True,
        )
        stdout = result.stdout.decode("utf-8-sig", errors="replace").replace("\x00", "")
        stderr = result.stderr.decode("utf-8-sig", errors="replace").replace("\x00", "")
        if result.returncode != 0:
            raise RuntimeError(f"native conversion failed ({result.returncode}): {stderr[-2000:]}")
        rust_report = json.loads(stdout)
        if rust_report.get("converted") != 1 or rust_report.get("failed") != 0:
            raise AssertionError(f"unexpected Rust conversion result: {rust_report}")
        history_path = Path(rust_report["run_history_path"])
        history = json.loads(history_path.read_text(encoding="utf-8"))
        conversion_event = next(
            (
                event
                for event in history.get("events", [])
                if event.get("phase") == "convert_file" and event.get("status") == "completed"
            ),
            None,
        )
        if not conversion_event or conversion_event.get("actual_engine") != "calamine + rust_xlsxwriter":
            raise AssertionError(f"unexpected conversion engine history: {history}")
        history_text = history_path.read_text(encoding="utf-8")
        if FIXTURE.name.lower() in history_text.lower():
            raise AssertionError("developer run history contains the input file name")
        output_path = Path(rust_report["items"][0]["output"])
        comparison = compare_values(FIXTURE, output_path)
        source_unchanged = FIXTURE.read_bytes() == source_before
        if not source_unchanged:
            raise AssertionError("the synthetic XLS source was modified")

        iso_source = output_dir / "iso-date.xlsm"
        iso_workbook = Workbook()
        iso_workbook.iso_dates = True
        iso_sheet = iso_workbook.active
        iso_sheet.title = "日期"
        iso_sheet["A1"] = datetime(2024, 2, 29, 12, 30)
        iso_workbook.save(iso_source)
        iso_source_before = iso_source.read_bytes()
        iso_result = subprocess.run(
            [
                str(RUST_CLI),
                "convert-excel",
                "--config",
                str(config),
                "--to",
                "xlsx",
                "--engine",
                "native",
                "--output-dir",
                str(output_dir),
                "--log-dir",
                str(output_dir / "logs"),
                "--log-view",
                "developer",
                str(iso_source),
            ],
            check=False,
            capture_output=True,
        )
        iso_stdout = iso_result.stdout.decode("utf-8-sig", errors="replace").replace("\x00", "")
        if iso_result.returncode != 2:
            raise AssertionError(
                f"ISO-typed date input must fail closed; exit={iso_result.returncode}, "
                f"stdout={iso_stdout[-1000:]!r}"
            )
        iso_report = json.loads(iso_stdout)
        if iso_report.get("converted") != 0 or iso_report.get("failed") != 1:
            raise AssertionError(f"unexpected ISO date refusal result: {iso_report}")
        iso_output = output_dir / "iso-date.xlsx"
        if iso_output.exists():
            raise AssertionError("ISO-typed date refusal published an output workbook")
        if iso_source.read_bytes() != iso_source_before:
            raise AssertionError("the ISO date source workbook was modified")
        iso_failure_detail = str(iso_report["items"][0].get("detail", ""))
        if "ISO-typed date/time" not in iso_failure_detail or "cell type" not in iso_failure_detail:
            raise AssertionError(f"ISO date refusal detail is incomplete: {iso_failure_detail!r}")

    report = {
        "schema_version": 1,
        "suite": "pytools.conversion.native_excel_value_only",
        "date": date.today().isoformat(),
        "status": "passed_synthetic_python_rust_value_projection_and_type_guard",
        "environment": {
            "os": platform.platform(),
            "rust_entry": "Windows-native pytools convert-excel --engine native",
            "python_baseline": f"xlrd {xlrd.__version__} source reader + openpyxl {openpyxl.__version__} output reader",
        },
        "input": {
            "format": "xls",
            "kind": "synthetic smoke workbook",
            "fixture": "pytools-core/tests/fixtures/timeline/wide_summary_legacy.xls",
        },
        "output": {
            "format": "xlsx",
            "requested_engine": "native",
            "actual_engine": conversion_event["actual_engine"],
            "actual_components": conversion_event.get("components"),
            "run_history_redacted": True,
            "source_unchanged": source_unchanged,
        },
        "comparison": comparison,
        "type_guard": {
            "input": "synthetic .xlsm containing an ISO-typed Excel date generated with openpyxl iso_dates=True",
            "rust_result": "conversion refused before publishing because rust_xlsxwriter output would change the cell type to text",
            "converted": iso_report["converted"],
            "failed": iso_report["failed"],
            "output_published": False,
            "source_unchanged": True,
        },
        "unverified": [
            "Real business XLS workbooks and configurations.",
            "Formula text/cache parity against legacy XLS conversion, numeric serial date/time cells, booleans, error cells and empty-string cells.",
            "Styles, comments, merged regions, dimensions, names, print settings, validations, conditional formatting, drawings and macros are intentionally not preserved by this value-only route.",
            "Windows desktop file dialog, glibc 2.28, Kylin V10/UOS V20 and offline deb delivery.",
        ],
    }
    REPORT.parent.mkdir(parents=True, exist_ok=True)
    REPORT.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(
        json.dumps(
            {"status": report["status"], "report": str(REPORT), "comparison": comparison},
            ensure_ascii=False,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
