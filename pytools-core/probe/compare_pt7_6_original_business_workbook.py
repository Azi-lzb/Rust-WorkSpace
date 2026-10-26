from __future__ import annotations

import argparse
import importlib
import json
import logging
import re
import subprocess
import sys
from datetime import date, datetime, time, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook


HEADERS = [
    "是否启用",
    "工作簿关键字",
    "工作表关键字",
    "排除关键字",
    "提取区域",
    "多区域拼接方式",
    "输出文件夹",
    "文件名",
    "备注",
]
AREA_SPEC = "A1:XFD1048576"


def project_cell(cell: object) -> dict[str, str] | None:
    value = cell.value
    if value is None:
        return None
    if cell.data_type == "e":
        return {"type": "error", "value": str(value)}
    if isinstance(value, datetime):
        return {"type": "datetime", "value": value.isoformat(timespec="milliseconds")}
    if isinstance(value, date):
        return {"type": "date", "value": value.isoformat()}
    if isinstance(value, time):
        return {"type": "time", "value": value.isoformat(timespec="milliseconds")}
    if isinstance(value, timedelta):
        return {"type": "timedelta", "value": str(value)}
    return {"type": type(value).__name__, "value": str(value)}


def project_workbook(path: Path) -> dict[str, object]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        sheets = []
        for worksheet in workbook.worksheets:
            cells = []
            for row in worksheet.iter_rows():
                for cell in row:
                    projected = project_cell(cell)
                    if projected is not None:
                        cells.append(
                            {
                                "cell": cell.coordinate,
                                **projected,
                                "number_format": cell.number_format,
                            }
                        )
            sheets.append({"name": worksheet.title, "cells": cells})
        return {"sheet_names": workbook.sheetnames, "sheets": sheets}
    finally:
        workbook.close()


def create_rust_config(path: Path, output_dir: Path, output_folder: str, output_name: str) -> None:
    workbook = Workbook()
    global_config = workbook.active
    global_config.title = "全局配置"
    global_config.append(["配置项", "值"])
    global_config.append(["输入目录", ""])
    global_config.append(["输出目录", str(output_dir)])
    global_config.append(["日志目录", str(output_dir / "logs")])

    extraction = workbook.create_sheet("工作表提取")
    extraction.append(HEADERS)
    extraction.append(
        [
            "是",
            "",
            "",
            "",
            AREA_SPEC,
            "",
            output_folder,
            output_name,
            "PT-7-6 original PyTools business workbook",
        ]
    )
    workbook.save(path)
    workbook.close()


def projection_summary(projection: dict[str, object]) -> dict[str, object]:
    sheets = projection["sheets"]
    type_counts: dict[str, int] = {}
    error_cells = 0
    non_empty_cells = 0
    for sheet in sheets:
        cells = sheet["cells"]
        non_empty_cells += len(cells)
        for cell in cells:
            cell_type = cell["type"]
            type_counts[cell_type] = type_counts.get(cell_type, 0) + 1
            if cell_type == "error":
                error_cells += 1
    return {
        "worksheet_count": len(sheets),
        "non_empty_cells": non_empty_cells,
        "cell_type_counts": dict(sorted(type_counts.items())),
        "error_cells": error_cells,
    }


def compare_projections(
    python_projection: dict[str, object], rust_projection: dict[str, object]
) -> tuple[bool, list[str]]:
    differences = []
    if python_projection["sheet_names"] != rust_projection["sheet_names"]:
        return False, ["worksheet names/order"]
    for python_sheet, rust_sheet in zip(
        python_projection["sheets"], rust_projection["sheets"], strict=True
    ):
        if python_sheet["name"] != rust_sheet["name"]:
            differences.append("worksheet name/order")
            continue
        python_cells = {cell["cell"]: cell for cell in python_sheet["cells"]}
        rust_cells = {cell["cell"]: cell for cell in rust_sheet["cells"]}
        if python_cells.keys() != rust_cells.keys():
            differences.append(f"{python_sheet['name']}: coordinates")
        for address in python_cells.keys() & rust_cells.keys():
            if python_cells[address] != rust_cells[address]:
                differences.append(f"{python_sheet['name']}!{address}")
    return not differences, differences[:20]


def find_run_log(log_dir: Path) -> dict[str, object]:
    candidates = []
    for path in log_dir.rglob("*.run.json"):
        try:
            report = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            continue
        if report.get("feature_id") == "PT-7-6":
            candidates.append(report)
    if not candidates:
        raise RuntimeError("Rust menu did not create a PT-7-6 developer run log")
    return candidates[-1]


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--pytools-cli", type=Path, required=True)
    parser.add_argument("--python-project", type=Path, required=True)
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--output-root", type=Path, required=True)
    parser.add_argument("--report", type=Path, required=True)
    args = parser.parse_args()

    if not args.source.is_file():
        raise FileNotFoundError(args.source)
    if args.output_root.exists() and any(args.output_root.iterdir()):
        raise FileExistsError(f"refusing to reuse non-empty output directory {args.output_root}")
    args.output_root.mkdir(parents=True, exist_ok=True)
    source_before = args.source.read_bytes()

    sys.dont_write_bytecode = True
    source_python = args.python_project.resolve() / "PyCode"
    sys.path.insert(0, str(source_python))
    legacy = importlib.import_module("core.feature_1_6_extract_sheets")

    python_output = args.output_root / "python-output"
    python_logs = args.output_root / "python-logs"
    rust_output = args.output_root / "rust-output"
    config_path = args.output_root / "rust-menu-config.xlsx"
    output_folder = "PT-7-6区域提取"
    output_name = "1.报表分口径_区域提取.xlsx"
    task = {
        "row_no": 2,
        "workbook_keywords": "",
        "sheet_keywords": "",
        "exclude_keywords": "",
        "area_spec": AREA_SPEC,
        "stitch_mode": "",
        "output_folder": output_folder,
        "output_name": output_name,
        "remark": "PT-7-6 original PyTools business workbook",
    }

    try:
        python_stats = legacy.run_extract_sheets(
            [args.source],
            [task],
            python_output,
            python_logs,
            append_existing=False,
            default_output_name="unused.xlsx",
        )
        create_rust_config(config_path, rust_output, output_folder, output_name)
        rust_process = subprocess.run(
            [
                str(args.pytools_cli.resolve()),
                "menu",
                "--config",
                str(config_path),
                "--log-view",
                "developer",
            ],
            input=f"7\n6\nn\n{args.source.resolve()}\n",
            capture_output=True,
            encoding="utf-8",
            errors="replace",
            check=False,
        )
        if rust_process.returncode != 0:
            raise RuntimeError(f"Rust PyTools menu exited with {rust_process.returncode}")

        python_path = python_output / output_folder / output_name
        rust_path = rust_output / output_folder / output_name
        if not python_path.is_file() or not rust_path.is_file():
            raise FileNotFoundError("one or both Python/Rust extraction outputs are missing")
        python_projection = project_workbook(python_path)
        rust_projection = project_workbook(rust_path)
        exact_match, difference_locations = compare_projections(
            python_projection, rust_projection
        )
        python_summary = projection_summary(python_projection)
        rust_summary = projection_summary(rust_projection)
        run_log = find_run_log(rust_output / "logs")
        events = run_log.get("events", [])
        merge_events = [
            event for event in events if event.get("phase") == "merge_value_compatibility"
        ]
        patch_events = [event for event in events if event.get("phase") == "ooxml_error_patch"]
        if not merge_events:
            raise AssertionError("Rust menu did not record merged-cell compatibility handling")
        merge_match = re.search(
            r"忽略 (\d+) 个合并子单元格值", str(merge_events[-1].get("user_message", ""))
        )
        if merge_match is None:
            raise AssertionError("Rust menu merge compatibility event omitted its cell count")
        source_unchanged = args.source.read_bytes() == source_before

        if python_stats.get("failed_files") != 0:
            raise AssertionError("Python baseline reported failed source files")
        if run_log.get("status") != "succeeded":
            raise AssertionError("Rust menu run log did not report success")
        if python_summary["error_cells"] == 0:
            raise AssertionError("the selected original business sample no longer contains errors")
        if not patch_events or patch_events[-1].get("status") != "succeeded":
            raise AssertionError("Rust menu did not report successful Excel error-cell preservation")
        if not exact_match:
            raise AssertionError(
                f"Python/Rust projection differs at {len(difference_locations)} sampled locations"
            )
        if not source_unchanged:
            raise AssertionError("the original PyTools source workbook changed")

        report = {
            "schema_version": 1,
            "status": "passed",
            "feature": "PT-7-6 original PyTools business workbook area extraction",
            "sample_kind": "original PyTools business workbook; read-only source",
            "source_basename": args.source.name,
            "source_bytes": len(source_before),
            "area_spec": AREA_SPEC,
            "entry": "Rust PyTools CLI menu → 7 → 6",
            "python_baseline": "PyCode/core/feature_1_6_extract_sheets.py::run_extract_sheets",
            "python_stats": {
                key: value for key, value in python_stats.items() if key != "paths"
            },
            "rust_run_status": run_log.get("status"),
            "rust_engines": sorted(
                {
                    event.get("actual_engine")
                    for event in events
                    if isinstance(event, dict) and event.get("actual_engine")
                }
            ),
            "rust_merge_value_compatibility": {
                "actual_engine": merge_events[-1].get("actual_engine"),
                "stored_merged_child_values_suppressed": int(merge_match.group(1)),
            },
            "rust_error_patch_succeeded": bool(
                patch_events and patch_events[-1].get("status") == "succeeded"
            ),
            "rust_error_cells_preserved": rust_summary["error_cells"],
            "python_projection": python_summary,
            "rust_projection": rust_summary,
            "exact_output_projection_match": exact_match,
            "difference_locations": difference_locations,
            "source_unchanged": source_unchanged,
            "output_location": "local F: drive test folder; excluded from Git",
            "scope_note": (
                "This closes one original PyTools workbook comparison for the configured full-area "
                "value-extraction task. It does not complete all PT-7-6 templates, Windows manual "
                "acceptance, or Linux/deb validation."
            ),
        }
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(
            json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
        )
        print(json.dumps(report, ensure_ascii=False, indent=2))
    finally:
        logging.shutdown()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
