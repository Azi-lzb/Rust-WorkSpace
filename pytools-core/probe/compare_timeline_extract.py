"""Compare PT-1-2 Rust extraction with the read-only PyCode implementation."""

from __future__ import annotations

import json
import os
import subprocess
import sys
import tempfile
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

WORKSPACE = Path(__file__).resolve().parents[2]
PYTOOLS_SOURCE = Path(r"C:\Users\AZI\Desktop\VibeCode\4.pytools-main")
sys.dont_write_bytecode = True
sys.path.insert(0, str(PYTOOLS_SOURCE))

from PyCode.core.config_xlsx import (  # noqa: E402
    PATH_MAP_COLS,
    TIMELINE_COLS,
    load_global,
    load_path_maps,
    load_timeline_rules,
)
from PyCode.core.timeline_396 import run_timeline_slim  # noqa: E402


def make_source(path: Path, month: int) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "本外币"
    sheet.append([f"2026年{month}月本外币存款", None, None])
    sheet.append(["项目", "本月余额", "比上月"])
    sheet.append(["活期存款", 100 + month, 5])
    sheet.append(["活期存款", 110 + month, 6])
    sheet.append(
        ["定期存款", 200 + month, datetime(2026, month, 15, 12, 30, 0, 123000)]
    )
    sheet.append(["活期存款", "文本余额", True])
    sheet.append(["合计", 410 + month * 3, 18])

    appendix = workbook.create_sheet("本外币附表")
    appendix.append([f"2026年{month}月附表", None, None])
    appendix.append(["项目", "本月余额", "比上月"])
    appendix.append(["活期存款", 20 + month, 2])
    workbook.save(path)


def make_config(
    path: Path,
    input_dir: Path,
    output_dir: Path,
    log_dir: Path,
    target_path: Path | None = None,
    error_policy: str = "continue",
) -> None:
    import pandas as pd

    global_rows = [
        ("输入目录", str(input_dir), ""),
        ("输出目录", str(output_dir), ""),
        ("日志目录", str(log_dir), ""),
        ("错误策略", error_policy, ""),
        ("默认编码", "utf-8", ""),
        ("源文件扩展名", ".xlsx", ""),
        ("时序规则命中策略", "all_match", ""),
    ]
    rule = {column: "" for column in TIMELINE_COLS}
    rule.update(
        {
            "是否启用": "是",
            "规则名称": "活期提取",
            "工作簿关键字": "存款表",
            "工作表关键字": "本外币",
            "行头列": 1,
            "列表头行": "2",
            "必含列头": "本月余额",
            "必含行头": "活期",
            "数据起始行": 3,
            "数据起始列": 2,
            "跳过关键字": "合计",
            "目标工作簿路径": str(target_path) if target_path else "",
            "目标工作表": "追加目标" if target_path else "",
            "启用目标写入": "是" if target_path else "否",
        }
    )
    duplicate_rule = dict(rule)
    duplicate_rule["规则名称"] = "重复规则"
    path_map = [
        "是",
        "活期标准化",
        "",
        "存款表",
        "本外币",
        "行头",
        "精确",
        "活期存款",
        "活期",
        "",
    ]
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        pd.DataFrame(global_rows, columns=["键", "值", "备注"]).to_excel(
            writer, sheet_name="全局配置", index=False
        )
        pd.DataFrame([rule, duplicate_rule], columns=TIMELINE_COLS).to_excel(
            writer, sheet_name="时序提取规则", index=False
        )
        pd.DataFrame([path_map], columns=PATH_MAP_COLS).to_excel(
            writer, sheet_name="路径标准化映射", index=False
        )


def normalize_value(value):
    if isinstance(value, (date, datetime)):
        return [type(value).__name__, value.isoformat()]
    return value


def workbook_projection(path: Path) -> dict:
    workbook = load_workbook(path, data_only=False)
    sheets = []
    for sheet in workbook.worksheets:
        cells = []
        for row in sheet.iter_rows():
            row_cells = []
            for cell in row:
                if cell.value is None and not cell.has_style:
                    row_cells.append(None)
                else:
                    row_cells.append(
                        {
                            "value": normalize_value(cell.value),
                            "type": cell.data_type,
                            "bold": cell.font.bold,
                            "number_format": cell.number_format,
                        }
                    )
            cells.append(row_cells)
        sheets.append(
            {
                "name": sheet.title,
                "state": sheet.sheet_state,
                "dimension": sheet.calculate_dimension(),
                "freeze_panes": str(sheet.freeze_panes) if sheet.freeze_panes else None,
                "cells": cells,
            }
        )
    return {"sheetnames": workbook.sheetnames, "sheets": sheets}


def run() -> dict:
    rust_cli = Path(
        os.environ.get("PYTOOLS_RUST_CLI", WORKSPACE / "target" / "debug" / "pytools.exe")
    )
    if not rust_cli.is_file():
        raise FileNotFoundError(f"Rust CLI is missing: {rust_cli}")

    with tempfile.TemporaryDirectory(
        prefix="pytools-pt12-parity-", ignore_cleanup_errors=True
    ) as raw_temp:
        root = Path(raw_temp)
        source_dir = root / "sources"
        source_dir.mkdir()
        source_a = source_dir / "本外币存款表_202604.xlsx"
        source_b = source_dir / "本外币存款表_202605.xlsx"
        broken_source = source_dir / "本外币存款表_202606.xlsx"
        make_source(source_a, 4)
        make_source(source_b, 5)
        broken_source.write_bytes(b"synthetic invalid xlsx")
        config = root / "config" / "config.xlsx"
        python_output = root / "python-output"
        rust_output = root / "rust-output"
        menu_output = root / "menu-output"
        logs = root / "logs"
        make_config(config, source_dir, menu_output, logs)

        global_config = load_global(config)
        global_config.output_dir = python_output
        rules = load_timeline_rules(config)
        path_maps = load_path_maps(config)
        python_path = run_timeline_slim(
            rules,
            path_maps,
            global_config,
            source_paths=[source_a, source_b, broken_source],
        )
        if python_path is None or not python_path.is_file():
            raise AssertionError("PyCode did not produce the expected timeline output")

        command = [
            str(rust_cli),
            "timeline-extract",
            "--config",
            str(config),
            "--output-dir",
            str(rust_output),
            "--log-dir",
            str(logs),
            "--log-view",
            "developer",
            str(source_a),
            str(source_b),
            str(broken_source),
        ]
        rust_run = subprocess.run(
            command,
            check=False,
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
        )
        if rust_run.returncode != 0:
            raise AssertionError(
                f"Rust CLI exited {rust_run.returncode}: {rust_run.stderr}\n{rust_run.stdout}"
            )
        rust_result = json.loads(rust_run.stdout)
        rust_path = Path(rust_result["path"])
        python_projection = workbook_projection(python_path)
        rust_projection = workbook_projection(rust_path)
        if python_projection != rust_projection:
            raise AssertionError(
                json.dumps(
                    {"python": python_projection, "rust": rust_projection},
                    ensure_ascii=False,
                    default=str,
                )
            )

        event_lines = [
            line for line in rust_run.stderr.splitlines() if line.startswith("{")
        ]
        developer_events = [json.loads(line) for line in event_lines]
        run_history_path = Path(rust_result["run_log_path"])
        history = json.loads(run_history_path.read_text(encoding="utf-8"))
        if history["events"] != developer_events:
            raise AssertionError("CLI developer event stream differs from persisted history")
        if history["feature_id"] != "PT-1-2" or history["status"] != "partial":
            raise AssertionError("unexpected PT-1-2 run-history status")
        rows = len(python_projection["sheets"][0]["cells"]) - 1
        if rows != 20:
            raise AssertionError(f"expected 20 extracted rows, got {rows}")

        menu_input = f"1\n2\n{source_a};{source_b};{broken_source}\n0\n"
        menu_run = subprocess.run(
            [str(rust_cli), "menu", "--config", str(config), "--log-view", "off"],
            input=menu_input,
            check=False,
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
        )
        if menu_run.returncode != 0:
            raise AssertionError(
                f"menu exited {menu_run.returncode}: {menu_run.stderr}\n{menu_run.stdout}"
            )
        if "选择时序提取源文件" not in menu_run.stdout:
            raise AssertionError("the real menu did not route through PT-1-2 source selection")
        result_field = menu_run.stdout.find('"configured_target_rules"')
        marker = menu_run.stdout.rfind("{", 0, result_field) if result_field >= 0 else -1
        if marker < 0:
            raise AssertionError(
                "menu output did not include the timeline extraction result:\n"
                + menu_run.stdout
                + "\n"
                + menu_run.stderr
            )
        menu_result, _ = json.JSONDecoder().raw_decode(menu_run.stdout[marker:])
        menu_projection = workbook_projection(Path(menu_result["path"]))
        if python_projection != menu_projection:
            raise AssertionError("menu 1→2 result differs from the PyCode output")
        if menu_result["failed_files"] != 1:
            raise AssertionError("menu path did not preserve continue-on-error behavior")

        target_path = root / "existing-target.xlsx"
        target_workbook = Workbook()
        target_workbook.active.title = "保留工作表"
        target_workbook.active["A1"] = "must remain unchanged"
        target_workbook.save(target_path)
        target_before = workbook_projection(target_path)
        target_config = root / "target-config" / "config.xlsx"
        make_config(target_config, source_dir, root / "target-output", logs, target_path)
        deferred_run = subprocess.run(
            [
                str(rust_cli),
                "timeline-extract",
                "--config",
                str(target_config),
                "--log-view",
                "off",
                str(source_a),
                str(source_b),
            ],
            check=False,
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
        )
        if deferred_run.returncode != 0:
            raise AssertionError(f"target-contract check failed: {deferred_run.stderr}")
        deferred_result = json.loads(deferred_run.stdout)
        if not deferred_result["target_writes_deferred"]:
            raise AssertionError("configured target write was not surfaced as deferred")
        if workbook_projection(target_path) != target_before:
            raise AssertionError("deferred target write modified its target workbook")

        fail_fast_output = root / "fail-fast-output"
        fail_fast_logs = root / "fail-fast-logs"
        fail_fast_config = root / "fail-fast-config" / "config.xlsx"
        make_config(
            fail_fast_config,
            source_dir,
            fail_fast_output,
            fail_fast_logs,
            error_policy="fail_fast",
        )
        fail_fast_global = load_global(fail_fast_config)
        fail_fast_rules = load_timeline_rules(fail_fast_config)
        fail_fast_maps = load_path_maps(fail_fast_config)
        try:
            run_timeline_slim(
                fail_fast_rules,
                fail_fast_maps,
                fail_fast_global,
                source_paths=[broken_source, source_a],
            )
        except Exception:
            pass
        else:
            raise AssertionError("PyCode fail_fast baseline did not stop on the first error")
        python_fail_fast_outputs = list(fail_fast_output.glob("*.xlsx"))
        if python_fail_fast_outputs:
            raise AssertionError("PyCode fail_fast unexpectedly wrote an output workbook")

        fail_fast_run = subprocess.run(
            [
                str(rust_cli),
                "timeline-extract",
                "--config",
                str(fail_fast_config),
                "--log-view",
                "developer",
                str(broken_source),
                str(source_a),
            ],
            check=False,
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
        )
        if fail_fast_run.returncode == 0:
            raise AssertionError("Rust fail_fast did not return a failure status")
        fail_fast_history_paths = list(fail_fast_logs.glob("PT-1-2_*.run.json"))
        if len(fail_fast_history_paths) != 1:
            raise AssertionError("Rust fail_fast did not persist exactly one run history")
        fail_fast_history = json.loads(
            fail_fast_history_paths[0].read_text(encoding="utf-8")
        )
        if fail_fast_history["status"] != "failed":
            raise AssertionError("Rust fail_fast history did not record failure")
        if fail_fast_history["progress"]["processed_inputs"] != 1:
            raise AssertionError("Rust fail_fast processed inputs beyond the first failure")
        if list(fail_fast_output.glob("*.xlsx")):
            raise AssertionError("Rust fail_fast unexpectedly wrote an output workbook")

        return {
            "schema_version": 1,
            "suite": "pytools.timeline.extract_python_rust_parity",
            "status": "passed",
            "feature_id": "PT-1-2",
            "platform": "windows-x64-native",
            "inputs": 3,
            "valid_inputs": 2,
            "matched_rule_sheets": rust_result["matched_rule_sheets"],
            "output_rows": rust_result["rows"],
            "full_workbook_projection_equal": True,
            "checked": [
                "sheet name, cell values and types",
                "numeric, text, boolean, and datetime extraction values",
                "header bold and cell number formats",
                "worksheet dimensions and freeze panes",
                "duplicate-rule global deduplication",
                "row suffix and path mapping",
                "developer event stream and persisted history",
                "menu 1→2 manual multi-file selection and continue-on-error behavior",
                "configured target append is surfaced as deferred and leaves the target untouched",
                "fail_fast stops at the first failed workbook and does not create output",
            ],
            "rust_engines": history["actual_engines"],
            "elapsed_ms": history["total_elapsed_ms"],
            "deferred_target_write_rules": rust_result["configured_target_rules"],
            "differences": [],
            "notes": [
                "Synthetic workbooks only; no source-project workbooks were modified.",
                "Configured target append behavior remains outside the confirmed PT-1-2 target-write contract.",
            ],
        }


if __name__ == "__main__":
    report = run()
    output = WORKSPACE / "pytools-core" / "probe" / "reports" / "timeline_extract_parity_20261025.json"
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
