"""Synthetic Python/Rust semantic parity probe for the selected dedup operations."""

from __future__ import annotations

import contextlib
import datetime as dt
import json
import logging
import shlex
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment

sys.dont_write_bytecode = True
WORKSPACE = Path(__file__).resolve().parents[2]
PYTOOLS_SOURCE = Path(r"C:\Users\AZI\Desktop\VibeCode\4.pytools-main")


def wsl_path(path: Path) -> str:
    return f"/mnt/{path.drive.rstrip(':').lower()}/{path.as_posix()[2:].lstrip('/')}"


def from_wsl_path(value: str) -> Path:
    normalized = value.replace("\\", "/")
    if not normalized.startswith("/mnt/"):
        raise ValueError(f"Rust output outside WSL mounted drive: {value}")
    return Path(f"{normalized[5].upper()}:\\{normalized[7:].replace('/', chr(92))}")


def make_source(path: Path) -> None:
    wb = Workbook()
    data = wb.active
    data.title = "Data"
    data.append(["ID", "Amount", "Tag"])
    data.append(["k1", 10, "alpha"])
    data.append(["k1", 20, "beta"])
    data.append(["k2", 20, "gamma"])
    data.append(["k3", 20, "delta"])
    data.append([None, 20, "blank-id"])

    other = wb.create_sheet("Other")
    other.append(["Code", "Name"])
    other.append(["x", "one"])
    other.append(["x", "one"])
    other.append(["y", "two"])
    other.append([None, "nonblank-row"])

    comments = wb.create_sheet("Commented")
    comments.append(["源数据", "Key", "Value"])
    comments["A1"].comment = Comment("本表为源数据", "probe")
    comments["B1"].comment = Comment("标识列", "probe")
    comments.append(["row-1", "dup", "first"])
    comments.append(["row-2", "dup", "duplicate"])
    comments.append(["row-3", None, "blank-key"])
    comments.append(["row-4", "keep", "first"])
    comments.append(["row-5", "keep", "duplicate"])

    ignored = wb.create_sheet("Ignored")
    ignored.append(["ID", "Value"])
    ignored.append(["x", 1])
    ignored.append(["x", 1])
    ignored["A1"].comment = Comment("不是源数据", "probe")
    wb.save(path)


def make_config(path: Path, empty_source_task: bool = False) -> None:
    wb = Workbook()
    global_sheet = wb.active
    global_sheet.title = "全局配置"
    global_sheet.append(["配置项", "值"])
    global_sheet.append(["输入目录", "input"])
    global_sheet.append(["输出目录", "rust_output"])
    global_sheet.append(["日志目录", "logs"])

    tasks = wb.create_sheet("去重追加数据配置")
    tasks.append([
        "是否启用", "任务名", "源数据工作簿", "源数据工作表", "标识列序号",
        "目标工作簿", "目标工作表", "执行模式", "备注",
    ])
    rows = [
        ["是", "Data-ID", "source.xlsx", "Data", "1", "rust/target.xlsx", "Append", "追加", ""],
        ["是", "Data-Amount", "source.xlsx", "Data", "2", "rust/target.xlsx", "Append", "追加", ""],
        ["是", "Data-ID-duplicate", "source.xlsx", "Data", "1", "rust/target.xlsx", "Append", "追加", ""],
        ["是", "Other-full-row", "source.xlsx", "Other", "", "rust/target.xlsx", "OtherAppend", "追加", ""],
        ["是", "Missing-sheet", "source.xlsx", "Missing", "1", "rust/target.xlsx", "MissingAppend", "追加", ""],
    ]
    if empty_source_task:
        rows.append(["是", "Missing-source", "", "Data", "1", "rust/target.xlsx", "Append", "追加", ""])
    for row in rows:
        tasks.append(row)
    tasks.append(["否", "Disabled", "source.xlsx", "Data", "1", "rust/target.xlsx", "Append", "追加", ""])
    wb.save(path)
    wb.close()


def make_existing_append_source(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Source"
    ws.append(["ID", "Amount", "Derived"])
    ws.append(["A", 10, "=B2*2"])
    ws.append(["A", 20, "=B3*2"])
    ws.append(["B", 10, "=B4*2"])
    ws.append([None, 30, "=B5*2"])
    wb.save(path)
    wb.close()


def make_existing_append_target(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Combined"
    ws.append(["ID", "Amount", "Derived"])
    ws.append(["Z", 99, "=B2*2"])
    wb.save(path)
    wb.close()


def make_existing_append_config(path: Path, source: str, target: str) -> None:
    wb = Workbook()
    global_sheet = wb.active
    global_sheet.title = "全局配置"
    global_sheet.append(["配置项", "值"])
    global_sheet.append(["输入目录", "input"])
    global_sheet.append(["输出目录", "rust_output"])
    global_sheet.append(["日志目录", "logs"])
    tasks = wb.create_sheet("去重追加数据配置")
    tasks.append([
        "是否启用", "任务名", "源数据工作簿", "源数据工作表", "标识列序号",
        "目标工作簿", "目标工作表", "执行模式", "备注",
    ])
    tasks.append(["是", "Existing-target", source, "Source", "1", target, "Combined", "追加", ""])
    wb.save(path)
    wb.close()


def rust_json(arguments: list[str], allow_nonzero: bool = False) -> object:
    encoded = [
        shlex.quote(wsl_path(Path(value.removeprefix("WINPATH:"))))
        if value.startswith("WINPATH:")
        else shlex.quote(value)
        for value in arguments
    ]
    command = (
        f"cd {shlex.quote(wsl_path(WORKSPACE))} && cargo run --quiet -p pytools-core --bin pytools -- "
        + " ".join(encoded)
    )
    result = subprocess.run(["wsl", "-e", "sh", "-lc", command], capture_output=True, check=False)
    stdout = result.stdout.decode("utf-8-sig", errors="replace").replace("\x00", "")
    stderr = result.stderr.decode("utf-8-sig", errors="replace").replace("\x00", "")
    if result.returncode != 0 and not allow_nonzero:
        raise RuntimeError(f"Rust CLI failed ({result.returncode}):\n{stdout}\n{stderr}")
    decoder = json.JSONDecoder()
    for index, character in enumerate(stdout):
        if character not in "[{":
            continue
        try:
            value, _ = decoder.raw_decode(stdout[index:])
            return value
        except json.JSONDecodeError:
            continue
    raise RuntimeError(f"Rust CLI emitted no JSON: {stdout}\n{stderr}")


def cli_path(path: Path) -> str:
    return "WINPATH:" + str(path)


def rust_comment_plan(source: Path, output: Path, action: str) -> dict[str, object]:
    result = rust_json([
        "dedup-comment-plan",
        "--input", cli_path(source),
        "--output", cli_path(output),
        "--action", action,
    ])
    assert isinstance(result, dict)
    return result


def rust_config_plan(config: Path, output: Path, action: str) -> list[dict[str, object]]:
    result = rust_json([
        "dedup-config-plan",
        "--config", cli_path(config),
        "--output-dir", cli_path(output),
        "--action", action,
    ])
    assert isinstance(result, list)
    return result


def rust_precheck(config: Path) -> dict[str, object]:
    result = rust_json(["dedup-precheck", "--config", cli_path(config)], allow_nonzero=True)
    assert isinstance(result, dict)
    return result


def rust_append(config: Path, output: Path) -> dict[str, object]:
    result = rust_json(["dedup-append-plan", "--config", cli_path(config), "--output-dir", cli_path(output)])
    assert isinstance(result, list) and len(result) == 1
    plan = result[0]
    plan_path = from_wsl_path(str(plan["plan_file"]))
    executed = rust_json(["dedup-append-execute", "--plan", cli_path(plan_path)])
    assert isinstance(executed, dict)
    plan["execution_result"] = executed
    return plan


def rust_existing_append(config: Path, output: Path, capability_path: Path) -> tuple[dict[str, object], Path, dict[str, object]]:
    result = rust_json(["dedup-append-plan", "--config", cli_path(config), "--output-dir", cli_path(output)])
    assert isinstance(result, list) and len(result) == 1
    report = result[0]
    plan_path = from_wsl_path(str(report["plan_file"]))
    plan_document = json.loads(plan_path.read_text(encoding="utf-8"))
    edit_plan = plan_document["edit_plan"]
    platform = "synthetic-pt-2-5-existing-target-probe"
    capability = {
        "engine": "umya_spreadsheet",
        "status": "passed",
        "format": edit_plan["format"],
        "operations": sorted({operation["operation"] for operation in edit_plan["operations"]}),
        "supported_features": edit_plan["workbook"]["features_present"],
        "preserved": edit_plan["preserve"],
        "platform": platform,
        "evidence": {
            "sample_id": "synthetic-pt-2-5-existing-target",
            "report_path": "pytools-core/probe/compare_dedup_tools.py",
            "adapter_version": "umya-spreadsheet-3.1.1",
            "engine_version": "umya-spreadsheet-3.1.1",
            "platform": platform,
        },
    }
    capability_path.write_text(json.dumps(capability, ensure_ascii=False, indent=2), encoding="utf-8")
    executed = rust_json([
        "dedup-append-execute",
        "--plan", cli_path(plan_path),
        "--capabilities", cli_path(capability_path),
        "--platform", platform,
    ])
    assert isinstance(executed, dict)
    candidate = from_wsl_path(str(executed["output"]))
    return report, candidate, executed


def normalize(value: object) -> object:
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.isoformat()
    return value


def workbook_projection(path: Path) -> list[dict[str, object]]:
    wb = load_workbook(path, data_only=False)
    result = []
    for ws in wb.worksheets:
        result.append({
            "name": ws.title,
            "rows": [[normalize(value) for value in row] for row in ws.iter_rows(values_only=True)],
        })
    wb.close()
    return result


def red_fill_cells(path: Path) -> set[tuple[str, str]]:
    wb = load_workbook(path)
    cells = {
        (ws.title, cell.coordinate)
        for ws in wb.worksheets
        for row in ws.iter_rows()
        for cell in row
        if cell.fill.fill_type == "solid"
        and cell.fill.fgColor.type == "rgb"
        and cell.fill.fgColor.rgb in ("FFFF0000", "00FF0000")
    }
    wb.close()
    return cells


def fill_plan_cells(plan: dict[str, object]) -> set[tuple[str, str]]:
    edit_plan = plan.get("plan") or {}
    return {
        (str(operation["sheet"]), str(operation["cell"]))
        for operation in edit_plan.get("operations", [])
        if operation.get("operation") == "set_cell_fill_color"
    }


def delete_plan_to_candidate(source: Path, output: Path, operations: list[dict[str, object]]) -> Path:
    shutil.copy2(source, output)
    wb = load_workbook(output)
    for operation in operations:
        if operation.get("operation") == "delete_rows":
            wb[str(operation["sheet"])].delete_rows(
                int(operation["first_row"]), int(operation["count"])
            )
    wb.save(output)
    wb.close()
    return output


def assert_delete_plan_matches_python(
    source: Path, python_result: Path, operations: list[dict[str, object]], temp: Path, label: str
) -> None:
    candidate = delete_plan_to_candidate(source, temp, operations)
    expected = workbook_projection(python_result)
    actual = workbook_projection(candidate)
    if expected != actual:
        raise AssertionError(json.dumps({"label": label, "python": expected, "rust_plan_applied_to_copy": actual}, ensure_ascii=False, indent=2))


def make_python_tasks(source: Path, target: Path):
    sys.path.insert(0, str(PYTOOLS_SOURCE / "PyCode"))
    from core.dedup_311 import DedupTask

    return [
        DedupTask(2, True, source, "Data", [1], target, "Append", "追加", "Data-ID"),
        DedupTask(3, True, source, "Data", [2], target, "Append", "追加", "Data-Amount"),
        DedupTask(4, True, source, "Data", [1], target, "Append", "追加", "Data-ID-duplicate"),
        DedupTask(5, True, source, "Other", [], target, "OtherAppend", "追加", "Other-full-row"),
        DedupTask(6, True, source, "Missing", [1], target, "MissingAppend", "追加", "Missing-sheet"),
    ]


def main() -> int:
    if not PYTOOLS_SOURCE.is_dir():
        raise FileNotFoundError(PYTOOLS_SOURCE)
    sys.path.insert(0, str(PYTOOLS_SOURCE / "PyCode"))
    from core.dedup_311 import (
        DedupTask,
        run_append_by_config,
        run_check_by_comment,
        run_check_by_config,
        run_delete_by_comment,
        run_delete_by_config,
        run_precheck_by_config,
    )

    with tempfile.TemporaryDirectory(prefix="pytools-dedup-parity-", ignore_cleanup_errors=True) as temp:
        root = Path(temp)
        (root / "PyCode").mkdir()
        (root / "input").mkdir()
        source = root / "source.xlsx"
        make_source(source)

        # 2-1 and 2-2: comment-based duplicate marks and delete plans.
        py_comment_mark = root / "python_comment_mark.xlsx"
        py_comment_delete = root / "python_comment_delete.xlsx"
        shutil.copy2(source, py_comment_mark)
        shutil.copy2(source, py_comment_delete)
        with contextlib.redirect_stdout(None):
            python_comment_check = run_check_by_comment(py_comment_mark, root / "python_logs")
            python_comment_delete_report = run_delete_by_comment(py_comment_delete, root / "python_logs")
        comment_mark = rust_comment_plan(source, root / "candidate_comment_mark.xlsx", "check")
        comment_delete = rust_comment_plan(source, root / "candidate_comment_delete.xlsx", "delete")
        if comment_mark.get("duplicate_rows") != python_comment_check["marked_rows"]:
            raise AssertionError("comment check duplicate-row count differs")
        if fill_plan_cells(comment_mark) != red_fill_cells(py_comment_mark):
            raise AssertionError("comment check fill plan differs from old Python marking")
        comment_delete_ops = (comment_delete.get("plan") or {}).get("operations", [])
        if comment_delete.get("duplicate_rows") != python_comment_delete_report["deleted_rows"]:
            raise AssertionError("comment delete count differs")
        assert_delete_plan_matches_python(
            source, py_comment_delete, comment_delete_ops, root / "rust_comment_delete_semantics.xlsx", "comment delete"
        )

        # 2-3, 2-4, and 2-5: one config contains distinct and duplicate tasks.
        config = root / "PyCode" / "config.xlsx"
        make_config(config)
        rust_output = root / "rust_plans"
        config_mark_reports = rust_config_plan(config, rust_output, "check")
        config_delete_reports = rust_config_plan(config, rust_output, "delete")
        if len(config_mark_reports) != 1 or len(config_delete_reports) != 1:
            raise AssertionError("expected one grouped config-dedup report")
        config_mark_report = config_mark_reports[0]["report"]
        config_delete_report = config_delete_reports[0]["report"]

        py_config_source_mark = root / "python_config_mark.xlsx"
        py_config_source_delete = root / "python_config_delete.xlsx"
        shutil.copy2(source, py_config_source_mark)
        shutil.copy2(source, py_config_source_delete)
        py_target_unused = root / "python_unused_target.xlsx"
        tasks_mark = make_python_tasks(py_config_source_mark, py_target_unused)
        tasks_delete = make_python_tasks(py_config_source_delete, py_target_unused)
        with contextlib.redirect_stdout(None):
            python_config_check = run_check_by_config(tasks_mark, root / "python_logs")
            python_config_delete_report = run_delete_by_config(tasks_delete, root / "python_logs")
        if int(config_mark_report["duplicate_rows"]) != int(python_config_check["marked_rows"]):
            raise AssertionError("config check duplicate-row count differs")
        if fill_plan_cells(config_mark_report) != red_fill_cells(py_config_source_mark):
            raise AssertionError("config check fill plan differs from old Python marking")
        if int(config_delete_report["duplicate_rows"]) != int(python_config_delete_report["deleted_rows"]):
            raise AssertionError("config delete duplicate-row count differs")
        assert_delete_plan_matches_python(
            source,
            py_config_source_delete,
            (config_delete_report.get("plan") or {}).get("operations", []),
            root / "rust_config_delete_semantics.xlsx",
            "config delete",
        )

        # 2-5: the absent target is created as a candidate; compare its rows to Python.
        python_target = root / "python_target" / "target.xlsx"
        tasks_append = make_python_tasks(source, python_target)
        with contextlib.redirect_stdout(None):
            python_append_report = run_append_by_config(tasks_append, root / "python_logs")
        append_report = rust_append(config, rust_output)
        candidate_path = from_wsl_path(str(append_report["execution_result"]["candidate"]))
        if int(append_report["appended_rows"]) != int(python_append_report["appended_rows"]):
            raise AssertionError(
                f"append row count differs: Python={python_append_report['appended_rows']}, "
                f"Rust={append_report['appended_rows']}; report={append_report}"
            )
        if int(append_report["duplicate_rows_filtered"]) != int(python_append_report["deleted_rows"]):
            raise AssertionError(
                f"append filtered-row count differs: Python={python_append_report['deleted_rows']}, "
                f"Rust={append_report['duplicate_rows_filtered']}; report={append_report}"
            )
        if workbook_projection(python_target) != workbook_projection(candidate_path):
            raise AssertionError(json.dumps({
                "python_append": workbook_projection(python_target),
                "rust_candidate": workbook_projection(candidate_path),
            }, ensure_ascii=False, indent=2))

        # 2-5 existing target: compare the legacy in-place save with the Rust
        # ExcelModifier candidate. A test-only capability is scoped to this
        # generated workbook; it is not written to production capability data.
        existing_source = root / "append_existing_source.xlsx"
        python_existing_target = root / "python_existing_target" / "target.xlsx"
        rust_existing_target = root / "rust_existing_target" / "target.xlsx"
        python_existing_target.parent.mkdir(parents=True)
        rust_existing_target.parent.mkdir(parents=True)
        make_existing_append_source(existing_source)
        make_existing_append_target(python_existing_target)
        make_existing_append_target(rust_existing_target)
        rust_existing_target_before = rust_existing_target.read_bytes()
        existing_config = root / "PyCode" / "config_existing_target.xlsx"
        make_existing_append_config(
            existing_config,
            "append_existing_source.xlsx",
            "rust_existing_target/target.xlsx",
        )
        existing_tasks = [
            DedupTask(
                2,
                True,
                existing_source,
                "Source",
                [1],
                python_existing_target,
                "Combined",
                "追加",
                "Existing-target",
            )
        ]
        python_target_before = workbook_projection(python_existing_target)
        with contextlib.redirect_stdout(None):
            python_existing_append_report = run_append_by_config(existing_tasks, root / "python_logs")
        rust_existing_report, rust_existing_candidate, rust_existing_execution = rust_existing_append(
            existing_config,
            root / "rust_existing_output",
            root / "existing_target_test_capability.json",
        )
        if workbook_projection(python_existing_target) != workbook_projection(rust_existing_candidate):
            raise AssertionError(json.dumps({
                "python_existing_target": workbook_projection(python_existing_target),
                "rust_existing_candidate": workbook_projection(rust_existing_candidate),
            }, ensure_ascii=False, indent=2))
        if int(rust_existing_report["appended_rows"]) != int(python_existing_append_report["appended_rows"]):
            raise AssertionError("existing-target append row count differs from Python")
        if int(rust_existing_report["duplicate_rows_filtered"]) != int(python_existing_append_report["deleted_rows"]):
            raise AssertionError("existing-target filtered-row count differs from Python")
        if rust_existing_target.read_bytes() != rust_existing_target_before:
            raise AssertionError("Rust existing target was unexpectedly modified in place")
        if not python_target_before or not rust_existing_execution.get("output"):
            raise AssertionError("existing-target candidate execution did not produce output")

        # 2-6: precheck includes missing worksheet and missing source cases.
        precheck_config = root / "PyCode" / "config_precheck.xlsx"
        make_config(precheck_config, empty_source_task=True)
        precheck_result = rust_precheck(precheck_config)
        tasks_precheck = make_python_tasks(source, py_target_unused)
        tasks_precheck.append(
            type(tasks_precheck[0])(7, True, None, "Data", [1], None, None, "", "Missing-source")
        )
        python_precheck = run_precheck_by_config(tasks_precheck, root / "python_logs")
        rust_counts = (int(precheck_result["ok"]), int(precheck_result["warn"]), int(precheck_result["fail"]))
        python_counts = (int(python_precheck["ok"]), int(python_precheck["warn"]), int(python_precheck["fail"]))
        if rust_counts != python_counts:
            raise AssertionError(f"dedup precheck counts differ: Python={python_counts}, Rust={rust_counts}")

        logging.shutdown()
        print(json.dumps({
            "status": "passed",
            "comment_check": {"duplicate_rows": python_comment_check["marked_rows"], "fill_cells": len(fill_plan_cells(comment_mark))},
            "comment_delete": {"deleted_rows": python_comment_delete_report["deleted_rows"]},
            "config_check": {"duplicate_rows": python_config_check["marked_rows"], "merged_tasks": config_mark_report["merged_tasks"]},
            "config_delete": {"deleted_rows": python_config_delete_report["deleted_rows"], "merged_tasks": config_delete_report["merged_tasks"]},
            "append": {
                "appended_rows": python_append_report["appended_rows"],
                "filtered_rows": python_append_report["deleted_rows"],
                "target_sheets": [sheet["name"] for sheet in workbook_projection(candidate_path)],
                "existing_target_modified": False,
            },
            "append_existing_target": {
                "appended_rows": python_existing_append_report["appended_rows"],
                "filtered_rows": python_existing_append_report["deleted_rows"],
                "target_sheets": [sheet["name"] for sheet in workbook_projection(rust_existing_candidate)],
                "rust_target_input_unchanged": True,
                "candidate_engine": rust_existing_execution["engine"],
                "python_target_before_rows": len(python_target_before[0]["rows"]),
            },
            "precheck": {"ok": python_counts[0], "warn": python_counts[1], "fail": python_counts[2]},
        }, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
