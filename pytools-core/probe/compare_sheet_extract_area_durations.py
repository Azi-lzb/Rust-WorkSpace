from __future__ import annotations

import argparse
import importlib
import json
import logging
import subprocess
import sys
import tempfile
from datetime import date, datetime, time, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook


HEADERS = [
    "是否启用",
    "工作簿关键字",
    "工作表关键字",
    "排除关键字",
    "提取区域",
    "多区域拼接方式",
    "输出文件夹",
    "文件名",
    "备注",
]


def cell_value(value: object) -> dict[str, str] | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return {"type": "datetime", "value": value.isoformat(timespec="milliseconds")}
    if isinstance(value, date):
        return {"type": "date", "value": value.isoformat()}
    if isinstance(value, time):
        return {"type": "time", "value": value.isoformat(timespec="milliseconds")}
    if isinstance(value, timedelta):
        return {"type": "timedelta", "value": str(value)}
    return {"type": type(value).__name__, "value": str(value)}


def project_workbook(path: Path) -> dict[str, object]:
    workbook = load_workbook(path, data_only=True)
    try:
        sheets = []
        for worksheet in workbook.worksheets:
            cells = []
            for row in worksheet.iter_rows():
                for cell in row:
                    projected = cell_value(cell.value)
                    if projected is not None:
                        cells.append(
                            {
                                "cell": cell.coordinate,
                                **projected,
                                "number_format": cell.number_format,
                            }
                        )
            sheets.append({"name": worksheet.title, "cells": cells})
        return {"sheet_names": workbook.sheetnames, "sheets": sheets}
    finally:
        workbook.close()


def create_source(path: Path) -> list[timedelta]:
    durations = [
        timedelta(days=1, hours=2, minutes=3, seconds=4, microseconds=123456),
        timedelta(hours=27, minutes=12, seconds=1, microseconds=2500),
        timedelta(microseconds=500),
        timedelta(microseconds=1500),
        timedelta(microseconds=2500),
        timedelta(microseconds=3500),
        timedelta(microseconds=-500),
        timedelta(microseconds=-1500),
        timedelta(days=-2, hours=3, minutes=4, seconds=5, microseconds=123456),
    ]
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "时长源"
    for row, value in enumerate(durations, start=1):
        worksheet.cell(row=row, column=1, value=value)
    workbook.save(path)
    workbook.close()
    return durations


def create_rust_config(path: Path, output_dir: Path, row_count: int) -> None:
    workbook = Workbook()
    global_config = workbook.active
    global_config.title = "全局配置"
    global_config.append(["配置项", "值"])
    global_config.append(["输入目录", ""])
    global_config.append(["输出目录", str(output_dir)])
    global_config.append(["日志目录", str(output_dir / "logs")])

    extraction = workbook.create_sheet("工作表提取")
    extraction.append(HEADERS)
    extraction.append(
        [
            "是",
            "",
            "",
            "",
            f"A1:A{row_count}",
            "",
            "时长提取",
            "duration-output.xlsx",
            "elapsed duration parity",
        ]
    )
    workbook.save(path)
    workbook.close()


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--pytools-cli", type=Path, required=True)
    parser.add_argument("--python-project", type=Path, required=True)
    parser.add_argument("--report", type=Path)
    args = parser.parse_args()

    probe_root = Path(__file__).resolve().parent
    report_path = args.report or (
        probe_root / "reports" / "sheet_extract_elapsed_duration_parity_20261026.json"
    )
    source_python = args.python_project.resolve() / "PyCode"
    sys.path.insert(0, str(source_python))
    legacy = importlib.import_module("core.feature_1_6_extract_sheets")

    with tempfile.TemporaryDirectory(prefix="pytools-duration-parity-") as temporary:
        root = Path(temporary)
        source = root / "elapsed-durations.xlsx"
        durations = create_source(source)
        source_bytes = source.read_bytes()
        python_output = root / "python-output"
        python_logs = root / "python-logs"
        rust_output = root / "rust-output"
        config_path = root / "config.xlsx"
        task = {
            "row_no": 2,
            "workbook_keywords": "",
            "sheet_keywords": "",
            "exclude_keywords": "",
            "area_spec": f"A1:A{len(durations)}",
            "stitch_mode": "",
            "output_folder": "时长提取",
            "output_name": "duration-output.xlsx",
            "remark": "elapsed duration parity",
        }

        try:
            python_result = legacy.run_extract_sheets(
                [source],
                [task],
                python_output,
                python_logs,
                append_existing=False,
                default_output_name="unused.xlsx",
            )
            create_rust_config(config_path, rust_output, len(durations))
            rust_process = subprocess.run(
                [
                    str(args.pytools_cli.resolve()),
                    "extract-sheets-area",
                    "--config",
                    str(config_path),
                    "--output-dir",
                    str(rust_output),
                    str(source),
                ],
                capture_output=True,
                text=True,
                check=False,
            )
            if rust_process.returncode != 0:
                raise RuntimeError(
                    f"Rust CLI failed ({rust_process.returncode}): {rust_process.stderr}"
                )
            rust_stats = json.loads(rust_process.stdout)
            python_projection = project_workbook(
                python_output / "时长提取" / "duration-output.xlsx"
            )
            rust_projection = project_workbook(
                rust_output / "时长提取" / "duration-output.xlsx"
            )
            source_unchanged = source.read_bytes() == source_bytes
            if not source_unchanged:
                raise AssertionError("synthetic source workbook changed")
            if python_projection != rust_projection:
                raise AssertionError(
                    "Python/Rust output mismatch:\n"
                    + json.dumps(
                        {"python": python_projection, "rust": rust_projection},
                        ensure_ascii=False,
                        indent=2,
                    )
                )
            if python_result["failed_files"] != 0 or rust_stats["failed_files"] != 0:
                raise AssertionError("the duration extraction task reported a failed source")

            report = {
                "schema_version": 1,
                "status": "passed",
                "feature": "PT-7-6 area extraction elapsed duration subset",
                "sample_kind": "synthetic XLSX generated with openpyxl timedelta values",
                "source_python": "PyCode/core/feature_1_6_extract_sheets.py::run_extract_sheets",
                "rust_entry": "pytools extract-sheets-area",
                "python_stats": {
                    key: value for key, value in python_result.items() if key != "paths"
                },
                "rust_stats": {
                    key: value for key, value in rust_stats.items() if key != "paths"
                },
                "output_projection": python_projection,
                "source_unchanged": source_unchanged,
                "scope_note": (
                    "Synthetic parity only. Real business workbooks, time-only cells, ISO-typed "
                    "durations, non-XLSX inputs, and platform delivery remain unverified."
                ),
            }
            report_path.parent.mkdir(parents=True, exist_ok=True)
            report_path.write_text(
                json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
            )
            print(json.dumps(report, ensure_ascii=False, indent=2))
        finally:
            logging.shutdown()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
