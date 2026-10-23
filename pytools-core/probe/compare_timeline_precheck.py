"""Compare timeline configuration precheck counts with the read-only Python baseline."""

from __future__ import annotations

import contextlib
import io
import json
import re
import shlex
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import Workbook

sys.dont_write_bytecode = True
WORKSPACE = Path(__file__).resolve().parents[2]
PYTOOLS_SOURCE = Path(r"C:\Users\AZI\Desktop\VibeCode\4.pytools-main")

TIMELINE_HEADERS = [
    "是否启用", "规则名称", "工作簿关键字", "工作表关键字", "行头列", "列表头行",
    "必含列头", "必含行头", "数据起始行", "数据结束行", "数据起始列", "数据结束列",
    "跳过关键字", "目标工作簿路径", "目标工作表",
]
TIMELINE_OPTIONAL = ["启用目标写入", "set区域", "目标去重列", "是否输出行序号"]
PATH_HEADERS = [
    "是否启用", "映射名称", "适用规则名", "工作簿关键字", "工作表关键字",
    "作用对象", "匹配方式", "原始路径", "标准路径", "备注",
]


def write_config(path: Path, variant: str) -> None:
    wb = Workbook()
    global_sheet = wb.active
    global_sheet.title = "全局配置"
    global_sheet.append(["配置项", "值"])
    for key, value in (
        ("输入目录", "input"),
        ("输出目录", "output"),
        ("日志目录", "logs"),
        ("源文件扩展名", ".xlsx"),
    ):
        global_sheet.append([key, value])

    timeline = wb.create_sheet("时序提取规则")
    timeline.append(TIMELINE_HEADERS + TIMELINE_OPTIONAL)
    if variant == "detailed_errors":
        timeline.append([
            "是", "DUP", "sample", "Report", "", "0,foo", "", "必含行头",
            "abc", "2", "bad", "A", "", "", "", "是", "别名@A0", "", "否",
        ])
        timeline.append([
            "是", "DUP", "sample", "Report", "A", "2", "", "", 5, "", "B", "C",
            "", "target.xlsx", "Append", "是", "", "", "否",
        ])
    else:
        timeline.append([
            "是", "Rule", "sample", "Report", "A", "2", "", "", 5, "", "B", "C",
            "", "target.xlsx", "Append", "是", "", "", "否",
        ])

    path_map = wb.create_sheet("路径标准化映射")
    headers = PATH_HEADERS.copy()
    if variant == "missing_required_column":
        headers.remove("备注")
    path_map.append(headers)
    if variant == "detailed_errors":
        path_map.append(["是", "", "MissingRule", "sample", "Report", "bad", "bad", "", "", ""])
    else:
        path_map.append(["是", "Map", "Rule", "sample", "Report", "两者", "精确", "old", "new", ""])
    wb.save(path)


def python_counts(config: Path) -> tuple[int, int]:
    sys.path.insert(0, str(PYTOOLS_SOURCE / "PyCode"))
    from core.precheck_399 import run_precheck

    output = io.StringIO()
    with contextlib.redirect_stdout(output):
        run_precheck(config)
    match = re.search(r"=== 预校验完成: 错误 (\d+) 项，警告 (\d+) 项 ===", output.getvalue())
    if not match:
        raise AssertionError(f"legacy precheck summary missing:\n{output.getvalue()}")
    return int(match.group(1)), int(match.group(2))


def rust_counts(config: Path) -> tuple[int, int]:
    config_wsl = f"/mnt/{config.drive.rstrip(':').lower()}/{config.as_posix()[2:].lstrip('/')}"
    workspace_wsl = f"/mnt/{WORKSPACE.drive.rstrip(':').lower()}/{WORKSPACE.as_posix()[2:].lstrip('/')}"
    command = (
        f"cd {shlex.quote(workspace_wsl)} && cargo run --quiet -p pytools-core --bin pytools -- "
        f"timeline-precheck --config {shlex.quote(config_wsl)}"
    )
    result = subprocess.run(["wsl", "-e", "sh", "-lc", command], capture_output=True, check=False)
    stdout = result.stdout.decode("utf-8-sig", errors="replace").replace("\x00", "")
    for index in range(len(stdout)):
        if stdout[index] == "{":
            try:
                payload = json.loads(stdout[index:])
                return int(payload["errors"]), int(payload["warnings"])
            except (json.JSONDecodeError, KeyError, TypeError):
                continue
    raise RuntimeError(f"Rust precheck emitted no JSON report: {stdout}")


def main() -> int:
    sys.dont_write_bytecode = True
    if not PYTOOLS_SOURCE.is_dir():
        raise FileNotFoundError(PYTOOLS_SOURCE)
    results = {}
    for variant in ("valid", "detailed_errors", "missing_required_column"):
        with tempfile.TemporaryDirectory(prefix=f"pytools-precheck-{variant}-", ignore_cleanup_errors=True) as temp:
            root = Path(temp)
            (root / "PyCode").mkdir()
            (root / "input").mkdir()
            config = root / "PyCode" / "config.xlsx"
            write_config(config, variant)
            expected = python_counts(config)
            actual = rust_counts(config)
            if expected != actual:
                raise AssertionError(f"{variant}: Python {expected}, Rust {actual}")
            results[variant] = {"errors": actual[0], "warnings": actual[1]}
    print(json.dumps({"status": "passed", "scenarios": results}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
