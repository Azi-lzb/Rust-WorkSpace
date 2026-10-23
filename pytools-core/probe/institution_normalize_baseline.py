#!/usr/bin/env python3
"""Generate a synthetic xlsx and source-Python baseline for institution normalization."""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill


def dump_value(value):
    if isinstance(value, bool):
        return {"type": "boolean", "value": value}
    if isinstance(value, int):
        return {"type": "integer", "value": value}
    if isinstance(value, float):
        return {"type": "number", "value": value}
    if isinstance(value, str):
        return {"type": "string", "value": value}
    raise TypeError(f"unsupported fixture value: {value!r} ({type(value).__name__})")


def make_fixture(cases, output: Path):
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet_data in cases["sheets"]:
        worksheet = workbook.create_sheet(sheet_data["name"])
        for row_index, row in enumerate(sheet_data["rows"], 1):
            for column_index, value in enumerate(row, 1):
                if value is None:
                    continue
                if isinstance(value, str) and value.startswith("="):
                    worksheet.cell(row_index, column_index).value = value
                else:
                    worksheet.cell(row_index, column_index).value = value

    for item in cases["comments"]:
        workbook[item["sheet"]][item["cell"]].comment = Comment(item["text"], item["author"])
    for item in cases["merges"]:
        workbook[item["sheet"]].merge_cells(item["range"])
    for item in cases["style_cells"]:
        workbook[item["sheet"]][item["cell"]].fill = PatternFill(
            fill_type="solid", fgColor=item["fill"]
        )
    for item in cases["column_widths"]:
        workbook[item["sheet"]].column_dimensions[item["column"]].width = item["width"]
    for item in cases["row_heights"]:
        workbook[item["sheet"]].row_dimensions[item["row"]].height = item["height"]
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    workbook.close()


def dump_workbook(path: Path):
    workbook = load_workbook(path, data_only=False)
    sheets = []
    for worksheet in workbook.worksheets:
        values = {}
        formulas = {}
        comments = {}
        for row in worksheet.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                if cell.data_type == "f":
                    formulas[cell.coordinate] = str(cell.value)[1:]
                else:
                    values[cell.coordinate] = dump_value(cell.value)
                if cell.comment is not None:
                    comments[cell.coordinate] = cell.comment.text
        sheets.append(
            {
                "name": worksheet.title,
                "values": values,
                "formulas": formulas,
                "formula_status": "read_by_openpyxl_baseline",
                "merges": sorted(str(item) for item in worksheet.merged_cells.ranges),
                "merge_status": "read_by_openpyxl_baseline",
                "comments": comments,
                "comments_status": "read_by_openpyxl_baseline",
                "conditional_formatting": [],
                "conditional_formatting_status": "fixture_has_none",
                "data_validations": [],
                "data_validations_status": "fixture_has_none",
            }
        )
    workbook.close()
    return sheets


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source-root", required=True)
    parser.add_argument(
        "--cases",
        default=str(
            Path(__file__).resolve().parents[1]
            / "tests"
            / "fixtures"
            / "institution_normalize_cases.json"
        ),
    )
    parser.add_argument(
        "--fixture",
        default=str(
            Path(__file__).resolve().parents[1]
            / "tests"
            / "fixtures"
            / "institution_normalize.xlsx"
        ),
    )
    parser.add_argument(
        "--output",
        default=str(
            Path(__file__).resolve().parent
            / "reports"
            / "institution_normalize_baseline.json"
        ),
    )
    args = parser.parse_args()
    source_root = Path(args.source_root).resolve()
    cases = json.loads(Path(args.cases).read_text(encoding="utf-8"))
    fixture = Path(args.fixture)
    make_fixture(cases, fixture)

    sys.path.insert(0, str(source_root))
    from PyCode.core.feature_1_2_normalize_institution import _process_sheet

    mapping = {key: value for key, value in cases["mapping"]}
    keys_lower = [(key.lower(), key) for key in mapping]
    with TemporaryDirectory(prefix="pytools-normalize-baseline-") as temporary:
        baseline_path = Path(temporary) / "baseline.xlsx"
        baseline_path.write_bytes(fixture.read_bytes())
        workbook = load_workbook(baseline_path, data_only=False)
        branch_sheets = mapped_sheets = mapped_rows = 0
        for worksheet in workbook.worksheets:
            if "分机构" not in worksheet.title:
                continue
            branch_sheets += 1
            changed = _process_sheet(worksheet, mapping, keys_lower)
            if changed:
                mapped_sheets += 1
                mapped_rows += changed
        if mapped_rows:
            workbook.save(baseline_path)
        workbook.close()
        expected_sheets = dump_workbook(baseline_path)

    report = {
        "schema_version": 3,
        "reader": "openpyxl-baseline",
        "elapsed_ms": 0,
        "suite": "pytools.institution_normalize",
        "baseline": "original PyCode.core.feature_1_2_normalize_institution._process_sheet",
        "source_python": sys.version.split()[0],
        "source_root": str(source_root),
        "mapping": cases["mapping"],
        "sheets": expected_sheets,
        "stats": {
            "files": 1,
            "branch_sheets": branch_sheets,
            "mapped_sheets": mapped_sheets,
            "mapped_rows": mapped_rows,
            "failed_files": 0,
            "mapping_empty": False,
        },
        "fixture_metadata": {
            "style_cells": cases["style_cells"],
            "column_widths": cases["column_widths"],
            "row_heights": cases["row_heights"],
        },
        "status": "基线已生成，等待 Rust 对拍",
    }
    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"suite": report["suite"], "stats": report["stats"]}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
