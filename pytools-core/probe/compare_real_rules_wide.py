"""Real-rules × synthetic-data wide-summary parity.

Uses the 15 real timeline rules and 146 real path maps from the user's VBA
toolbox workbook (tests/fixtures/vba_config) through a bridge config (those
two sheets copied verbatim plus a synthetic global sheet), driving three
synthetic source workbooks crafted to hit different real rules. Not
real-sample acceptance: rules are real, source data is synthetic.
"""

import sys
import json
import tempfile
from pathlib import Path

sys.dont_write_bytecode = True
PROBE = Path(r"C:/Users/AZI/Desktop/VibeCode/pytools-dev/pytools-core/probe")
sys.path.insert(0, str(PROBE))

import compare_wide_summary as cws  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402

VBA_BOOK = Path(
    r"C:/Users/AZI/Desktop/VibeCode/pytools-dev/pytools-core/tests/fixtures/vba_config/VbaV2.4.2.xlsm"
)


def make_bridge_config(path: Path, output: Path) -> None:
    source = load_workbook(VBA_BOOK, data_only=False)
    bridge = Workbook()
    default = bridge.active
    bridge.remove(default)
    global_sheet = bridge.create_sheet("全局配置")
    global_sheet.append(["配置项", "值"])
    for key, value in (
        ("输入目录", cws.rust_path_argument(output / "unused_input")),
        ("输出目录", cws.rust_path_argument(output)),
        ("日志目录", cws.rust_path_argument(output / "logs")),
        ("错误策略", "fail_fast"),
        ("源文件扩展名", ".xlsx"),
        ("时序规则命中策略", "all_match"),
    ):
        global_sheet.append([key, value])
    for sheet_name in ("时序提取规则", "路径标准化映射"):
        src = source[sheet_name]
        dst = bridge.create_sheet(sheet_name)
        for row in src.iter_rows():
            for cell in row:
                if cell.value is not None:
                    dst.cell(row=cell.row, column=cell.column, value=cell.value)
    source.close()
    path.parent.mkdir(parents=True, exist_ok=True)
    bridge.save(path)


def make_monthly_source(path: Path, *, include_data: bool = True) -> None:
    """Hits 中外资本外币分地区: 双列头 38/39, 行头 B 列, 数据 40-66 有界."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "中外资本外币分地区"
    sheet["A1"] = "某省金融机构月报（合成）"
    sheet["B38"] = "地区"
    for column, header in ((3, "存款"), (4, "存款"), (5, "贷款"), (6, "票据")):
        sheet.cell(row=38, column=column, value=header)
    for column, header in ((3, "余额"), (4, "增量"), (5, "本息"), (6, "余额")):
        sheet.cell(row=39, column=column, value=header)
    if include_data:
        institutions = ["甲市", "乙市", "丙市", "丁市"]
        for index in range(8):
            row = 40 + index
            sheet.cell(row=row, column=2, value=institutions[index % len(institutions)])
            for column in range(3, 7):
                sheet.cell(row=row, column=column, value=100.5 + index * column)
    # bounded data region ends at 66; leave 48-66 empty like a real report tail
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def make_inclusive_source(path: Path) -> None:
    """Hits 普惠分地区: 行头 A 列, 列头行 2, 数据行 3+, 列 2+."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "普惠贷款分地区表"
    sheet["A2"] = "地区"
    sheet["B2"] = "余额"
    sheet["C2"] = "户数"
    for index in range(6):
        row = 3 + index
        sheet.cell(row=row, column=1, value=f"{['甲', '乙', '丙'][index % 3]}地区")
        sheet.cell(row=row, column=2, value=200 + index)
        sheet.cell(row=row, column=3, value=10 + index)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def make_survey_source(path: Path) -> None:
    """Hits 分地区统计研究调研问卷: 三列拼接行头, 列头行 1/2, 数据列 4+."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet["A1"] = "A列标题"
    sheet["B1"] = "B列标题"
    sheet["C1"] = "C列标题"
    sheet["D1"] = "问卷项一"
    sheet["E1"] = "问卷项二"
    sheet["A2"] = "a2"
    sheet["B2"] = "b2"
    sheet["C2"] = "c2"
    sheet["D2"] = "回答"
    sheet["E2"] = "金额"
    rows = (
        ("广东", "深圳", "甲行", "是", 11.5),
        ("广东", "广州", "乙行", "否", 22.5),
        ("浙江", "杭州", "丙行", "是", 33.5),
    )
    for index, (province, city, bank, answer, amount) in enumerate(rows):
        row = 3 + index
        sheet.cell(row=row, column=1, value=province)
        sheet.cell(row=row, column=2, value=city)
        sheet.cell(row=row, column=3, value=bank)
        sheet.cell(row=row, column=4, value=answer)
        sheet.cell(row=row, column=5, value=amount)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def python_side(config: Path, sources: list[Path], row_suffix: bool):
    import logging

    sys.path.insert(0, r"C:/Users/AZI/Desktop/VibeCode/4.pytools-main")
    from PyCode.core.config_xlsx import load_global, load_path_maps, load_timeline_rules
    from PyCode.core.wide_397_398 import run_wide_summary

    try:
        return run_wide_summary(
            load_timeline_rules(config),
            load_path_maps(config),
            load_global(config),
            row_suffix,
            source_paths=sources,
        )
    finally:
        logging.shutdown()


def run_short_data_boundary_case(root: Path, row_suffix: bool) -> dict[str, object]:
    case = root / ("short_data_suffix" if row_suffix else "short_data_no_suffix")
    source = case / "sources" / "20260609_月报空数据段.xlsx"
    make_monthly_source(source, include_data=False)
    config = case / "config.xlsx"
    output = case / "output"
    make_bridge_config(config, output)

    python_output = python_side(config, [source], row_suffix)
    rust_output, plans = cws.run_rust_sources(config, [source], row_suffix)
    if plans:
        raise AssertionError(f"short data-range case unexpectedly produced plans: {plans}")
    if python_output is None or rust_output is None:
        if python_output is not None or rust_output is not None:
            raise AssertionError(
                f"short data-range output presence differs: python={python_output}, rust={rust_output}"
            )
        output_projection = "both skipped output"
    else:
        cws.assert_parity(python_output, rust_output)
        output_projection = "exact"
    return {
        "row_suffix": row_suffix,
        "configured_data_rows": "40-66",
        "source_last_row": 39,
        "python_rust_output_projection": output_projection,
        "rust_panicked_or_failed": False,
    }


def main() -> int:
    import datetime as dt

    with tempfile.TemporaryDirectory(
        prefix="pytools-real-rules-parity-", ignore_cleanup_errors=True
    ) as temp:
        root = Path(temp)
        sources = []
        monthly = root / "sources" / "20260609_月报一批分地区.xlsx"
        make_monthly_source(monthly)
        inclusive = root / "sources" / "20260609_绿色普惠分地区.xlsx"
        make_inclusive_source(inclusive)
        survey = root / "sources" / "2026分地区统计研究调研问卷.xlsx"
        make_survey_source(survey)
        sources = [monthly, inclusive, survey]

        results = []
        short_data_boundary_results = []
        for row_suffix in (False, True):
            case = root / ("suffix" if row_suffix else "no_suffix")
            output = case / "output"
            config = case / "config.xlsx"
            make_bridge_config(config, output)
            python_output = python_side(config, sources, row_suffix)
            rust_output, plans = cws.run_rust_sources(config, sources, row_suffix)
            if python_output is None or rust_output is None:
                raise AssertionError(
                    f"wide summary did not save: python={python_output}, rust={rust_output}"
                )
            if plans:
                raise AssertionError(f"target writes disabled, got plans: {plans}")
            cws.assert_parity(python_output, rust_output)
            projection = cws.normalized(rust_output)
            results.append(
                {
                    "row_suffix": row_suffix,
                    "python_rust_summary": "exact",
                    "sheet_names": projection["sheet_names"],
                    "dimensions": projection["dimensions"],
                }
            )
            short_data_boundary_results.append(
                run_short_data_boundary_case(root, row_suffix)
            )

        report = {
            "schema_version": 1,
            "suite": "pytools.real_rules_synthetic_wide_parity",
            "date": dt.date.today().isoformat(),
            "status": "passed_exact",
            "config_source": {
                "workbook": "tests/fixtures/vba_config/VbaV2.4.2.xlsm",
                "origin": "https://github.com/Azi-lzb/VBA_Excel_Tools cea6dfa",
                "note": "real 15 timeline rules and 146 path maps via a bridge config; global sheet synthetic",
            },
            "sources": "synthetic workbooks crafted to hit 中外资本外币分地区 / 普惠分地区 / 分地区统计研究调研问卷",
            "modes": results,
            "short_data_before_configured_start": short_data_boundary_results,
            "limitations": [
                "Rules are real business configuration; source workbooks are synthetic.",
                "Not real-sample acceptance; outputs were temporary.",
            ],
        }
        report_path = PROBE / "reports" / "real_rules_wide_parity_20261025.json"
        report_path.write_text(
            json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
        )
        print(
            "REAL_RULES_RESULT="
            + json.dumps(
                {"status": report["status"], "modes": results, "report": str(report_path)},
                ensure_ascii=False,
            )
        )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
