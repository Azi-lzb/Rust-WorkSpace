"""Generate a synthetic XLSX fixture and capture the read-only Python baseline.

This probe imports the original PyCode implementation only for comparison.
It writes all generated artifacts under pytools-core, never into the source tree.
"""
from __future__ import annotations

import hashlib
import json
import logging
import os
import shutil
import sys
import tempfile
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import PatternFill
from openpyxl.worksheet.datavalidation import DataValidation


WORKSPACE = Path(__file__).resolve().parents[2]
PYTOOLS_ROOT = WORKSPACE / "pytools-core"
FIXTURE = PYTOOLS_ROOT / "tests" / "fixtures" / "fx_header" / "source.xlsx"
REPORT = PYTOOLS_ROOT / "probe" / "reports" / "fx_header_python_baseline.json"
SOURCE_ROOT = Path(
    os.environ.get(
        "PYTOOLS_SOURCE_ROOT",
        r"C:\Users\AZI\Desktop\VibeCode\4.pytools-main",
    )
)
SOURCE_MODULE = SOURCE_ROOT / "PyCode" / "core" / "feature_1_4_fx_header.py"

if not SOURCE_MODULE.is_file() and Path("/mnt/c/Users/AZI/Desktop/VibeCode/4.pytools-main").is_dir():
    SOURCE_ROOT = Path("/mnt/c/Users/AZI/Desktop/VibeCode/4.pytools-main")
    SOURCE_MODULE = SOURCE_ROOT / "PyCode" / "core" / "feature_1_4_fx_header.py"

sys.path.insert(0, str(SOURCE_ROOT))
from PyCode.core.feature_1_4_fx_header import run_fx_header_fix  # noqa: E402


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def make_fixture(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    fx = workbook.active
    fx.title = "外汇报表"
    fx.append(["指标", "金额", "公式", "说明"])
    fx.append(["本期", 12.5, "=B2*2", "单位：万元"])
    fx.merge_cells("C3:D3")
    fx["C3"] = "合并保留"
    fx["A1"].fill = PatternFill(fill_type="solid", fgColor="FFF2CC")
    fx["A1"].comment = Comment("原有批注保留", "probe")
    fx.column_dimensions["D"].width = 18
    fx.page_setup.orientation = "landscape"
    fx.page_setup.fitToWidth = 1
    fx.sheet_properties.pageSetUpPr.fitToPage = True
    fx.print_area = "A1:D3"
    validation = DataValidation(type="list", formula1='"A,B"', allow_blank=True)
    fx.add_data_validation(validation)
    validation.add("A2")
    fx.conditional_formatting.add(
        "B2:B3",
        CellIsRule(
            operator="greaterThan",
            formula=["10"],
            fill=PatternFill(fill_type="solid", fgColor="E2F0D9"),
        ),
    )

    fx.oddHeader.left.text = "单位：万元"
    fx.oddHeader.center.text = "外汇报表 单位：万元"
    fx.oddHeader.right.text = "右侧万元"
    fx.evenHeader.left.text = "偶数页：万元"
    fx.firstHeader.center.text = "首页不替换"

    blank_fx = workbook.create_sheet("外汇空白")
    blank_fx["A1"] = "空工作表也参与匹配计数"
    blank_fx.oddHeader.left.text = "单位：万美元"

    ordinary = workbook.create_sheet("普通报表")
    ordinary["A1"] = "不得修改"
    ordinary.oddHeader.left.text = "普通报表：万元"
    workbook.save(path)
    workbook.close()


def header_values(path: Path) -> dict:
    workbook = load_workbook(path, data_only=False)
    result = {}
    try:
        for worksheet in workbook.worksheets:
            result[worksheet.title] = {}
            for variant in ("odd", "even", "first"):
                header = getattr(worksheet, f"{variant}Header")
                result[worksheet.title][variant] = {
                    section: getattr(header, section).text
                    for section in ("left", "center", "right")
                }
    finally:
        workbook.close()
    return result


def preservation_values(path: Path) -> dict:
    workbook = load_workbook(path, data_only=False)
    worksheet = workbook["外汇报表"]
    try:
        return {
            "cell_b2": worksheet["B2"].value,
            "formula_c2": worksheet["C2"].value,
            "merged_ranges": sorted(str(item) for item in worksheet.merged_cells.ranges),
            "comment_a1": worksheet["A1"].comment.text,
            "column_d_width": worksheet.column_dimensions["D"].width,
            "orientation": worksheet.page_setup.orientation,
            "fit_to_width": worksheet.page_setup.fitToWidth,
            "print_area": str(worksheet.print_area),
            "data_validation_count": len(worksheet.data_validations.dataValidation),
            "conditional_formatting_count": len(worksheet.conditional_formatting),
        }
    finally:
        workbook.close()


def main() -> None:
    if not SOURCE_MODULE.is_file():
        raise SystemExit(f"read-only Python baseline source not found: {SOURCE_MODULE}")

    make_fixture(FIXTURE)
    fixture_digest = sha256(FIXTURE)
    with tempfile.TemporaryDirectory(prefix="pytools-fx-header-baseline-") as temporary:
        work = Path(temporary)
        target = work / FIXTURE.name
        shutil.copy2(FIXTURE, target)
        try:
            stats = run_fx_header_fix([target], work / "logs")
        finally:
            # The legacy logger caches its file handler globally. Close it so
            # Windows can remove the temporary comparison directory.
            logger = logging.getLogger("PyCode.feature_1_4_fx_header")
            for handler in tuple(logger.handlers):
                logger.removeHandler(handler)
                handler.close()
        report = {
            "schema_version": 1,
            "feature": "ordinary_fx_header_unit_replacement",
            "source_project": "PyCode",
            "source_file": str(SOURCE_MODULE),
            "source_sha256": sha256(SOURCE_MODULE),
            "fixture": str(FIXTURE.relative_to(WORKSPACE)).replace("\\", "/"),
            "fixture_sha256": fixture_digest,
            "python_baseline_stats": stats,
            "headers_after": header_values(target),
            "preserved_workbook_properties": preservation_values(target),
        }

    REPORT.parent.mkdir(parents=True, exist_ok=True)
    REPORT.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(REPORT)
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
