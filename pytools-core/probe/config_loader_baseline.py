"""Generate a workbook from shared vectors and run the read-only Python baseline."""

from __future__ import annotations

import argparse
import json
import sys
from dataclasses import asdict, is_dataclass
from pathlib import Path

from openpyxl import Workbook


def json_value(value, base: Path):
    if is_dataclass(value):
        return json_value(asdict(value), base)
    if isinstance(value, dict):
        return {key: json_value(item, base) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [json_value(item, base) for item in value]
    if isinstance(value, Path):
        return str(value).replace(str(base), "$BASE").replace("\\", "/")
    return value


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source-root", required=True, type=Path)
    parser.add_argument(
        "--cases",
        type=Path,
        default=Path(__file__).parents[1] / "tests/fixtures/config_loader_cases.json",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=Path(__file__).parents[1] / "probe/reports/config_loader_baseline.json",
    )
    parser.add_argument(
        "--fixture",
        type=Path,
        default=Path(__file__).parents[1] / "tests/fixtures/config_loader.xlsx",
    )
    args = parser.parse_args()
    source_root = args.source_root.resolve()
    sys.path.insert(0, str(source_root))
    from PyCode.core.config_xlsx import (  # noqa: PLC0415
        load_global,
        load_global_value,
        load_path_maps,
        load_timeline_rules,
    )

    cases = json.loads(args.cases.read_text(encoding="utf-8"))
    config_path = args.fixture.resolve()
    config_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet_name, rows in cases["sheets"].items():
        worksheet = workbook.create_sheet(sheet_name)
        for row_index, row in enumerate(rows, start=1):
            for column_index, value in enumerate(row, start=1):
                if value is not None:
                    worksheet.cell(row_index, column_index, value)
    workbook.save(config_path)

    base = config_path.parent.parent
    global_config = load_global(config_path)
    result = {
        "global": json_value(global_config, base),
        "global_value_existing": load_global_value(config_path, "错误策略"),
        "global_value_missing": load_global_value(config_path, "missing", "fallback"),
        "timeline_rules": json_value(load_timeline_rules(config_path), base),
        "path_maps": json_value(load_path_maps(config_path), base),
    }

    report = {
        "schema_version": 1,
        "suite": "pytools.config_loader",
        "source": "PyCode/core/config_xlsx.py",
        "shared_vectors": "tests/fixtures/config_loader_cases.json",
        "python_baseline": {"status": "保留", "passed": 1, "total": 1},
        "result": result,
        "limitations": [
            "Covers global settings, timeline rules, and path maps only.",
            "Uses a synthetic workbook generated from shared vectors; does not establish full workflow or real-template parity.",
        ],
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
