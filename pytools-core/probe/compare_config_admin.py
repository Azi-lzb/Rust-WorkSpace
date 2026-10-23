"""Compare Rust config init/view with the read-only Python config baseline."""

from __future__ import annotations

import hashlib
import json
import shlex
import subprocess
import sys
import tempfile
from datetime import date, datetime, time
from pathlib import Path

from openpyxl import load_workbook

sys.dont_write_bytecode = True
WORKSPACE = Path(__file__).resolve().parents[2]
PYTOOLS_SOURCE = Path(r"C:\Users\AZI\Desktop\VibeCode\4.pytools-main")


def wsl_path(path: Path) -> str:
    return f"/mnt/{path.drive.rstrip(':').lower()}/{path.as_posix()[2:].lstrip('/')}"


def rust_command(arguments: list[str]) -> subprocess.CompletedProcess[bytes]:
    workspace = wsl_path(WORKSPACE)
    args = " ".join(shlex.quote(value) for value in arguments)
    command = f"cd {shlex.quote(workspace)} && cargo run --quiet -p pytools-core --bin pytools -- {args}"
    return subprocess.run(["wsl", "-e", "sh", "-lc", command], capture_output=True, check=False)


def parse_json(result: subprocess.CompletedProcess[bytes]) -> object:
    stdout = result.stdout.decode("utf-8-sig", errors="replace").replace("\x00", "")
    for index, char in enumerate(stdout):
        if char in "[{":
            try:
                return json.loads(stdout[index:])
            except json.JSONDecodeError:
                continue
    raise RuntimeError(f"Rust command emitted no JSON (exit={result.returncode}): {stdout}")


def normalize(value: object) -> object:
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return value


def workbook_projection(path: Path) -> dict[str, object]:
    wb = load_workbook(path, data_only=False)
    sheets = []
    authors = []
    for ws in wb.worksheets:
        cells = {
            cell.coordinate: normalize(cell.value)
            for row in ws.iter_rows()
            for cell in row
            if cell.value is not None
        }
        comments = {
            cell.coordinate: cell.comment.text
            for row in ws.iter_rows()
            for cell in row
            if cell.comment is not None
        }
        comment_authors = {
            cell.coordinate: cell.comment.author
            for row in ws.iter_rows()
            for cell in row
            if cell.comment is not None
        }
        max_row = max(
            [cell.row for row in ws.iter_rows() for cell in row if cell.value is not None]
            + [int("".join(filter(str.isdigit, address))) for address in comments]
            + [0]
        )
        max_col = max(
            [cell.column for row in ws.iter_rows() for cell in row if cell.value is not None]
            + [column_number(address) for address in comments]
            + [0]
        )
        rows = [
            [cells.get(f"{column_letters(column)}{row}") for column in range(1, max_col + 1)]
            for row in range(1, max_row + 1)
        ]
        sheets.append({"name": ws.title, "rows": rows, "comments": comments})
        authors.append({"name": ws.title, "comments": comment_authors})
    wb.close()
    return {"sheet_names": [sheet["name"] for sheet in sheets], "sheets": sheets, "comment_authors": authors}


def column_number(address: str) -> int:
    letters = "".join(filter(str.isalpha, address))
    result = 0
    for letter in letters:
        result = result * 26 + ord(letter.upper()) - ord("A") + 1
    return result


def column_letters(number: int) -> str:
    result = ""
    while number:
        number, remainder = divmod(number - 1, 26)
        result = chr(ord("A") + remainder) + result
    return result


def main() -> int:
    sys.dont_write_bytecode = True
    source_file = PYTOOLS_SOURCE / "PyCode" / "core" / "config_init.py"
    if not source_file.is_file():
        raise FileNotFoundError(source_file)
    sys.path.insert(0, str(PYTOOLS_SOURCE / "PyCode"))
    from core.config_init import initialize_or_repair_config

    with tempfile.TemporaryDirectory(prefix="pytools-config-admin-", ignore_cleanup_errors=True) as temp:
        root = Path(temp)
        (root / "Python" / "config").mkdir(parents=True)
        (root / "Rust" / "config").mkdir(parents=True)
        python_path = root / "Python" / "config" / "config.xlsx"
        rust_path = root / "Rust" / "config" / "config.xlsx"
        python_result = initialize_or_repair_config(python_path)
        rust_init = rust_command(["config-init", "--path", wsl_path(rust_path)])
        if rust_init.returncode != 0:
            raise RuntimeError(rust_init.stderr.decode("utf-8", errors="replace"))
        rust_init_result = parse_json(rust_init)
        expected = workbook_projection(python_path)
        actual = workbook_projection(rust_path)
        if expected != actual:
            if expected["sheet_names"] != actual["sheet_names"]:
                raise AssertionError(f"sheet order differs: Python={expected['sheet_names']}; Rust={actual['sheet_names']}")
            for expected_sheet, actual_sheet in zip(expected["sheets"], actual["sheets"]):
                if expected_sheet["rows"] != actual_sheet["rows"]:
                    for row_index, (python_row, rust_row) in enumerate(zip(expected_sheet["rows"], actual_sheet["rows"]), start=1):
                        if python_row != rust_row:
                            raise AssertionError(f"{expected_sheet['name']} row {row_index} differs: Python={python_row}; Rust={rust_row}")
                    raise AssertionError(f"{expected_sheet['name']} row dimensions differ")
                if expected_sheet["comments"] != actual_sheet["comments"]:
                    addresses = sorted(set(expected_sheet["comments"]) | set(actual_sheet["comments"]))
                    address = next(address for address in addresses if expected_sheet["comments"].get(address) != actual_sheet["comments"].get(address))
                    raise AssertionError(f"{expected_sheet['name']}!{address} comment differs: Python={expected_sheet['comments'].get(address)!r}; Rust={actual_sheet['comments'].get(address)!r}")
            if expected["comment_authors"] != actual["comment_authors"]:
                for python_sheet, rust_sheet in zip(expected["comment_authors"], actual["comment_authors"]):
                    if python_sheet != rust_sheet:
                        raise AssertionError(f"comment author differs: Python={python_sheet}; Rust={rust_sheet}")
            raise AssertionError("workbook projection differs")

        rust_view_result = rust_command(["config-view", "--config", wsl_path(rust_path)])
        if rust_view_result.returncode != 0:
            raise RuntimeError(rust_view_result.stderr.decode("utf-8", errors="replace"))
        rust_view = parse_json(rust_view_result)
        if rust_view["sheets"] != actual["sheets"]:
            raise AssertionError("config-view JSON differs from the initialized Rust workbook projection")

        before = hashlib.sha256(rust_path.read_bytes()).hexdigest()
        retry = rust_command(["config-init", "--path", wsl_path(rust_path)])
        after = hashlib.sha256(rust_path.read_bytes()).hexdigest()
        if retry.returncode == 0 or before != after:
            raise AssertionError("config-init must reject an existing file without changing its bytes")

        real_config = PYTOOLS_SOURCE / "PyCode" / "config" / "config.xlsx"
        real_config_hash_before = hashlib.sha256(real_config.read_bytes()).hexdigest()
        real_view_result = rust_command(["config-view", "--config", wsl_path(real_config)])
        if real_view_result.returncode != 0:
            raise RuntimeError(real_view_result.stderr.decode("utf-8", errors="replace"))
        real_view = parse_json(real_view_result)
        real_projection = workbook_projection(real_config)
        real_view_sheets = real_view["sheets"]
        if real_view_sheets != real_projection["sheets"]:
            if len(real_view_sheets) != len(real_projection["sheets"]):
                raise AssertionError(f"real config sheet count differs: Python={len(real_projection['sheets'])}; Rust={len(real_view_sheets)}")
            for python_sheet, rust_sheet in zip(real_projection["sheets"], real_view_sheets):
                if python_sheet["name"] != rust_sheet["name"]:
                    raise AssertionError(f"real config sheet order differs: Python={python_sheet['name']}; Rust={rust_sheet['name']}")
                if python_sheet["rows"] != rust_sheet["rows"]:
                    for row_index, (python_row, rust_row) in enumerate(zip(python_sheet["rows"], rust_sheet["rows"]), start=1):
                        if python_row != rust_row:
                            for column, (python_value, rust_value) in enumerate(zip(python_row, rust_row), start=1):
                                if python_value != rust_value:
                                    raise AssertionError(f"real config {python_sheet['name']}!{column_letters(column)}{row_index} differs: Python={python_value!r}; Rust={rust_value!r}")
                    raise AssertionError(f"real config {python_sheet['name']} row dimensions differ")
                if python_sheet["comments"] != rust_sheet["comments"]:
                    addresses = sorted(set(python_sheet["comments"]) | set(rust_sheet["comments"]))
                    address = next(address for address in addresses if python_sheet["comments"].get(address) != rust_sheet["comments"].get(address))
                    raise AssertionError(f"real config {python_sheet['name']}!{address} comment differs: Python={python_sheet['comments'].get(address)!r}; Rust={rust_sheet['comments'].get(address)!r}")
            raise AssertionError("real config projection differs")
        real_config_hash_after = hashlib.sha256(real_config.read_bytes()).hexdigest()
        if real_config_hash_before != real_config_hash_after:
            raise AssertionError("read-only Rust config-view changed the source config bytes")

        report = {
            "status": "passed_synthetic_parity",
            "scope": ["5-1 new configuration initialization", "5-2 read-only configuration view"],
            "python_init": {"created_sheets": python_result["created_sheets"], "repaired_sheets": python_result["repaired_sheets"]},
            "rust_init": rust_init_result,
            "comparison": {"sheet_order": "exact", "cell_values": "exact", "header_comment_text": "exact", "header_comment_author": "exact", "config_view_rows_and_comments": "exact"},
            "existing_file_protection": "Rust config-init rejected the second invocation; SHA-256 unchanged",
            "real_source_config_view": {"status": "passed_read_only_projection", "sheets": len(real_view_sheets), "sha256_before_after": real_config_hash_before},
            "source_python_sha256": hashlib.sha256(source_file.read_bytes()).hexdigest(),
            "limitations": ["Initialization parity uses a synthetic new configuration; the real source config was only read for config-view parity and was not modified.", "Config repair is a plan-only operation and its ExcelModifier execution remains unverified.", "Target Linux/deb behavior was not validated."],
        }
        output = WORKSPACE / "pytools-core" / "probe" / "reports" / "config_admin_parity_20261023.json"
        output.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        print(json.dumps(report, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
