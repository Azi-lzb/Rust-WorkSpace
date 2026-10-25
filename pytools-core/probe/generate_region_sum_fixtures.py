"""Create deterministic synthetic workbooks for the region-sum feature."""

from __future__ import annotations

import json
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

from openpyxl import Workbook


FIXTURE_DIR = Path(__file__).parents[1] / "tests" / "fixtures" / "region_sum"
MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


def save_book(path: Path, sheets: dict[str, list[list[object]]]) -> None:
    book = Workbook()
    book.remove(book.active)
    for title, rows in sheets.items():
        sheet = book.create_sheet(title)
        for row in rows:
            sheet.append(row)
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)
    book.close()


def save_cached_formula_book(path: Path) -> None:
    """Create a synthetic source whose total cells have cached formula values."""
    book = Workbook()
    book.remove(book.active)
    formula_results = [75, 60]
    for title in ("本期分地区", "同期分地区"):
        sheet = book.create_sheet(title)
        sheet.append(["地区总分校验"])
        sheet.append(["序号", "地区", "指标甲"])
        sheet.append([None, "惠州市", "=SUM(C4:C9)"])
        for index, name in enumerate(
            ("惠阳区", "博罗县", "惠东县", "龙门县", "惠州市本部", "大亚湾"),
            start=1,
        ):
            sheet.append([index, name, 10])
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)
    book.close()

    # openpyxl writes formulas without cached values; populate the stored
    # OOXML values so both Python data_only and Calamine read the same result.
    patched_path = path.with_name(f"{path.stem}.patched{path.suffix}")
    ET.register_namespace("", MAIN_NS)
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(
        patched_path, "w", zipfile.ZIP_DEFLATED
    ) as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename in ("xl/worksheets/sheet1.xml", "xl/worksheets/sheet2.xml"):
                sheet_index = int(
                    item.filename.removesuffix(".xml").removeprefix("xl/worksheets/sheet")
                )
                root = ET.fromstring(data)
                formula_cell = root.find(f".//{{{MAIN_NS}}}c[@r='C3']")
                if formula_cell is None:
                    raise RuntimeError(f"formula cell C3 missing from {item.filename}")
                cached_value = formula_cell.find(f"{{{MAIN_NS}}}v")
                if cached_value is None:
                    cached_value = ET.SubElement(formula_cell, f"{{{MAIN_NS}}}v")
                cached_value.text = str(formula_results[sheet_index - 1])
                data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            target.writestr(item, data)
    patched_path.replace(path)


def main() -> None:
    save_book(
        FIXTURE_DIR / "balanced.xlsx",
        {
            "本期分地区": [
                ["惠州市分地区数据"],
                ["序号", "地区", "指标甲", "指标乙", "指标丙"],
                [None, "惠州市", 100, 1000, 0],
                [1, "惠阳区", 10, 100, 0],
                [2, "博罗县", 20, 200, 0],
                [3, "惠东县", 15, 150, 0],
                [4, "龙门县", 10, 100, 0],
                [5, "惠州市本部", 20, 200, 0],
                [6, "大亚湾", 25, 250, 0],
                [None, "惠州市", 999, 999, 999],
            ]
        },
    )
    save_book(
        FIXTURE_DIR / "mismatch.xlsx",
        {
            "分地区汇总": [
                ["地区总分校验"],
                ["序号", "地区", "金额", "比例", "非数值"],
                [None, "惠州市", 100, 1000, "bad"],
                [1, "惠阳区", 40, 500, "bad"],
                [2, "博罗县", 50, 500, "bad"],
            ]
        },
    )
    save_book(FIXTURE_DIR / "no_target.xlsx", {"普通表": [["地区", "金额"], ["惠州市", 100]]})
    save_book(
        FIXTURE_DIR / "non_numeric_missing.xlsx",
        {
            "年度分地区": [
                ["年度数据"],
                ["序号", "地区", "金额"],
                [None, "惠州市", "bad"],
                [1, "惠阳区", "bad"],
            ]
        },
    )
    (FIXTURE_DIR / "broken.xlsx").write_bytes(b"not a workbook")
    cases = {
        "schema_version": 1,
        "run_files": [
            "balanced.xlsx",
            "mismatch.xlsx",
            "no_target.xlsx",
            "non_numeric_missing.xlsx",
            "broken.xlsx",
        ],
        "expected_stats": {"files": 4, "sheets": 3, "errors": 1, "failed_files": 1},
    }
    (FIXTURE_DIR / "cases.json").write_text(
        json.dumps(cases, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )

    cached_fixture_dir = FIXTURE_DIR / "cached_formula"
    save_cached_formula_book(cached_fixture_dir / "cached_formula_total.xlsx")
    cached_cases = {
        "schema_version": 1,
        "run_files": ["cached_formula_total.xlsx"],
        "expected_stats": {"files": 1, "sheets": 2, "errors": 1, "failed_files": 0},
    }
    (cached_fixture_dir / "cases.json").write_text(
        json.dumps(cached_cases, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    print(json.dumps({"fixture_dir": str(FIXTURE_DIR), "case_count": len(cases["run_files"])}, ensure_ascii=False))


if __name__ == "__main__":
    main()
