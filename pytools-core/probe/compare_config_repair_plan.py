"""Compare Python config repair with the Rust plan and gated XLSX modifier execution."""

from __future__ import annotations

import hashlib
import json
import shlex
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment

sys.dont_write_bytecode = True
WORKSPACE = Path(__file__).resolve().parents[2]
PYTOOLS_SOURCE = Path(r"C:\Users\AZI\Desktop\VibeCode\4.pytools-main")


def wsl_path(path: Path) -> str:
    return f"/mnt/{path.drive.rstrip(':').lower()}/{path.as_posix()[2:].lstrip('/')}"


def create_incomplete_config(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "全局配置"
    ws.append(["wrong-key", "wrong-value", "wrong-note"])
    ws.append(["输入目录", "custom-input", "user note"])
    ws.append(["输出目录", None, None])
    ws.append(["错误策略", "first-policy", None])
    ws.append(["错误策略", None, None])

    rename = wb.create_sheet("重命名配置")
    for column in range(1, 18):
        rename.cell(1, column, f"wrong-{column}")
    rename["A2"] = "user alias"
    rename["L1"].comment = Comment("obsolete header note", "user")
    rename["L2"].comment = Comment("obsolete cell note", "user")

    archive = wb.create_sheet("归档类型配置")
    archive.append(["wrong-1", "wrong-2", "wrong-3", "wrong-4", "wrong-5", "wrong-6"])

    sort = wb.create_sheet("文件排序配置")
    sort.append(["是否启用", "排序前缀", "匹配关键词", "备注"])
    sort.append(["否", "777", "keep-me", "existing data prevents examples"])

    for obsolete in ("校验规则生成配置", "校验规则映射表"):
        wb.create_sheet(obsolete)["A1"] = "remove-me"
    custom = wb.create_sheet("Custom user sheet")
    custom["A1"] = "keep custom data"
    custom["B3"].comment = Comment("keep custom comment", "user")
    wb.save(path)
    wb.close()


def workbook_projection(path: Path) -> dict[str, object]:
    wb = load_workbook(path, data_only=False)
    sheets = []
    for ws in wb.worksheets:
        max_row = max(ws.max_row, 1)
        max_column = max(ws.max_column, 1)
        values = [
            [ws.cell(row, column).value for column in range(1, max_column + 1)]
            for row in range(1, max_row + 1)
        ]
        comments = {
            cell.coordinate: {"text": cell.comment.text, "author": cell.comment.author}
            for row in ws.iter_rows()
            for cell in row
            if cell.comment is not None
        }
        sheets.append({"name": ws.title, "values": values, "comments": comments})
    wb.close()
    return {"sheet_names": [sheet["name"] for sheet in sheets], "sheets": sheets}


def rust_repair_plan(source: Path, candidate: Path) -> dict[str, object]:
    workspace = wsl_path(WORKSPACE)
    command = (
        f"cd {shlex.quote(workspace)} && cargo run --quiet -p pytools-core --bin pytools -- "
        f"config-repair-plan --input {shlex.quote(wsl_path(source))} "
        f"--output {shlex.quote(wsl_path(candidate))}"
    )
    result = subprocess.run(["wsl", "-e", "sh", "-lc", command], capture_output=True, check=False)
    stdout = result.stdout.decode("utf-8-sig", errors="replace").replace("\x00", "")
    if result.returncode != 0:
        raise RuntimeError(f"Rust config repair plan failed: {stdout}\n{result.stderr.decode('utf-8', errors='replace')}")
    for index, char in enumerate(stdout):
        if char == "{":
            try:
                return json.loads(stdout[index:])
            except json.JSONDecodeError:
                continue
    raise RuntimeError(f"Rust config repair emitted no JSON: {stdout}")


def rust_execute_repair(
    source: Path, candidate: Path, capability_path: Path, plan: dict[str, object]
) -> dict[str, object]:
    platform = "synthetic-pt-5-1-config-repair-probe"
    capability = {
        "engine": "openpyxl_worker",
        "status": "passed",
        "format": plan["format"],
        "operations": sorted({operation["operation"] for operation in plan["operations"]}),
        "supported_features": plan["workbook"]["features_present"],
        "preserved": plan["preserve"],
        "platform": platform,
        "evidence": {
            "sample_id": "synthetic-pt-5-1-config-repair",
            "report_path": "pytools-core/probe/compare_config_repair_plan.py",
            "adapter_version": "openpyxl-worker-1",
            "engine_version": "openpyxl 3.1.5",
            "platform": platform,
        },
    }
    capability_path.write_text(
        json.dumps(capability, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    workspace = wsl_path(WORKSPACE)
    package_dir = wsl_path(WORKSPACE / "target" / "libexec" / "pytools" / "openpyxl")
    runtime_root = wsl_path(WORKSPACE / ".test-openpyxl-venv")
    worker_source = wsl_path(WORKSPACE / "crates" / "excel-editor" / "python" / "openpyxl_worker.py")
    command = "\n".join([
        "set -eu",
        f"workspace={shlex.quote(workspace)}",
        f"package_dir={shlex.quote(package_dir)}",
        f"runtime_root={shlex.quote(runtime_root)}",
        f"worker_source={shlex.quote(worker_source)}",
        "if [ -e \"$package_dir\" ] || [ -L \"$package_dir\" ]; then echo 'test runtime layout already exists' >&2; exit 73; fi",
        "mkdir -p \"$package_dir/worker\"",
        "ln -s \"$runtime_root\" \"$package_dir/runtime\"",
        "cp \"$worker_source\" \"$package_dir/worker/openpyxl_worker.py\"",
        "cleanup() { expected=\"$(readlink -f \"$workspace/target\")/libexec/pytools/openpyxl\"; actual=\"$(readlink -f \"$package_dir\")\"; [ \"$actual\" = \"$expected\" ] || { echo 'refusing to clean unexpected runtime path' >&2; return 1; }; rm -f -- \"$package_dir/runtime\" \"$package_dir/worker/openpyxl_worker.py\"; rmdir \"$package_dir/worker\" \"$package_dir\"; }",
        "trap cleanup EXIT",
        f"cd {shlex.quote(workspace)}",
        "cargo run --quiet -p pytools-core --bin pytools -- "
        f"config-repair-plan --input {shlex.quote(wsl_path(source))} "
        f"--output {shlex.quote(wsl_path(candidate))} "
        f"--capabilities {shlex.quote(wsl_path(capability_path))} "
        f"--platform {shlex.quote(platform)}",
    ])
    result = subprocess.run(["wsl", "-e", "sh", "-lc", command], capture_output=True, check=False)
    stdout = result.stdout.decode("utf-8-sig", errors="replace").replace("\x00", "")
    stderr = result.stderr.decode("utf-8-sig", errors="replace").replace("\x00", "")
    if result.returncode != 0:
        raise RuntimeError(f"Rust config repair execution failed: {stdout}\n{stderr}")
    for index, char in enumerate(stdout):
        if char == "{":
            try:
                return json.loads(stdout[index:])
            except json.JSONDecodeError:
                continue
    raise RuntimeError(f"Rust config repair execution emitted no JSON: {stdout}\n{stderr}")


def apply_repair_plan_to_copy(source: Path, output: Path, plan: dict[str, object]) -> None:
    shutil.copy2(source, output)
    wb = load_workbook(output)
    for operation in plan.get("operations", []):
        kind = operation["operation"]
        if kind == "create_sheet":
            if operation["name"] not in wb.sheetnames:
                wb.create_sheet(operation["name"])
        elif kind == "delete_sheet":
            if operation["name"] in wb.sheetnames:
                wb.remove(wb[operation["name"]])
        elif kind == "set_cell_value":
            value = operation["value"]
            cell_value = value.get("value")
            wb[operation["sheet"]][operation["cell"]] = cell_value
        elif kind == "set_comment":
            cell = wb[operation["sheet"]][operation["cell"]]
            if operation.get("text") is None:
                cell.comment = None
            else:
                cell.comment = Comment(operation["text"], operation.get("author") or "")
        else:
            raise AssertionError(f"unexpected operation in config repair plan: {operation}")
    wb.save(output)
    wb.close()


def main() -> int:
    sys.dont_write_bytecode = True
    if not PYTOOLS_SOURCE.is_dir():
        raise FileNotFoundError(PYTOOLS_SOURCE)
    sys.path.insert(0, str(PYTOOLS_SOURCE / "PyCode"))
    from core.config_init import initialize_or_repair_config

    with tempfile.TemporaryDirectory(prefix="pytools-config-repair-", ignore_cleanup_errors=True) as temp:
        root = Path(temp)
        original = root / "incomplete.xlsx"
        python_copy = root / "python_repaired.xlsx"
        plan_copy = root / "rust_plan_applied_copy.xlsx"
        candidate = root / "candidate.xlsx"
        capability_path = root / "test-capability.json"
        create_incomplete_config(original)
        original_hash = hashlib.sha256(original.read_bytes()).hexdigest()
        shutil.copy2(original, python_copy)
        python_result = initialize_or_repair_config(python_copy)
        plan_result = rust_repair_plan(original, candidate)
        if plan_result.get("status") == "already_complete":
            raise AssertionError("incomplete fixture was reported complete")
        plan = plan_result.get("plan", plan_result)
        apply_repair_plan_to_copy(original, plan_copy, plan)
        execution = rust_execute_repair(original, candidate, capability_path, plan)
        if execution.get("execution", {}).get("engine") != "openpyxl_worker":
            raise AssertionError(f"unexpected backend execution: {execution.get('execution')}")
        expected = workbook_projection(python_copy)
        actual_plan = workbook_projection(plan_copy)
        actual_engine = workbook_projection(candidate)
        if expected != actual_plan or expected != actual_engine:
            raise AssertionError(json.dumps({
                "python": expected,
                "rust_plan_applied": actual_plan,
                "rust_openpyxl_candidate": actual_engine,
            }, ensure_ascii=False, indent=2))
        if hashlib.sha256(original.read_bytes()).hexdigest() != original_hash:
            raise AssertionError("config-repair-plan changed the source workbook")
        if not candidate.is_file():
            raise AssertionError("ExcelModifier did not create the validated candidate workbook")
        report = {
            "status": "passed_semantics_and_openpyxl_execution_synthetic",
            "scope": "5-1 existing configuration repair plan and capability-gated execution",
            "environment": {
                "python_baseline": "Original PyTools .venv; Windows Python 3.12.10, openpyxl 3.1.5",
                "rust_cli": "Ubuntu on WSL2 x86_64; rustc 1.99.1, cargo 1.99.1",
                "openpyxl_worker_runtime": "workspace isolated test venv; openpyxl 3.1.5",
                "windows_native_rust_validated": False,
                "kylin_v10_validated": False,
                "uos_v20_validated": False,
                "glibc_2_28_validated": False,
                "offline_deb_validated": False,
            },
            "python_repair_counts": {"created_sheets": python_result["created_sheets"], "repaired_sheets": python_result["repaired_sheets"]},
            "rust_plan_operations": len(plan["operations"]),
            "comparison": ["sheet order", "cell values", "comment text and author"],
            "execution": {
                "engine": execution["execution"]["engine"],
                "candidate_created": True,
                "temporary_test_capability_only": True,
                "production_capability_created": False,
                "default_route_changed": False,
            },
            "engine_evaluation": {
                "umya_spreadsheet": {
                    "status": "rejected_for_this_synthetic_operation_batch",
                    "observation": "Python readback of the Umya candidate did not retain expected repaired header comments; parity failed.",
                    "production_capability_created": False,
                },
                "openpyxl_worker": {
                    "status": "passed_synthetic_and_real_configuration_copy_parity",
                    "selection_is_default": False,
                },
            },
            "scenarios": ["wrong headers/comments", "missing sheets and global defaults", "duplicate/blank global keys", "default examples only for empty sheets", "obsolete sheet removal", "rename extra cells", "unrelated user sheet/data/comments preserved"],
            "synthetic_input_sha256_unchanged": original_hash,
            "python_baseline_source_sha256": hashlib.sha256(
                (PYTOOLS_SOURCE / "PyCode" / "core" / "config_init.py").read_bytes()
            ).hexdigest(),
            "production_capability_created": False,
            "default_route_changed": False,
            "limitations": [
                "Synthetic XLSX only; the real user configuration and target Linux package are not validated.",
                "The probe compares sheet order, cell values, and comment text/author; style, merge, dimensions, and broader OOXML preservation are asserted separately or remain unverified.",
                "The openpyxl capability is generated only for this synthetic plan and is not production evidence.",
            ],
        }
        report_path = WORKSPACE / "pytools-core" / "probe" / "reports" / "config_repair_execution_20261023.json"

        real_config = PYTOOLS_SOURCE / "PyCode" / "config" / "config.xlsx"
        if not real_config.is_file():
            raise FileNotFoundError(real_config)
        real_config_hash = hashlib.sha256(real_config.read_bytes()).hexdigest()
        real_python_copy = root / "real_python_repaired.xlsx"
        real_rust_input = root / "real_rust_input.xlsx"
        real_candidate = root / "real_rust_candidate.xlsx"
        shutil.copy2(real_config, real_python_copy)
        shutil.copy2(real_config, real_rust_input)
        real_python_result = initialize_or_repair_config(real_python_copy)
        real_plan_result = rust_repair_plan(real_rust_input, real_candidate)
        if real_plan_result.get("status") == "already_complete":
            raise AssertionError("real config Python repair reported changes but Rust returned already_complete")
        real_plan = real_plan_result.get("plan", real_plan_result)
        real_execution = rust_execute_repair(
            real_rust_input,
            real_candidate,
            root / "real-test-capability.json",
            real_plan,
        )
        if real_execution.get("execution", {}).get("engine") != "openpyxl_worker":
            raise AssertionError(f"unexpected real config backend: {real_execution.get('execution')}")
        real_expected = workbook_projection(real_python_copy)
        real_actual = workbook_projection(real_candidate)
        if real_expected != real_actual:
            raise AssertionError(json.dumps({
                "python_real_config": real_expected,
                "rust_real_config_candidate": real_actual,
            }, ensure_ascii=False, indent=2))
        if hashlib.sha256(real_config.read_bytes()).hexdigest() != real_config_hash:
            raise AssertionError("read-only PyTools source configuration changed")
        if hashlib.sha256(real_rust_input.read_bytes()).hexdigest() != real_config_hash:
            raise AssertionError("Rust config repair changed its source copy")
        report["real_configuration"] = {
            "source": "4.pytools-main/PyCode/config/config.xlsx (read-only; only temporary copies were modified)",
            "source_sha256_unchanged": real_config_hash,
            "python_repair_counts": {
                "created_sheets": real_python_result["created_sheets"],
                "repaired_sheets": real_python_result["repaired_sheets"],
            },
            "rust_plan_operations": len(real_plan["operations"]),
            "engine": real_execution["execution"]["engine"],
            "python_rust_projection_equal": True,
            "rust_input_copy_unchanged": True,
            "temporary_test_capability_only": True,
        }
        report["status"] = "passed_synthetic_and_real_config_execution_parity"
        report["overall_production_acceptance"] = "unverified"
        report["limitations"] = [
            "The real sample is the PyTools configuration workbook, not a business data workbook.",
            "Value/sheet/comment projections were compared; full preservation of every real OOXML object and target Linux package behavior remains unverified.",
            "The capability records were generated only for temporary test copies and are not production evidence.",
        ]
        report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        print(json.dumps(report, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
