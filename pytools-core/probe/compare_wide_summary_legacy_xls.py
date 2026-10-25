from __future__ import annotations

import json
import logging
import shutil
import subprocess
import sys
import tempfile
import time
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook

REPO = Path(__file__).resolve().parents[2]
PYTOOLS_SOURCE = Path(r"C:\Users\AZI\Desktop\VibeCode\4.pytools-main")
FIXTURE = REPO / "pytools-core" / "tests" / "fixtures" / "timeline" / "wide_summary_legacy.xls"
RUST_CLI = REPO / "target" / "debug" / "pytools.exe"
REPORT = REPO / "pytools-core" / "probe" / "reports" / "wide_summary_legacy_xls_metadata_parity.json"

sys.path.insert(0, str(PYTOOLS_SOURCE / "PyCode"))
from core.config_xlsx import GlobalConfig, TimelineRule
from core.wide_397_398 import run_wide_summary


def make_config(path: Path, root: Path) -> None:
    app = root / "app"
    input_dir = app / "input"
    output_dir = app / "output"
    log_dir = app / "logs"
    input_dir.mkdir(parents=True, exist_ok=True)
    output_dir.mkdir(parents=True, exist_ok=True)
    log_dir.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    global_sheet = workbook.active
    global_sheet.title = "全局配置"
    global_sheet.append(["配置项", "值"])
    for key, value in (
        ("输入目录", str(input_dir)),
        ("输出目录", str(output_dir)),
        ("日志目录", str(log_dir)),
        ("错误策略", "fail_fast"),
        ("源文件扩展名", ".xlsx"),
        ("时序规则命中策略", "all_match"),
    ):
        global_sheet.append([key, value])

    timeline = workbook.create_sheet("时序提取规则")
    timeline.append([
        "是否启用", "规则名称", "工作簿关键字", "工作表关键字", "行头列",
        "列表头行", "必含列头", "必含行头", "数据起始行", "数据结束行",
        "数据起始列", "数据结束列", "跳过关键字", "目标工作簿路径",
        "目标工作表", "启用目标写入", "set区域", "目标去重列", "是否输出行序号",
    ])
    timeline.append([
        "是", "WideRule", "input", "Data", "A", "1", "", "", 2, "",
        "B", "B", "", "", "", "否", "", "", "是",
    ])

    path_maps = workbook.create_sheet("路径标准化映射")
    path_maps.append([
        "是否启用", "映射名称", "适用规则名", "工作簿关键字", "工作表关键字",
        "作用对象", "匹配方式", "原始路径", "标准路径", "备注",
    ])
    workbook.save(path)


def make_python_rule() -> TimelineRule:
    return TimelineRule(
        enabled=True,
        name="WideRule",
        wb_keyword="input",
        sheet_keyword="Data",
        row_header_col=1,
        col_header_row=1,
        row_header_cols=[1],
        row_header_col_specified=True,
        col_header_rows=[1],
        data_row_start=2,
        data_col_start=2,
        data_col_end=2,
        set_items=[],
        output_row_no=True,
        target_wb_path=None,
        target_sheet="",
        target_write_enabled=False,
    )


def close_python_log_handlers() -> None:
    for logger in list(logging.root.manager.loggerDict.values()):
        if isinstance(logger, logging.Logger):
            for handler in logger.handlers[:]:
                logger.removeHandler(handler)
                handler.flush()
                handler.close()
    logging.shutdown()


def cleanup_temp(path: Path) -> None:
    for _ in range(30):
        try:
            shutil.rmtree(path)
            return
        except PermissionError:
            time.sleep(0.05)
    raise PermissionError(f"could not release synthetic test directory: {path}")


def run_case(command: str) -> dict[str, object]:
    root = Path(tempfile.mkdtemp(prefix="pytools-wide-xls-parity-"))
    try:
        app = root / "app"
        input_dir = app / "input"
        source = input_dir / "input.xls"
        input_dir.mkdir(parents=True)
        shutil.copyfile(FIXTURE, source)
        source_before = source.read_bytes()

        config = app / "config" / "config.xlsx"
        config.parent.mkdir(parents=True)
        make_config(config, root)

        python_global = GlobalConfig(
            input_dir=input_dir,
            output_dir=root / "python-output",
            log_dir=root / "python-logs",
            error_policy="fail_fast",
            source_exts=[".xlsx"],
        )
        try:
            python_summary = run_wide_summary(
                [make_python_rule()],
                [],
                python_global,
                command == "wide-summary-row-suffix",
                source_paths=[source],
            )
        finally:
            close_python_log_handlers()

        rust = subprocess.run(
            [
                str(RUST_CLI),
                command,
                "--config",
                str(config),
                "--log-view",
                "developer",
                str(source),
            ],
            capture_output=True,
            check=False,
        )
        stdout = rust.stdout.decode("utf-8-sig", errors="replace").replace("\x00", "")
        stderr = rust.stderr.decode("utf-8-sig", errors="replace").replace("\x00", "")
        if rust.returncode != 0:
            raise RuntimeError(f"{command} exited {rust.returncode}: {stderr[-1500:]}")
        rust_report = json.loads(stdout)
        history_path = Path(rust_report["run_log_path"])
        history = json.loads(history_path.read_text(encoding="utf-8"))

        if python_summary is not None:
            raise AssertionError(f"synthetic smoke XLS unexpectedly produced Python output: {python_summary}")
        if rust_report.get("output") is not None:
            raise AssertionError(f"Rust summary output differs from Python no-output case: {rust_report}")
        if history.get("status") != "skipped":
            raise AssertionError(f"Rust should skip output when no wide cells were extracted: {history}")
        if history.get("total_inputs") != 1 or history.get("completed_inputs") != 1:
            raise AssertionError(f"Rust did not complete the single XLS input: {history}")
        if history.get("failed_inputs") != 0 or history.get("matched_rule_sheets") != 1:
            raise AssertionError(f"Rust input/rule state differs from expected: {history}")
        if history.get("extracted_cells") != 0:
            raise AssertionError(f"unexpected extracted cells in smoke workbook: {history}")
        events = history.get("events", [])
        metadata = next(
            (
                event
                for event in events
                if event.get("phase") == "elapsed_duration_metadata"
            ),
            None,
        )
        if not metadata or metadata.get("status") != "skipped":
            raise AssertionError(f"BIFF metadata skip event is missing: {history}")
        if "格式识别未验证" not in metadata.get("user_message", ""):
            raise AssertionError(f"BIFF limitation is not visible in the event: {metadata}")
        if source.read_bytes() != source_before:
            raise AssertionError("the synthetic XLS source was modified")

        return {
            "entry": command,
            "python_summary": "none",
            "rust_summary": "none",
            "python_cells": 0,
            "rust_matched_sheets": history["matched_rule_sheets"],
            "rust_extracted_cells": history["extracted_cells"],
            "rust_completed_inputs": history["completed_inputs"],
            "rust_failed_inputs": history["failed_inputs"],
            "metadata_event": metadata["status"],
            "source_unchanged": True,
        }
    finally:
        close_python_log_handlers()
        cleanup_temp(root)


def main() -> int:
    if not RUST_CLI.is_file():
        raise FileNotFoundError(f"build the native PyTools CLI first: {RUST_CLI}")
    if not FIXTURE.is_file():
        raise FileNotFoundError(f"synthetic legacy XLS fixture missing: {FIXTURE}")

    cases = [
        run_case("wide-summary"),
        run_case("wide-summary-row-suffix"),
    ]
    report = {
        "schema_version": 1,
        "suite": "pytools.timeline.wide_summary.legacy_xls_metadata_parity",
        "date": datetime.now().date().isoformat(),
        "status": "passed_synthetic_python_rust_skip_parity",
        "baseline": "PyCode/core/wide_397_398.py::run_wide_summary",
        "rust_entry": "Windows-native pytools CLI wide-summary / wide-summary-row-suffix",
        "input": {
            "format": "xls",
            "kind": "synthetic smoke workbook",
            "fixture": "pytools-core/tests/fixtures/timeline/wide_summary_legacy.xls",
            "sheet": "Data",
            "dimensions": [2, 2],
        },
        "comparison": {
            "summary_output": "both implementations returned no summary for this smoke fixture",
            "rust_input_processing": "one input completed, one rule sheet matched, zero extractable wide cells",
            "elapsed_duration_metadata": "explicitly skipped for BIFF .xls; format classification remains unverified",
            "source_unchanged": all(case["source_unchanged"] for case in cases),
            "cases": cases,
        },
        "limitations": [
            "The synthetic smoke XLS contains no extractable wide data cells; this verifies the metadata failure regression and no-output parity, not a nonempty XLS summary.",
            "Real business workbooks and XLS elapsed-time/date formats remain unverified.",
            "This does not change the XLS production capability or claim full PT-1-3/PT-1-4 format acceptance.",
        ],
    }
    REPORT.parent.mkdir(parents=True, exist_ok=True)
    REPORT.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"status": report["status"], "report": str(REPORT), "cases": len(cases)}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())