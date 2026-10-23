"""Run read-only pytools operation 4-7 against isolated synthetic workbooks."""

from __future__ import annotations

import argparse
import hashlib
import json
import logging
import sys
import tempfile
from pathlib import Path

from openpyxl import Workbook


WORKSPACE = Path(__file__).resolve().parents[2]
CASES_PATH = WORKSPACE / "pytools-core/tests/fixtures/content_unrename_cases.json"
SOURCE_RELATIVE = Path("PyCode/core/convert_33_34_35_37.py")


def make_config(path: Path, case: dict[str, object], *, include_sort_sheet: bool) -> None:
    workbook = Workbook()
    rename_sheet = workbook.active
    rename_sheet.title = "重命名配置"
    rename_sheet.cell(1, 15, "内容重命名启用")
    rename_sheet.cell(1, 16, "内容重命名取值规则")
    rename_sheet.cell(1, 17, "备注")
    rename_sheet.cell(2, 15, "Y")
    rename_sheet.cell(2, 16, case["content_spec"])
    rename_sheet.cell(2, 17, "4-7 对拍")

    if include_sort_sheet:
        sort_sheet = workbook.create_sheet("文件排序配置")
        sort_sheet.append(["是否启用", "排序前缀", "匹配关键词", "备注"])
        for rule in case["sort_rows"]:
            sort_sheet.append(
                [rule["enabled"], rule["sort_prefix"], rule["keywords"], rule["remark"]]
            )

    workbook.save(path)
    workbook.close()


def build_inputs(input_dir: Path, case: dict[str, object]) -> list[Path]:
    for name in case.get("existing_names", []):
        (input_dir / name).write_bytes(b"occupied target")

    missing = set(case.get("missing_names", []))
    directories = set(case.get("directory_names", []))
    corrupt = set(case.get("corrupt_names", []))
    workbooks = case["workbooks"]
    paths = []
    for name in case["input_names"]:
        path = input_dir / name
        paths.append(path)
        if name in missing:
            continue
        if name in directories:
            path.mkdir()
            (path / "marker.txt").write_text("directory", encoding="utf-8")
        elif name in corrupt:
            path.write_bytes(b"not a workbook package")
        else:
            workbook_data = workbooks.get(name)
            if workbook_data is None:
                path.write_bytes(b"not a workbook package")
                continue
            workbook = Workbook()
            default_sheet = workbook.active
            default_sheet.title = next(iter(workbook_data))
            for sheet_index, (sheet_name, cells) in enumerate(workbook_data.items()):
                sheet = default_sheet if sheet_index == 0 else workbook.create_sheet(sheet_name)
                sheet.title = sheet_name
                for address, value in cells.items():
                    sheet[address] = value
            workbook.save(path)
            workbook.close()
    return paths


def close_source_loggers() -> None:
    for name, value in list(logging.root.manager.loggerDict.items()):
        if not name.endswith("4_7_batch_unrename_by_content"):
            continue
        if not isinstance(value, logging.Logger):
            continue
        for handler in list(value.handlers):
            handler.close()
            value.removeHandler(handler)


def run_case(function, root: Path, case: dict[str, object], *, include_sort_sheet: bool) -> dict[str, object]:
    input_dir = root / "input"
    input_dir.mkdir(parents=True)
    config_path = root / "config.xlsx"
    make_config(config_path, case, include_sort_sheet=include_sort_sheet)
    files = build_inputs(input_dir, case)
    original_bytes = {
        path.name: path.read_bytes()
        for path in files
        if path.is_file()
    }
    try:
        result = function(config_path, files, root / "logs")
    finally:
        close_source_loggers()

    final_entries = sorted(path.name for path in input_dir.iterdir())
    return {
        "stats": {
            key: result[key]
            for key in ("ok", "skip", "spec", "config_row", "config_remark")
        },
        "final_entries": final_entries,
        "original_bytes_sha256": {
            name: hashlib.sha256(data).hexdigest().upper()
            for name, data in sorted(original_bytes.items())
        },
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--source-root",
        type=Path,
        default=Path(r"C:\Users\AZI\Desktop\VibeCode\4.pytools-main"),
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=WORKSPACE / "pytools-core/probe/reports/content_unrename_47_python_baseline.json",
    )
    args = parser.parse_args()

    sys.dont_write_bytecode = True
    sys.path.insert(0, str(args.source_root))
    from PyCode.core.convert_33_34_35_37 import run_batch_unrename_by_content_47

    cases = json.loads(CASES_PATH.read_text(encoding="utf-8"))
    with tempfile.TemporaryDirectory(prefix="pytools-content-unrename-47-") as raw_root:
        root = Path(raw_root)
        main_result = run_case(
            run_batch_unrename_by_content_47,
            root / "main-case",
            cases["operation47"],
            include_sort_sheet=True,
        )
        legacy_result = run_case(
            run_batch_unrename_by_content_47,
            root / "legacy-case",
            cases["legacy_without_sort_sheet"],
            include_sort_sheet=False,
        )

    source_file = args.source_root / SOURCE_RELATIVE
    payload = {
        "source": {
            "project": str(args.source_root),
            "module": SOURCE_RELATIVE.as_posix(),
            "sha256": hashlib.sha256(source_file.read_bytes()).hexdigest().upper(),
            "operation": "run_batch_unrename_by_content_47",
        },
        "operation47": main_result,
        "legacy_without_sort_sheet": legacy_result,
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(args.output)


if __name__ == "__main__":
    main()
