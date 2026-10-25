from __future__ import annotations

import datetime as dt
import json
import sys
from pathlib import Path
from typing import Any


def cell_value(value: Any) -> dict[str, Any] | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return {"type": "boolean", "value": value}
    if isinstance(value, (int, float)):
        number = int(value) if float(value).is_integer() else float(value)
        return {"type": "number", "value": number}
    return {"type": "string", "value": str(value)}


def main() -> None:
    if len(sys.argv) != 4:
        raise SystemExit("usage: submission_check_baseline.py <pytools-source-root> <work-dir> <report.json>")
    source_root = Path(sys.argv[1]).resolve()
    work_root = Path(sys.argv[2]).resolve()
    report_path = Path(sys.argv[3]).resolve()
    work_root.mkdir(parents=True, exist_ok=True)
    sys.path.insert(0, str(source_root))

    from openpyxl import Workbook, load_workbook
    from PyCode.core.convert_33_34_35_37 import run_submission_check_48

    config = work_root / "config.xlsx"
    config_book = Workbook()
    sheet = config_book.active
    sheet.title = "提交检查配置"
    sheet.append(["是否启用", "日期关键词", "县区关键词", "机构关键词", "备用关键词", "备注"])
    sheet.append(["是", "2026", "Area", "Bank", "", "base rule"])
    sheet.append(["是", "2026", "Area", "Bank", "multi", "second rule"])
    sheet.append(["是", "2026", "Missing", "Bank", "", "missing rule"])
    sheet.append(["否", "disabled", "Area", "Bank", "", "ignored rule"])
    sheet.append(
        ["是", dt.datetime(2026, 6, 30, 1, 2, 3, 123000), "Area", "Bank", "", "date keyword"]
    )
    config_book.save(config)

    sources = work_root / "sources"
    sources.mkdir(exist_ok=True)
    files = [
        sources / "2026_Area_Bank_base.txt",
        sources / "2026_Area_Bank_multi.txt",
        sources / "misc_submission.txt",
    ]
    for index, path in enumerate(files):
        path.write_text(f"synthetic payload {index}", encoding="utf-8")

    output_dir = work_root / "output"
    log_dir = work_root / "logs"
    output_dir.mkdir(exist_ok=True)
    log_dir.mkdir(exist_ok=True)
    summary = run_submission_check_48(config, files, output_dir, log_dir)
    workbook = load_workbook(summary["output"], data_only=True, read_only=False)
    normalized_sheets = []
    for ws in workbook.worksheets:
        cells: dict[str, dict[str, Any]] = {}
        for row in ws.iter_rows():
            for cell in row:
                value = cell_value(cell.value)
                if value is None:
                    continue
                if value["type"] == "string":
                    value["value"] = value["value"].replace(str(work_root), "<ROOT>")
                cells[cell.coordinate] = value
        normalized_sheets.append({
            "name": ws.title,
            "max_row": ws.max_row,
            "max_column": ws.max_column,
            "cells": cells,
            "header_bold": [bool(ws.cell(1, column).font.bold) for column in range(1, ws.max_column + 1)],
        })
    payload = {
        "baseline": "PyCode.run_submission_check_48 synthetic contract",
        "python": sys.version.split()[0],
        "statistics": {key: int(summary[key]) for key in ("rules", "files", "submitted", "missing", "unmatched", "multi")},
        "output_workbook": {"sheets": normalized_sheets},
    }
    report_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"wrote {report_path.name}")


if __name__ == "__main__":
    main()
