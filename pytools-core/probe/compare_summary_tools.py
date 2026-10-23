"""Synthetic Python/Rust parity probe for pytools summary tools 6-1 and 6-2."""

from __future__ import annotations

import contextlib
import datetime as dt
import hashlib
import json
import logging
import math
import shlex
import subprocess
import sys
import tempfile
import xml.etree.ElementTree as ET
from pathlib import Path
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill

sys.dont_write_bytecode = True
WORKSPACE = Path(__file__).resolve().parents[2]
PYTOOLS_SOURCE = Path(r"C:\Users\AZI\Desktop\VibeCode\4.pytools-main")


def wsl_path(path: Path) -> str:
    return f"/mnt/{path.drive.rstrip(':').lower()}/{path.as_posix()[2:].lstrip('/')}"


def from_wsl_path(value: str) -> Path:
    normalized = value.replace("\\", "/")
    if not normalized.startswith("/mnt/"):
        raise ValueError(f"Rust output outside WSL mounted drive: {value}")
    return Path(f"{normalized[5].upper()}:\\{normalized[7:].replace('/', chr(92))}")


def make_source(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Branch_A"
    ws["A1"] = "机构"
    ws["B1"] = "余额"
    ws["C1"] = "余额"
    ws["D1"] = "华东"
    ws["A2"] = "甲"
    ws["B2"] = 10.5
    ws["C2"] = 20
    ws["A3"] = "乙"
    ws["B3"] = dt.date(2026, 6, 30)
    ws["C3"] = True
    ws["A4"] = "丙"
    ws["B4"] = "text"
    ws["C4"] = 0
    for column, error in enumerate(
        ("#DIV/0!", "#N/A", "#NAME?", "#NULL!", "#NUM!", "#REF!", "#VALUE!"), start=5
    ):
        ws.cell(row=4, column=column).value = error
    ws["F6"] = "merged value"
    ws.merge_cells("F6:G6")
    ws["E8"] = "=1+1"
    ws["F8"] = "=1+2"
    ws["H10"] = "dimension tail"
    ws["H10"].fill = PatternFill(fill_type="solid", fgColor="FFFF00")

    matching = wb.create_sheet("模板")
    matching["A1"] = "sheet"
    matching["B2"] = 5
    matching["C2"] = 6
    matching["D1"] = "华南"
    wb.save(path)



def patch_formula_cache(path: Path, address: str, cached_value: str) -> None:
    main_ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    ET.register_namespace("", main_ns)
    with ZipFile(path, "r") as source_archive:
        entries = [
            (info, source_archive.read(info.filename))
            for info in source_archive.infolist()
        ]

    worksheet_found = False
    patched_entries = []
    for info, payload in entries:
        if info.filename == "xl/worksheets/sheet1.xml":
            root = ET.fromstring(payload)
            cell = root.find(f".//{{{main_ns}}}c[@r='{address}']")
            if cell is None:
                raise AssertionError(f"formula cell {address} is missing from source")
            formula = cell.find(f"{{{main_ns}}}f")
            if formula is None:
                raise AssertionError(f"formula text is missing from {address}")
            cached = cell.find(f"{{{main_ns}}}v")
            if cached is None:
                cached = ET.SubElement(cell, f"{{{main_ns}}}v")
            cached.text = cached_value
            payload = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            worksheet_found = True
        patched_entries.append((info, payload))

    if not worksheet_found:
        raise AssertionError("source worksheet part is missing")
    with ZipFile(path, "w") as output_archive:
        for info, payload in patched_entries:
            output_archive.writestr(info, payload)

    formula_book = load_workbook(path, data_only=False, read_only=False)
    value_book = load_workbook(path, data_only=True, read_only=False)
    try:
        formula_ws = formula_book["Branch_A"]
        value_ws = value_book["Branch_A"]
        if formula_ws[address].value != "=1+1" or value_ws[address].value != 2:
            raise AssertionError("cached formula fixture did not round trip as expected")
        if formula_ws["F8"].value != "=1+2" or value_ws["F8"].value is not None:
            raise AssertionError("missing formula cache fixture did not round trip as expected")
    finally:
        formula_book.close()
        value_book.close()

def make_second_source(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Extra"
    ws["A2"] = "第二个工作簿"
    ws["C2"] = 42
    wb.create_sheet("Empty")
    wb.save(path)
    if path.suffix.lower() == ".xlsm":
        mark_macro_enabled(path)


def mark_macro_enabled(path: Path) -> None:
    content_types_ns = "http://schemas.openxmlformats.org/package/2006/content-types"
    with ZipFile(path, "r") as archive:
        entries = [(info, archive.read(info.filename)) for info in archive.infolist()]
    content_types_found = False
    updated_entries = []
    ET.register_namespace("", content_types_ns)
    for info, payload in entries:
        if info.filename == "[Content_Types].xml":
            root = ET.fromstring(payload)
            workbook_type = root.find(
                f"{{{content_types_ns}}}Override[@PartName='/xl/workbook.xml']"
            )
            if workbook_type is None:
                raise AssertionError("workbook content type is missing")
            workbook_type.set(
                "ContentType", "application/vnd.ms-excel.sheet.macroEnabled.main+xml"
            )
            payload = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            content_types_found = True
        updated_entries.append((info, payload))
    if not content_types_found:
        raise AssertionError("package content types are missing")
    with ZipFile(path, "w") as archive:
        for info, payload in updated_entries:
            archive.writestr(info, payload)


def make_comment_template(path: Path, duplicate_row_boundaries: bool = False) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "模板"
    ws["A1"] = "单位"
    ws["A2"] = "区域起"
    ws["A4"] = "区域止"
    ws["A2"].comment = Comment("行区域1", "probe")
    if duplicate_row_boundaries:
        ws["A3"].comment = Comment("行区域1", "probe")
    ws["A4"].comment = Comment("行区域#1", "probe")
    if duplicate_row_boundaries:
        ws["A5"].comment = Comment("行区域#1", "probe")
    ws["B1"] = "余额"
    ws["C1"] = "余额"
    ws["B1"].comment = Comment("列区域1", "probe")
    ws["C1"].comment = Comment("列区域#1", "probe")
    ws["D1"] = "区域"
    ws["D1"].comment = Comment("set(区域)", "probe")
    wb.save(path)


def python_baselines(
    sources: list[Path], comment_source: Path, template: Path, output: Path
) -> tuple[Path | None, Path | None]:
    sys.path.insert(0, str(PYTOOLS_SOURCE / "PyCode"))
    from core.summary_321_322 import run_summary_by_comment, run_summary_by_usedrange

    output.mkdir(parents=True, exist_ok=True)
    with contextlib.redirect_stdout(None):
        used = run_summary_by_usedrange(sources, output / "python_used", output / "python_logs")
        comments = run_summary_by_comment(
            template, [comment_source], output / "python_comments", output / "python_logs"
        )
    logging.shutdown()
    return (
        Path(used["path"]) if used.get("saved") else None,
        Path(comments["path"]) if comments.get("saved") else None,
    )


def python_used_range_only(sources: list[Path], output: Path) -> Path | None:
    sys.path.insert(0, str(PYTOOLS_SOURCE / "PyCode"))
    from core.summary_321_322 import run_summary_by_usedrange

    output.mkdir(parents=True, exist_ok=True)
    with contextlib.redirect_stdout(None):
        result = run_summary_by_usedrange(sources, output / "used", output / "logs")
    logging.shutdown()
    return Path(result["path"]) if result.get("saved") else None


def python_comments_only(template: Path, sources: list[Path], output: Path) -> Path | None:
    sys.path.insert(0, str(PYTOOLS_SOURCE / "PyCode"))
    from core.summary_321_322 import run_summary_by_comment

    output.mkdir(parents=True, exist_ok=True)
    with contextlib.redirect_stdout(None):
        result = run_summary_by_comment(
            template, sources, output / "comments", output / "logs"
        )
    logging.shutdown()
    return Path(result["path"]) if result.get("saved") else None


def python_nocom_used_range(sources: list[Path], output: Path) -> Path | None:
    sys.path.insert(0, str(PYTOOLS_SOURCE / "NoComPyTools"))
    from nocompytools.core.summary_321_322 import run_summary_by_usedrange

    with contextlib.redirect_stdout(None):
        result = run_summary_by_usedrange(sources, output / "nocom_used", output / "nocom_logs")
    logging.shutdown()
    return Path(result["path"]) if result.get("saved") else None


def run_rust(command_name: str, args: list[Path], output: Path) -> tuple[Path | None, dict[str, object]]:
    cli = " ".join(shlex.quote(wsl_path(path)) for path in args)
    command = (
        f"cd {shlex.quote(wsl_path(WORKSPACE))} && cargo run --quiet -p pytools-core --bin pytools -- "
        f"{command_name} --output-dir {shlex.quote(wsl_path(output))} {cli}"
    )
    completed = subprocess.run(["wsl", "-e", "sh", "-lc", command], check=True, capture_output=True)
    stdout = completed.stdout.decode("utf-8-sig", errors="replace").replace("\x00", "")
    stderr = completed.stderr.decode("utf-8-sig", errors="replace").replace("\x00", "")
    for index, character in enumerate(stdout):
        if character != "{":
            continue
        try:
            result = json.loads(stdout[index:])
            path = result.get("path")
            return (from_wsl_path(path) if path else None), result
        except json.JSONDecodeError:
            continue
    raise RuntimeError(f"Rust summary CLI emitted no JSON: {stdout}\n{stderr}")


def run_rust_comments(template: Path, source: Path, output: Path) -> tuple[Path | None, dict[str, object]]:
    command = (
        f"cd {shlex.quote(wsl_path(WORKSPACE))} && cargo run --quiet -p pytools-core --bin pytools -- "
        f"summary-comments --template {shlex.quote(wsl_path(template))} "
        f"--output-dir {shlex.quote(wsl_path(output))} {shlex.quote(wsl_path(source))}"
    )
    completed = subprocess.run(["wsl", "-e", "sh", "-lc", command], check=True, capture_output=True)
    stdout = completed.stdout.decode("utf-8-sig", errors="replace").replace("\x00", "")
    stderr = completed.stderr.decode("utf-8-sig", errors="replace").replace("\x00", "")
    for index, character in enumerate(stdout):
        if character != "{":
            continue
        try:
            result = json.loads(stdout[index:])
            path = result.get("path")
            return (from_wsl_path(path) if path else None), result
        except json.JSONDecodeError:
            continue
    raise RuntimeError(f"Rust comment-summary CLI emitted no JSON: {stdout}\n{stderr}")


def normalize(value: object) -> object:
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.isoformat()
    if isinstance(value, float) and not math.isfinite(value):
        return str(value)
    return value


def workbook_projection(path: Path) -> tuple[list[dict[str, object]], list[dict[str, object]]]:
    wb = load_workbook(path, data_only=False)
    sheets = []
    for ws in wb.worksheets:
        cells = [cell for row in ws.iter_rows() for cell in row]
        values = {
            cell.coordinate: normalize(cell.value)
            for cell in cells
            if cell.value is not None
        }
        types = {
            cell.coordinate: cell.data_type
            for cell in cells
            if cell.value is not None
        }
        styles = {
            cell.coordinate: {
                "number_format": cell.number_format,
                "bold": bool(cell.font.bold),
                "fill_type": cell.fill.fill_type,
                "fill_color": cell.fill.fgColor.rgb if cell.fill.fgColor.type == "rgb" else None,
            }
            for cell in cells
            if cell.value is not None
        }
        setup_properties = ws.sheet_properties.pageSetUpPr
        rows = [
            [normalize(cell.value) for cell in row]
            for row in ws.iter_rows()
        ]
        sheets.append({
            "name": ws.title,
            "rows": rows,
            "values": values,
            "types": types,
            "styles": styles,
            "freeze_panes": ws.freeze_panes,
            "header_bold": [ws.cell(1, column).font.bold for column in range(1, ws.max_column + 1)],
            "merged_ranges": sorted(str(item) for item in ws.merged_cells.ranges),
            "formula_cells": {
                cell.coordinate: cell.value
                for cell in cells
                if cell.data_type == "f"
            },
            "defined_names": sorted(wb.defined_names.keys()),
            "print_area": str(ws.print_area) if ws.print_area else None,
            "page_setup": {
                "orientation": ws.page_setup.orientation or "portrait",
                "fitToPage": setup_properties.fitToPage if setup_properties else None,
                "fitToWidth": ws.page_setup.fitToWidth,
                "fitToHeight": ws.page_setup.fitToHeight,
                "paperSize": ws.page_setup.paperSize,
                "scale": ws.page_setup.scale,
            },
            "page_margins": {
                "left": ws.page_margins.left,
                "right": ws.page_margins.right,
                "top": ws.page_margins.top,
                "bottom": ws.page_margins.bottom,
                "header": ws.page_margins.header,
                "footer": ws.page_margins.footer,
            },
            "horizontal_centered": ws.print_options.horizontalCentered,
            "vertical_centered": ws.print_options.verticalCentered,
        })
    wb.close()
    return sheets, [
        {"name": sheet["name"], "rows": len(sheet["rows"]), "columns": max((len(row) for row in sheet["rows"]), default=0)}
        for sheet in sheets
    ]


def assert_equivalent(
    expected: Path | None,
    actual: Path | None,
    label: str,
    compare_page_settings: bool = True,
) -> None:
    if expected is None or actual is None:
        if expected is not actual:
            raise AssertionError(f"{label}: saved status differs Python={expected}, Rust={actual}")
        return
    expected_sheets, expected_dimensions = workbook_projection(expected)
    actual_sheets, actual_dimensions = workbook_projection(actual)
    if not compare_page_settings:
        omitted_fields = {
            "page_setup",
            "page_margins",
            "horizontal_centered",
            "vertical_centered",
        }
        expected_sheets = [
            {key: value for key, value in sheet.items() if key not in omitted_fields}
            for sheet in expected_sheets
        ]
        actual_sheets = [
            {key: value for key, value in sheet.items() if key not in omitted_fields}
            for sheet in actual_sheets
        ]
    if expected_sheets != actual_sheets or expected_dimensions != actual_dimensions:
        field_differences = []
        for expected_sheet, actual_sheet in zip(expected_sheets, actual_sheets):
            field_differences.extend(
                f"{expected_sheet['name']}.{field}"
                for field in expected_sheet
                if expected_sheet.get(field) != actual_sheet.get(field)
            )
            if expected_sheet.get("styles") != actual_sheet.get("styles"):
                expected_styles = expected_sheet.get("styles", {})
                actual_styles = actual_sheet.get("styles", {})
                style_cells = sorted(set(expected_styles) | set(actual_styles))
                style_differences = {
                    cell: {"python": expected_styles.get(cell), "rust": actual_styles.get(cell)}
                    for cell in style_cells
                    if expected_styles.get(cell) != actual_styles.get(cell)
                }
            else:
                style_differences = {}
        if expected_dimensions != actual_dimensions:
            field_differences.append("worksheet_dimensions")
        raise AssertionError(json.dumps({
            "label": label,
            "different_fields": field_differences,
            "style_differences": style_differences,
            "python_page_setup": expected_sheets[0].get("page_setup"),
            "rust_page_setup": actual_sheets[0].get("page_setup"),
            "python_page_margins": expected_sheets[0].get("page_margins"),
            "rust_page_margins": actual_sheets[0].get("page_margins"),
            "python_dimensions": expected_dimensions,
            "rust_dimensions": actual_dimensions,
        }, ensure_ascii=False, default=str, indent=2))


def source_fingerprints(paths: list[Path]) -> dict[str, str]:
    return {
        path.name: hashlib.sha256(path.read_bytes()).hexdigest()
        for path in paths
    }


def projection_difference_fields(expected: Path, actual: Path) -> list[str]:
    expected_sheets, expected_dimensions = workbook_projection(expected)
    actual_sheets, actual_dimensions = workbook_projection(actual)
    differences = []
    for expected_sheet, actual_sheet in zip(expected_sheets, actual_sheets):
        differences.extend(
            f"{expected_sheet['name']}.{field}"
            for field in expected_sheet
            if expected_sheet.get(field) != actual_sheet.get(field)
        )
    if expected_dimensions != actual_dimensions:
        differences.append("worksheet_dimensions")
    return sorted(differences)




def projection_difference_summary(expected: Path, actual: Path) -> dict[str, object]:
    expected_sheets, expected_dimensions = workbook_projection(expected)
    actual_sheets, actual_dimensions = workbook_projection(actual)
    cell_differences: dict[str, dict[str, object]] = {}
    for expected_sheet, actual_sheet in zip(expected_sheets, actual_sheets):
        for field in ("values", "types", "styles"):
            expected_cells = expected_sheet.get(field, {})
            actual_cells = actual_sheet.get(field, {})
            coordinates = sorted(set(expected_cells) | set(actual_cells))
            changed = [
                coordinate
                for coordinate in coordinates
                if expected_cells.get(coordinate) != actual_cells.get(coordinate)
            ]
            if changed:
                cell_differences[f"{expected_sheet['name']}.{field}"] = {
                    "count": len(changed),
                    "coordinates_sample": changed[:12],
                    "examples": [
                        {
                            "cell": coordinate,
                            "pycode": expected_cells.get(coordinate),
                            "nocom_pytools": actual_cells.get(coordinate),
                        }
                        for coordinate in changed[:6]
                    ],
                }
    return {
        "difference_fields": projection_difference_fields(expected, actual),
        "cell_differences": cell_differences,
        "pycode_dimensions": expected_dimensions,
        "nocom_pytools_dimensions": actual_dimensions,
    }

def main() -> int:
    if not PYTOOLS_SOURCE.is_dir():
        raise FileNotFoundError(PYTOOLS_SOURCE)
    with tempfile.TemporaryDirectory(prefix="pytools-summary-parity-", ignore_cleanup_errors=True) as temp:
        root = Path(temp)
        source = root / "sample_2026-06-30.xlsx"
        second_source = root / "sample_2026-07-31.xlsm"
        template = root / "comment_template.xlsx"
        duplicate_marker_template = root / "comment_template_duplicate_marker.xlsx"
        unsupported_error_source = WORKSPACE / "pytools-core/tests/fixtures/summary_used_range/getting_data_error.xlsx"
        make_source(source)
        patch_formula_cache(source, "E8", "2")
        make_second_source(second_source)
        make_comment_template(template)
        make_comment_template(duplicate_marker_template, duplicate_row_boundaries=True)
        sources = [source, second_source]
        fingerprints_before = source_fingerprints(
            [*sources, template, duplicate_marker_template, unsupported_error_source]
        )
        python_used, python_comments = python_baselines(sources, source, template, root)
        nocom_used = python_nocom_used_range(sources, root)
        rust_used, used_report = run_rust("summary-used-range", sources, root / "rust_used")
        rust_comments, comments_report = run_rust_comments(template, source, root / "rust_comments")
        assert_equivalent(python_used, rust_used, "6-1 used-range summary")
        assert_equivalent(
            python_comments,
            rust_comments,
            "6-2 comment summary",
            compare_page_settings=False,
        )
        python_duplicate_comments = python_comments_only(
            duplicate_marker_template, [source], root / "python_duplicate_comments"
        )
        rust_duplicate_comments, duplicate_comments_report = run_rust_comments(
            duplicate_marker_template, source, root / "rust_duplicate_comments"
        )
        if python_duplicate_comments is None:
            raise AssertionError("PyCode did not save the duplicate-marker comment summary")
        assert_equivalent(
            python_duplicate_comments,
            rust_duplicate_comments,
            "6-2 duplicate row-region marker behavior",
            compare_page_settings=False,
        )
        duplicate_sheets, _ = workbook_projection(rust_duplicate_comments)
        duplicate_source_rows = [
            duplicate_sheets[0]["values"].get(f"F{row}") for row in range(2, 8)
        ]
        if duplicate_comments_report.get("rows") != 6 or duplicate_source_rows != [3, 4, 5, 3, 4, 5]:
            raise AssertionError(
                "duplicate row-region markers should select the final start marker: "
                f"rows={duplicate_comments_report.get('rows')}, values={duplicate_source_rows}"
            )
        python_unknown_error = python_used_range_only(
            [unsupported_error_source], root / "python_unknown_error"
        )
        if python_unknown_error is None:
            raise AssertionError("PyCode did not save the unknown-error summary")
        unknown_error_sheets, _ = workbook_projection(python_unknown_error)
        python_unknown_cell = unknown_error_sheets[0]
        if (
            python_unknown_cell["values"].get("D2") != "#DATA!"
            or python_unknown_cell["types"].get("D2") != "s"
        ):
            raise AssertionError("PyCode unknown-error behavior changed from string output")
        try:
            run_rust(
                "summary-used-range",
                [unsupported_error_source],
                root / "rust_unknown_error",
            )
        except subprocess.CalledProcessError as error:
            rust_unknown_error_stderr = error.stderr.decode(
                "utf-8-sig", errors="replace"
            )
            rust_unknown_error_marker = "Unsupported cell error value '#DATA!'"
            if rust_unknown_error_marker not in rust_unknown_error_stderr:
                raise AssertionError(
                    "Rust CLI failed for the unknown error for an unexpected reason: "
                    + rust_unknown_error_stderr[-1600:]
                )
        else:
            raise AssertionError("Rust CLI unexpectedly accepted the unknown XLSX error")
        fingerprints_after = source_fingerprints(
            [*sources, template, duplicate_marker_template, unsupported_error_source]
        )
        if fingerprints_before != fingerprints_after:
            raise AssertionError("a source workbook or template was modified")

        nocom_matches_rust = False
        nocom_difference_fields = ["not_saved"]
        if nocom_used is not None and rust_used is not None:
            nocom_difference_fields = projection_difference_fields(nocom_used, rust_used)
            nocom_matches_rust = not nocom_difference_fields

        if used_report.get("rows") != 14 or used_report.get("files") != 2:
            raise AssertionError(f"unexpected used-range summary report: {used_report}")
        if used_report.get("execution", {}).get("ooxml_error_cells_patched") != 7:
            raise AssertionError(f"expected seven preserved Excel error values: {used_report}")
        report = {
            "schema_version": 1,
            "suite": "pytools.summary_used_range_python_rust_parity",
            "date": dt.date.today().isoformat(),
            "status": "passed_synthetic_cli_parity",
            "source_python_file": "PyCode/core/summary_321_322.py::run_summary_by_usedrange",
            "real_sample": False,
            "inputs": {
                "files": [path.name for path in sources],
                "sha256_before_after_equal": True,
                "source_workbooks": 2,
            },
            "scenario": {
                "rows": used_report["rows"],
                "features": [
                    "actual Rust summary-used-range CLI",
                    "multiple workbooks and worksheets in selection order",
                    "macro-free XLSM source with macro-enabled workbook content type",
                    "empty worksheet produces the default-bound row seen in PyCode",
                    "format-only used-range tail",
                    "merged range used as an extent",
                    "cached formula values and missing formula-cache semantics",
                    "date, boolean, numeric, text and blank semantics",
                    "seven common Excel error values preserved as error-typed cells",
                ],
            },
            "comparison": {
                "cell_values_and_types": "exact",
                "cell_styles": "exact for populated output cells",
                "worksheet_names_and_dimensions": "exact",
                "freeze_panes": "exact",
                "header_bold_flags": "exact",
                "formulas_merges_defined_names_print_area_page_setup": "exact (default portrait normalized)",
            },
            "rust_execution": used_report.get("execution"),
            "unknown_cell_error_probe": {
                "input": "synthetic XLSX t=e cell with value #DATA!",
                "pycode": {"saved": True, "output_value": "#DATA!", "output_type": "s"},
                "rust_cli": {
                    "status": "failed_closed",
                    "message_contains": "Unsupported cell error value '#DATA!'",
                    "reader_engine": "calamine",
                },
                "note": "Calamine 0.37 does not parse this worksheet error token; PyTools does not use an OOXML value fallback.",
            },
            "comment_summary_probe": {
                "baseline_rows": comments_report.get("rows"),
                "duplicate_row_region_marker": {
                    "scenario": "duplicate 行区域1 starts at A2 and A3; duplicate 行区域#1 ends at A4 and A5",
                    "python": {
                        "saved": True,
                        "rows": 6,
                        "source_row_numbers": [3, 4, 5, 3, 4, 5],
                    },
                    "rust_cli": {
                        "saved": duplicate_comments_report.get("saved"),
                        "rows": duplicate_comments_report.get("rows"),
                        "source_row_numbers": duplicate_source_rows,
                    },
                    "comparison": "exact",
                    "rule": "last matching start/end marker in row-major order, matching PyCode overwrite behavior",
                },
            },
            "python_variant_scope": {
                "nocom_pytools_output_matches_rust_for_this_xlsx_sample": nocom_matches_rust,
                "nocom_pytools_difference_fields": nocom_difference_fields,
                "nocom_pytools_vs_pycode": (
                    projection_difference_summary(python_used, nocom_used)
                    if python_used is not None and nocom_used is not None
                    else {"status": "not_saved"}
                ),
                "format_difference": "This probe covers xlsx and macro-free xlsm. VBA-bearing xlsm and NoComPyTools-only xls/xlsb remain outside the validated Rust/PyCode set.",
            },
            "limitations": [
                "Synthetic XLSX and macro-free XLSM fixtures only; no authorized real summary workbook was present.",
                "XLSM was tested only as a macro-free synthetic package; VBA-bearing XLSM, XLS, XLSB, Kylin/UOS, glibc 2.28 and deb packaging were not validated.",
                "NoComPyTools and PyCode input-format scopes differ; no production scope was inferred from the NoCom-only formats.",
                "Output is a new data workbook; source workbook styles, formulas and unrelated OOXML parts are not copied by contract.",
                "A synthetic XLSX t=e #DATA! cell is written as a string by PyCode; Calamine 0.37 rejects it. Supporting this value requires a coordinated shared-reader capability; no OOXML fallback was added.",
            ],
        }
        report_day = dt.date.today().strftime("%Y%m%d")
        report_path = Path(__file__).resolve().parent / "reports" / f"summary_used_range_parity_{report_day}.json"
        report_path.parent.mkdir(parents=True, exist_ok=True)
        report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        print(json.dumps({
            "status": report["status"],
            "used_range_rows_rust": used_report.get("rows"),
            "used_range_files_rust": used_report.get("files"),
            "error_cells_patched": (used_report.get("execution") or {}).get("ooxml_error_cells_patched"),
            "nocom_pytools_matches_rust": nocom_matches_rust,
            "comment_summary_rows_rust": comments_report.get("rows"),
            "comment_duplicate_marker_rows_rust": duplicate_comments_report.get("rows"),
            "report": str(report_path),
        }, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
