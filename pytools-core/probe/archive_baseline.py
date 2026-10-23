"""Create the Python baseline for archive preview/copy behavior in isolation."""

from __future__ import annotations

import argparse
import hashlib
import json
import logging
import sys
import tempfile
from pathlib import Path

from openpyxl import Workbook, load_workbook


WORKSPACE = Path(__file__).resolve().parents[2]
CASES_PATH = WORKSPACE / "pytools-core/tests/fixtures/archive_cases.json"
CONFIG_PATH = WORKSPACE / "pytools-core/tests/fixtures/archive_type_config.xlsx"
SOURCE_RELATIVE = Path("PyCode/core/convert_33_34_35_37.py")


def create_config(cases: dict[str, object]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "归档类型配置"
    headers = ["是否启用", "匹配关键词", "归档文件夹", "备注", "需要归档的后缀", "排除归档的后缀"]
    for column, value in enumerate(headers, start=1):
        sheet.cell(1, column).value = value
    for row_number, row in enumerate(cases["archive_type_rules"], start=2):
        for column, value in enumerate(row, start=1):
            sheet.cell(row_number, column).value = value
    sheet["E2"] = cases["include_extensions"]
    sheet["F2"] = cases["exclude_extensions"]
    CONFIG_PATH.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(CONFIG_PATH)
    workbook.close()


def create_inputs(root: Path, cases: dict[str, object], existing_target: str | None = None) -> tuple[Path, Path]:
    input_dir = root / "input"
    input_dir.mkdir(parents=True)
    for relative in cases["input_names"]:
        path = input_dir / relative
        if path.name.startswith("~$"):
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_bytes(b"temporary")
        else:
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_bytes(f"fixture:{relative}".encode("utf-8"))
    missing = input_dir / cases["missing_name"]
    if existing_target:
        target = root / "target" / existing_target
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_bytes(b"occupied target")
    return input_dir, missing


def relative(path: object, root: Path) -> str:
    text = str(path or "")
    if not text:
        return ""
    return Path(text).resolve().relative_to(root.resolve()).as_posix()


def report_rows(path: Path, root: Path, sheet_name: str, is_plan: bool) -> list[dict[str, str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[sheet_name]
        out = []
        for values in sheet.iter_rows(min_row=2, values_only=True):
            if all(value is None for value in values):
                continue
            values = ["" if value is None else str(value) for value in values]
            if is_plan:
                out.append(
                    {
                        "source": relative(values[0], root),
                        "file_name": values[1],
                        "mode": values[2],
                        "category": values[3],
                        "target_dir": relative(values[4], root),
                        "target": relative(values[5], root),
                        "status": values[6],
                        "note": values[7],
                    }
                )
            else:
                out.append(
                    {
                        "source": relative(values[0], root),
                        "planned_target": relative(values[1], root),
                        "actual_target": relative(values[2], root),
                        "status": values[3],
                        "note": values[4],
                    }
                )
        return out
    finally:
        workbook.close()


def stats(result: dict[str, object], copy: bool = False) -> dict[str, object]:
    if copy:
        return {key: int(result.get(key, 0)) for key in ("rows", "copied", "skip", "failed", "files", "invalid", "planned", "rules", "skipped_ext")}
    return {key: int(result.get(key, 0)) for key in ("files", "invalid", "planned", "rules", "skipped_ext")}


def setup_python(source_root: Path) -> None:
    sys.dont_write_bytecode = True
    sys.path.insert(0, str(source_root))


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source-root", type=Path, default=Path(r"C:\Users\AZI\Desktop\VibeCode\4.pytools-main"))
    parser.add_argument("--output", type=Path, default=WORKSPACE / "pytools-core/probe/reports/archive_python_baseline.json")
    args = parser.parse_args()
    setup_python(args.source_root)
    from PyCode.core.convert_33_34_35_37 import (
        run_archive_by_date_copy,
        run_archive_by_date_preview,
        run_archive_by_type_copy,
        run_archive_by_type_preview,
        run_archive_plan_copy,
    )

    cases = json.loads(CASES_PATH.read_text(encoding="utf-8"))
    create_config(cases)
    source_hash = hashlib.sha256((args.source_root / SOURCE_RELATIVE).read_bytes()).hexdigest()
    report: dict[str, object] = {"schema_version": 1, "suite": "pytools.archive.python_baseline", "source_sha256": source_hash}

    with tempfile.TemporaryDirectory(prefix="pytools-archive-python-") as raw_root:
        root = Path(raw_root)
        input_dir, missing = create_inputs(root / "date_preview", cases, cases["existing_targets"]["date"])
        result = run_archive_by_date_preview([input_dir, missing], root / "date_preview/target", root / "date_preview/out", root / "date_preview/logs")
        report["date_preview"] = {"stats": stats(result), "rows": report_rows(Path(result["output"]), root / "date_preview", "归档计划", True)}

        input_dir, missing = create_inputs(root / "date_filtered", cases)
        result = run_archive_by_date_preview([input_dir, missing], root / "date_filtered/target", root / "date_filtered/out", root / "date_filtered/logs", CONFIG_PATH)
        report["date_filtered_preview"] = {"stats": stats(result), "rows": report_rows(Path(result["output"]), root / "date_filtered", "归档计划", True)}

        input_dir, missing = create_inputs(root / "type_preview", cases, cases["existing_targets"]["type"])
        result = run_archive_by_type_preview(CONFIG_PATH, [input_dir, missing], root / "type_preview/target", root / "type_preview/out", root / "type_preview/logs")
        report["type_preview"] = {"stats": stats(result), "rows": report_rows(Path(result["output"]), root / "type_preview", "归档计划", True)}

        input_dir, _ = create_inputs(root / "date_copy", cases, cases["existing_targets"]["date"])
        result = run_archive_by_date_copy([input_dir], root / "date_copy/target", root / "date_copy/out", root / "date_copy/logs")
        report["date_copy"] = {"stats": stats(result, copy=True), "rows": report_rows(Path(result["output"]), root / "date_copy", "归档执行结果", False)}

        input_dir, _ = create_inputs(root / "type_copy", cases, cases["existing_targets"]["type"])
        result = run_archive_by_type_copy(CONFIG_PATH, [input_dir], root / "type_copy/target", root / "type_copy/out", root / "type_copy/logs")
        report["type_copy"] = {"stats": stats(result, copy=True), "rows": report_rows(Path(result["output"]), root / "type_copy", "归档执行结果", False)}

        input_dir, missing = create_inputs(root / "plan_copy", cases, cases["existing_targets"]["date"])
        preview = run_archive_by_date_preview([input_dir, missing], root / "plan_copy/target", root / "plan_copy/plan", root / "plan_copy/logs")
        result = run_archive_plan_copy(Path(preview["output"]), root / "plan_copy/out", root / "plan_copy/logs")
        report["plan_copy"] = {"stats": stats(result, copy=True), "rows": report_rows(Path(result["output"]), root / "plan_copy", "归档执行结果", False)}

        configs = logging.Logger.manager.loggerDict
        for logger_name, logger in list(configs.items()):
            if logger_name.startswith("PyCode.archive_") and isinstance(logger, logging.Logger):
                for handler in list(logger.handlers):
                    handler.close()
                    logger.removeHandler(handler)

    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"archive baseline written: {args.output}")


if __name__ == "__main__":
    main()
