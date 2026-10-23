"""Run the read-only PyCode baseline and export result/workbook semantics."""

from __future__ import annotations

import argparse
import hashlib
import json
import logging
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


def typed(value: object) -> dict[str, object]:
    if isinstance(value, bool):
        return {"type": "boolean", "value": value}
    if isinstance(value, int):
        return {"type": "integer", "value": value}
    if isinstance(value, float):
        return {"type": "number", "value": value}
    if isinstance(value, str):
        return {"type": "string", "value": value}
    raise TypeError(f"unexpected output cell type: {type(value).__name__}")


def color_rgb(color: object) -> str | None:
    if color is None or getattr(color, "type", None) != "rgb":
        return None
    value = getattr(color, "rgb", None)
    return str(value)[-6:].upper() if value else None


def inspect_output(path: Path) -> dict[str, object]:
    book = load_workbook(path, data_only=True, read_only=False)
    formula_book = load_workbook(path, data_only=False, read_only=False)
    try:
        sheet = book["核对结果"]
        formula_sheet = formula_book["核对结果"]
        cells: dict[str, object] = {}
        styles: dict[str, object] = {}
        formulas: dict[str, str] = {}
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                cells[cell.coordinate] = typed(cell.value)
                styles[cell.coordinate] = {
                    "bold": bool(cell.font.bold),
                    "font_color": color_rgb(cell.font.color),
                    "fill_color": color_rgb(cell.fill.fgColor)
                    if cell.fill.fill_type == "solid"
                    else None,
                    "horizontal_alignment": cell.alignment.horizontal,
                }
        for row in formula_sheet.iter_rows():
            for cell in row:
                if cell.data_type == "f":
                    formulas[cell.coordinate] = str(cell.value)
        widths = {
            get_column_letter(column): sheet.column_dimensions[get_column_letter(column)].width
            for column in range(1, 7)
        }
        return {
            "sheet_name": sheet.title,
            "max_row": sheet.max_row,
            "max_column": sheet.max_column,
            "cells": cells,
            "formulas": formulas,
            "column_widths": widths,
            "styles": styles,
        }
    finally:
        book.close()
        formula_book.close()


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source-root", required=True, type=Path)
    parser.add_argument(
        "--fixture-dir",
        type=Path,
        default=Path(__file__).parents[1] / "tests" / "fixtures" / "region_sum",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=Path(__file__).parents[1] / "probe" / "reports" / "region_sum_python_baseline.json",
    )
    args = parser.parse_args()
    sys.dont_write_bytecode = True
    source_root = args.source_root.resolve()
    fixture_dir = args.fixture_dir.resolve()
    cases = json.loads((fixture_dir / "cases.json").read_text(encoding="utf-8"))
    input_paths = [fixture_dir / name for name in cases["run_files"]]
    input_hashes_before = {path.name: sha256(path) for path in input_paths}

    sys.path.insert(0, str(source_root))
    from PyCode.core.feature_1_5_region_sum_check import run_region_sum_check  # noqa: PLC0415

    with tempfile.TemporaryDirectory(prefix="pytools-region-sum-") as temporary:
        base = Path(temporary)
        try:
            result = run_region_sum_check(input_paths, base / "out", base / "logs")
            output_path = Path(result["path"])
            output = inspect_output(output_path)
            statistics = {key: result[key] for key in ("files", "sheets", "errors", "failed_files")}
        finally:
            logger = logging.getLogger("PyCode.feature_1_5_region_sum_check")
            for handler in list(logger.handlers):
                try:
                    handler.flush()
                    handler.close()
                finally:
                    logger.removeHandler(handler)

    input_hashes_after = {path.name: sha256(path) for path in input_paths}
    if input_hashes_before != input_hashes_after:
        raise RuntimeError("baseline mutated an input fixture")

    source_files = {
        "feature_1_5_region_sum_check.py": source_root
        / "PyCode"
        / "core"
        / "feature_1_5_region_sum_check.py",
        "logger.py": source_root / "PyCode" / "core" / "logger.py",
    }
    report = {
        "schema_version": 1,
        "suite": "pytools.region_sum_check",
        "source": "PyCode/core/feature_1_5_region_sum_check.py::run_region_sum_check",
        "source_python_version": sys.version.split()[0],
        "source_sha256": {name: sha256(path) for name, path in source_files.items()},
        "fixtures": cases["run_files"],
        "fixture_sha256": input_hashes_before,
        "scope": "Synthetic XLSX inputs; normal non-COM implementation only.",
        "python_baseline": {"status": "保留", "case_count": len(input_paths)},
        "statistics": statistics,
        "output_workbook": output,
        "limitations": [
            "No real business workbook, XLSM/XLT input, target Linux platform, or menu/CLI integration was tested.",
            "The source feature's COM variant is outside this migration slice; no Office engine is used.",
        ],
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"suite": report["suite"], "statistics": statistics, "source_sha256": report["source_sha256"]}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
