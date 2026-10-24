"""Create the synthetic PT-7-6 full-sheet baseline with the read-only PyCode helper."""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill


DEFAULT_PYCODE_ROOT = Path(
    r"C:\Users\AZI\Desktop\VibeCode\4.pytools-main\PyCode"
)
FEATURE_DIR = Path(__file__).resolve().parents[1]
FIXTURE_DIR = FEATURE_DIR / "tests" / "fixtures" / "sheet_extract_full_menu"
REPORT_PATH = FEATURE_DIR / "probe" / "reports" / "sheet_extract_full_menu_python_baseline.json"


def create_source(path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "报表"

    worksheet["A1"] = "汇总标题"
    worksheet["A1"].font = Font(name="等线", bold=True, color="FF17365D")
    worksheet["A1"].fill = PatternFill("solid", fgColor="FFD9EAF7")
    worksheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    worksheet.merge_cells("A1:C1")

    worksheet["A2"] = "项目"
    worksheet["B2"] = "金额"
    worksheet["C2"] = "比例"
    worksheet["A3"] = "贷款"
    worksheet["B3"] = 1250.5
    worksheet["B3"].number_format = "#,##0.00"
    worksheet["C3"] = 0.25
    worksheet["C3"].number_format = "0.0%"
    worksheet["D5"] = "备注位置"
    worksheet["D5"].comment = Comment("synthetic comment", "PyTools baseline")

    # A styled blank extends the cell rectangle beyond the populated values.
    worksheet["E8"].fill = PatternFill("solid", fgColor="FFE2F0D9")
    worksheet.row_dimensions[12].height = 26
    worksheet.row_dimensions[12].hidden = True
    worksheet.row_dimensions[12].outlineLevel = 1
    worksheet.column_dimensions["G"].width = 20
    worksheet.column_dimensions["G"].hidden = True
    worksheet.column_dimensions["G"].outlineLevel = 2
    workbook.save(path)


def _cell_projection(cell: Any) -> dict[str, object] | None:
    value = cell.value
    if value is None:
        return None
    if isinstance(value, bool):
        kind = "boolean"
    elif isinstance(value, int):
        kind = "integer"
    elif isinstance(value, float):
        kind = "number"
    else:
        kind = "string"
        value = str(value)
    return {"type": kind, "value": value}


def create_python_expected(source: Path, expected: Path, pycode_root: Path) -> dict[str, object]:
    sys.dont_write_bytecode = True
    sys.path.insert(0, str(pycode_root))
    from core.feature_1_6_extract_sheets import _copy_full_sheet_with_style

    source_workbook = load_workbook(source, data_only=True, read_only=False)
    source_sheet = source_workbook["报表"]
    output_workbook = Workbook()
    output_sheet = output_workbook.active
    output_sheet.title = "报表"
    copied = _copy_full_sheet_with_style(source_sheet, output_sheet)
    if not copied:
        raise RuntimeError("PyCode full-sheet helper did not copy the synthetic worksheet")
    output_workbook.save(expected)

    check_workbook = load_workbook(expected, data_only=True, read_only=False)
    check_sheet = check_workbook["报表"]
    if check_sheet.max_row != source_sheet.max_row or check_sheet.max_column != source_sheet.max_column:
        raise RuntimeError("PyCode full-sheet result dimensions differ from its source")
    if list(map(str, check_sheet.merged_cells.ranges)) != list(map(str, source_sheet.merged_cells.ranges)):
        raise RuntimeError("PyCode full-sheet result merge ranges differ from its source")
    if check_sheet["D5"].comment is not None:
        raise RuntimeError("ordinary PyCode extraction unexpectedly copied a comment")

    style_read_error = None
    try:
        output_style_checks = {
            "A1_bold": check_sheet["A1"].font.bold,
            "A1_fill": check_sheet["A1"].fill.fgColor.rgb,
            "B3_number_format": check_sheet["B3"].number_format,
            "C3_number_format": check_sheet["C3"].number_format,
            "E8_fill": check_sheet["E8"].fill.fgColor.rgb,
        }
    except (IndexError, KeyError) as error:
        output_style_checks = None
        style_read_error = f"{type(error).__name__}: {error}"

    source_style_checks = {
        "A1_bold": source_sheet["A1"].font.bold,
        "A1_fill": source_sheet["A1"].fill.fgColor.rgb,
        "B3_number_format": source_sheet["B3"].number_format,
        "C3_number_format": source_sheet["C3"].number_format,
        "E8_fill": source_sheet["E8"].fill.fgColor.rgb,
    }

    return {
        "name": check_sheet.title,
        "sheet": check_sheet.title,
        "max_row": check_sheet.max_row,
        "max_column": check_sheet.max_column,
        "non_empty_cells": sum(
            cell.value is not None
            for row in check_sheet.iter_rows()
            for cell in row
        ),
        "merged_ranges": [str(value) for value in check_sheet.merged_cells.ranges],
        "comment_copied": check_sheet["D5"].comment is not None,
        "extended_row_dimension": {
            "row": 12,
            "height": check_sheet.row_dimensions[12].height,
            "hidden": check_sheet.row_dimensions[12].hidden,
            "outline_level": check_sheet.row_dimensions[12].outlineLevel,
        },
        "extended_column_dimension": {
            "column": "G",
            "width": check_sheet.column_dimensions["G"].width,
            "hidden": check_sheet.column_dimensions["G"].hidden,
            "outline_level": check_sheet.column_dimensions["G"].outlineLevel,
        },
        "values": [
            [_cell_projection(check_sheet.cell(row=row, column=column))
             for column in range(1, check_sheet.max_column + 1)]
            for row in range(1, check_sheet.max_row + 1)
        ],
        "style_checks": {
            "python_output": output_style_checks,
            "python_output_read_error": style_read_error,
            "source_workbook_expected": source_style_checks,
        },
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--pycode-root", type=Path, default=DEFAULT_PYCODE_ROOT)
    args = parser.parse_args()

    FIXTURE_DIR.mkdir(parents=True, exist_ok=True)
    REPORT_PATH.parent.mkdir(parents=True, exist_ok=True)
    source = FIXTURE_DIR / "source.xlsx"
    expected = FIXTURE_DIR / "python_expected.xlsx"
    create_source(source)
    worksheet = create_python_expected(source, expected, args.pycode_root)
    report = {
        "schema_version": 1,
        "status": "partial",
        "suite": "pytools.sheet_extract.full_sheet_menu_python_baseline",
        "baseline": "PyCode/core/feature_1_6_extract_sheets.py::_copy_full_sheet_with_style",
        "formula_mode": "data_only",
        "ordinary_comment_variant": "comments are not copied",
        "source_tree_modified": False,
        "fixture": "tests/fixtures/sheet_extract_full_menu/source.xlsx",
        "expected": "tests/fixtures/sheet_extract_full_menu/python_expected.xlsx",
        "limitations": [
            "The legacy helper copies private style arrays across workbooks; openpyxl cannot read the emitted style table (IndexError).",
            "Value, merge, comment, and row/column dimension behavior is captured from the helper output; Rust style checks use source-workbook properties because the Python output style table is unreadable.",
        ],
        "worksheet": worksheet,
    }
    REPORT_PATH.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
