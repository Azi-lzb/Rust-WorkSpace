"""Execute PT-3-3 against the read-only PyTools baseline and a Rust candidate."""

from __future__ import annotations

import contextlib
import hashlib
import io
import json
import logging
import math
import shlex
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.styles.protection import Protection

sys.dont_write_bytecode = True
WORKSPACE = Path(__file__).resolve().parents[2]
PYTOOLS_SOURCE = Path(r"C:\Users\AZI\Desktop\VibeCode\4.pytools-main")
PLATFORM = "synthetic-pt-3-3-print-config-probe"


def wsl_path(path: Path) -> str:
    return f"/mnt/{path.drive.rstrip(':').lower()}/{path.as_posix()[2:].lstrip('/')}"


def make_source(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Source"
    sheet["A1"] = "Report"
    sheet["A1"].font = Font(bold=True)
    sheet.merge_cells("A1:B1")
    sheet["A2"] = "A"
    sheet["B2"] = 12.5
    sheet["B2"].number_format = "#,##0.00"
    sheet["B2"].fill = PatternFill(fill_type="solid", fgColor="FF0099CC")
    sheet["B2"].border = Border(top=Side(style="thin", color="FF112233"))
    sheet["B2"].alignment = Alignment(horizontal="right", wrap_text=True)
    sheet["B2"].protection = Protection(locked=False)
    sheet["A2"].comment = Comment("source note", "probe-author")
    sheet.row_dimensions[2].height = 24
    sheet.column_dimensions["A"].width = 18
    workbook.save(path)
    workbook.close()


def make_rust_config(path: Path, source_wsl: str, target_wsl: str, mode: int) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "打印配置"
    headers = [
        "是否启用",
        "打印模式",
        "源工作簿",
        "源工作表",
        "源工作表打印区域",
        "目标工作簿",
        "目标工作表",
        "FitToPagesWide",
        "FitToPagesTall",
        "打印方向",
        "水平居中",
        "垂直居中",
        "不输出批注",
        "零值不输出",
        "备注",
    ]
    for column, value in enumerate(headers, start=1):
        sheet.cell(1, column, value)
    values = [
        True,
        mode,
        source_wsl,
        "Source",
        "A1:B2",
        target_wsl,
        "Printed",
        1,
        1,
        "横向",
        True,
        True,
        False,
        False,
        "",
    ]
    for column, value in enumerate(values, start=1):
        sheet.cell(2, column, value)
    workbook.save(path)
    workbook.close()


def projection(path: Path) -> dict[str, Any]:
    workbook = load_workbook(path, data_only=False)
    sheets = []
    for sheet in workbook.worksheets:
        page_setup_properties = sheet.sheet_properties.pageSetUpPr
        cells = {}
        for address in ("A1", "B1", "A2", "B2"):
            cell = sheet[address]
            cells[address] = {
                "value": cell.value,
                "type": cell.data_type,
                "bold": bool(cell.font.bold),
                "fill_type": cell.fill.fill_type,
                "fill_rgb": cell.fill.fgColor.rgb,
                "number_format": cell.number_format,
                "border_top_style": cell.border.top.style,
                "border_top_color": None
                if cell.border.top.color is None
                else cell.border.top.color.rgb,
                "alignment_horizontal": cell.alignment.horizontal,
                "alignment_wrap_text": cell.alignment.wrap_text,
                "protection_locked": cell.protection.locked,
                "comment": None
                if cell.comment is None
                else {"text": cell.comment.text, "author": cell.comment.author},
            }
        footer = sheet.oddFooter.center.text or ""
        if footer.startswith("&C"):
            footer = footer[2:]
        sheets.append(
            {
                "name": sheet.title,
                "cells": cells,
                "merges": sorted(str(item) for item in sheet.merged_cells.ranges),
                "row_2_height": sheet.row_dimensions[2].height,
                "column_A_width": sheet.column_dimensions["A"].width,
                "page": {
                    "orientation": sheet.page_setup.orientation,
                    "fit_to_page": None
                    if page_setup_properties is None
                    else page_setup_properties.fitToPage,
                    "fit_to_width": sheet.page_setup.fitToWidth,
                    "fit_to_height": sheet.page_setup.fitToHeight,
                    "scale": sheet.page_setup.scale,
                    "paper_size": sheet.page_setup.paperSize,
                    "print_area": str(sheet.print_area),
                    "center_horizontal": sheet.print_options.horizontalCentered,
                    "center_vertical": sheet.print_options.verticalCentered,
                    "footer_center": footer,
                },
            }
        )
    workbook.close()
    return {"sheet_names": [sheet["name"] for sheet in sheets], "sheets": sheets}


def differences(left: Any, right: Any, path: str = "root") -> list[str]:
    if isinstance(left, bool) or isinstance(right, bool):
        return [] if left == right and type(left) is type(right) else [path]
    if isinstance(left, (int, float)) and isinstance(right, (int, float)):
        return [] if math.isclose(left, right, rel_tol=1e-15, abs_tol=1e-15) else [path]
    if type(left) is not type(right):
        return [path]
    if isinstance(left, dict):
        if left.keys() != right.keys():
            return [path]
        return [
            item
            for key in left
            for item in differences(left[key], right[key], f"{path}.{key}")
        ]
    if isinstance(left, list):
        if len(left) != len(right):
            return [path]
        return [
            item
            for index, (left_item, right_item) in enumerate(zip(left, right))
            for item in differences(left_item, right_item, f"{path}[{index}]")
        ]
    return [] if left == right else [path]


def run_rust_cli(arguments: list[str]) -> dict[str, Any]:
    workspace = wsl_path(WORKSPACE)
    command = " ".join(shlex.quote(argument) for argument in arguments)
    shell = f"cd {shlex.quote(workspace)} && cargo run --quiet --offline --locked -p pytools-core --bin pytools -- {command}"
    result = subprocess.run(
        ["wsl.exe", "-d", "Ubuntu", "--", "bash", "-lc", shell],
        capture_output=True,
        check=False,
    )
    stdout = result.stdout.decode("utf-8-sig", errors="replace").replace("\x00", "")
    stderr = result.stderr.decode("utf-8", errors="replace").replace("\x00", "")
    if result.returncode:
        raise RuntimeError(f"Rust CLI failed ({result.returncode}):\n{stdout}\n{stderr}")
    try:
        return json.loads(stdout)
    except json.JSONDecodeError as error:
        raise RuntimeError(f"Rust CLI emitted invalid JSON:\n{stdout}\n{stderr}") from error


def main() -> int:
    sys.path.insert(0, str(PYTOOLS_SOURCE / "PyCode"))
    from core.print_310 import run_print_config_all_modes

    with tempfile.TemporaryDirectory(prefix="pytools-pt-3-3-") as temporary:
        root = Path(temporary)
        source = root / "source.xlsx"
        source.parent.mkdir(parents=True, exist_ok=True)
        make_source(source)
        source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
        mode_results = {}
        for mode in (1, 2, 3):
            mode_root = root / f"mode-{mode}"
            python_dir = mode_root / "python"
            rust_dir = mode_root / "rust"
            python_dir.mkdir(parents=True)
            rust_dir.mkdir(parents=True)
            python_target = python_dir / "requested.xlsx"
            rust_target = rust_dir / "requested.xlsx"
            python_logs = python_dir / "logs"
            config = rust_dir / "config.xlsx"
            rust_output = rust_dir / "plans"
            capability_file = rust_dir / "test-capability.json"
            plan_file = rust_output / "print-target-0.plan.json"
            make_rust_config(config, wsl_path(source), wsl_path(rust_target), mode)
            config_hash = hashlib.sha256(config.read_bytes()).hexdigest()

            task = {
                "row_no": 2,
                "enabled": True,
                "mode": mode,
                "source_wb": source,
                "source_ws": "Source",
                "source_ranges": "A1:B2",
                "target_wb": python_target,
                "target_ws": "Printed",
                "fit_wide": 1,
                "fit_tall": 1,
                "orientation": "横向",
                "center_h": True,
                "center_v": True,
                "drop_comments": False,
                "hide_zero": False,
            }
            with contextlib.redirect_stdout(io.StringIO()):
                python_stats = run_print_config_all_modes([task], python_logs)
            logging.shutdown()
            if not python_target.is_file() or python_stats["task_ok"] != 1 or python_stats["task_skip"]:
                raise AssertionError(f"Python mode {mode} baseline did not produce output: {python_stats}")

            plan_response = run_rust_cli(
                [
                    "print-config-plan",
                    "--config",
                    wsl_path(config),
                    "--output-dir",
                    wsl_path(rust_output),
                ]
            )
            if len(plan_response) != 1 or plan_response[0]["execution"] != "plan_only":
                raise AssertionError(f"unexpected Rust mode {mode} plan response: {plan_response}")
            plan = json.loads(plan_file.read_text(encoding="utf-8"))
            if plan["output"] != wsl_path(rust_target):
                raise AssertionError(f"Rust mode {mode} plan did not target the configured new file: {plan['output']}")
            capability = {
                "engine": "umya_spreadsheet",
                "status": "passed",
                "format": plan["format"],
                "operations": sorted({operation["operation"] for operation in plan["operations"]}),
                "supported_features": plan["workbook"]["features_present"],
                "preserved": plan["preserve"],
                "platform": PLATFORM,
                "evidence": {
                    "sample_id": f"synthetic-pt-3-3-print-config-mode-{mode}",
                    "report_path": "pytools-core/probe/compare_print_config_execution.py",
                    "adapter_version": "umya-spreadsheet-3.1.1",
                    "engine_version": "umya-spreadsheet-3.1.1",
                    "platform": PLATFORM,
                },
            }
            capability_file.write_text(json.dumps(capability, ensure_ascii=False), encoding="utf-8")
            rust_response = run_rust_cli(
                [
                    "print-config-plan",
                    "--config",
                    wsl_path(config),
                    "--output-dir",
                    wsl_path(rust_output),
                    "--capabilities",
                    wsl_path(capability_file),
                    "--platform",
                    PLATFORM,
                ]
            )
            if not rust_target.is_file() or rust_response[0]["execution"]["engine"] != "umya_spreadsheet":
                raise AssertionError(f"Rust mode {mode} execution did not create configured target: {rust_response}")

            python_result = projection(python_target)
            rust_result = projection(rust_target)
            mismatch_paths = differences(python_result, rust_result)
            if mismatch_paths:
                with zipfile.ZipFile(rust_target) as package:
                    page_xml = {
                        name: package.read(name).decode("utf-8", errors="replace")
                        for name in package.namelist()
                        if name.startswith("xl/worksheets/sheet") and name.endswith(".xml")
                    }
                raise AssertionError(
                    f"Python/Rust print mode {mode} output differs:\n"
                    + "Mismatched paths: " + ", ".join(mismatch_paths) + "\n"
                    + json.dumps({"python": python_result, "rust": rust_result}, ensure_ascii=True, indent=2)
                    + "\nRust worksheet XML:\n"
                    + json.dumps(page_xml, ensure_ascii=True, indent=2)
                )
            if hashlib.sha256(config.read_bytes()).hexdigest() != config_hash:
                raise AssertionError(f"Rust mode {mode} print operation changed the configuration workbook")
            mode_results[str(mode)] = {
                "python_stats": python_stats,
                "rust_engine": rust_response[0]["execution"]["engine"],
                "target_created_at_configured_path": True,
                "projection": python_result,
            }

        if hashlib.sha256(source.read_bytes()).hexdigest() != source_hash:
            raise AssertionError("print operation changed the source workbook")

        report = {
            "status": "passed_synthetic_python_rust_execution_parity",
            "feature": "PT-3-3 configured print preparation, modes 1/2/3",
            "environment": {
                "python": sys.version.split()[0],
                "openpyxl": __import__("openpyxl").__version__,
                "rust_execution": "Ubuntu WSL2 x86_64; cargo CLI invoked from Windows Python probe",
                "target_linux_deb_validated": False,
            },
            "sample": "synthetic XLSX generated by this probe; no user business workbook was available",
            "source_sha256": source_hash,
            "comparison": [
                "worksheet identity/order",
                "cell values and types",
                "font/fill/border/alignment/protection/number format projections",
                "merged ranges",
                "comment text and author",
                "row height and column width",
                "orientation, fit settings, scale, paper size, print area, centering and footer",
            ],
            "float_comparison": "recursive relative and absolute tolerance 1e-15, per rules/13 §9.1",
            "modes": mode_results,
            "limitations": [
                "Modes 1/2/3 each use one synthetic task and one target workbook; multiple tasks per target are not covered.",
                "Capability is temporary test evidence and is not registered for production routing.",
                "No real PyTools print job, Windows native Rust build, Kylin/UOS, glibc 2.28 or offline deb was validated.",
                "The user's existing-target rejection contract is covered separately by the Rust CLI integration test.",
            ],
        }
        report_path = WORKSPACE / "pytools-core" / "probe" / "reports" / "print_config_execution_parity_20261023.json"
        report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        print(
            json.dumps(
                {
                    "status": report["status"],
                    "feature": report["feature"],
                    "modes": sorted(mode_results),
                    "report": str(report_path),
                },
                ensure_ascii=False,
            )
        )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
