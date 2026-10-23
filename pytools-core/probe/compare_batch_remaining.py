"""Compare remaining batch-operation plans with the read-only PyCode baseline.

The Rust production CLI emits ExcelEditPlan only. This probe applies the small
tested operation subset to a temporary copy with openpyxl solely to compare
plan semantics; it does not register or execute a production ExcelModifier
capability.
"""

from __future__ import annotations

import json
import hashlib
import math
import shlex
import shutil
import subprocess
import sys
import tempfile
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.utils.cell import coordinate_to_tuple

sys.dont_write_bytecode = True
WORKSPACE = Path(__file__).resolve().parents[2]
SOURCE = Path(r"C:\Users\AZI\Desktop\VibeCode\4.pytools-main")
sys.path.insert(0, str(SOURCE / "PyCode"))

from core.feature_1_1_split_village_bank import run_split_village_bank
from core.feature_1_3_remove_foreign_bank import run_remove_foreign_bank
from core.feature_1_8_adjust_rural_loan import run_adjust_rural_loan
from core.feature_1_9_adjust_rural_loan_quarterly import run_adjust_rural_loan_quarterly


def wsl_path(path: Path) -> str:
    return f"/mnt/{path.drive.rstrip(':').lower()}/{path.as_posix()[2:].lstrip('/')}"


def run_rust(command: str, args: list[tuple[str, Path]]) -> dict:
    workspace = wsl_path(WORKSPACE)
    command_line = "cd " + shlex.quote(workspace) + " && cargo run --quiet -p pytools-core --bin pytools -- "
    command_line += shlex.quote(command) + " " + " ".join(
        f"{shlex.quote(option)} {shlex.quote(wsl_path(path))}" for option, path in args
    )
    result = subprocess.run(["wsl", "-e", "sh", "-lc", command_line], capture_output=True, check=False)
    stdout = result.stdout.decode("utf-8-sig", errors="replace").replace("\x00", "")
    if result.returncode != 0:
        raise RuntimeError(
            f"Rust {command} failed (exit={result.returncode}): {stdout}\n"
            f"{result.stderr.decode('utf-8', errors='replace')}"
        )
    start = stdout.find("{")
    if start < 0:
        raise RuntimeError(f"Rust {command} emitted no JSON: {stdout}")
    return json.loads(stdout[start:])


def new_workbook(path: Path, sheets: dict[str, list[list[object]]]) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets.items():
        sheet = workbook.create_sheet(name)
        for row in rows:
            sheet.append(row)
    workbook.save(path)
    workbook.close()


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def apply_plan(plan: dict, source: Path, destination: Path) -> None:
    workbook = load_workbook(source, data_only=False)
    for operation in plan["operations"]:
        kind = operation["operation"]
        if kind in ("delete_rows", "insert_rows"):
            sheet = workbook[operation["sheet"]]
            if kind == "delete_rows":
                sheet.delete_rows(operation["first_row"], operation["count"])
            else:
                sheet.insert_rows(operation["first_row"], operation["count"])
        elif kind == "set_cell_value":
            sheet = workbook[operation["sheet"]]
            row, column = coordinate_to_tuple(operation["cell"])
            value = operation["value"]
            value_type = value["type"]
            raw = value.get("value")
            if value_type == "decimal":
                raw = float(Decimal(raw))
            elif value_type == "blank":
                raw = None
            sheet.cell(row, column).value = raw
        else:
            raise AssertionError(f"probe does not apply untested operation kind {kind}")
    workbook.save(destination)
    workbook.close()


def workbook_projection(path: Path) -> dict:
    workbook = load_workbook(path, data_only=False)
    output = {}
    for sheet in workbook.worksheets:
        rows = []
        for row in sheet.iter_rows():
            rows.append([cell.value for cell in row])
        while rows and all(value is None for value in rows[-1]):
            rows.pop()
        rows = [row[: max((i + 1 for i, value in enumerate(row) if value is not None), default=0)] for row in rows]
        output[sheet.title] = rows
    workbook.close()
    return output


def assert_projection_equal(label: str, expected: Path, actual: Path) -> None:
    left = workbook_projection(expected)
    right = workbook_projection(actual)
    if not values_equal(left, right):
        raise AssertionError(f"{label}: workbook value projection differs\nPython={left}\nRust-plan={right}")


def values_equal(left: object, right: object) -> bool:
    if isinstance(left, bool) or isinstance(right, bool):
        return left is right
    if isinstance(left, (int, float)) and isinstance(right, (int, float)):
        return math.isclose(float(left), float(right), rel_tol=1e-15, abs_tol=1e-15)
    if isinstance(left, list) and isinstance(right, list):
        return len(left) == len(right) and all(values_equal(a, b) for a, b in zip(left, right))
    if isinstance(left, dict) and isinstance(right, dict):
        return left.keys() == right.keys() and all(values_equal(left[key], right[key]) for key in left)
    return left == right


def village_case(root: Path) -> dict:
    current = root / "village-current.xlsx"
    previous = root / "village-previous.xlsx"
    target = root / "village-target.xlsx"
    python_target = root / "village-python.xlsx"
    rust_target = root / "village-rust-plan.xlsx"
    names = ["村镇银行本外币", "村镇银行人民币", "村镇银行外汇"]
    cur_sheets = {}
    prev_sheets = {}
    for index, name in enumerate(names, 1):
        cur_rows = [[None] * 89 for _ in range(6)]
        prev_rows = [[None] * 89 for _ in range(6)]
        for row in (5, 6):
            cur_rows[row - 1][2:5] = [index * row * 10, f"存款-{index}-{row}", row]
            prev_rows[row - 1][2] = index * row * 5
            cur_rows[row - 1][86:89] = [index * row * 20, f"贷款-{index}-{row}", row + 1]
            prev_rows[row - 1][86] = index * row * 10
        cur_sheets[name] = cur_rows
        prev_sheets[name] = prev_rows
    new_workbook(current, cur_sheets)
    new_workbook(previous, prev_sheets)
    new_workbook(
        target,
        {
            "广东分机构本外币": [["机构"], ["其他"], ["村镇银行"], ["尾行"]],
            "广东分机构人民币": [["机构"], ["村镇银行"]],
            "广东分机构外汇": [["机构"], ["村镇银行"]],
            "广东分机构未知": [["机构"], ["村镇银行"]],
            "无村镇行分机构本外币": [["机构"], ["普通银行"]],
            "非分机构工作表": [["村镇银行"]],
        },
    )
    import logging

    logger = logging.getLogger("batch-probe-village")
    logger.addHandler(logging.NullHandler())
    shutil.copy2(target, python_target)
    python_stats = run_split_village_bank(current, previous, [python_target], root / "logs-village")
    response = run_rust("village-split-plan", [
        ("--current", current), ("--previous", previous), ("--input", target), ("--output", rust_target)
    ])
    plan = response["plan"]
    apply_plan(plan, target, rust_target)
    assert_projection_equal("7-1 village split", python_target, rust_target)
    report = response
    if python_stats["filled_sheets"] != report["matched_sheets"]:
        raise AssertionError(f"7-1 matched sheet count differs: {python_stats} vs {report}")
    return {
        "python": python_stats,
        "rust_plan": {"matched_sheets": report["matched_sheets"], "changed_cells": report["changed_cells"], "warnings": report["warnings"]},
        "output_projection": "equal",
    }


def foreign_case(root: Path) -> dict:
    source = root / "foreign-source.xlsx"
    python_output = root / "foreign-source(金融局).xlsx"
    rust_output = root / "foreign-rust-plan.xlsx"
    config = root / "foreign-config.xlsx"
    new_workbook(
        source,
        {
            "甲分机构本外币": [["机构"], ["本地行"], ["外资甲"], ["外资乙"], ["外资丙"], ["尾行"]],
            "乙分机构": [["机构"], ["外资甲"], ["本地行"]],
            "非目标工作表": [["外资甲"]],
        },
    )
    source_sha256 = sha256(source)
    new_workbook(config, {"机构映射表": [["序号", "机构名称", "是否为外资行"], [1, "外资甲", 1], [2, "外资乙", 1]]})
    import logging

    logger = logging.getLogger("batch-probe-foreign")
    logger.addHandler(logging.NullHandler())
    python_stats = run_remove_foreign_bank([source], {"外资甲", "外资乙"}, root / "logs-foreign")
    response = run_rust("remove-foreign-bank-plan", [
        ("--config", config), ("--input", source), ("--output", rust_output)
    ])
    apply_plan(response["plan"], source, rust_output)
    assert_projection_equal("7-3 foreign bank removal", python_output, rust_output)
    if sha256(source) != source_sha256:
        raise AssertionError("7-3 modified the original source workbook")
    deleted_by_plan = sum(
        operation["count"]
        for operation in response["plan"]["operations"]
        if operation["operation"] == "delete_rows"
    )
    if python_stats["deleted_rows"] != deleted_by_plan:
        raise AssertionError(f"7-3 deleted row count differs: {python_stats} vs {response}")
    return {
        "python": python_stats,
        "rust_plan": {"matched_sheets": response["matched_sheets"], "deleted_rows": deleted_by_plan},
        "output_projection": "equal",
        "source_unchanged": True,
        "source_sha256": source_sha256,
    }


def make_institution_sheet(rows: list[list[object]]) -> list[list[object]]:
    return [["序号", "机构名称", "涉农贷款余额", "旧值"]] * 2 + rows


def loan_case(root: Path) -> dict:
    current = root / "loan-current.xlsx"
    previous = root / "loan-previous.xlsx"
    python_current = root / "loan-python.xlsx"
    rust_current = root / "loan-rust-plan.xlsx"
    new_workbook(
        current,
        {
            "惠州市涉农贷款分机构": make_institution_sheet([[1, "甲行", 14.123456789, 999], [2, "本期新增", 4.0, 999]]),
            "本外币涉农贷款分地区": [["地区", "涉农贷款比上月"], ["惠州市", 5.123456789]],
        },
    )
    new_workbook(
        previous,
        {"惠州市涉农贷款分机构": make_institution_sheet([[1, "甲行", 9.0, 888], [2, "上期独有", 3.0, 888]])},
    )
    shutil.copy2(current, python_current)
    import logging

    logger = logging.getLogger("batch-probe-loan")
    logger.addHandler(logging.NullHandler())
    python_summary = run_adjust_rural_loan(python_current, previous, root / "logs-loan")
    response = run_rust("rural-loan-plan", [("--current", current), ("--previous", previous), ("--output", rust_current)])
    apply_plan(response["edit"]["plan"], current, rust_current)
    assert_projection_equal("7-7 rural loan", python_current, rust_current)
    rust_summary = response["summary"]
    for key in ("matched", "consistent", "unmatched_current", "previous_only"):
        if python_summary[key] != rust_summary[key]:
            raise AssertionError(f"7-7 summary {key} differs: {python_summary} vs {rust_summary}")
    for key in ("fill_total", "area_total"):
        if not math.isclose(python_summary[key], rust_summary[key], rel_tol=1e-15, abs_tol=1e-15):
            raise AssertionError(f"7-7 summary {key} differs: {python_summary} vs {rust_summary}")
    return {"python": python_summary, "rust_plan": rust_summary, "output_projection": "equal"}


def quarterly_case(root: Path) -> dict:
    current = root / "quarter-current.xlsx"
    previous = root / "quarter-previous.xlsx"
    python_current = root / "quarter-python.xlsx"
    rust_current = root / "quarter-rust-plan.xlsx"
    new_workbook(
        current,
        {
            "惠州市涉农贷款分机构": make_institution_sheet([[1, "甲行", 14.123456789, 999], [2, "本期新增", 4.0, 999]]),
            "江门市涉农贷款分机构": make_institution_sheet([[1, "江门行", 8.0, 777]]),
            "涉农贷款分地区统计表": [["地区", "涉农贷款比上月"], ["惠州市", 5.123456789], ["江门市", 8.0]],
        },
    )
    new_workbook(
        previous,
        {
            "惠州市涉农贷款分机构": make_institution_sheet([[1, "甲行", 9.0, 888], [2, "上期独有", 3.0, 888]]),
            "江门市涉农贷款分机构": make_institution_sheet([[1, "江门行", 4.0, 666]]),
        },
    )
    shutil.copy2(current, python_current)
    import logging

    logger = logging.getLogger("batch-probe-quarterly")
    logger.addHandler(logging.NullHandler())
    python_summary = run_adjust_rural_loan_quarterly(python_current, previous, root / "logs-quarterly")
    response = run_rust("quarterly-rural-loan-plan", [("--current", current), ("--previous", previous), ("--output", rust_current)])
    apply_plan(response["edit"]["plan"], current, rust_current)
    assert_projection_equal("7-8 quarterly rural loan", python_current, rust_current)
    rust_summary = response["summary"]
    py_regions = python_summary["regions"]
    rust_regions = rust_summary["regions"]
    if len(py_regions) != len(rust_regions):
        raise AssertionError(f"7-8 region counts differ: {python_summary} vs {rust_summary}")
    for py_region, rust_region in zip(py_regions, rust_regions):
        for key in ("region", "sheet", "matched", "consistent", "unmatched_current", "previous_only", "error"):
            if py_region[key] != rust_region[key]:
                raise AssertionError(f"7-8 region {key} differs: {python_summary} vs {rust_summary}")
        for key in ("fill_total", "area_total"):
            if py_region[key] is None or rust_region[key] is None:
                if py_region[key] is not rust_region[key]:
                    raise AssertionError(f"7-8 region {key} differs: {python_summary} vs {rust_summary}")
            elif not math.isclose(py_region[key], rust_region[key], rel_tol=1e-15, abs_tol=1e-15):
                raise AssertionError(f"7-8 region {key} differs: {python_summary} vs {rust_summary}")
    for key in ("matched", "consistent", "error"):
        if python_summary[key] != rust_summary[key]:
            raise AssertionError(f"7-8 summary {key} differs: {python_summary} vs {rust_summary}")
    for key in ("fill_total", "area_total"):
        if not math.isclose(python_summary[key], rust_summary[key], rel_tol=1e-15, abs_tol=1e-15):
            raise AssertionError(f"7-8 summary {key} differs: {python_summary} vs {rust_summary}")
    return {"python": python_summary, "rust_plan": rust_summary, "output_projection": "equal"}


def quarterly_error_cases(root: Path) -> dict:
    current = root / "quarter-errors-current.xlsx"
    previous = root / "quarter-errors-previous.xlsx"
    python_current = root / "quarter-errors-python.xlsx"
    rust_current = root / "quarter-errors-rust-plan.xlsx"
    new_workbook(
        current,
        {
            "惠州市涉农贷款分机构": make_institution_sheet([[1, "甲行", 14.123456789, 999]]),
            "江门市涉农贷款分机构": make_institution_sheet([[1, "江门行", 8.0, 777]]),
            "珠海市涉农贷款分机构": make_institution_sheet([[1, "重复行", 4.0, 777], [2, "重复行", 5.0, 888]]),
            "深圳市涉农贷款分机构": make_institution_sheet([[1, "深行", 5.0, 999]]),
            "涉农贷款分地区统计表": [
                ["地区", "涉农贷款比上月"],
                ["惠州市", 5.123456789],
                ["江门市", 8.0],
                ["珠海市", 0.0],
            ],
        },
    )
    new_workbook(
        previous,
        {
            "惠州市涉农贷款分机构": make_institution_sheet([[1, "甲行", 9.0, 888]]),
            "珠海市涉农贷款分机构": make_institution_sheet([[1, "重复行", 3.0, 999]]),
            "深圳市涉农贷款分机构": make_institution_sheet([[1, "深行", 3.0, 999]]),
        },
    )
    shutil.copy2(current, python_current)
    import logging

    logger = logging.getLogger("batch-probe-quarterly-errors")
    logger.addHandler(logging.NullHandler())
    python_summary = run_adjust_rural_loan_quarterly(python_current, previous, root / "logs-quarter-errors")
    response = run_rust("quarterly-rural-loan-plan", [
        ("--current", current), ("--previous", previous), ("--output", rust_current)
    ])
    apply_plan(response["edit"]["plan"], current, rust_current)
    assert_projection_equal("7-8 quarterly rural loan edge cases", python_current, rust_current)
    rust_summary = response["summary"]
    if len(python_summary["regions"]) != len(rust_summary["regions"]):
        raise AssertionError(f"7-8 edge-case region count differs: {python_summary} vs {rust_summary}")
    for py_region, rust_region in zip(python_summary["regions"], rust_summary["regions"]):
        for key in ("region", "sheet", "matched", "consistent", "unmatched_current", "previous_only", "error"):
            if py_region[key] != rust_region[key]:
                raise AssertionError(f"7-8 edge-case region {key} differs: {python_summary} vs {rust_summary}")
        for key in ("fill_total", "area_total"):
            if py_region[key] is None or rust_region[key] is None:
                if py_region[key] is not rust_region[key]:
                    raise AssertionError(f"7-8 edge-case region {key} differs: {python_summary} vs {rust_summary}")
            elif not math.isclose(py_region[key], rust_region[key], rel_tol=1e-15, abs_tol=1e-15):
                raise AssertionError(f"7-8 edge-case region {key} differs: {python_summary} vs {rust_summary}")
    for key in ("matched", "consistent", "error"):
        if python_summary[key] != rust_summary[key]:
            raise AssertionError(f"7-8 edge-case summary {key} differs: {python_summary} vs {rust_summary}")
    for key in ("fill_total", "area_total"):
        if not math.isclose(python_summary[key], rust_summary[key], rel_tol=1e-15, abs_tol=1e-15):
            raise AssertionError(f"7-8 edge-case summary {key} differs: {python_summary} vs {rust_summary}")
    errors = [region["error"] for region in rust_summary["regions"]]
    if errors != [
        None,
        "上期文件缺少江门市分机构表",
        "分机构表存在重复机构名称：重复行",
        "“涉农贷款分地区统计表”未找到地区行：深圳市",
    ]:
        raise AssertionError(f"7-8 expected per-region error contract differs: {errors}")
    if rust_summary["consistent"]:
        raise AssertionError("7-8 incomplete/error regions must make the overall result inconsistent")
    return {
        "python": python_summary,
        "rust_plan": rust_summary,
        "output_projection": "equal",
        "error_contract": "equal",
    }


def main() -> int:
    with tempfile.TemporaryDirectory(prefix="pytools-batch-remaining-", ignore_cleanup_errors=True) as temp:
        root = Path(temp)
        result = {
            "schema_version": 1,
            "status": "passed_synthetic_plan_semantics",
            "cases": {
                "7-1 village split": village_case(root),
                "7-3 remove foreign banks": foreign_case(root),
                "7-7 rural loan adjustment": loan_case(root),
                "7-8 quarterly rural loan adjustment": quarterly_case(root),
                "7-8 missing/duplicate region cases": quarterly_error_cases(root),
            },
            "scope_limits": [
                "All inputs and outputs were synthetic temporary workbooks.",
                "Python ran only as a read-only source baseline; no original source workbook was written.",
                "Rust plans were applied to temporary copies by the probe's openpyxl helper, not through ExcelModifier.",
                "No production capability or default route was changed.",
                "Real templates, GUI integration, Windows native Rust, target Linux/deb, and file-preservation metadata are not covered.",
            ],
        }
        output = WORKSPACE / "pytools-core" / "probe" / "reports" / "batch_remaining_plan_parity_20261023_r2.json"
        output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        print(json.dumps(result, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
