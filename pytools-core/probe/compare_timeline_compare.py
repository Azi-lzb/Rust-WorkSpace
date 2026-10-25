"""Synthetic Python/Rust parity probe for pytools timeline operation 1-1."""

from __future__ import annotations

import json
import logging
import os
import subprocess
import sys
import tempfile
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.utils.datetime import CALENDAR_MAC_1904

sys.dont_write_bytecode = True
WORKSPACE = Path(__file__).resolve().parents[2]
PYTOOLS_SOURCE = Path(os.environ.get("PYTOOLS_SOURCE_ROOT", WORKSPACE.parent / "4.pytools-main"))


def make_book(path: Path, changed: bool = False, sheet_title: str = "Report Data") -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title
    ws["A2"] = "机构"
    ws["B2"] = "存款"
    ws["C2"] = "贷款"
    ws["B3"] = "定期" if changed else "活期"
    ws["C3"] = "公司"
    ws["A5"] = "甲支行"
    ws["B5"] = 11
    ws["C5"] = 21
    ws["A6"] = "丙支行" if changed else "乙支行"
    ws["B6"] = 12
    ws["C6"] = 22
    wb.save(path)


def make_date_header_book(path: Path, value: datetime, date_1904: bool = False) -> None:
    make_book(path)
    wb = load_workbook(path)
    if date_1904:
        wb.epoch = CALENDAR_MAC_1904
    ws = wb["Report Data"]
    ws["B3"] = value
    ws["B3"].number_format = "yyyy-mm-dd"
    wb.save(path)
    wb.close()


def make_config(
    path: Path,
    output: Path,
    inputs: Path,
    workbook_keyword: str = "SAMPLE",
    sheet_keyword: str = "report,data",
) -> None:
    wb = Workbook()
    global_sheet = wb.active
    global_sheet.title = "全局配置"
    global_sheet.append(["配置项", "值"])
    for key, value in (
        ("输入目录", str(inputs)),
        ("输出目录", str(output)),
        ("日志目录", str(output / "logs")),
        ("错误策略", "continue"),
        ("源文件扩展名", ".xlsx"),
        ("时序规则命中策略", "all_match"),
    ):
        global_sheet.append([key, value])
    sheet = wb.create_sheet("时序提取规则")
    sheet.append([
        "是否启用", "规则名称", "工作簿关键字", "工作表关键字", "行头列", "列表头行",
        "数据起始行", "数据结束行", "数据起始列", "数据结束列",
    ])
    sheet.append(["是", "CompareRule", workbook_keyword, sheet_keyword, "A", "2,3", 5, "", "B", "C"])
    wb.save(path)


def python_baseline(template: Path, sources: list[Path], config: Path, output: Path) -> Path:
    sys.path.insert(0, str(PYTOOLS_SOURCE / "PyCode"))
    from core.compare_393 import run_compare
    from core.config_xlsx import GlobalConfig, load_timeline_rules
    from core.io_excel import clear_sheet_cache

    output.mkdir(parents=True, exist_ok=True)
    global_config = GlobalConfig(
        input_dir=template.parent,
        output_dir=output,
        log_dir=output / "logs",
        error_policy="continue",
        source_exts=[".xlsx"],
        timeline_rule_match_mode="all_match",
    )
    try:
        return run_compare(template, sources, load_timeline_rules(config), global_config)
    finally:
        clear_sheet_cache()
        logging.shutdown()


def run_rust(template: Path, sources: list[Path], config: Path) -> Path:
    executable = "pytools.exe" if sys.platform == "win32" else "pytools"
    binary = WORKSPACE / "target" / "debug" / executable
    if not binary.is_file():
        raise FileNotFoundError(f"build the native debug CLI first: {binary}")
    command = [
        str(binary),
        "timeline-compare",
        "--config",
        str(config),
        "--template",
        str(template),
        *(str(path) for path in sources),
    ]
    completed = subprocess.run(
        command,
        check=True,
        capture_output=True,
        text=True,
        encoding="utf-8",
    )
    stdout = completed.stdout
    stderr = completed.stderr
    for line in reversed(stdout.splitlines()):
        try:
            payload = json.loads(line)
            return Path(payload["output"])
        except (json.JSONDecodeError, KeyError):
            continue
    raise RuntimeError(f"Rust CLI did not return an output path: {stdout}\n{stderr}")


def _color_projection(color):
    if color is None:
        return None
    return [color.type, color.rgb, color.indexed, color.theme, color.tint, color.auto]


def _style_projection(cell) -> dict[str, object]:
    font = cell.font
    fill = cell.fill
    border = cell.border
    alignment = cell.alignment
    protection = cell.protection
    return {
        "font": [font.name, font.sz, font.bold, font.italic, font.underline, _color_projection(font.color)],
        "fill": [fill.fill_type, _color_projection(fill.fgColor), _color_projection(fill.bgColor)],
        "alignment": [
            alignment.horizontal,
            alignment.vertical,
            alignment.text_rotation,
            alignment.wrap_text,
            alignment.shrink_to_fit,
            alignment.indent,
        ],
        "border": [
            [getattr(border, side).style, _color_projection(getattr(border, side).color)]
            for side in ("left", "right", "top", "bottom")
        ],
        "protection": [protection.locked, protection.hidden],
        "number_format": cell.number_format,
    }



def _projection_without_styles(projection: dict[str, object]) -> dict[str, object]:
    sheets = []
    for sheet in projection["sheets"]:
        sheets.append(
            {
                **sheet,
                "rows": [
                    [None if cell is None else cell[:2] for cell in row]
                    for row in sheet["rows"]
                ],
            }
        )
    return {**projection, "sheets": sheets}


def _effective_excel_style(style: dict[str, object]) -> dict[str, object]:
    """Normalize equivalent implicit and explicit XLSX defaults for display parity."""
    normalized = dict(style)
    font = list(style["font"])
    if font[0] is None:
        font[0] = "Calibri"
    if font[1] is None:
        font[1] = 11.0
    font_color = font[5]
    if font_color is None or (font_color[0] == "theme" and font_color[3] == 1):
        font[5] = "default_text_color"
    normalized["font"] = font
    normalized["border"] = [
        [
            side[0],
            "automatic_border_color"
            if side[1] is None or (side[1][0] == "auto" and side[1][5] is True)
            else side[1],
        ]
        for side in style["border"]
    ]
    return normalized


def _style_mismatch_cells(expected: dict[str, object], actual: dict[str, object]) -> list[str]:
    mismatches = []
    for expected_sheet, actual_sheet in zip(expected["sheets"], actual["sheets"]):
        for row_index, (expected_row, actual_row) in enumerate(zip(expected_sheet["rows"], actual_sheet["rows"])):
            for column_index, (expected_cell, actual_cell) in enumerate(zip(expected_row, actual_row)):
                if (
                    expected_cell is not None
                    and actual_cell is not None
                    and _effective_excel_style(expected_cell[2])
                    != _effective_excel_style(actual_cell[2])
                ):
                    column = chr(ord("A") + column_index)
                    mismatches.append(f"{expected_sheet['title']}!{column}{row_index + 1}")
    return mismatches


def workbook_projection(path: Path) -> dict[str, object]:
    wb = load_workbook(path, data_only=False)
    sheets = []
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows():
            cells = [
                None
                if cell.value is None and not cell.has_style
                else [cell.value, cell.data_type, _style_projection(cell)]
                for cell in row
            ]
            while cells and cells[-1] is None:
                cells.pop()
            rows.append(cells)
        while rows and not rows[-1]:
            rows.pop()
        setup = ws.sheet_properties.pageSetUpPr
        sheets.append(
            {
                "title": ws.title,
                "dimension": ws.calculate_dimension(),
                "rows": rows,
                "merged_ranges": sorted(str(item) for item in ws.merged_cells.ranges),
                "freeze_panes": str(ws.freeze_panes) if ws.freeze_panes is not None else None,
                "page_setup": [
                    ws.page_setup.orientation,
                    ws.page_setup.paperSize,
                    setup.fitToPage if setup is not None else None,
                    ws.page_setup.fitToWidth,
                    ws.page_setup.fitToHeight,
                    ws.page_setup.scale,
                ],
                "print_area": str(ws.print_area) if ws.print_area else None,
            }
        )
    projection = {
        "sheet_names": [sheet["title"] for sheet in sheets],
        "sheets": sheets,
        "defined_names": sorted(wb.defined_names),
    }
    wb.close()
    return projection

def real_config_sample_availability() -> dict[str, object]:
    sys.path.insert(0, str(PYTOOLS_SOURCE / "PyCode"))
    from core.config_xlsx import load_global, load_timeline_rules
    from core.io_excel import list_sheet_names
    from core.rule_engine import match_all_keywords

    config = PYTOOLS_SOURCE / "PyCode" / "config" / "config.xlsx"
    if not config.is_file():
        return {"status": "source_config_missing"}
    global_config = load_global(config)
    rules = load_timeline_rules(config)

    def matching_count(files: list[Path]) -> int:
        count = 0
        for path in files:
            try:
                sheets = list_sheet_names(path)
            except Exception:
                continue
            if any(
                match_all_keywords(path.name, rule.wb_keyword)
                and any(match_all_keywords(sheet, rule.sheet_keyword) for sheet in sheets)
                for rule in rules
            ):
                count += 1
        return count

    extensions = set(global_config.source_exts)
    configured_files = (
        sorted(
            path
            for path in global_config.input_dir.rglob("*")
            if path.is_file() and path.suffix.lower() in extensions and not path.name.startswith("~$")
        )
        if global_config.input_dir.is_dir()
        else []
    )
    available_assets = sorted((PYTOOLS_SOURCE / "xlsx").rglob("*.xlsx"))
    return {
        "status": "real_config_loaded",
        "configured_input_dir_exists": global_config.input_dir.is_dir(),
        "configured_input_workbooks": len(configured_files),
        "configured_input_rule_matches": matching_count(configured_files),
        "available_xlsx_assets": len(available_assets),
        "available_xlsx_rule_matches": matching_count(available_assets),
    }


def main() -> int:
    if not PYTOOLS_SOURCE.is_dir():
        raise FileNotFoundError(PYTOOLS_SOURCE)
    with tempfile.TemporaryDirectory(prefix="pytools-timeline-compare-", ignore_cleanup_errors=True) as temp:
        root = Path(temp)
        inputs = root / "inputs"
        inputs.mkdir()
        template = root / "template.xlsx"
        unchanged = inputs / "sample_unchanged.xlsx"
        changed = inputs / "sample_changed.xlsx"
        mixed_case = inputs / "sample_case_changed.xlsx"
        unmatched = inputs / "unmatched.xlsx"
        make_book(template)
        make_book(unchanged)
        make_book(changed, changed=True)
        make_book(unmatched, changed=True)
        make_book(mixed_case, changed=True, sheet_title="REPORT DATA")
        config = root / "config.xlsx"
        python_output = root / "python_output"
        rust_output = root / "rust_output"
        make_config(config, rust_output, inputs)
        sources = [unchanged, changed, unmatched, mixed_case, template]
        expected = python_baseline(template, sources, config, python_output)
        actual = run_rust(template, sources, config)
        expected_projection = workbook_projection(expected)
        actual_projection = workbook_projection(actual)
        expected_values_types = _projection_without_styles(expected_projection)
        actual_values_types = _projection_without_styles(actual_projection)
        if expected_values_types != actual_values_types:
            raise AssertionError(
                json.dumps(
                    {"expected": expected_values_types, "actual": actual_values_types},
                    ensure_ascii=False,
                    default=str,
                    indent=2,
                )
            )
        style_mismatches = _style_mismatch_cells(expected_projection, actual_projection)
        if style_mismatches:
            raise AssertionError(f"style differences: {style_mismatches}")
        header_cells = [
            cell
            for sheet in actual_projection["sheets"]
            for cell in (sheet["rows"][0] if sheet["rows"] else [])
            if cell is not None
        ]
        if len(header_cells) != 17:
            raise AssertionError(f"expected 17 header cells, found {len(header_cells)}")
        for cell in header_cells:
            style = cell[2]
            if (
                style["font"][2] is not True
                or style["alignment"][:2] != ["center", "top"]
                or any(side[0] != "thin" for side in style["border"])
            ):
                raise AssertionError(f"unexpected header style at {cell[0]!r}: {style!r}")

        date_template = root / "date_template.xlsx"
        date_source = inputs / "DATE_source.xlsx"
        make_date_header_book(date_template, datetime(2024, 1, 31))
        make_date_header_book(date_source, datetime(2024, 2, 29))
        date_config = root / "date_config.xlsx"
        make_config(date_config, rust_output, inputs, workbook_keyword="DATE")
        expected_dates = python_baseline(
            date_template, [date_source], date_config, root / "python_date_output"
        )
        actual_dates = run_rust(date_template, [date_source], date_config)
        expected_date_projection = _projection_without_styles(workbook_projection(expected_dates))
        actual_date_projection = _projection_without_styles(workbook_projection(actual_dates))
        if expected_date_projection != actual_date_projection:
            raise AssertionError(
                json.dumps(
                    {
                        "case": "date-valued multi-row column header",
                        "expected": expected_date_projection,
                        "actual": actual_date_projection,
                    },
                    ensure_ascii=False,
                    default=str,
                    indent=2,
                )
            )
        epoch_template = root / "epoch_template.xlsx"
        epoch_source = inputs / "EPOCH_source.xlsx"
        same_date = datetime(2024, 1, 31)
        make_date_header_book(epoch_template, same_date)
        make_date_header_book(epoch_source, same_date, date_1904=True)
        epoch_config = root / "epoch_config.xlsx"
        make_config(epoch_config, rust_output, inputs, workbook_keyword="EPOCH")
        expected_epoch = python_baseline(
            epoch_template, [epoch_source], epoch_config, root / "python_epoch_output"
        )
        actual_epoch = run_rust(epoch_template, [epoch_source], epoch_config)
        expected_epoch_projection = _projection_without_styles(workbook_projection(expected_epoch))
        actual_epoch_projection = _projection_without_styles(workbook_projection(actual_epoch))
        if expected_epoch_projection != actual_epoch_projection:
            raise AssertionError(
                json.dumps(
                    {
                        "case": "same calendar date stored with 1900 and 1904 epochs",
                        "expected": expected_epoch_projection,
                        "actual": actual_epoch_projection,
                    },
                    ensure_ascii=False,
                    default=str,
                    indent=2,
                )
            )
        report = {
            "schema_version": 1,
            "suite": "pytools.timeline.header_path_compare_parity",
            "date": "2026-10-26",
            "status": "passed_synthetic_semantics_header_style_and_dates",
            "entry": "timeline-compare CLI -> run_header_path_compare",
            "baseline": "PyCode/core/compare_393.py::run_compare",
            "rust_cli": f"{sys.platform} native debug",
            "style_comparison": "effective Excel style; implicit Calibri 11/theme text and omitted/automatic border colors are equivalent",
            "real_config_sample_availability": real_config_sample_availability(),
            "coverage": {
                "case_insensitive_workbook_keyword": "passed",
                "case_insensitive_sheet_keyword": "passed",
                "ascii_comma_sheet_keyword_delimiter": "passed",
                "nonmatching_workbook_keyword": "passed",
                "same_template_source_skipped": "passed",
                "output_cells_values_types_and_structure": "exact",
                "output_header_style": "exact",
                "date_valued_column_header_projection": "exact",
                "mixed_1900_1904_date_epoch_projection": "exact",
            },
            "sheet_names": expected_projection["sheet_names"],
            "matched_sources": 3,
            "selected_sources_including_unmatched": 4,
            "header_style_cells_checked": len(header_cells),
            "header_style_mismatch_cells": len(style_mismatches),
            "style_gap": None,
        }
        report_path = WORKSPACE / "pytools-core" / "probe" / "reports" / "timeline_compare_parity_20261026.json"
        report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2) + chr(10), encoding="utf-8")
        print(json.dumps(report, ensure_ascii=False))

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
