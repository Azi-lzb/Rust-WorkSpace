"""Run read-only pytools operation 4-8 against isolated synthetic files."""

from __future__ import annotations

import argparse
import hashlib
import json
import logging
import sys
import tempfile
from collections import Counter
from pathlib import Path

from openpyxl import Workbook


WORKSPACE = Path(__file__).resolve().parents[2]
CASES_PATH = WORKSPACE / "pytools-core/tests/fixtures/local_filename_map_48_cases.json"
SOURCE_RELATIVE = Path("PyCode/core/convert_33_34_35_37.py")


def sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest().upper()


def mappings(case: dict[str, object]) -> list[tuple[str, str]]:
    return [
        (
            rule["source"],
            rule.get("target", "x" * int(rule.get("target_repeat", 0))),
        )
        for rule in case["mapping"]
    ]


def make_config(path: Path, case: dict[str, object]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "重命名配置"
    sheet.cell(1, 13, "原文件名片段")
    sheet.cell(1, 14, "新文件名片段")
    for row, (source, target) in enumerate(mappings(case), start=2):
        sheet.cell(row, 13, source)
        sheet.cell(row, 14, target)
    workbook.save(path)
    workbook.close()


def build_inputs(input_dir: Path, case: dict[str, object]) -> list[Path]:
    for name in case.get("existing_names", []):
        (input_dir / name).write_bytes(b"occupied target")
    missing = set(case.get("missing_names", []))
    directories = set(case.get("directory_names", []))
    paths = []
    for name in case["input_names"]:
        path = input_dir / name
        paths.append(path)
        if name in missing:
            continue
        if name in directories:
            path.mkdir()
            (path / "marker.txt").write_bytes(b"directory")
        else:
            path.write_bytes(f"content for {name}".encode("utf-8"))
    return paths


def close_source_loggers() -> None:
    for name, value in list(logging.root.manager.loggerDict.items()):
        if not name.endswith("4_8_batch_rename_local_map"):
            continue
        if not isinstance(value, logging.Logger):
            continue
        for handler in list(value.handlers):
            handler.close()
            value.removeHandler(handler)


def run_case(function, root: Path, case: dict[str, object]) -> dict[str, object]:
    input_dir = root / "input"
    input_dir.mkdir(parents=True)
    config = root / "config.xlsx"
    make_config(config, case)
    files = build_inputs(input_dir, case)
    before = {
        path.name: sha256(path.read_bytes())
        for path in files
        if path.is_file()
    }
    try:
        stats = function(config, files, root / "logs")
    finally:
        close_source_loggers()

    final_entries = sorted(path.name for path in input_dir.iterdir())
    after = {
        path.name: sha256(path.read_bytes())
        for path in input_dir.iterdir()
        if path.is_file()
    }
    before_counts = Counter(before.values())
    after_counts = Counter(after.values())
    return {
        "stats": stats,
        "final_entries": final_entries,
        "input_file_sha256_by_name": dict(sorted(before.items())),
        "final_file_sha256_by_name": dict(sorted(after.items())),
        "input_file_bytes_preserved": all(
            after_counts[digest] >= count for digest, count in before_counts.items()
        ),
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--source-root",
        type=Path,
        default=Path(r"C:\Users\AZI\Desktop\VibeCode\4.pytools-main"),
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=WORKSPACE / "pytools-core/probe/reports/local_filename_map_48_python_baseline.json",
    )
    args = parser.parse_args()

    sys.dont_write_bytecode = True
    sys.path.insert(0, str(args.source_root))
    from PyCode.core.convert_33_34_35_37 import run_batch_rename_local_map_48

    cases = json.loads(CASES_PATH.read_text(encoding="utf-8"))
    with tempfile.TemporaryDirectory(prefix="pytools-local-map-48-") as raw_root:
        result = run_case(
            run_batch_rename_local_map_48,
            Path(raw_root) / "main-case",
            cases["operation48"],
        )

    source_file = args.source_root / SOURCE_RELATIVE
    payload = {
        "source": {
            "project": str(args.source_root),
            "module": SOURCE_RELATIVE.as_posix(),
            "sha256": sha256(source_file.read_bytes()),
            "operation": "run_batch_rename_local_map_48",
        },
        "operation48": result,
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(args.output)


if __name__ == "__main__":
    main()
