"""Compare the Rust 6-1/6-2 summary CLI with the read-only PyCode baseline.

The script creates only synthetic workbooks under a temporary directory. It
imports the original PyCode implementation as a baseline and never writes to
the source project.
"""

from __future__ import annotations

import argparse
import json
import math
import shlex
import shutil
import subprocess
import sys
import tempfile
from datetime import date, datetime, time, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill


def create_used_range_inputs(root: Path) -> list[Path]:
    first = root / "used_first.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "First"
    sheet["B2"] = "left"
    sheet["D4"] = 12
    sheet["E5"] = date(2024, 6, 30)
    sheet["F6"] = True
    sheet["J8"].fill = PatternFill(fill_type="solid", fgColor="FFFF00")
    sheet.merge_cells("H9:I10")
    sheet["H9"] = "merged"
    workbook.create_sheet("Empty")
    workbook.save(first)

    second = root / "used_second.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Second"
    sheet["A1"] = "head"
    sheet["C3"] = 3.5
    workbook.save(second)
    return [first, second]


def mark_region(sheet, start_cell: str, end_cell: str, keyword: str, number: int) -> None:
    sheet[start_cell].comment = Comment(f"{keyword}{number}", "probe")
    sheet[end_cell].comment = Comment(f"{keyword}#{number}", "probe")


def create_comment_template(path: Path) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, top, bottom, set_value in [
        ("Zeta", "上层", "金额", "甲机构"),
        ("Alpha", "年份", "余额", "乙机构"),
        ("模板", "指标", "本期", "默认机构"),
    ]:
        sheet = workbook.create_sheet(name)
        sheet["A1"] = "row-start"
        sheet["A3"] = "row-end"
        sheet["B1"] = top
        sheet["B3"] = bottom
        sheet["C1"] = set_value
        mark_region(sheet, "A1", "A3", "行区域", 1)
        mark_region(sheet, "B1", "B3", "列区域", 1)
        sheet["C1"].comment = Comment("set(机构)", "probe")
    workbook.save(path)


def create_comment_sources(root: Path) -> list[Path]:
    first = root / "comment_first.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Zeta"
    sheet["B1"] = "分组"
    sheet["B2"] = "甲"
    sheet["B3"] = 11
    sheet["C1"] = "甲机构"
    sheet["C2"] = "ignored"
    sheet["C3"] = "ignored"
    workbook.create_sheet("Alpha")
    workbook["Alpha"]["B1"] = "年份"
    workbook["Alpha"]["B2"] = 2024
    workbook["Alpha"]["B3"] = 2025
    workbook["Alpha"]["C1"] = "乙机构"
    workbook.save(first)

    second = root / "comment_fallback.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "NotInTemplate"
    sheet["B1"] = "项目"
    sheet["B2"] = "期初"
    sheet["B3"] = "期末"
    sheet["C1"] = "丙机构"
    workbook.save(second)
    return [first, second]


def scalar(value):
    if isinstance(value, (datetime, date, time)):
        return {"date_type": type(value).__name__, "value": value.isoformat()}
    if isinstance(value, timedelta):
        return {"date_type": "timedelta", "seconds": value.total_seconds()}
    if isinstance(value, float):
        if not math.isfinite(value):
            return {"float": repr(value)}
        return {"float": value}
    return value


def snapshot(path: Path) -> dict:
    workbook = load_workbook(path, data_only=True, read_only=False)
    result = {"sheet_order": workbook.sheetnames, "sheets": {}}
    for sheet in workbook.worksheets:
        values = [
            [scalar(sheet.cell(row, col).value) for col in range(1, sheet.max_column + 1)]
            for row in range(1, sheet.max_row + 1)
        ]
        header_bold = [
            sheet.cell(1, col).font.bold for col in range(1, sheet.max_column + 1)
        ] if sheet.max_row else []
        result["sheets"][sheet.title] = {
            "values": values,
            "rows": sheet.max_row,
            "columns": sheet.max_column,
            "header_bold": header_bold,
            "freeze_panes": sheet.freeze_panes,
        }
    workbook.close()
    return result


def equal_value(left, right) -> bool:
    if isinstance(left, dict) and isinstance(right, dict):
        if set(left) != set(right):
            return False
        if "float" in left and "float" in right:
            a, b = left["float"], right["float"]
            if isinstance(a, (int, float)) and isinstance(b, (int, float)):
                return abs(a - b) <= 1e-15 * max(1.0, abs(a), abs(b))
        return all(equal_value(left[key], right[key]) for key in left)
    if isinstance(left, list) and isinstance(right, list):
        return len(left) == len(right) and all(equal_value(a, b) for a, b in zip(left, right))
    if isinstance(left, (int, float)) and not isinstance(left, bool) and isinstance(right, (int, float)) and not isinstance(right, bool):
        return abs(left - right) <= 1e-15 * max(1.0, abs(left), abs(right))
    return left == right


def compare_snapshots(expected: dict, actual: dict) -> list[str]:
    differences = []
    if expected["sheet_order"] != actual["sheet_order"]:
        differences.append("sheet order")
    for name in expected["sheet_order"]:
        if name not in actual["sheets"]:
            differences.append(f"missing sheet {name}")
            continue
        for field in ("values", "rows", "columns", "header_bold", "freeze_panes"):
            if not equal_value(expected["sheets"][name][field], actual["sheets"][name][field]):
                differences.append(f"{name}.{field}")
    return differences


def call_rust(
    workspace: Path,
    operation: str,
    output_dir: Path,
    inputs: list[Path],
    template: Path | None = None,
):
    command = [
        "cargo", "run", "--quiet", "--manifest-path", str(workspace / "Cargo.toml"),
        "-p", "pytools-core", "--bin", "pytools", "--", operation,
        "--output-dir", str(output_dir),
    ]
    if template is not None:
        command.extend(["--template", str(template)])
    command.extend(str(path) for path in inputs)
    if shutil.which("cargo"):
        completed = subprocess.run(command, cwd=workspace, capture_output=True)
    elif shutil.which("wsl"):
        def linux_path(path: Path) -> str:
            raw = str(path.resolve()).replace("\\", "/")
            if len(raw) >= 2 and raw[1] == ":":
                return f"/mnt/{raw[0].lower()}{raw[2:]}"
            raise RuntimeError(f"cannot map Windows path into WSL: {path}")

        wsl_command = (
            f"cd {shlex.quote(linux_path(workspace))} && "
            + shlex.join([
                "cargo", "run", "--quiet", "--manifest-path", linux_path(workspace / "Cargo.toml"),
                "-p", "pytools-core", "--bin", "pytools", "--", operation,
                "--output-dir", linux_path(output_dir),
                *( ["--template", linux_path(template)] if template is not None else [] ),
                *(linux_path(path) for path in inputs),
            ])
        )
        completed = subprocess.run(
            ["wsl", "-e", "sh", "-lc", wsl_command], capture_output=True
        )
    else:
        raise RuntimeError("Rust toolchain not found (need cargo or WSL cargo)")
    stdout = completed.stdout.decode("utf-8", errors="replace").replace("\x00", "")
    stderr = completed.stderr.decode("utf-8", errors="replace").replace("\x00", "")
    if completed.returncode:
        raise RuntimeError(
            f"Rust command failed ({completed.returncode}): {stderr.strip()}"
        )
    for line in reversed(stdout.splitlines()):
        try:
            return json.loads(line)
        except json.JSONDecodeError:
            continue
    raise RuntimeError(f"Rust command did not print a JSON result: {stdout[-1000:]}")


def newest_output(directory: Path) -> Path:
    candidates = sorted(directory.glob("*.xlsx"), key=lambda path: path.stat().st_mtime_ns)
    if not candidates:
        raise RuntimeError(f"no output workbook in {directory}")
    return candidates[-1]


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workspace", type=Path, required=True)
    parser.add_argument("--source-root", type=Path, required=True)
    parser.add_argument("--report", type=Path, required=True)
    args = parser.parse_args()

    sys.path.insert(0, str(args.source_root))
    from PyCode.core.summary_321_322 import run_summary_by_comment, run_summary_by_usedrange

    with tempfile.TemporaryDirectory(prefix="pytools-summary-parity-") as temporary:
        root = Path(temporary)
        inputs = create_used_range_inputs(root)
        used_python_dir = root / "used-python"
        used_rust_dir = root / "used-rust"
        used_python = run_summary_by_usedrange(inputs, used_python_dir, root / "logs")
        used_rust = call_rust(args.workspace, "summary-used-range", used_rust_dir, inputs)
        used_differences = compare_snapshots(
            snapshot(Path(used_python["path"])), snapshot(newest_output(used_rust_dir))
        )

        template = root / "summary_template.xlsx"
        create_comment_template(template)
        comment_inputs = create_comment_sources(root)
        comment_python_dir = root / "comment-python"
        comment_rust_dir = root / "comment-rust"
        comment_python = run_summary_by_comment(
            template, comment_inputs, comment_python_dir, root / "logs"
        )
        comment_rust = call_rust(
            args.workspace, "summary-comments", comment_rust_dir, comment_inputs, template
        )
        comment_differences = compare_snapshots(
            snapshot(Path(comment_python["path"])), snapshot(newest_output(comment_rust_dir))
        )

        report = {
            "schema_version": 1,
            "suite": "pytools.summary.operations",
            "baseline": "read-only PyCode.core.summary_321_322",
            "scope": "synthetic XLSX; 6-1 used-range summary and 6-2 comment-template summary",
            "used_range": {
                "python_rows": used_python.get("rows"),
                "rust_rows": used_rust.get("rows"),
                "differences": used_differences,
                "status": "passed" if not used_differences else "failed",
            },
            "comment_summary": {
                "python_rows": comment_python.get("rows"),
                "rust_rows": comment_rust.get("rows"),
                "differences": comment_differences,
                "status": "passed" if not comment_differences else "failed",
            },
        }
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        print(json.dumps(report, ensure_ascii=False, indent=2))
        return 0 if report["used_range"]["status"] == report["comment_summary"]["status"] == "passed" else 1


if __name__ == "__main__":
    raise SystemExit(main())
