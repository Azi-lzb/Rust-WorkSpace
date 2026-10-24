"""Compare the PT-7-1 Python behavior with the real Rust menu execution."""

from __future__ import annotations

import argparse
import contextlib
import io
import json
import math
import sys
import tempfile
from pathlib import Path

from openpyxl import Workbook, load_workbook

sys.dont_write_bytecode = True
WORKSPACE = Path(__file__).resolve().parents[2]
DEFAULT_SOURCE = Path(r"C:\Users\AZI\Desktop\VibeCode\4.pytools-main")


def make_workbook(path: Path, sheets: dict[str, list[list[object]]]) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets.items():
        sheet = workbook.create_sheet(name)
        for row in rows:
            sheet.append(row)
    workbook.save(path)
    workbook.close()


def workbook_projection(path: Path) -> dict[str, list[list[object]]]:
    workbook = load_workbook(path, data_only=False)
    projection = {}
    for sheet in workbook.worksheets:
        rows = []
        for row in sheet.iter_rows():
            values = [cell.value for cell in row]
            while values and values[-1] is None:
                values.pop()
            rows.append(values)
        while rows and not rows[-1]:
            rows.pop()
        projection[sheet.title] = rows
    workbook.close()
    return projection


def values_equal(left: object, right: object) -> bool:
    if isinstance(left, bool) or isinstance(right, bool):
        return left is right
    if isinstance(left, (int, float)) and isinstance(right, (int, float)):
        return math.isclose(float(left), float(right), rel_tol=1e-15, abs_tol=1e-15)
    if isinstance(left, list) and isinstance(right, list):
        return len(left) == len(right) and all(values_equal(a, b) for a, b in zip(left, right))
    if isinstance(left, dict) and isinstance(right, dict):
        return left.keys() == right.keys() and all(values_equal(left[key], right[key]) for key in left)
    return left == right


def projection_differences(
    python_projection: dict[str, list[list[object]]],
    rust_projection: dict[str, list[list[object]]],
) -> list[dict[str, object]]:
    differences = []
    for sheet in sorted(set(python_projection) | set(rust_projection)):
        python_rows = python_projection.get(sheet, [])
        rust_rows = rust_projection.get(sheet, [])
        differing_rows = [
            index + 1
            for index in range(max(len(python_rows), len(rust_rows)))
            if not values_equal(
                python_rows[index] if index < len(python_rows) else None,
                rust_rows[index] if index < len(rust_rows) else None,
            )
        ]
        if differing_rows:
            differences.append(
                {
                    "sheet": sheet,
                    "python_rows": len(python_rows),
                    "rust_rows": len(rust_rows),
                    "differing_row_numbers": differing_rows,
                }
            )
    return differences


def run_json(command: list[str], *, input_text: str | None = None) -> tuple[dict, str]:
    import subprocess

    result = subprocess.run(
        command,
        input=input_text.encode("utf-8") if input_text is not None else None,
        capture_output=True,
        check=False,
    )
    stdout = result.stdout.decode("utf-8-sig", errors="replace")
    stderr = result.stderr.decode("utf-8-sig", errors="replace")
    if result.returncode != 0:
        raise RuntimeError(
            f"command failed ({result.returncode}): {command[0]}\n"
            f"stdout={stdout}\nstderr={stderr}"
        )
    marker = '{\n  "feature_id": "PT-7-1"'
    start = stdout.find(marker)
    if start < 0:
        start = stdout.find("{")
    if start < 0:
        raise RuntimeError(f"command emitted no JSON: {stdout}")
    report, _ = json.JSONDecoder().raw_decode(stdout[start:])
    return report, stderr


def run_probe(pytools: Path, source: Path) -> dict:
    sys.path.insert(0, str(source / "PyCode"))
    from core.feature_1_1_split_village_bank import run_split_village_bank

    platform = "synthetic-pt-7-1-umya-menu-probe"
    with tempfile.TemporaryDirectory(prefix="pytools-pt71-menu-") as temporary:
        root = Path(temporary)
        current = root / "current.xlsx"
        previous = root / "previous.xlsx"
        target = root / "target.xlsx"
        python_target = root / "python-target.xlsx"
        candidate = root / "rust-candidate.xlsx"
        capability_path = root / "test-capability.json"
        config_path = root / "config.xlsx"
        output_dir = root / "output"
        log_dir = root / "logs"

        names = ["村镇银行本外币", "村镇银行人民币", "村镇银行外汇"]
        current_sheets = {}
        previous_sheets = {}
        for index, name in enumerate(names, 1):
            current_rows = [[None] * 89 for _ in range(6)]
            previous_rows = [[None] * 89 for _ in range(6)]
            for row in (5, 6):
                current_rows[row - 1][2:5] = [index * row * 10, f"存款-{index}-{row}", row]
                previous_rows[row - 1][2] = index * row * 5
                current_rows[row - 1][86:89] = [index * row * 20, f"贷款-{index}-{row}", row + 1]
                previous_rows[row - 1][86] = index * row * 10
            current_sheets[name] = current_rows
            previous_sheets[name] = previous_rows
        make_workbook(current, current_sheets)
        make_workbook(previous, previous_sheets)
        make_workbook(
            target,
            {
                "广东分机构本外币": [["机构"], ["其他"], ["村镇银行"], ["尾行"]],
                "广东分机构人民币": [["机构"], ["村镇银行"]],
                "广东分机构外汇": [["机构"], ["村镇银行"]],
                "广东分机构未知": [["机构"], ["村镇银行"]],
                "无村镇行分机构本外币": [["机构"], ["普通银行"]],
                "非分机构工作表": [["村镇银行"]],
            },
        )
        current_before = current.read_bytes()
        previous_before = previous.read_bytes()
        target_before = target.read_bytes()
        config = Workbook()
        global_sheet = config.active
        global_sheet.title = "全局配置"
        global_sheet.append(["配置项", "值"])
        global_sheet.append(["输入目录", str(root / "input")])
        global_sheet.append(["输出目录", str(output_dir)])
        global_sheet.append(["日志目录", str(log_dir)])
        config.save(config_path)
        config.close()

        python_target.write_bytes(target_before)
        import logging

        try:
            with contextlib.redirect_stderr(io.StringIO()):
                python_stats = run_split_village_bank(current, previous, [python_target], log_dir)
        finally:
            logger = logging.getLogger("PyCode.feature_1_1_split_village_bank")
            for handler in list(logger.handlers):
                handler.close()
                logger.removeHandler(handler)

        plan_command = [
            str(pytools),
            "village-split-plan",
            "--current",
            str(current),
            "--previous",
            str(previous),
            "--input",
            str(target),
            "--output",
            str(candidate),
        ]
        plan_report, _ = run_json(plan_command)
        plan = plan_report["plan"]
        capability = {
            "engine": "umya_spreadsheet",
            "status": "passed",
            "format": plan["format"],
            "operations": [operation["operation"] for operation in plan["operations"]],
            "supported_features": plan["workbook"]["features_present"],
            "preserved": plan["preserve"],
            "platform": platform,
            "evidence": {
                "sample_id": "synthetic-pt-7-1-village-split-menu",
                "report_path": "pytools-core/probe/compare_village_split_menu.py",
                "adapter_version": "umya-spreadsheet-3.1.1",
                "engine_version": "umya-spreadsheet-3.1.1",
                "platform": platform,
            },
        }
        capability_path.write_text(json.dumps(capability, ensure_ascii=False, indent=2), encoding="utf-8")
        menu_input = f"7\n1\n{current}\n{previous}\n1\n{target}\n0\n0\n"
        rust_report, stderr = run_json(
            [
                str(pytools),
                "menu",
                "--config",
                str(config_path),
                "--capabilities",
                str(capability_path),
                "--platform",
                platform,
            ],
            input_text=menu_input,
        )
        rust_projection = workbook_projection(target)
        python_projection = workbook_projection(python_target)
        differences = projection_differences(python_projection, rust_projection)
        projection_equal = not differences
        if current.read_bytes() != current_before or previous.read_bytes() != previous_before:
            raise AssertionError("read-only current/prior period workbooks changed")
        if target.read_bytes() == target_before:
            raise AssertionError("Rust menu did not commit the target workbook")
        item = rust_report["items"][0]
        if python_stats["filled_sheets"] != item["matched_sheets"]:
            raise AssertionError("Python/Rust filled sheet counts differ")
        if item["status"] != "committed":
            raise AssertionError("Rust menu did not report a committed target")
        return {
            "status": "passed_synthetic_menu_execution_parity"
            if projection_equal
            else "failed_synthetic_menu_execution_parity",
            "python_version": sys.version.split()[0],
            "openpyxl_version": __import__("openpyxl").__version__,
            "python": {
                "files": python_stats["files"],
                "branch_sheets": python_stats["branch_sheets"],
                "filled_sheets": python_stats["filled_sheets"],
                "skipped_sheets": python_stats["skipped_sheets"],
                "failed_files": python_stats["failed_files"],
            },
            "rust": {
                "feature_id": rust_report["feature_id"],
                "files": rust_report["files"],
                "matched_sheets": item["matched_sheets"],
                "changed_cells": item["changed_cells"],
                "failed_files": rust_report["failed_files"],
                "engine": item["execution"].get("engine"),
                "status": item["status"],
            },
            "output_projection_equal": projection_equal,
            "projection_differences": differences,
            "current_and_previous_unchanged": True,
            "target_committed": True,
            "report_contains_raw_cells": False,
            "rust_stderr_nonempty": bool(stderr.strip()),
        }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--pytools-exe", required=True, type=Path)
    parser.add_argument("--python-source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument(
        "--report-path",
        type=Path,
        default=WORKSPACE / "pytools-core/probe/reports/village_split_menu_parity_20261024.json",
    )
    arguments = parser.parse_args()
    if not arguments.pytools_exe.is_file():
        raise SystemExit(f"Rust CLI executable not found: {arguments.pytools_exe}")
    report = {
        "schema_version": 1,
        "suite": "pytools.PT-7-1.real_menu_python_rust_parity",
        "date": "2026-10-24",
        "platform": "Windows native",
        "sample_kind": "synthetic",
        **run_probe(arguments.pytools_exe, arguments.python_source),
    }
    arguments.report_path.parent.mkdir(parents=True, exist_ok=True)
    arguments.report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    if report["status"] != "passed_synthetic_menu_execution_parity":
        raise SystemExit(1)


if __name__ == "__main__":
    main()
