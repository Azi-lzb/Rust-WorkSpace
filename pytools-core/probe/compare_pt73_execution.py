"""Compare the PT-7-3 Python baseline with the Rust batch execution route.

All workbooks are generated in a temporary directory. The Python source
function is imported read-only. The Rust CLI receives a temporary, exact-match
capability solely to exercise the already-tested Umya execution route; this
probe does not create production capability evidence or change routing.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import logging
import shlex
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import openpyxl
from openpyxl import Workbook, load_workbook

sys.dont_write_bytecode = True
WORKSPACE = Path(__file__).resolve().parents[2]
SOURCE = Path(r"C:\Users\AZI\Desktop\VibeCode\4.pytools-main")
sys.path.insert(0, str(SOURCE / "PyCode"))

from core.feature_1_3_remove_foreign_bank import run_remove_foreign_bank


TEST_PLATFORM = "synthetic-pt-7-3-umya-execution-probe"
FOREIGN_BANKS = {"外资甲", "外资乙"}


def wsl_path(path: Path) -> str:
    return f"/mnt/{path.drive.rstrip(':').lower()}/{path.as_posix()[2:].lstrip('/')}"


def run_rust(command: str, arguments: list[str | Path]) -> dict:
    workspace = wsl_path(WORKSPACE)
    rendered = [shlex.quote(command)]
    for argument in arguments:
        if isinstance(argument, Path):
            rendered.append(shlex.quote(wsl_path(argument)))
        else:
            rendered.append(shlex.quote(argument))
    command_line = (
        "cd "
        + shlex.quote(workspace)
        + " && cargo run --offline --locked --quiet -p pytools-core --bin pytools -- "
        + " ".join(rendered)
    )
    result = subprocess.run(
        ["wsl", "-e", "sh", "-lc", command_line], capture_output=True, check=False
    )
    stdout = result.stdout.decode("utf-8-sig", errors="replace").replace("\x00", "")
    if result.returncode != 0:
        raise RuntimeError(
            f"Rust {command} failed (exit={result.returncode}): {stdout}\n"
            f"{result.stderr.decode('utf-8', errors='replace')}"
        )
    start = stdout.find("{")
    if start < 0:
        raise RuntimeError(f"Rust {command} emitted no JSON: {stdout}")
    return json.loads(stdout[start:])


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def make_config(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "机构映射表"
    for row in (
        ("序号", "机构名称", "是否为外资行"),
        (1, "外资甲", 1),
        (2, "外资乙", 1),
        (3, "普通行", 0),
    ):
        sheet.append(row)
    workbook.save(path)
    workbook.close()


def make_input(path: Path, include_foreign: bool) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)

    first = workbook.create_sheet("甲分机构本外币")
    first.append(["机构", "余额"])
    first.append(["普通甲", 10])
    first.append(["外资甲" if include_foreign else "普通己", 20])
    first.append(["普通乙", 30])
    first.append(["外资乙" if include_foreign else "普通庚", 40])
    first.append(["普通丙", 50])
    first.append(["尾行", 60])

    second = workbook.create_sheet("乙分机构人民币")
    second.append(["机构"])
    second.append(["外资乙" if include_foreign else "普通戊"])
    second.append(["普通丁"])

    other = workbook.create_sheet("非目标工作表")
    other.append(["机构"])
    other.append(["外资甲"])

    workbook.save(path)
    workbook.close()


def workbook_projection(path: Path) -> dict:
    workbook = load_workbook(path, data_only=False)
    projection = {}
    for sheet in workbook.worksheets:
        rows = []
        for row in sheet.iter_rows():
            rows.append([cell.value for cell in row])
        while rows and all(value is None for value in rows[-1]):
            rows.pop()
        rows = [
            row[: max((index + 1 for index, value in enumerate(row) if value is not None), default=0)]
            for row in rows
        ]
        projection[sheet.title] = rows
    workbook.close()
    return projection


def capability_for(plan: dict) -> dict:
    platform = TEST_PLATFORM
    return {
        "engine": "umya_spreadsheet",
        "status": "passed",
        "format": plan["format"],
        "operations": sorted({operation["operation"] for operation in plan["operations"]}),
        "supported_features": sorted(plan["workbook"]["features_present"]),
        "preserved": sorted(plan["preserve"]),
        "platform": platform,
        "evidence": {
            "sample_id": "synthetic-pt-7-3-batch-execution",
            "report_path": "pytools-core/probe/compare_pt73_execution.py",
            "adapter_version": "umya-spreadsheet-3.1.1",
            "engine_version": "umya-spreadsheet-3.1.1",
            "platform": platform,
        },
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--report-path",
        type=Path,
        default=(
            WORKSPACE
            / "pytools-core"
            / "probe"
            / "reports"
            / "PT-7-3_execution_parity_20261023.json"
        ),
        help="write the machine-readable parity report to this path",
    )
    args = parser.parse_args()

    source_hashes = {}
    baseline_hashes = {}
    with tempfile.TemporaryDirectory(prefix="pt-7-3-execution-") as temporary:
        root = Path(temporary)
        config = root / "config.xlsx"
        sources = [root / "changed.xlsx", root / "unchanged.xlsx"]
        python_sources = [root / "python-changed.xlsx", root / "python-unchanged.xlsx"]
        make_config(config)
        make_input(sources[0], include_foreign=True)
        make_input(sources[1], include_foreign=False)
        for source, python_source in zip(sources, python_sources, strict=True):
            source_hashes[source.name] = sha256(source)
            shutil.copy2(source, python_source)

        logger = logging.getLogger("pt-7-3-execution-parity")
        logger.addHandler(logging.NullHandler())
        python_summary = run_remove_foreign_bank(
            python_sources, FOREIGN_BANKS, root / "python-logs"
        )
        logging.shutdown()
        python_outputs = [
            source.with_name(f"{source.stem}(金融局){source.suffix}")
            for source in python_sources
        ]
        baseline_hashes = {output.name: sha256(output) for output in python_outputs}

        plan_only = run_rust(
            "remove-foreign-bank", ["--config", config, *sources]
        )
        if plan_only.get("mode") != "plan_only":
            raise AssertionError(f"unexpected no-capability mode: {plan_only}")
        if plan_only.get("planned_files") != python_summary["ok_files"]:
            raise AssertionError(f"planned file count differs: {plan_only} vs {python_summary}")
        if any(
            source.with_name(f"{source.stem}(金融局){source.suffix}").exists()
            for source in sources
        ):
            raise AssertionError("plan-only mode created an output workbook")

        first_plan = plan_only["items"][0].get("plan")
        if not first_plan:
            raise AssertionError(f"changed workbook did not produce a plan: {plan_only}")
        capability = root / "test-only-capability.json"
        capability.write_text(
            json.dumps(capability_for(first_plan), ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )

        execution = run_rust(
            "remove-foreign-bank",
            [
                "--config",
                config,
                "--capabilities",
                capability,
                "--platform",
                TEST_PLATFORM,
                *sources,
            ],
        )
        rust_outputs = [
            source.with_name(f"{source.stem}(金融局){source.suffix}")
            for source in sources
        ]
        if execution.get("mode") != "execute":
            raise AssertionError(f"unexpected execution mode: {execution}")
        for key in ("files", "ok_files", "branch_sheets", "deleted_rows", "failed_files"):
            if execution.get(key) != python_summary.get(key):
                raise AssertionError(
                    f"PT-7-3 {key} differs: Rust report={execution!r}, "
                    f"Python summary={python_summary!r}"
                )
        for python_output, rust_output in zip(python_outputs, rust_outputs, strict=True):
            if workbook_projection(python_output) != workbook_projection(rust_output):
                raise AssertionError(
                    f"workbook projection differs: {python_output.name} vs {rust_output.name}"
                )
        if execution["items"][0]["execution"].get("engine") != "umya_spreadsheet":
            raise AssertionError(f"unexpected execution engine: {execution['items'][0]}")
        if execution["items"][1]["execution"].get("operation") != "copied":
            raise AssertionError(f"unchanged workbook was not byte-copied: {execution['items'][1]}")

        for source in sources:
            if sha256(source) != source_hashes[source.name]:
                raise AssertionError(f"Rust or Python modified the original input: {source.name}")
        unchanged_copy_is_identical = rust_outputs[1].read_bytes() == sources[1].read_bytes()
        if not unchanged_copy_is_identical:
            raise AssertionError("unchanged workbook output is not a byte-identical source copy")

        report = {
            "schema_version": 1,
            "suite": "pytools.pt_7_3_batch_execution_parity",
            "status": "passed_synthetic_execution_parity",
            "feature_id": "PT-7-3",
            "sample_kind": "synthetic_xlsx",
            "host": "Windows 11 x64",
            "python": {
                "version": sys.version.split()[0],
                "openpyxl": openpyxl.__version__,
                "baseline_module": str(
                    SOURCE / "PyCode/core/feature_1_3_remove_foreign_bank.py"
                ),
                "baseline_sha256": sha256(
                    SOURCE / "PyCode/core/feature_1_3_remove_foreign_bank.py"
                ),
                "summary": python_summary,
            },
            "rust": {
                "command": "pytools remove-foreign-bank",
                "execution": execution,
                "engine": "umya_spreadsheet",
                "platform_capability": TEST_PLATFORM,
            },
            "comparison": {
                "counts": "equal",
                "sheet_order_and_cell_values": "equal",
                "original_inputs_unchanged": True,
                "unchanged_workbook_output_byte_copy": unchanged_copy_is_identical,
                "execution_capability": "temporary_test_only",
                "production_capability_created": False,
                "default_route_changed": False,
            },
            "synthetic_source_sha256": source_hashes,
            "python_output_sha256": baseline_hashes,
            "limitations": [
                "No real pytools workbook or configuration was available; production acceptance remains unverified.",
                "This probe compares business values, sheet order, counts and source safety; it does not qualify OOXML object preservation.",
                "Formula, comment, merge and cross-sheet reference behavior remains covered only by the focused Rust integration test and must be compared with an Excel COM oracle for this feature.",
                "Windows-native Rust, Kylin/UOS, glibc 2.28 and offline deb packaging were not verified.",
            ],
        }
        output = args.report_path
        output.parent.mkdir(parents=True, exist_ok=True)
        output.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        print(json.dumps(report, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
