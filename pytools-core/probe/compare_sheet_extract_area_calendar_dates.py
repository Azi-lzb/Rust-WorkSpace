from __future__ import annotations

import argparse
import importlib
import json
import logging
import subprocess
import sys
import tempfile
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook


HEADERS = [
    "是否启用",
    "工作簿关键字",
    "工作表关键字",
    "排除关键字",
    "提取区域",
    "多区域拼接方式",
    "输出文件夹",
    "文件名",
    "备注",
]


def cell_value(value: object) -> dict[str, str] | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return {"type": "datetime", "value": value.isoformat(timespec="milliseconds")}
    if isinstance(value, date):
        return {"type": "date", "value": value.isoformat()}
    if isinstance(value, time):
        return {"type": "time", "value": value.isoformat(timespec="milliseconds")}
    if isinstance(value, timedelta):
        return {"type": "timedelta", "value": str(value)}
    return {"type": type(value).__name__, "value": str(value)}


def project_workbook(path: Path) -> dict[str, object]:
    workbook = load_workbook(path, data_only=True)
    try:
        sheets = []
        for worksheet in workbook.worksheets:
            cells = []
            for row in worksheet.iter_rows():
                for cell in row:
                    projected = cell_value(cell.value)
                    if projected is not None:
                        cells.append(
                            {
                                "cell": cell.coordinate,
                                **projected,
                                "number_format": cell.number_format,
                            }
                        )
            sheets.append({"name": worksheet.title, "cells": cells})
        return {"sheet_names": workbook.sheetnames, "sheets": sheets}
    finally:
        workbook.close()


def create_rust_config(path: Path, output_dir: Path) -> None:
    workbook = Workbook()
    global_config = workbook.active
    global_config.title = "全局配置"
    global_config.append(["配置项", "值"])
    global_config.append(["输入目录", ""])
    global_config.append(["输出目录", str(output_dir)])
    global_config.append(["日志目录", str(output_dir / "logs")])

    extraction = workbook.create_sheet("工作表提取")
    extraction.append(HEADERS)
    extraction.append(
        [
            "是",
            "",
            "",
            "",
            "E1",
            "",
            "日期提取",
            "date-output.xlsx",
            "calendar datetime parity",
        ]
    )
    workbook.save(path)
    workbook.close()


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--pytools-cli", type=Path, required=True)
    parser.add_argument("--python-project", type=Path, required=True)
    parser.add_argument("--fixture", type=Path)
    parser.add_argument("--report", type=Path)
    args = parser.parse_args()

    probe_root = Path(__file__).resolve().parent
    pytools_root = probe_root.parent
    fixture = args.fixture or (
        pytools_root / "tests" / "fixtures" / "sheet_extract_area" / "省级月度报表_A.xlsx"
    )
    report_path = args.report or (
        probe_root / "reports" / "sheet_extract_calendar_dates_parity_20261026.json"
    )
    source_python = args.python_project.resolve() / "PyCode"
    sys.path.insert(0, str(source_python))
    legacy = importlib.import_module("core.feature_1_6_extract_sheets")
    source_bytes = fixture.read_bytes()

    with tempfile.TemporaryDirectory(prefix="pytools-calendar-parity-") as temporary:
        root = Path(temporary)
        python_output = root / "python-output"
        python_logs = root / "python-logs"
        rust_output = root / "rust-output"
        config_path = root / "config.xlsx"
        task = {
            "row_no": 2,
            "workbook_keywords": "",
            "sheet_keywords": "",
            "exclude_keywords": "",
            "area_spec": "E1",
            "stitch_mode": "",
            "output_folder": "日期提取",
            "output_name": "date-output.xlsx",
            "remark": "calendar datetime parity",
        }
        python_result = legacy.run_extract_sheets(
            [fixture],
            [task],
            python_output,
            python_logs,
            append_existing=False,
            default_output_name="unused.xlsx",
        )
        create_rust_config(config_path, rust_output)
        rust_process = subprocess.run(
            [
                str(args.pytools_cli.resolve()),
                "extract-sheets-area",
                "--config",
                str(config_path),
                "--output-dir",
                str(rust_output),
                str(fixture),
            ],
            capture_output=True,
            text=True,
            check=False,
        )
        if rust_process.returncode != 0:
            raise RuntimeError(
                f"Rust CLI failed ({rust_process.returncode}): {rust_process.stderr}"
            )
        rust_stats = json.loads(rust_process.stdout)
        python_path = python_output / "日期提取" / "date-output.xlsx"
        rust_path = rust_output / "日期提取" / "date-output.xlsx"
        python_projection = project_workbook(python_path)
        rust_projection = project_workbook(rust_path)
        source_unchanged = fixture.read_bytes() == source_bytes
        if not source_unchanged:
            raise AssertionError("read-only source fixture changed")
        if python_projection != rust_projection:
            raise AssertionError(
                "Python/Rust output mismatch:\n"
                + json.dumps(
                    {"python": python_projection, "rust": rust_projection},
                    ensure_ascii=False,
                    indent=2,
                )
            )
        if python_result["failed_files"] != 0 or rust_stats["failed_files"] != 0:
            raise AssertionError("the date extraction task reported a failed source")

        report = {
            "schema_version": 1,
            "status": "passed",
            "feature": "PT-7-6 area extraction calendar datetime subset",
            "fixture": "tests/fixtures/sheet_extract_area/省级月度报表_A.xlsx",
            "source_python": "PyCode/core/feature_1_6_extract_sheets.py::run_extract_sheets",
            "rust_entry": "pytools extract-sheets-area",
            "python_stats": {key: value for key, value in python_result.items() if key != "paths"},
            "rust_stats": {key: value for key, value in rust_stats.items() if key != "paths"},
            "output_projection": python_projection,
            "source_unchanged": source_unchanged,
            "scope_note": (
                "Synthetic fixture only. Time-only values on Excel epoch days, elapsed durations, "
                "and real business samples remain unverified and fail closed."
            ),
        }
        report_path.parent.mkdir(parents=True, exist_ok=True)
        report_path.write_text(
            json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
        )
        print(json.dumps(report, ensure_ascii=False, indent=2))
        logging.shutdown()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())