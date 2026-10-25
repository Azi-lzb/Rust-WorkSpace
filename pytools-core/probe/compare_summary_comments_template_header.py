from __future__ import annotations

import datetime as dt
import json
import logging
import subprocess
import sys
import tempfile
from pathlib import Path

import openpyxl
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment

WORKSPACE = Path(__file__).resolve().parents[2]
PYTOOLS_SOURCE = Path(r"C:\Users\AZI\Desktop\VibeCode\4.pytools-main")
RUST_CLI = WORKSPACE / "target" / "debug" / "pytools.exe"
REPORT = (
    Path(__file__).resolve().parent
    / "reports"
    / f"summary_comments_template_header_parity_{dt.date.today():%Y%m%d}.json"
)

sys.path.insert(0, str(WORKSPACE / "pytools-core" / "probe"))
sys.path.insert(0, str(PYTOOLS_SOURCE / "PyCode"))

from compare_summary_tools import (  # noqa: E402
    assert_equivalent,
    projection_difference_fields,
    source_fingerprints,
)
from core.summary_321_322 import run_summary_by_comment  # noqa: E402


def create_template(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Branch_A"
    sheet["A2"] = "row start"
    sheet["A2"].comment = Comment("行区域1", "probe")
    sheet["A4"] = "row end"
    sheet["A4"].comment = Comment("行区域#1", "probe")
    sheet["B1"] = '="余额"'
    sheet["B1"].comment = Comment("列区域1", "probe")
    sheet["C1"].comment = Comment("列区域#1", "probe")
    sheet["D1"] = True
    sheet["D1"].comment = Comment("列区域2", "probe")
    sheet["D2"].comment = Comment("列区域#2", "probe")
    workbook.save(path)
    workbook.close()


def create_source(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Branch_A"
    for row in range(2, 5):
        sheet.cell(row, 2).value = row * 10
        sheet.cell(row, 3).value = row * 100
    workbook.save(path)
    workbook.close()


def main() -> int:
    if sys.platform != "win32":
        raise RuntimeError("this parity probe targets the Windows native CLI")
    if not PYTOOLS_SOURCE.is_dir():
        raise FileNotFoundError(PYTOOLS_SOURCE)
    if not RUST_CLI.is_file():
        raise FileNotFoundError(RUST_CLI)
    rust_toolchain = subprocess.run(
        ["rustc", "--version"], check=True, capture_output=True, text=True, encoding="utf-8"
    ).stdout.split()[1]

    with tempfile.TemporaryDirectory(prefix="pytools-summary-comments-header-") as temp:
        root = Path(temp)
        template = root / "template.xlsx"
        source = root / "source.xlsx"
        create_template(template)
        create_source(source)
        before = source_fingerprints([template, source])

        logging.disable(logging.CRITICAL)
        try:
            python_result = run_summary_by_comment(
                template, [source], root / "python", root / "logs"
            )
        finally:
            logging.shutdown()
            logging.disable(logging.NOTSET)
        if not python_result.get("saved"):
            raise AssertionError("PyCode did not save the synthetic PT-6-2 output")

        completed = subprocess.run(
            [
                str(RUST_CLI),
                "summary-comments",
                "--template",
                str(template),
                "--output-dir",
                str(root / "rust"),
                "--log-view",
                "off",
                str(source),
            ],
            check=True,
            capture_output=True,
            text=True,
            encoding="utf-8",
        )
        rust_result = json.loads(completed.stdout)
        python_output = Path(python_result["path"])
        rust_output = Path(rust_result["path"])
        python_dimensions, rust_dimensions = assert_equivalent(
            python_output,
            rust_output,
            "PT-6-2 formula and blank template header semantics",
        )
        differences = projection_difference_fields(python_output, rust_output)
        if differences:
            raise AssertionError(f"workbook projection differs: {differences}")

        output_book = load_workbook(rust_output, data_only=False)
        try:
            output_sheet = output_book["Branch_A"]
            formula_header = output_sheet["C1"]
            blank_header = output_sheet["D1"]
            boolean_header = output_sheet["E1"]
            if (
                formula_header.data_type != "f"
                or blank_header.value != "C"
                or boolean_header.value != "True"
            ):
                raise AssertionError("Rust template header semantics changed")
        finally:
            output_book.close()

        if before != source_fingerprints([template, source]):
            raise AssertionError("the generated source input changed")
        if rust_result.get("rows") != 3:
            raise AssertionError("summary row count differs from the three-row input region")

        report = {
            "schema_version": 1,
            "suite": "pytools.summary_comments_template_header_parity",
            "date": dt.date.today().isoformat(),
            "status": "passed_python_rust_projection_parity",
            "data_origin": "synthetic XLSX contract fixture and read-only PyCode baseline",
            "environment": {
                "platform": "Windows x86_64",
                "rust_target": "x86_64-pc-windows-msvc",
                "rust_toolchain": rust_toolchain,
                "python_version": sys.version.split()[0],
                "openpyxl_version": openpyxl.__version__,
                "pytools_branch": "pytools-dev",
            },
            "entry_route": "pytools summary-comments",
            "baseline": "PyCode/core/summary_321_322.py::run_summary_by_comment",
            "input": {"format": "xlsx", "sample_count": 2},
            "behaviors": [
                "formula template header remains formula text in the output workbook",
                "blank template header falls back to its Excel column letter",
                "boolean template header uses Python string capitalization",
            ],
            "comparison": {
                "workbook_projection": "exact",
                "different_fields": differences,
                "fields": [
                    "sheet names and dimensions",
                    "row values and cell types",
                    "non-empty cell styles",
                    "freeze panes and header bold state",
                    "merged ranges, formula text, and defined names",
                    "print area, page setup, margins, and centering",
                ],
            },
            "execution": {
                "status": "succeeded",
                "summary_rows": rust_result["rows"],
                "output_dimensions": rust_dimensions,
                "python_dimensions": python_dimensions,
                "reader_engine": rust_result["execution"]["reader_engine"],
                "writer_engine": rust_result["execution"]["writer_engine"],
            },
            "safety": {
                "source_unchanged": True,
                "outputs_temporary_only": True,
                "report_contains_input_filenames_or_cell_values": False,
            },
            "limitations": [
                "This evidence covers a synthetic template and does not substitute for a valid real PT-6-2 business template.",
                "Linux targets and offline deb acceptance were not tested.",
            ],
        }

    REPORT.parent.mkdir(parents=True, exist_ok=True)
    REPORT.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"status": report["status"], "report": str(REPORT)}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
