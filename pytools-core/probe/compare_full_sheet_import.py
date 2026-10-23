from __future__ import annotations

import argparse
import hashlib
import json
import sys
import tempfile
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Protection


def add_source_project() -> None:
    root = Path(r"C:\Users\AZI\Desktop\VibeCode\4.pytools-main")
    if str(root) not in sys.path:
        sys.path.insert(0, str(root))


def make_source_fixture(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Source"
    sheet["A1"] = "primary"
    sheet["A1"].font = Font(bold=True)
    sheet["A1"].fill = PatternFill(fill_type="solid", fgColor="FFFFCC00")
    sheet["A1"].alignment = Alignment(
        horizontal="center", vertical="center", wrap_text=True
    )
    sheet["A1"].protection = Protection(
        locked=False, hidden=False
    )
    sheet["A1"].number_format = "0.00"
    sheet["C3"] = "=1+1"
    sheet.merge_cells("A1:B1")
    sheet.row_dimensions[8].height = 22.0
    sheet.row_dimensions[8].hidden = True
    sheet.row_dimensions[8].outlineLevel = 2
    sheet.column_dimensions["D"].width = 24.0
    sheet.column_dimensions["D"].hidden = True
    sheet.column_dimensions["D"].outlineLevel = 1
    workbook.save(path)
    workbook.close()

    # openpyxl has no public cached-formula writer. Patch only the generated
    # test fixture's C3 cached value through a namespace-aware XML parser.
    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    cell_tag = f"{{{namespace}}}c"
    value_tag = f"{{{namespace}}}v"
    temporary = path.with_suffix(".patched.xlsx")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(
        temporary, "w", compression=zipfile.ZIP_DEFLATED
    ) as patched:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                root = ET.fromstring(data)
                cell = root.find(f".//{cell_tag}[@r='C3']")
                if cell is None or cell.find(f"{{{namespace}}}f") is None:
                    raise RuntimeError("expected formula cell C3 in fixture XML")
                value = cell.find(value_tag)
                if value is None:
                    value = ET.SubElement(cell, value_tag)
                value.text = "2"
                data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            patched.writestr(info, data)
    temporary.replace(path)

    target = Workbook()
    target.active.title = "Source"
    target.save(path.with_name("target.xlsx"))
    target.close()


def _meaningful_dimensions(sheet):
    rows = {}
    for index, dimension in sheet.row_dimensions.items():
        values = {
            "height": dimension.height,
            "hidden": bool(dimension.hidden),
            "outline": int(dimension.outlineLevel or 0),
        }
        if values["height"] is not None or values["hidden"] or values["outline"]:
            rows[int(index)] = values
    columns = {}
    for key, dimension in sheet.column_dimensions.items():
        values = {
            "width": dimension.width,
            "hidden": bool(dimension.hidden),
            "outline": int(dimension.outlineLevel or 0),
        }
        if values["width"] is not None or values["hidden"] or values["outline"]:
            columns[str(key)] = values
    return rows, columns


def _snapshot(path: Path):
    workbook = load_workbook(path, data_only=True, read_only=False)
    result = {}
    for sheet in workbook.worksheets:
        cells = {}
        for row in range(1, sheet.max_row + 1):
            for column in range(1, sheet.max_column + 1):
                cell = sheet.cell(row, column)
                style = {
                    "number_format": cell.number_format,
                    "font": {
                        "name": cell.font.name,
                        "size": cell.font.sz,
                        "bold": cell.font.bold,
                        "italic": cell.font.italic,
                        "color_type": cell.font.color.type if cell.font.color else None,
                        "color_rgb": cell.font.color.rgb if cell.font.color and cell.font.color.type == "rgb" else None,
                    },
                    "fill": {
                        "fill_type": cell.fill.fill_type,
                        "fg_type": cell.fill.fgColor.type,
                        "fg_rgb": cell.fill.fgColor.rgb if cell.fill.fgColor.type == "rgb" else None,
                    },
                    "alignment": {
                        "horizontal": cell.alignment.horizontal,
                        "vertical": cell.alignment.vertical,
                        "wrap_text": cell.alignment.wrap_text,
                        "text_rotation": cell.alignment.text_rotation,
                    },
                    "protection": {
                        "locked": cell.protection.locked,
                        "hidden": cell.protection.hidden,
                    },
                }
                cells[cell.coordinate] = {"value": cell.value, "style": style}
        rows, columns = _meaningful_dimensions(sheet)
        result[sheet.title] = {
            "max_row": sheet.max_row,
            "max_column": sheet.max_column,
            "cells": cells,
            "merges": sorted(str(value) for value in sheet.merged_cells.ranges),
            "row_dimensions": rows,
            "column_dimensions": columns,
            "comment_cells": sorted(
                cell.coordinate
                for row in sheet.iter_rows()
                for cell in row
                if cell.comment is not None
            ),
        }
    workbook.close()
    return result


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", required=True, type=Path)
    parser.add_argument("--prepare-source", action="store_true")
    parser.add_argument("--rust-output", type=Path)
    parser.add_argument("--expected-output", type=Path)
    parser.add_argument("--json-output", type=Path)
    args = parser.parse_args()

    if args.prepare_source:
        args.source.parent.mkdir(parents=True, exist_ok=True)
        make_source_fixture(args.source)
        print(json.dumps({"source_fixture": str(args.source), "created": True}, ensure_ascii=False))
        return 0
    if args.rust_output is None or args.expected_output is None:
        parser.error("--rust-output and --expected-output are required unless --prepare-source is set")

    add_source_project()
    from PyCode.core.feature_1_6_extract_sheets import (
        _copy_full_sheet_with_style,
        _remove_default_blank,
    )

    before = hashlib.sha256(args.source.read_bytes()).hexdigest()
    source_workbook = load_workbook(args.source, data_only=True, read_only=False)
    expected_workbook = Workbook()
    source_sheet = source_workbook.worksheets[0]
    target_sheet = expected_workbook.create_sheet(title=source_sheet.title)
    copied = _copy_full_sheet_with_style(source_sheet, target_sheet)
    if not copied:
        raise SystemExit("source helper unexpectedly skipped the worksheet")
    _remove_default_blank(expected_workbook)
    expected_workbook.save(args.expected_output)
    source_workbook.close()
    expected_workbook.close()
    after = hashlib.sha256(args.source.read_bytes()).hexdigest()

    # Compare against the source worksheet properties that the legacy helper
    # copies. The helper's own output is also generated for inspection, but
    # openpyxl currently reports an out-of-range style index for this styled
    # cross-workbook copy, so it cannot be used as a valid style oracle.
    reference = _snapshot(args.source)
    actual = _snapshot(args.rust_output)
    legacy_style_readable = True
    legacy_style_error = None
    try:
        legacy_output = load_workbook(args.expected_output, data_only=True, read_only=False)
        for worksheet in legacy_output.worksheets:
            for row in worksheet.iter_rows():
                for cell in row:
                    _ = cell.font
        legacy_output.close()
    except Exception as error:  # record source behavior without hiding candidate differences
        legacy_style_readable = False
        legacy_style_error = f"{type(error).__name__}: {error}"
    differences = []
    for sheet_name in sorted(set(reference) | set(actual)):
        if sheet_name not in reference or sheet_name not in actual:
            differences.append({"sheet": sheet_name, "kind": "sheet_presence"})
            continue
        left = reference[sheet_name]
        right = actual[sheet_name]
        for key in ("max_row", "max_column", "cells", "merges", "row_dimensions", "column_dimensions", "comment_cells"):
            if left[key] != right[key]:
                differences.append({"sheet": sheet_name, "kind": key})

    report = {
        "source": args.source.name,
        "python_legacy_output": args.expected_output.name,
        "rust_output": args.rust_output.name,
        "source_sha256_before": before,
        "source_sha256_after": after,
        "source_unchanged": before == after,
        "sheets_compared": len(set(reference) | set(actual)),
        "difference_count": len(differences),
        "differences": differences,
        "legacy_helper_output_style_readable": legacy_style_readable,
        "legacy_helper_output_style_error": legacy_style_error,
        "comparison_oracle": "source worksheet properties copied by the legacy helper; its generated styled cross-workbook output is separately checked and is not used as the style oracle when unreadable",
        "compared_properties": [
            "cell values read with data_only=True",
            "cell number formats, font, fill, alignment and protection",
            "merged ranges",
            "row height, hidden state and outline level",
            "column width, hidden state and outline level",
            "comment cell positions",
        ],
        "limitations": [
            "synthetic XLSX only",
            "direct Umya candidate output; does not register a production capability",
            "does not validate every OOXML package preservation item",
        ],
    }
    encoded = json.dumps(report, ensure_ascii=False, indent=2)
    print(encoded)
    if args.json_output:
        args.json_output.parent.mkdir(parents=True, exist_ok=True)
        args.json_output.write_text(encoded + "\n", encoding="utf-8")
    return 0 if report["source_unchanged"] and not differences else 1


if __name__ == "__main__":
    raise SystemExit(main())
