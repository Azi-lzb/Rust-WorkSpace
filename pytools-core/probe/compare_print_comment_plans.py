"""Compare Rust print-plan structure with 3-1/3-2 Python workbook outputs."""

from __future__ import annotations

import hashlib
import json
import logging
import math
import os
import shlex
import shutil
import subprocess
import sys
import tempfile
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Border, Font, PatternFill, Side
from openpyxl.utils import column_index_from_string, get_column_letter

sys.dont_write_bytecode = True
WORKSPACE = Path(__file__).resolve().parents[2]
PYTOOLS_SOURCE = Path(r"C:\Users\AZI\Desktop\VibeCode\4.pytools-main")
EXPECTED_RANGES = ["B2:D4", "F2:G3"]
PLATFORM = "synthetic-pt-3-1-3-2-comment-print-probe"
REPORT_DATE = os.environ.get("PYTOOLS_REPORT_DATE", date.today().strftime("%Y%m%d"))
if len(REPORT_DATE) != 8 or not REPORT_DATE.isdigit():
    raise ValueError("PYTOOLS_REPORT_DATE must be an eight-digit YYYYMMDD value")


def wsl_path(path: Path) -> str:
    return f"/mnt/{path.drive.rstrip(':').lower()}/{path.as_posix()[2:].lstrip('/')}"


def make_source(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Marked"
    ws["A1"] = "Report"
    ws["A1"].comment = Comment("打印输出 横向", "probe")
    for row, values in enumerate(((0, "A", "B"), (2, "C", "D"), (3, "E", "F")), start=2):
        for column, value in enumerate(values, start=2):
            ws.cell(row, column, value)
    ws["B2"].comment = Comment("打印区域 1", "probe")
    ws["D4"].comment = Comment("打印区域#1", "probe")
    ws["F2"] = "second"
    ws["G3"] = "block"
    ws["F2"].comment = Comment("打印区域 2", "probe")
    ws["G3"].comment = Comment("打印区域#2", "probe")
    ws["B2"].font = Font(bold=True)
    ws["C3"].fill = PatternFill("solid", fgColor="FFCC00")
    ws["D4"].border = Border(bottom=Side(style="thin"))
    ws.merge_cells("C3:D3")
    ws.row_dimensions[2].height = 22
    ws.column_dimensions["B"].width = 14
    plain = wb.create_sheet("Plain")
    plain["A1"] = "not printed"
    wb.save(path)
    wb.close()


def rust_plan(
    source: Path,
    output: Path,
    drop_comments: bool,
    hide_zero: bool,
    capability_path: Path | None = None,
) -> dict[str, object]:
    arguments = [
        "print-comment-plan",
        "--input",
        str(source),
        "--output",
        str(output),
        "--drop-comments",
        "true" if drop_comments else "false",
        "--hide-zero",
        "true" if hide_zero else "false",
    ]
    if capability_path is not None:
        arguments.extend(["--capabilities", str(capability_path), "--platform", PLATFORM])
    native_cli = os.environ.get("PYTOOLS_RUST_CLI")
    if native_cli:
        result = subprocess.run([native_cli, *arguments], capture_output=True, check=False)
    else:
        workspace = wsl_path(WORKSPACE)
        wsl_arguments = arguments.copy()
        wsl_arguments[2] = wsl_path(source)
        wsl_arguments[4] = wsl_path(output)
        if capability_path is not None:
            wsl_arguments[-3] = wsl_path(capability_path)
        command = (
            f"cd {shlex.quote(workspace)} && cargo run --quiet -p pytools-core --bin pytools -- "
            + " ".join(shlex.quote(argument) for argument in wsl_arguments)
        )
        result = subprocess.run(["wsl", "-e", "sh", "-lc", command], capture_output=True, check=False)
    stdout = result.stdout.decode("utf-8-sig", errors="replace").replace("\x00", "")
    if result.returncode != 0:
        raise RuntimeError(f"Rust print plan failed: {stdout}\n{result.stderr.decode('utf-8', errors='replace')}")
    for index, char in enumerate(stdout):
        if char == "{":
            try:
                return json.loads(stdout[index:])
            except json.JSONDecodeError:
                continue
    raise RuntimeError(f"Rust print plan emitted no JSON: {stdout}")


def python_output(source: Path, output_dir: Path, fast: bool) -> Path:
    sys.path.insert(0, str(PYTOOLS_SOURCE / "PyCode"))
    from core.print_310 import run_print_fast_by_comment, run_print_keep_by_comment

    log_dir = output_dir / "logs"
    output_dir.mkdir(parents=True)
    with __import__("contextlib").redirect_stdout(None):
        if fast:
            stats = run_print_fast_by_comment([source], output_dir, log_dir, hide_zero=True)
        else:
            stats = run_print_keep_by_comment([source], output_dir, log_dir, hide_zero=False)
    if stats != {"workbooks_hit": 1, "sheets_hit": 2, "saved_files": 1, "skipped": 0}:
        raise AssertionError(f"unexpected Python print stats: {stats}")
    outputs = list(output_dir.glob("*.xlsx"))
    if len(outputs) != 1:
        raise AssertionError(f"expected one Python output, got {outputs}")
    return outputs[0]


def normalized_column_dimensions(ws: object) -> dict[str, dict[str, object]]:
    dimensions = {}
    for key, dimension in ws.column_dimensions.items():
        first = dimension.min or column_index_from_string(key)
        last = dimension.max or first
        if dimension.width is None and not dimension.hidden:
            continue
        for column in range(first, last + 1):
            dimensions[get_column_letter(column)] = {
                "width": dimension.width,
                "hidden": dimension.hidden,
            }
    return dimensions


def workbook_projection(path: Path) -> list[dict[str, object]]:
    wb = load_workbook(path, data_only=False)
    result = []
    for ws in wb.worksheets:
        cells = {}
        comments = {}
        for row in ws.iter_rows():
            for cell in row:
                if cell.comment is not None:
                    comments[cell.coordinate] = cell.comment.text
                if cell.value is None and cell.comment is None and not cell.has_style:
                    continue
                cells[cell.coordinate] = {
                    "value": cell.value,
                    "type": cell.data_type,
                    "bold": bool(cell.font.bold),
                    "italic": bool(cell.font.italic),
                    "font_name": cell.font.name,
                    "font_size": cell.font.sz,
                    "fill_type": cell.fill.fill_type,
                    "fill_rgb": cell.fill.fgColor.rgb,
                    "border_top_style": None
                    if cell.border.top is None
                    else cell.border.top.style,
                    "border_bottom_style": None
                    if cell.border.bottom is None
                    else cell.border.bottom.style,
                    "alignment_horizontal": cell.alignment.horizontal,
                    "alignment_vertical": cell.alignment.vertical,
                    "alignment_wrap_text": cell.alignment.wrap_text,
                    "number_format": cell.number_format,
                    "protection_locked": cell.protection.locked,
                    "comment": None
                    if cell.comment is None
                    else {"text": cell.comment.text, "author": cell.comment.author},
                }
        page_setup_properties = ws.sheet_properties.pageSetUpPr
        result.append(
            {
                "name": ws.title,
                "max_row": ws.max_row,
                "max_column": ws.max_column,
                "print_area": str(ws.print_area),
                "orientation": ws.page_setup.orientation,
                "fit_to_width": ws.page_setup.fitToWidth,
                "fit_to_height": ws.page_setup.fitToHeight,
                "paper_size": ws.page_setup.paperSize,
                "footer": ws.oddFooter.center.text,
                "comments": comments,
                "zero_cell": ws["A1"].value,
                "cells": cells,
                "merges": sorted(str(item) for item in ws.merged_cells.ranges),
                "row_heights": {
                    str(index): dimension.height
                    for index, dimension in ws.row_dimensions.items()
                    if dimension.height is not None
                },
                "column_dimensions": normalized_column_dimensions(ws),
                "page": {
                    "orientation": ws.page_setup.orientation,
                    "fit_to_page": None
                    if page_setup_properties is None
                    else page_setup_properties.fitToPage,
                    "fit_to_width": ws.page_setup.fitToWidth,
                    "fit_to_height": ws.page_setup.fitToHeight,
                    "scale": ws.page_setup.scale,
                    "paper_size": ws.page_setup.paperSize,
                    "print_area": str(ws.print_area),
                    "center_horizontal": ws.print_options.horizontalCentered,
                    "center_vertical": ws.print_options.verticalCentered,
                    "footer_center": ws.oddFooter.center.text,
                },
            }
        )
    wb.close()
    return result


def differences(left: object, right: object, path: str = "root") -> list[str]:
    if isinstance(left, bool) or isinstance(right, bool):
        return [] if left == right and type(left) is type(right) else [path]
    if isinstance(left, (int, float)) and isinstance(right, (int, float)):
        return [] if math.isclose(left, right, rel_tol=1e-15, abs_tol=1e-15) else [path]
    if type(left) is not type(right):
        return [path]
    if isinstance(left, dict):
        if left.keys() != right.keys():
            return [path]
        return [
            item
            for key in left
            for item in differences(left[key], right[key], f"{path}.{key}")
        ]
    if isinstance(left, list):
        if len(left) != len(right):
            return [path]
        return [
            item
            for index, (left_item, right_item) in enumerate(zip(left, right))
            for item in differences(left_item, right_item, f"{path}[{index}]")
        ]
    return [] if left == right else [path]


def preserve_mismatch_artifacts(name: str, python_path: Path, rust_path: Path) -> None:
    if not os.environ.get("PYTOOLS_KEEP_PRINT_PROBE_ARTIFACTS"):
        return
    directory = Path(tempfile.mkdtemp(prefix="pytools-print-comment-mismatch-"))
    shutil.copy2(python_path, directory / f"{name}_python.xlsx")
    shutil.copy2(rust_path, directory / f"{name}_rust.xlsx")
    print(f"print probe mismatch artifacts: {directory}", file=sys.stderr)


def capability_for(plan: dict[str, object]) -> dict[str, object]:
    return {
        "engine": "umya_spreadsheet",
        "status": "passed",
        "format": plan["format"],
        "operations": sorted({operation["operation"] for operation in plan["operations"]}),
        "supported_features": plan["workbook"]["features_present"],
        "preserved": plan["preserve"],
        "platform": PLATFORM,
        "evidence": {
            "sample_id": "synthetic-pt-3-1-3-2-comment-print",
            "report_path": "pytools-core/probe/compare_print_comment_plans.py",
            "adapter_version": "umya-spreadsheet-3.1.1",
            "engine_version": "umya-spreadsheet-3.1.1",
            "platform": PLATFORM,
        },
    }


def verify_plan(plan: dict[str, object], expected_sheets: list[dict[str, object]], drop_comments: bool, hide_zero: bool) -> None:
    if plan.get("execution") != "plan_only":
        raise AssertionError(f"print plan unexpectedly executed: {plan.get('execution')}")
    if plan.get("matching_sheets") != 1 or plan.get("marked_ranges") != 2:
        raise AssertionError(f"marked range summary differs: {plan}")
    body = plan.get("plan")
    if not isinstance(body, dict):
        raise AssertionError("Rust did not return an edit plan")
    operations = body.get("operations", [])
    imports = [operation for operation in operations if operation.get("operation") == "import_range"]
    if [operation.get("source_range") for operation in imports] != EXPECTED_RANGES:
        raise AssertionError(f"Rust marked ranges differ: {[operation.get('source_range') for operation in imports]}")
    widths = [operation for operation in operations if operation.get("operation") == "set_column_width"]
    expected_widths = (
        [("Marked_1", column, 13.0) for column in range(1, 4)]
        + [("Marked_2", column, 13.0) for column in range(1, 3)]
        if not drop_comments
        else []
    )
    actual_widths = [
        (operation.get("sheet"), operation.get("column_index"), operation.get("width"))
        for operation in widths
    ]
    if actual_widths != expected_widths:
        raise AssertionError(f"Rust default-width plan differs: {actual_widths}")
    if [operation.get("target_sheet") for operation in imports] != [sheet["name"] for sheet in expected_sheets]:
        raise AssertionError("Rust target sheet order/names differ from Python")
    for operation in imports:
        if operation.get("target_start") != "A1" or not operation.get("copy_values") or operation.get("copy_formulas"):
            raise AssertionError(f"unexpected range import semantics: {operation}")
        if operation.get("style_mode") != ("basic" if drop_comments else "full"):
            raise AssertionError(f"unexpected style mode: {operation}")
        if operation.get("copy_comments") is not (not drop_comments):
            raise AssertionError(f"unexpected comment policy: {operation}")
        if operation.get("copy_merges") is not True:
            raise AssertionError("merge preservation is missing")
        if operation.get("hide_zero") is not hide_zero:
            raise AssertionError("hide-zero policy differs")
    pages = {}
    for operation in operations:
        if operation.get("operation") == "set_page_settings":
            page = pages.setdefault(operation["sheet"], {})
            for key, value in operation["settings"].items():
                if value is not None and (key != "clear_scale" or value):
                    page[key] = value
    for sheet in expected_sheets:
        settings = pages.get(sheet["name"], {})
        expected = {
            "orientation": "landscape",
            "fit_to_page": True,
            "fit_to_width": 1,
            "fit_to_height": 0,
            "clear_scale": True,
            "paper_size": 9,
            "print_area": sheet["print_area"].split("!")[-1].replace("$", ""),
            "horizontal_centered": False,
            "vertical_centered": False,
            "footer": sheet["footer"],
        }
        if settings != expected:
            raise AssertionError(f"{sheet['name']} print settings differ: expected={expected}, rust={settings}")


def main() -> int:
    if not PYTOOLS_SOURCE.is_dir():
        raise FileNotFoundError(PYTOOLS_SOURCE)
    with tempfile.TemporaryDirectory(prefix="pytools-print-comment-", ignore_cleanup_errors=True) as temp:
        root = Path(temp)
        source = root / "source.xlsx"
        make_source(source)
        source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
        python_keep = python_output(source, root / "python_keep", fast=False)
        python_fast = python_output(source, root / "python_fast", fast=True)
        keep_projection = workbook_projection(python_keep)
        fast_projection = workbook_projection(python_fast)
        if [sheet["name"] for sheet in keep_projection] != ["Marked_1", "Marked_2"]:
            raise AssertionError(f"unexpected Python output names: {keep_projection}")
        if [sheet["print_area"] for sheet in keep_projection] != ["'Marked_1'!$A$1:$C$3", "'Marked_2'!$A$1:$B$2"]:
            raise AssertionError(f"unexpected Python print areas: {keep_projection}")
        keep = rust_plan(source, root / "rust_keep.xlsx", drop_comments=False, hide_zero=False)
        fast = rust_plan(source, root / "rust_fast.xlsx", drop_comments=True, hide_zero=True)
        verify_plan(keep, keep_projection, drop_comments=False, hide_zero=False)
        verify_plan(fast, fast_projection, drop_comments=True, hide_zero=True)
        if len(keep_projection[0]["comments"]) == 0 or any(sheet["comments"] for sheet in fast_projection):
            raise AssertionError("Python keep/fast comment behavior was not represented by the fixture")
        if fast_projection[0]["zero_cell"] is not None:
            raise AssertionError("Python fast print did not hide the source zero value")
        if hashlib.sha256(source.read_bytes()).hexdigest() != source_hash:
            raise AssertionError("print plan probe changed source workbook bytes")
        report = {
            "status": "passed_plan_semantics_synthetic",
            "scope": ["3-1 keep-format print plan", "3-2 fast print plan"],
            "python": {
                "keep": {"sheets": len(keep_projection), "ranges": EXPECTED_RANGES, "comments_copied": True},
                "fast": {"sheets": len(fast_projection), "ranges": EXPECTED_RANGES, "comments_copied": False, "zero_hidden": True},
            },
            "rust_plan": {"keep_imports": len([op for op in keep["plan"]["operations"] if op.get("operation") == "import_range"]), "fast_imports": len([op for op in fast["plan"]["operations"] if op.get("operation") == "import_range"])},
            "comparison": ["marked ranges and order", "output sheet names", "style/comment/merge/zero policies", "orientation, fit settings, paper size, print area and footer"],
            "limitations": ["Rust output was not applied: ImportRange and page-setting capabilities remain gated by registered ExcelModifier evidence.", "Synthetic XLSX only; no real print template or XSLM was used.", "This validates plan semantics, not backend output fidelity or print rendering.", "Windows release build, Kylin/UOS and deb were not validated."],
        }
        report_path = WORKSPACE / "pytools-core" / "probe" / "reports" / f"print_comment_plan_parity_{REPORT_DATE}.json"
        report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

        executions = {}
        for name, response, python_path, drop_comments, hide_zero in (
            ("keep", keep, python_keep, False, False),
            ("fast", fast, python_fast, True, True),
        ):
            rust_path = root / f"rust_{name}.xlsx"
            capability_path = root / f"rust_{name}.capability.json"
            plan = response["plan"]
            capability_path.write_text(
                json.dumps(capability_for(plan), ensure_ascii=False, indent=2),
                encoding="utf-8",
            )
            executed = rust_plan(
                source,
                rust_path,
                drop_comments=drop_comments,
                hide_zero=hide_zero,
                capability_path=capability_path,
            )
            if executed["execution"].get("engine") != "umya_spreadsheet":
                raise AssertionError(f"{name} print did not execute through Umya: {executed}")
            python_result = workbook_projection(python_path)
            rust_result = workbook_projection(rust_path)
            mismatch_paths = differences(python_result, rust_result)
            if mismatch_paths:
                preserve_mismatch_artifacts(name, python_path, rust_path)
                raise AssertionError(f"Python/Rust {name} print differs at {mismatch_paths}")
            executions[name] = {
                "engine": executed["execution"]["engine"],
                "sheet_count": len(rust_result),
                "projection": rust_result,
            }
        if hashlib.sha256(source.read_bytes()).hexdigest() != source_hash:
            raise AssertionError("print execution changed source workbook bytes")
        execution_report = {
            "status": "passed_synthetic_python_rust_execution_parity",
            "scope": ["PT-3-1 keep-format print", "PT-3-2 fast print"],
            "sample": "synthetic XLSX only",
            "engine": "ExcelModifier -> UmyaSpreadsheet using temporary test-only capability",
            "float_comparison": "recursive relative and absolute tolerance 1e-15, per rules/13 §9.1",
            "column_dimension_comparison": "expand OOXML min/max spans into per-column effective width and hidden state",
            "source_unchanged": True,
            "comparison": [
                "worksheet order, identity, dimensions and values/types",
                "font, fill, border, alignment, number format and protection",
                "comments, merged ranges, row/column dimensions",
                "orientation, fit settings, scale, paper size, print area, centering and footer",
            ],
            "modes": executions,
            "limitations": [
                "Capability evidence was temporary and was not registered for production routing.",
                "This evaluates only the Umya modifier candidate for a newly generated workbook; it does not compare or select the rust_xlsxwriter new-workbook path.",
                "No real PyTools print template, Windows native Rust build, Kylin/UOS, glibc 2.28 or offline deb was validated.",
                "No PDF rendering was performed; PDF export remains out of scope.",
            ],
        }
        execution_report_path = WORKSPACE / "pytools-core" / "probe" / "reports" / f"print_comment_execution_parity_{REPORT_DATE}.json"
        execution_report_path.write_text(
            json.dumps(execution_report, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )
        logging.shutdown()
        print(
            json.dumps(
                {
                    "plan_report": str(report_path),
                    "execution_report": str(execution_report_path),
                    "status": execution_report["status"],
                    "modes": sorted(executions),
                },
                ensure_ascii=False,
            )
        )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
