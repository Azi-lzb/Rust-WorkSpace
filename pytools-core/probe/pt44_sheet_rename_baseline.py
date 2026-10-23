#!/usr/bin/env python3
"""Run the read-only PyCode PT-4-4 implementation on an isolated fixture copy."""

from __future__ import annotations

import argparse
import hashlib
import json
import logging
import shutil
import sys
import tempfile
from pathlib import Path

sys.dont_write_bytecode = True


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def projection(path: Path) -> dict[str, object]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, data_only=False, keep_links=True)
    try:
        first = workbook["RenamedA"]
        return {
            "sheet_names": workbook.sheetnames,
            "a1_values": {sheet.title: sheet["A1"].value for sheet in workbook.worksheets},
            "formula": workbook["A"]["B1"].value,
            "defined_names": [
                {"name": name.name, "reference": name.attr_text}
                for name in workbook.defined_names.values()
            ],
            "renamed_a1": {
                "fill": first["A1"].fill.fgColor.rgb,
                "bold": first["A1"].font.bold,
                "comment": first["A1"].comment.text if first["A1"].comment else None,
                "merged_ranges": sorted(str(item) for item in first.merged_cells.ranges),
                "row_1_height": first.row_dimensions[1].height,
                "column_b_width": first.column_dimensions["B"].width,
                "orientation": first.page_setup.orientation,
                "print_area": first.print_area,
            },
        }
    finally:
        workbook.close()


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source-root", type=Path, required=True)
    parser.add_argument("--workspace-root", type=Path, default=Path(__file__).resolve().parents[1])
    parser.add_argument("--report", type=Path, required=True)
    parser.add_argument("--expected-xlsx", type=Path)
    args = parser.parse_args()

    source = args.source_root / "PyCode" / "core" / "convert_33_34_35_37.py"
    sys.path.insert(0, str(args.source_root))
    from PyCode.core.convert_33_34_35_37 import run_batch_rename_sheet_37

    fixture_dir = args.workspace_root / "tests" / "fixtures" / "sheet_rename"
    config = fixture_dir / "config.xlsx"
    source_fixture = fixture_dir / "input.xlsx"
    source_fixture_before = sha256(source_fixture)
    with tempfile.TemporaryDirectory(prefix="pytools-pt44-baseline-") as temporary:
        root = Path(temporary)
        input_copy = root / "input.xlsx"
        input_copy.write_bytes(source_fixture.read_bytes())
        input_before = sha256(input_copy)
        config_before = sha256(config)
        result = run_batch_rename_sheet_37(config, [input_copy], root / "logs")
        output_projection = projection(input_copy)
        logger = logging.getLogger("PyCode.3_7_batch_rename_sheet")
        for handler in logger.handlers[:]:
            handler.flush()
            handler.close()
            logger.removeHandler(handler)
        expected_sha256 = None
        if args.expected_xlsx is not None:
            args.expected_xlsx.parent.mkdir(parents=True, exist_ok=True)
            shutil.copyfile(input_copy, args.expected_xlsx)
            expected_sha256 = sha256(args.expected_xlsx)
        report = {
            "suite": "pytools.pt_4_4_sheet_rename_python_baseline",
            "source_function": "PyCode.core.convert_33_34_35_37.run_batch_rename_sheet_37",
            "source_sha256": sha256(source),
            "python": sys.version.split()[0],
            "openpyxl": __import__("openpyxl").__version__,
            "pandas": __import__("pandas").__version__,
            "fixture_sha256": {
                "config.xlsx": sha256(config),
                "input.xlsx": sha256(source_fixture),
            },
            "result": result,
            "projection": output_projection,
            "expected_xlsx_sha256": expected_sha256,
            "safety": {
                "baseline_copy_changed": sha256(input_copy) != input_before,
                "source_fixture_unchanged": sha256(source_fixture) == source_fixture_before,
                "config_unchanged": sha256(config) == config_before,
            },
            "scope_limit": (
                "Synthetic XLSX only. The source normal branch supports XLSX/XLSM; this "
                "workspace candidate currently accepts only XLSX and needs real-template, "
                "XLSM/VBA, target-platform and deb verification."
            ),
        }
    args.report.parent.mkdir(parents=True, exist_ok=True)
    args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
