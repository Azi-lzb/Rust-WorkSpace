"""Read-only precheck parity using a temporary copy of the real pytools config."""

from __future__ import annotations

import contextlib
import hashlib
import io
import json
import re
import shlex
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

sys.dont_write_bytecode = True
WORKSPACE = Path(__file__).resolve().parents[2]
PYTOOLS_SOURCE = Path(r"C:\Users\AZI\Desktop\VibeCode\4.pytools-main")
SOURCE_CONFIG = PYTOOLS_SOURCE / "PyCode" / "config" / "config.xlsx"


def wsl_path(path: Path) -> str:
    return f"/mnt/{path.drive.rstrip(':').lower()}/{path.as_posix()[2:].lstrip('/')}"


def copy_with_safe_directories(destination: Path, root: Path) -> None:
    wb = load_workbook(SOURCE_CONFIG)
    ws = wb["全局配置"]
    safe_values = {
        "输入目录": "input",
        "输出目录": "output",
        "日志目录": "logs",
    }
    for row in range(2, ws.max_row + 1):
        key = str(ws.cell(row, 1).value or "").strip()
        if key in safe_values:
            ws.cell(row, 2).value = safe_values[key]
    wb.save(destination)
    wb.close()


def python_counts(config: Path) -> tuple[int, int]:
    sys.path.insert(0, str(PYTOOLS_SOURCE / "PyCode"))
    from core.precheck_399 import run_precheck

    output = io.StringIO()
    with contextlib.redirect_stdout(output):
        run_precheck(config)
    match = re.search(r"=== 预校验完成: 错误 (\d+) 项，警告 (\d+) 项 ===", output.getvalue())
    if not match:
        raise AssertionError(f"legacy precheck summary missing; exit text captured but not retained")
    return int(match.group(1)), int(match.group(2))


def rust_report(config: Path) -> dict[str, object]:
    workspace_wsl = wsl_path(WORKSPACE)
    config_wsl = wsl_path(config)
    command = (
        f"cd {shlex.quote(workspace_wsl)} && cargo run --quiet -p pytools-core --bin pytools -- "
        f"timeline-precheck --config {shlex.quote(config_wsl)}"
    )
    result = subprocess.run(["wsl", "-e", "sh", "-lc", command], capture_output=True, check=False)
    stdout = result.stdout.decode("utf-8-sig", errors="replace").replace("\x00", "")
    for index, character in enumerate(stdout):
        if character != "{":
            continue
        try:
            payload = json.loads(stdout[index:])
            return payload
        except (json.JSONDecodeError, KeyError, TypeError):
            continue
    raise RuntimeError("Rust precheck emitted no JSON report")


def main() -> int:
    if not SOURCE_CONFIG.is_file():
        raise FileNotFoundError(SOURCE_CONFIG)
    source_hash = hashlib.sha256(SOURCE_CONFIG.read_bytes()).hexdigest()
    with tempfile.TemporaryDirectory(prefix="pytools-real-config-precheck-", ignore_cleanup_errors=True) as temp:
        root = Path(temp)
        (root / "PyCode").mkdir()
        for name in ("input", "output", "logs"):
            (root / name).mkdir()
        config = root / "PyCode" / "config.xlsx"
        copy_with_safe_directories(config, root)
        expected = python_counts(config)
        actual_report = rust_report(config)
        actual = (int(actual_report["errors"]), int(actual_report["warnings"]))
        if expected != actual:
            safe_issues = [
                {
                    "level": issue.get("level"),
                    "location": str(issue.get("location", "")).replace(str(root), "<temporary-root>"),
                    "message": str(issue.get("message", "")).replace(str(root), "<temporary-root>"),
                }
                for issue in actual_report.get("issues", [])
            ]
            raise AssertionError(
                f"real configuration precheck count mismatch: Python={expected}, Rust={actual}, issues={safe_issues}"
            )
        print(json.dumps({
            "status": "passed",
            "source_configuration_sha256": source_hash,
            "python_errors": expected[0],
            "python_warnings": expected[1],
            "rust_errors": actual[0],
            "rust_warnings": actual[1],
            "source_configuration_modified": False,
            "configured_input_directories_replaced_in_temporary_copy": True
        }, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
