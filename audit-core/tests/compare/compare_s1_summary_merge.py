#!/usr/bin/env python3
"""Compare the Rust S1-F02 table-merge CLI with the read-only Python business flow."""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import subprocess
import sys
import tempfile
import time
from datetime import date, datetime
from pathlib import Path
from typing import Any

sys.dont_write_bytecode = True
SOURCE_CORE = Path(r"C:\Users\AZI\Desktop\VibeCode\3.open-data-audit\core")
sys.path.insert(0, str(SOURCE_CORE / "src"))

from base_audit.name_config import (  # noqa: E402
    WORKBOOK_TABLE_MERGE_FUNCTION,
    FeatureMapping,
)
from base_audit.region_summary import run_region_summaries  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def create_fixture(input_dir: Path, template_path: Path) -> list[Path]:
    nested = input_dir / "机构报送"
    nested.mkdir(parents=True)
    template = Workbook()
    template.active.title = "模板"
    template.save(template_path)

    def write(path: Path, sheets: dict[str, list[list[Any]]]) -> None:
        workbook = Workbook()
        workbook.remove(workbook.active)
        for sheet_name, rows in sheets.items():
            sheet = workbook.create_sheet(sheet_name)
            for row in rows:
                sheet.append(row)
        workbook.save(path)

    root_summary = input_dir / "上一版汇总.xlsx"
    write(root_summary, {"Data": [["旧表"], ["不应进入"]]})
    alpha = nested / "!Alpha_20260731.xlsx"
    write(
        alpha,
        {
            "Data": [["单位", "数值"], ["甲", 10], ["乙", 20]],
            "Stats": [["指标", "数值"], ["总数", 2]],
            "Empty": [],
        },
    )
    beta = nested / "Beta_20260731.xlsx"
    write(beta, {"Data": [["单位", "数值"], ["丙", 30], [None, None]]})
    gamma = nested / "Gamma_20260731.xlsx"
    write(gamma, {"Different": [["地区", "数量"], ["东", 5]]})
    return [root_summary, alpha, beta, gamma]


def canonical(value: Any) -> Any:
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, float) and math.isfinite(value):
        return value
    return value


def equal_value(left: Any, right: Any) -> bool:
    left, right = canonical(left), canonical(right)
    if isinstance(left, bool) or isinstance(right, bool):
        return left is right
    if isinstance(left, (int, float)) and isinstance(right, (int, float)):
        return left == right or abs(left - right) <= max(abs(left), abs(right), 1.0) * 1e-15
    return left == right


def compare_workbooks(python_path: Path, rust_path: Path) -> dict[str, Any]:
    python = load_workbook(python_path, data_only=False, read_only=False)
    rust = load_workbook(rust_path, data_only=False, read_only=False)
    try:
        sheet_name_mismatches = python.sheetnames != rust.sheetnames
        value_mismatches = 0
        style_mismatches = 0
        filter_mismatches = 0
        freeze_mismatches = 0
        dimension_mismatches = 0
        differing_cells: list[dict[str, Any]] = []
        style_details: list[dict[str, Any]] = []
        freeze_details: list[dict[str, Any]] = []
        width_details: list[dict[str, Any]] = []
        width_differences: list[float] = []
        for name in python.sheetnames:
            if name not in rust.sheetnames:
                continue
            expected, actual = python[name], rust[name]
            if (expected.max_row, expected.max_column) != (actual.max_row, actual.max_column):
                dimension_mismatches += 1
            rows = max(expected.max_row, actual.max_row)
            columns = max(expected.max_column, actual.max_column)
            for row in range(1, rows + 1):
                for column in range(1, columns + 1):
                    left, right = expected.cell(row, column), actual.cell(row, column)
                    if not equal_value(left.value, right.value):
                        value_mismatches += 1
                        differing_cells.append({
                            "sheet": name,
                            "cell": left.coordinate,
                            "python": repr(left.value),
                            "rust": repr(right.value),
                        })
                    if left.number_format != right.number_format:
                        style_mismatches += 1
                        style_details.append({
                            "sheet": name,
                            "cell": left.coordinate,
                            "python_number_format": left.number_format,
                            "rust_number_format": right.number_format,
                        })
                    if row == 1 and (
                        left.font.bold != right.font.bold
                        or left.fill.fill_type != right.fill.fill_type
                        or left.fill.fgColor.rgb != right.fill.fgColor.rgb
                    ):
                        style_mismatches += 1
                        style_details.append({
                            "sheet": name,
                            "cell": left.coordinate,
                            "python_header": [left.font.bold, left.fill.fill_type, left.fill.fgColor.rgb],
                            "rust_header": [right.font.bold, right.fill.fill_type, right.fill.fgColor.rgb],
                        })
            if expected.auto_filter.ref != actual.auto_filter.ref:
                filter_mismatches += 1
            if expected.freeze_panes != actual.freeze_panes:
                freeze_mismatches += 1
                freeze_details.append({
                    "sheet": name,
                    "python": str(expected.freeze_panes),
                    "rust": str(actual.freeze_panes),
                })
            for column in range(1, columns + 1):
                letter = expected.cell(1, column).column_letter
                left = expected.column_dimensions[letter].width
                right = actual.column_dimensions[letter].width
                width_differences.append(abs(left - right))
                width_details.append({
                    "sheet": name,
                    "column": letter,
                    "python": left,
                    "rust": right,
                    "delta": abs(left - right),
                })
        return {
            "status": "passed" if not any((
                sheet_name_mismatches,
                value_mismatches,
                style_mismatches,
                filter_mismatches,
                freeze_mismatches,
                dimension_mismatches,
            )) else "failed",
            "sheet_name_mismatches": int(sheet_name_mismatches),
            "python_sheet_names": python.sheetnames,
            "rust_sheet_names": rust.sheetnames,
            "dimension_mismatches": dimension_mismatches,
            "value_mismatches": value_mismatches,
            "style_mismatches": style_mismatches,
            "filter_mismatches": filter_mismatches,
            "freeze_mismatches": freeze_mismatches,
            "column_width_max_delta": max(width_differences, default=0.0),
            "column_widths_exact": all(delta == 0 for delta in width_differences),
            "sheets": len(python.sheetnames),
            "differing_cells": differing_cells,
            "style_details": style_details,
            "freeze_details": freeze_details,
            "column_width_details": width_details,
        }
    finally:
        python.close()
        rust.close()


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--audit-binary", type=Path, required=True)
    parser.add_argument("--report", type=Path)
    args = parser.parse_args()
    binary = args.audit_binary.resolve(strict=True)

    with tempfile.TemporaryDirectory(prefix="audit-s1-summary-merge-") as temporary:
        root = Path(temporary)
        input_dir = root / "input"
        output_dir = root / "outputs"
        input_dir.mkdir()
        output_dir.mkdir()
        template_path = root / "template.xlsx"
        input_files = create_fixture(input_dir, template_path)
        before = {path: sha256(path) for path in input_files}
        feature = FeatureMapping(
            name="汇总表合并",
            feature_type=WORKBOOK_TABLE_MERGE_FUNCTION,
            range_names=(),
            workbook_limited=False,
            workbook_keyword="",
            remark="",
        )

        python_path = None
        python_started = time.perf_counter()
        progress: list[str] = []
        result = run_region_summaries(
            template_path=template_path,
            input_dir=input_dir,
            output_dir=output_dir / "python",
            config_path=template_path,
            summary_features=(feature,),
            on_step=progress.append,
            engine_preference="自动",
        )
        python_elapsed_ms = round((time.perf_counter() - python_started) * 1000)
        python_path = result.output_path

        rust_path = output_dir / "rust.xlsx"
        rust_started = time.perf_counter()
        completed = subprocess.run(
            [
                str(binary),
                "s1-merge-workbook-tables",
                "--input-dir",
                str(input_dir),
                "--output-workbook",
                str(rust_path),
                "--developer-view",
            ],
            check=False,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True,
            encoding="utf-8",
        )
        rust_elapsed_ms = round((time.perf_counter() - rust_started) * 1000)
        if completed.returncode:
            raise RuntimeError(f"Rust CLI failed: {completed.stderr.strip()}")
        rust_payload = json.loads(completed.stdout)
        rust_progress = [
            json.loads(line)
            for line in completed.stderr.splitlines()
            if line.strip().startswith("{")
        ]

        comparison = compare_workbooks(python_path, rust_path)
        if comparison["status"] == "failed":
            status = "failed"
        elif comparison["column_widths_exact"]:
            status = "passed"
        else:
            status = "passed_with_column_width_variance"
        report = {
            "feature": "S1-F02 workbook-table merge",
            "status": status,
            "fixture": "synthetic",
            "python_engine": next(
                (line.split("：", 1)[1] for line in progress if line.startswith("已连接表格引擎：")),
                "not_reported",
            ),
            "rust_reader": rust_payload["summary"]["reader"],
            "rust_writer": rust_payload["summary"]["writer"],
            "python_elapsed_ms": python_elapsed_ms,
            "rust_cli_elapsed_ms": rust_elapsed_ms,
            "rust_events": len(rust_progress),
            "source_workbooks_seen": rust_payload["summary"]["source_workbooks"],
            "source_sheets_seen": rust_payload["summary"]["source_sheets"],
            "data_rows": rust_payload["summary"]["data_rows"],
            "input_files_unchanged": all(sha256(path) == digest for path, digest in before.items()),
            "comparison": comparison,
            "scope": "formula-free synthetic XLSX; exact values, sheet order, dimensions, number formats, header style, filter and freeze panes; column widths differ by engine; no real templates or Linux runtime",
        }
        encoded = json.dumps(report, ensure_ascii=False, indent=2) + "\n"
        if args.report:
            args.report.parent.mkdir(parents=True, exist_ok=True)
            args.report.write_text(encoded, encoding="utf-8")
        print(encoded, end="")
        return 0 if report["status"].startswith("passed") and report["input_files_unchanged"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
