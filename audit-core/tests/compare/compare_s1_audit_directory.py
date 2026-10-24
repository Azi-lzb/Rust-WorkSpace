#!/usr/bin/env python3
"""Compare the Rust S1 issue chain with the read-only Python workflow stages."""

from __future__ import annotations

import argparse
import dataclasses
import hashlib
import json
import math
import ntpath
import os
import subprocess
import sys
import tempfile
import time
from collections import Counter
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any
from urllib.parse import unquote

sys.dont_write_bytecode = True


def wsl_path(path: Path) -> str:
    drive = path.drive.rstrip(":").lower()
    if len(drive) != 1:
        raise ValueError(f"无法转换为 WSL 挂载路径：{path.drive}")
    relative = path.as_posix().split("/", 1)[1]
    return f"/mnt/{drive}/{relative}"


def canonical(value: Any) -> Any:
    if isinstance(value, Decimal):
        return value
    if isinstance(value, bool) or value is None:
        return value
    if isinstance(value, (int, float)):
        return Decimal(str(value))
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, dict):
        return {str(key): canonical(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [canonical(item) for item in value]
    return value


def normalized_issue(issue: dict[str, Any]) -> dict[str, Any]:
    row = canonical(issue)
    for field in ("source_file", "audit_file"):
        value = row.get(field)
        if value:
            row[field] = Path(str(value)).name.casefold()
    return row


def compare_issues(expected: list[Any], actual: list[dict[str, Any]]) -> dict[str, Any]:
    expected_rows = {
        issue.issue_id: normalized_issue(dataclasses.asdict(issue))
        for issue in expected
    }
    actual_rows = {
        issue["issue_id"]: normalized_issue(issue)
        for issue in actual
    }
    if len(expected_rows) != len(expected) or len(actual_rows) != len(actual):
        raise ValueError("Issue identity is not unique; refusing to hide duplicate records")
    left, right = set(expected_rows), set(actual_rows)
    differences: Counter[str] = Counter()
    for issue_id in left & right:
        expected_item = expected_rows[issue_id]
        actual_item = actual_rows[issue_id]
        for field in sorted(set(expected_item) | set(actual_item)):
            if expected_item.get(field) != actual_item.get(field):
                differences[field] += 1
    return {
        "expected_count": len(expected),
        "actual_count": len(actual),
        "missing_in_rust": len(left - right),
        "missing_in_python": len(right - left),
        "field_difference_counts": dict(sorted(differences.items())),
    }


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def invoke_rust(args: argparse.Namespace, manifest: Path, source: Path) -> tuple[dict[str, Any], str, int]:
    binary = args.audit_binary.resolve(strict=True)
    if os.name == "nt" and binary.suffix.lower() != ".exe":
        command = [
            "wsl.exe",
            "--distribution",
            args.wsl_distribution,
            "--user",
            "root",
            "--exec",
            wsl_path(binary),
        ]
        input_arg = wsl_path(args.input_dir.resolve(strict=True))
        template_arg = wsl_path(args.template_workbook.resolve(strict=True))
        template_dir_arg = (
            wsl_path(args.template_dir.resolve(strict=True)) if args.template_dir else None
        )
        manifest_arg = wsl_path(manifest.resolve(strict=True))
        source_arg = source.relative_to(args.input_dir.resolve(strict=True)).as_posix()
        audit_arg = wsl_path(args.audit_workbook.resolve(strict=True))
        mapping_arg = wsl_path(args.mapping_config.resolve(strict=True)) if args.mapping_config else None
        result_arg = wsl_path(args.result_workbook.resolve()) if args.result_workbook else None
        history_result_arg = (
            wsl_path(args.history_result_workbook.resolve())
            if args.history_result_workbook
            else None
        )
    else:
        command = [str(binary)]
        input_arg = str(args.input_dir.resolve(strict=True))
        template_arg = str(args.template_workbook.resolve(strict=True))
        template_dir_arg = (
            str(args.template_dir.resolve(strict=True)) if args.template_dir else None
        )
        manifest_arg = str(manifest.resolve(strict=True))
        source_arg = source.relative_to(args.input_dir.resolve(strict=True)).as_posix()
        audit_arg = str(args.audit_workbook.resolve(strict=True))
        mapping_arg = str(args.mapping_config.resolve(strict=True)) if args.mapping_config else None
        result_arg = str(args.result_workbook.resolve()) if args.result_workbook else None
        history_result_arg = (
            str(args.history_result_workbook.resolve()) if args.history_result_workbook else None
        )

    command.extend(
        [
            "s1-audit-directory",
            "--input-dir",
            input_arg,
            "--selected-file",
            source_arg,
            "--formula-manifest",
            manifest_arg,
            "--period",
            args.period,
            "--batch-id",
            args.batch_id,
            "--audit-time",
            args.audit_time,
            "--developer-view",
        ]
    )
    if template_dir_arg:
        command.extend(["--template-dir", template_dir_arg])
    else:
        command.extend(["--template-workbook", template_arg])
    if mapping_arg:
        command.extend(["--mapping-config", mapping_arg])
    if result_arg:
        command.extend(["--result-workbook", result_arg])
    if history_result_arg:
        command.extend(["--history-result-workbook", history_result_arg])
    started = time.perf_counter()
    completed = subprocess.run(
        command,
        check=False,
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
    )
    elapsed_ms = round((time.perf_counter() - started) * 1000)
    if completed.returncode:
        raise RuntimeError(
            "Rust S1 issue-chain CLI failed: "
            + completed.stderr.strip().splitlines()[-1]
        )
    output = json.loads(completed.stdout)
    return output, completed.stderr, elapsed_ms


def normalize_hyperlink_target(value: str) -> str:
    target = value
    if target.lower().startswith("file://"):
        target = unquote(target[len("file://") :])
        if target.startswith("/") and len(target) >= 3 and target[2] == ":":
            target = target[1:]
        elif not target.startswith("/") and "/" in target:
            target = "\\\\" + target.replace("/", "\\")
    if target.lower().startswith("/mnt/") and len(target) > 6:
        drive = target[5].upper()
        target = f"{drive}:\\{target[7:].replace('/', chr(92))}"
    if target.startswith("\\\\") or (len(target) >= 2 and target[1] == ":"):
        return ntpath.normcase(ntpath.normpath(target))
    return os.path.normcase(os.path.normpath(target))


def compare_result_workbooks(expected_path: Path, actual_path: Path) -> dict[str, Any]:
    from openpyxl import load_workbook

    expected_book = load_workbook(expected_path, data_only=False, read_only=False)
    actual_book = load_workbook(actual_path, data_only=False, read_only=False)
    try:
        expected = expected_book["本期审核结果"]
        actual = actual_book["本期审核结果"]
        mismatches: Counter[str] = Counter()
        if expected_book.sheetnames != actual_book.sheetnames:
            mismatches["sheet_names"] += 1
        if expected.freeze_panes != actual.freeze_panes:
            mismatches["freeze_panes"] += 1
        if expected.auto_filter.ref != actual.auto_filter.ref:
            mismatches["auto_filter"] += 1
        max_row = max(expected.max_row, actual.max_row)
        max_column = max(expected.max_column, actual.max_column)
        for row in range(1, max_row + 1):
            for column in range(1, max_column + 1):
                expected_cell = expected.cell(row, column)
                actual_cell = actual.cell(row, column)
                expected_value = expected_cell.value or ""
                actual_value = actual_cell.value or ""
                if expected_value != actual_value:
                    mismatches["cell_values"] += 1
                if expected_value and expected_cell.data_type != actual_cell.data_type:
                    mismatches["cell_types"] += 1
                if expected_value and expected_cell.number_format != actual_cell.number_format:
                    mismatches["number_formats"] += 1

        def effective_column_width(sheet, column_index: int) -> float:
            for dimension in sheet.column_dimensions.values():
                minimum = dimension.min
                maximum = dimension.max or minimum
                if minimum <= column_index <= maximum:
                    return dimension.width
            return sheet.sheet_format.defaultColWidth or 13.0

        for column_index in range(1, 14):
            expected_width = effective_column_width(expected, column_index)
            actual_width = effective_column_width(actual, column_index)
            if abs(expected_width - actual_width) > 0.01:
                mismatches["column_widths"] += 1

        for row in (1,):
            for column in range(1, 14):
                expected_cell = expected.cell(row, column)
                actual_cell = actual.cell(row, column)
                if expected_cell.font.bold != actual_cell.font.bold:
                    mismatches["header_font"] += 1
                if (expected_cell.fill.fgColor.rgb or "")[-6:] != (
                    actual_cell.fill.fgColor.rgb or ""
                )[-6:]:
                    mismatches["header_fill"] += 1

        def hyperlinks(sheet) -> dict[str, tuple[str, str, str]]:
            result = {}
            for row in sheet.iter_rows():
                for cell in row:
                    link = cell.hyperlink
                    if link is not None:
                        result[cell.coordinate] = (
                            normalize_hyperlink_target(link.target or ""),
                            link.location or "",
                            link.display or "",
                        )
            return result

        def color_signature(color):
            if color is None:
                return None
            if color.type == "rgb":
                return ("rgb", (color.rgb or "")[-6:].lower(), round(color.tint, 6))
            if color.type == "theme":
                return ("theme", color.theme, round(color.tint, 6))
            if color.type == "indexed":
                return ("indexed", color.indexed, round(color.tint, 6))
            return (color.type, str(color.value), round(color.tint, 6))

        def hyperlink_font_signature(font):
            return (
                font.name,
                font.sz,
                font.bold,
                font.italic,
                font.underline,
                color_signature(font.color),
            )

        expected_links = hyperlinks(expected)
        actual_links = hyperlinks(actual)
        for coordinate in set(expected_links) | set(actual_links):
            expected_link = expected_links.get(coordinate)
            actual_link = actual_links.get(coordinate)
            if expected_link is None or actual_link is None:
                mismatches["hyperlink_cells"] += 1
                continue
            if expected_link[0] != actual_link[0]:
                mismatches["hyperlink_targets"] += 1
            if expected_link[1] != actual_link[1]:
                mismatches["hyperlink_locations"] += 1
            if expected_link[2] != actual_link[2]:
                mismatches["hyperlink_display"] += 1
            if hyperlink_font_signature(expected[coordinate].font) != hyperlink_font_signature(
                actual[coordinate].font
            ):
                mismatches["hyperlink_styles"] += 1
        return {
            "status": "passed" if not mismatches else "failed",
            "rows": max_row - 1,
            "hyperlinks": len(actual_links),
            "mismatch_counts": dict(sorted(mismatches.items())),
        }
    finally:
        expected_book.close()
        actual_book.close()


def compare_history_result_workbooks(expected_path: Path, actual_path: Path) -> dict[str, Any]:
    from openpyxl import load_workbook

    expected_book = load_workbook(expected_path, data_only=False, read_only=False)
    actual_book = load_workbook(actual_path, data_only=False, read_only=False)
    try:
        mismatches: Counter[str] = Counter()
        if expected_book.sheetnames != actual_book.sheetnames:
            mismatches["sheet_names"] += 1
        expected = expected_book["核查表校验结果"] if "核查表校验结果" in expected_book else None
        actual = actual_book["核查表校验结果"] if "核查表校验结果" in actual_book else None
        if expected is None or actual is None:
            mismatches["history_sheet_missing"] += 1
            return {
                "status": "failed",
                "rows": 0,
                "mismatch_counts": dict(sorted(mismatches.items())),
            }

        if expected.freeze_panes != actual.freeze_panes:
            mismatches["freeze_panes"] += 1
        if expected.auto_filter.ref != actual.auto_filter.ref:
            mismatches["auto_filter"] += 1

        max_row = max(expected.max_row, actual.max_row)
        for row in range(1, max_row + 1):
            for column in range(1, 14):
                expected_cell = expected.cell(row, column)
                actual_cell = actual.cell(row, column)
                expected_value = expected_cell.value
                actual_value = actual_cell.value
                if expected_value is None or actual_value is None:
                    if expected_value is not actual_value:
                        mismatches["cell_values"] += 1
                elif (
                    isinstance(expected_value, bool)
                    or isinstance(actual_value, bool)
                    or isinstance(expected_value, (int, float))
                    and isinstance(actual_value, (int, float))
                ):
                    if isinstance(expected_value, bool) and isinstance(actual_value, bool):
                        if expected_value != actual_value:
                            mismatches["cell_values"] += 1
                    elif (
                        isinstance(expected_value, (int, float))
                        and not isinstance(expected_value, bool)
                        and isinstance(actual_value, (int, float))
                        and not isinstance(actual_value, bool)
                    ):
                        if not math.isclose(
                            float(expected_value), float(actual_value), rel_tol=1e-12, abs_tol=1e-12
                        ):
                            mismatches["cell_values"] += 1
                    else:
                        mismatches["cell_values"] += 1
                elif canonical(expected_value) != canonical(actual_value):
                    mismatches["cell_values"] += 1

                if expected_value is not None or actual_value is not None:
                    if expected_cell.data_type != actual_cell.data_type:
                        mismatches["cell_types"] += 1

        def effective_column_width(sheet, column_index: int) -> float:
            for dimension in sheet.column_dimensions.values():
                minimum = dimension.min
                maximum = dimension.max or minimum
                if minimum <= column_index <= maximum:
                    return dimension.width
            return sheet.sheet_format.defaultColWidth or 13.0

        for column_index in range(1, 14):
            if abs(
                effective_column_width(expected, column_index)
                - effective_column_width(actual, column_index)
            ) > 0.01:
                mismatches["column_widths"] += 1

        return {
            "status": "passed" if not mismatches else "failed",
            "rows": max_row - 1,
            "mismatch_counts": dict(sorted(mismatches.items())),
        }
    finally:
        expected_book.close()
        actual_book.close()


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source-root", required=True, type=Path)
    parser.add_argument("--audit-binary", required=True, type=Path)
    parser.add_argument("--input-dir", required=True, type=Path)
    parser.add_argument("--source-file", required=True, type=Path)
    parser.add_argument("--audit-workbook", required=True, type=Path)
    parser.add_argument("--template-workbook", required=True, type=Path)
    parser.add_argument("--template-dir", type=Path)
    parser.add_argument("--protected-original", action="append", type=Path, default=[])
    parser.add_argument("--mapping-config", type=Path)
    parser.add_argument("--result-workbook", type=Path)
    parser.add_argument("--history-result-workbook", type=Path)
    parser.add_argument("--rust-cli-peak-rss-file", type=Path)
    parser.add_argument("--wsl-distribution", default="Ubuntu")
    parser.add_argument("--period", required=True)
    parser.add_argument("--batch-id", default="s1-chain-parity")
    parser.add_argument("--audit-time", default="2026-10-24 12:00:00")
    parser.add_argument("--report", required=True, type=Path)
    parser.add_argument("--source-filename-alias-used", action="store_true")
    args = parser.parse_args()

    source = args.source_file.resolve(strict=True)
    audit = args.audit_workbook.resolve(strict=True)
    template = args.template_workbook.resolve(strict=True)
    input_dir = args.input_dir.resolve(strict=True)
    source.relative_to(input_dir)
    history_result_path = (
        args.history_result_workbook.resolve() if args.history_result_workbook else None
    )
    if history_result_path is not None and history_result_path.exists():
        raise ValueError("Rust S1 history result output path already exists")
    mapping_config = args.mapping_config.resolve(strict=True) if args.mapping_config else None
    protected_paths = [source, audit, template]
    protected_paths.extend(path.resolve(strict=True) for path in args.protected_original)
    before = {
        path: sha256_file(path) for path in protected_paths
    }

    sys.path.insert(0, str(args.source_root.resolve(strict=True) / "core" / "src"))
    from base_audit.conditional_evaluators import ConditionalFormatService
    from base_audit.name_config import (
        CONDITIONAL_FORMAT_EXTRACT_FUNCTION,
        FORMULA_COPY_FUNCTION,
        ISSUE_EXTRACT_FUNCTION,
        STRUCTURE_COMPARE_FUNCTION,
        features_of_type,
        load_feature_mappings,
    )
    from base_audit.native.openpyxl_workbook import (
        extract_issues,
        named_ranges,
        read_template,
    )
    from base_audit.service import _organisation_from_name
    from base_audit.history import classify_current_issues, merge_history
    from base_audit.native.result_writer import write_audit_summary
    from base_audit.native.history_io import write_history_xlsx
    from openpyxl import load_workbook

    mappings = load_feature_mappings(mapping_config, template)
    definition = read_template(
        template,
        formula_mappings=features_of_type(mappings, FORMULA_COPY_FUNCTION),
        structure_mappings=features_of_type(mappings, STRUCTURE_COMPARE_FUNCTION),
        extraction_mappings=features_of_type(mappings, ISSUE_EXTRACT_FUNCTION),
    )
    conditional_mappings = features_of_type(mappings, CONDITIONAL_FORMAT_EXTRACT_FUNCTION)
    template_book = load_workbook(template, read_only=True, data_only=False, keep_links=False)
    try:
        conditional_ranges = named_ranges(template_book, conditional_mappings)
        structure_ranges = definition.structure_ranges
    finally:
        template_book.close()

    org_code, org_name = _organisation_from_name(source)
    formula_issues = extract_issues(
        audit,
        definition.rules,
        period=args.period,
        batch_id=args.batch_id,
        org_code=org_code,
        org_name=org_name,
        source_file=source,
        extraction_ranges=definition.extraction_ranges,
    )
    for issue in formula_issues:
        issue.audit_time = args.audit_time

    conditional_issues, conditional_stats = ConditionalFormatService(mode="PYTHON").extract_issues(
        workbook_path=source,
        ranges=conditional_ranges,
        structure_ranges=structure_ranges,
        period=args.period,
        batch_id=args.batch_id,
        audit_time=args.audit_time,
        source_file=source,
    )
    formula_and_conditional = formula_issues + [
        dataclasses.replace(issue, audit_file=str(audit))
        for issue in conditional_issues
    ]
    python_merged_history = merge_history([], formula_and_conditional)
    python_current, python_resolved = classify_current_issues(
        args.period,
        formula_and_conditional,
        [],
        batch_id=args.batch_id,
        audit_time=args.audit_time,
    )

    audit_manifest_path = (
        wsl_path(audit)
        if os.name == "nt"
        and args.audit_binary.suffix.lower() != ".exe"
        else str(audit)
    )
    with tempfile.TemporaryDirectory(prefix="audit-s1-issue-chain-") as directory:
        manifest = Path(directory) / "formula-copies.json"
        manifest.write_text(
            json.dumps(
                {
                    "files": [
                        {
                            "source_file": source.relative_to(input_dir).as_posix(),
                            "audit_workbook": audit_manifest_path,
                        }
                    ]
                },
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )
        rust, rust_log, rust_elapsed_ms = invoke_rust(args, manifest, source)
    template_selection = rust.get("template_selection", {})
    expected_template_mode = "automatic" if args.template_dir else "explicit"
    if template_selection.get("mode") != expected_template_mode:
        raise ValueError("Rust S1 template selection mode differs from the requested mode")
    if template_selection.get("template_file") != template.name:
        raise ValueError("Rust S1 selected a different template from the Python baseline")
    rust_cli_peak_rss_kib = None
    if args.rust_cli_peak_rss_file:
        try:
            rust_cli_peak_rss_kib = int(
                args.rust_cli_peak_rss_file.read_text(encoding="ascii").strip()
            )
        except (OSError, ValueError) as error:
            raise ValueError("Could not read the measured Rust CLI peak RSS") from error
        if rust_cli_peak_rss_kib <= 0:
            raise ValueError("Measured Rust CLI peak RSS must be positive")

    result_workbook_comparison = None
    if args.result_workbook:
        if not args.result_workbook.is_file():
            raise ValueError("Rust S1 result workbook was not created")
        with tempfile.TemporaryDirectory(prefix="audit-s1-python-result-") as directory:
            python_result_workbook = Path(directory) / "python-result.xlsx"
            write_audit_summary(
                python_result_workbook,
                python_current,
                sheet_name="本期审核结果",
                mode="OPENPYXL",
            )
            result_workbook_comparison = compare_result_workbooks(
                python_result_workbook,
                args.result_workbook.resolve(strict=True),
            )
        writer_summary = rust.get("result_workbook")
        if not isinstance(writer_summary, dict):
            raise ValueError("Rust CLI did not report the S1 result-workbook execution")
        if writer_summary.get("issue_count") != len(rust.get("issues", [])):
            raise ValueError("Rust S1 result-workbook issue count differs from the CLI result")
    else:
        writer_summary = None

    history_result_workbook_comparison = None
    if args.history_result_workbook:
        with tempfile.TemporaryDirectory(prefix="audit-s1-python-history-result-") as directory:
            python_history_workbook = Path(directory) / "python-history-result.xlsx"
            write_history_xlsx(python_history_workbook, python_merged_history)
            history_result_workbook_comparison = compare_history_result_workbooks(
                python_history_workbook,
                history_result_path.resolve(strict=True),
            )
        history_writer_summary = rust.get("history_result_workbook")
        if not isinstance(history_writer_summary, dict):
            raise ValueError("Rust CLI did not report the S1 history result-workbook execution")
        if history_writer_summary.get("issue_count") != len(rust.get("merged_history", [])):
            raise ValueError("Rust S1 history result-workbook count differs from merged history")
    else:
        history_writer_summary = None

    formula_comparison = compare_issues(python_current, rust.get("issues", []))
    resolved_comparison = compare_issues(python_resolved, rust.get("resolved_issues", []))
    history_comparison = compare_issues(python_merged_history, rust.get("merged_history", []))
    actual_formula_count = rust["formula_summary"]["issues"]
    actual_conditional_count = rust["conditional_summary"]["issues"]
    actual_preflight = rust.get("structure_preflight")
    if actual_preflight and actual_preflight["files_accepted"] != 1:
        raise ValueError("Rust structure preflight did not accept the selected real source")
    if actual_formula_count != len(formula_issues):
        raise ValueError("Rust formula issue count differs from the Python formula baseline")
    if actual_conditional_count != len(conditional_issues):
        raise ValueError("Rust conditional issue count differs from the Python baseline")

    after = {
        path: sha256_file(path) for path in protected_paths
    }
    inputs_unchanged = before == after
    all_comparisons = [formula_comparison, resolved_comparison, history_comparison]
    mismatch_count = sum(
        item["missing_in_rust"]
        + item["missing_in_python"]
        + sum(item["field_difference_counts"].values())
        for item in all_comparisons
    )
    if result_workbook_comparison:
        mismatch_count += sum(result_workbook_comparison["mismatch_counts"].values())
    if history_result_workbook_comparison:
        mismatch_count += sum(history_result_workbook_comparison["mismatch_counts"].values())
    log_events = [
        json.loads(line)
        for line in rust_log.splitlines()
        if line.lstrip().startswith("{")
    ]
    actual_readers = sorted(
        {
            event.get("engine") or event.get("result_reader")
            for event in log_events
            if event.get("step")
            in {
                "formula_result_issue_extraction",
                "conditional_format_issue_extraction",
                "issue_merge_and_history_classification",
            }
            and (event.get("status") == "completed" or event.get("engine"))
        }
    )
    report = {
        "schema_version": 1,
        "suite": "audit.s1_issue_chain",
        "baseline": "Open-Data-Audit UOS/native DAG adapter formula/conditional Issue stages and history classifier",
        "status": "passed" if mismatch_count == 0 and inputs_unchanged else "failed",
        "python": {
            "formula_issues": len(formula_issues),
            "conditional_issues": len(conditional_issues),
            "conditional_evaluated_cells": conditional_stats.get("evaluated_cells", 0),
            "merged_issues": len(formula_and_conditional),
        },
        "rust": {
            "formula_issues": actual_formula_count,
            "conditional_issues": actual_conditional_count,
            "merged_issues": len(rust.get("issues", [])),
            "resolved_issues": len(rust.get("resolved_issues", [])),
            "readers": actual_readers,
            "cli_elapsed_ms": rust_elapsed_ms,
            "cli_peak_rss_kib": rust_cli_peak_rss_kib,
            "template_selection_mode": template_selection.get("mode"),
            "template_selection_matches_baseline": (
                template_selection.get("template_file") == template.name
            ),
            "result_workbook": (
                {
                    "engine": writer_summary.get("engine"),
                    "issue_count": writer_summary.get("issue_count"),
                    "hyperlink_count": writer_summary.get("hyperlink_count"),
                    "elapsed_ms": writer_summary.get("elapsed_ms"),
                }
                if writer_summary
                else None
            ),
            "history_result_workbook": (
                {
                    "engine": history_writer_summary.get("engine"),
                    "issue_count": history_writer_summary.get("issue_count"),
                    "elapsed_ms": history_writer_summary.get("elapsed_ms"),
                }
                if history_writer_summary
                else None
            ),
        },
        "comparison": {
            "current_issues": formula_comparison,
            "resolved_issues": resolved_comparison,
            "merged_history": history_comparison,
            "result_workbook": result_workbook_comparison,
            "history_result_workbook": history_result_workbook_comparison,
            "mismatch_count": mismatch_count,
        },
        "verification": {
            "all_workbooks_unchanged": inputs_unchanged,
            "formula_recalculation": "not performed; saved audit-copy values only",
            "raw_issue_values_or_paths_persisted": False,
            "python_conditional_metadata_source": "workflow.adapters.native_audit_handlers.handle_conditional_issue_extract",
            "windows_general_handler_metadata_behavior": "workflow.adapters.audit_handlers.handle_conditional_issue_extract passes source organization fields; S1 UOS/native path intentionally leaves them empty",
            "scope": "one read-only real source/audit-copy pair and one S1 template",
            "source_filename_alias_used_for_linux_component_limit": args.source_filename_alias_used,
            "rust_cli_peak_rss_scope": (
                "complete s1-audit-directory process, including result workbook output"
                if rust_cli_peak_rss_kib is not None
                else "not measured"
            ),
            "python_conditional_total_time_s": round(
                float(conditional_stats.get("total_time", 0.0)), 3
            ),
        },
    }
    args.report.parent.mkdir(parents=True, exist_ok=True)
    args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if report["status"] == "passed" else 1


if __name__ == "__main__":
    raise SystemExit(main())
