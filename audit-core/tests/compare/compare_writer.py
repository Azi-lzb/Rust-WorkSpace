"""Verify a rust_xlsxwriter-created workbook with the OpenPyXL baseline reader."""
from __future__ import annotations

import argparse
import json
import os
import re
import statistics
import subprocess
from pathlib import Path

from openpyxl import load_workbook


def norm(value):
    return re.sub(r"'([^']+)'!", r"\1!", str(value or "")).replace("$", "").replace("=", "").upper()


def wsl_path(path: Path) -> str:
    text = str(path.resolve())
    if len(text) >= 2 and text[1] == ":":
        return f"/mnt/{text[0].lower()}/{text[3:].replace(chr(92), '/') }"
    return text.replace("\\", "/")


def run_writer(binary: str, output: Path, wsl: bool) -> dict:
    prefix = ["wsl.exe", binary] if wsl else [binary]
    out_arg = wsl_path(output) if wsl else str(output)
    proc = subprocess.run(prefix + ["write-smoke", "--output", out_arg], check=True, capture_output=True, text=True, encoding="utf-8", errors="replace")
    return json.loads(proc.stdout)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--report", type=Path)
    parser.add_argument("--rust-bin", default=os.getenv("RUST_FASTPATH_PROBE"))
    parser.add_argument("--wsl", action="store_true", help="invoke Linux binary through wsl.exe")
    parser.add_argument("--runs", type=int, default=5)
    args = parser.parse_args()
    performance = {"status": "not_measured"}
    if args.rust_bin:
        if args.runs < 1:
            parser.error("--runs must be positive")
        run_writer(args.rust_bin, args.input, args.wsl)  # warm-up
        timings = [run_writer(args.rust_bin, args.input, args.wsl)["elapsed_ms"] for _ in range(args.runs)]
        performance = {"warmup_runs": 1, "measured_runs": args.runs, "rust_xlsxwriter_internal_ms": {"median": statistics.median(timings), "min": min(timings), "max": max(timings)}}
    checks = []
    try:
        wb = load_workbook(args.input, data_only=False)
        ws = wb.active
        checks.append({"name": "openpyxl_openability", "passed": True, "actual": wb.sheetnames})
    except Exception as exc:
        result = {"input": str(args.input), "checks": [{"name": "openpyxl_openability", "passed": False, "actual": repr(exc)}], "passed": 0, "total": 1, "rate": 0.0, "performance_ms": performance}
        print(json.dumps(result, ensure_ascii=False, indent=2))
        return

    def check(name, passed, actual):
        checks.append({"name": name, "passed": bool(passed), "actual": actual})

    check("sheet_names_and_order", wb.sheetnames == ["Synthetic"], wb.sheetnames)
    check("values", ws["A1"].value == 1 and ws["A2"].value == 2 and ws["A3"].value == 3 and ws["E5"].value == "named target", {"A1": ws["A1"].value, "A2": ws["A2"].value, "A3": ws["A3"].value, "E5": ws["E5"].value})
    check("formula_text", ws["D4"].data_type == "f" and norm(ws["D4"].value) == "SUM(A1:A3)", ws["D4"].value)
    check("merge", "A7:B7" in {str(x) for x in ws.merged_cells.ranges}, sorted(str(x) for x in ws.merged_cells.ranges))
    a6 = ws["A6"]
    fill = a6.fill.fgColor.rgb if a6.fill.fgColor.type == "rgb" else None
    check("style", a6.font.bold and fill in ("FFFFFF00", "FFFF00") and a6.number_format == "0.00", {"bold": a6.font.bold, "fill": fill, "number_format": a6.number_format})
    defined = wb.defined_names.get("SyntheticNamedCell")
    check("defined_name", defined is not None and norm(defined.attr_text) == "SYNTHETIC!E5", None if defined is None else defined.attr_text)
    col_width = ws.column_dimensions["I"].width
    check("row_column_dimensions", ws.row_dimensions[3].height == 30 and abs(col_width - 25) <= 1, {"row3": ws.row_dimensions[3].height, "colI": col_width, "requested_colI": 25, "tolerance": 1})
    check("print_area", norm(ws.print_area) == "SYNTHETIC!A1:D10", ws.print_area)
    passed = sum(c["passed"] for c in checks)
    result = {"input": str(args.input), "checks": checks, "passed": passed, "total": len(checks), "rate": passed / len(checks), "performance_ms": performance}
    serialized = json.dumps(result, ensure_ascii=False, indent=2)
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(serialized + "\n", encoding="utf-8")
    print(serialized)


if __name__ == "__main__":
    main()
