"""Compare Calamine JSON exports with the OpenPyXL parity baseline."""
from __future__ import annotations

import argparse
import json
import math
import os
import statistics
import subprocess
import tempfile
import time as timer
from datetime import date, datetime, time, timedelta
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.datetime import to_excel


PROJECT = Path(__file__).resolve().parents[2]


def wsl_path(path: Path) -> str:
    text = str(path.resolve())
    if len(text) >= 2 and text[1] == ":":
        return f"/mnt/{text[0].lower()}/{text[3:].replace(chr(92), '/') }"
    return text.replace("\\", "/")


def invoke_probe(binary: str, input_path: Path, output_path: Path, wsl: bool) -> None:
    prefix = ["wsl.exe", binary] if wsl else [binary]
    input_arg = wsl_path(input_path) if wsl else str(input_path)
    output_arg = wsl_path(output_path) if wsl else str(output_path)
    subprocess.run(prefix + ["read", "--input", input_arg, "--output", output_arg], check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)


def scalar(value, epoch, data_type=None):
    if isinstance(value, (datetime, date)):
        return {"type": "datetime_serial", "value": to_excel(value, epoch=epoch)}
    if isinstance(value, time):
        return {"type": "time", "value": value.isoformat()}
    if isinstance(value, timedelta):
        return {"type": "duration", "value": value.total_seconds()}
    if isinstance(value, bool):
        return {"type": "boolean", "value": value}
    if isinstance(value, int):
        return {"type": "integer", "value": value}
    if isinstance(value, float):
        if math.isnan(value):
            return {"type": "number", "value": "NaN"}
        return {"type": "number", "value": value}
    if isinstance(value, str):
        if data_type == "e":
            return {"type": "error", "value": value}
        return {"type": "string", "value": value}
    return None


def openpyxl_dump(path: Path) -> dict:
    wb = load_workbook(path, data_only=False, read_only=False)
    values_wb = load_workbook(path, data_only=True, read_only=False)
    epoch = values_wb.epoch
    sheets = []
    for ws in wb.worksheets:
        value_ws = values_wb[ws.title]
        values, formulas, comments = {}, {}, {}
        for row in ws.iter_rows():
            for cell in row:
                key = cell.coordinate
                value = scalar(value_ws[key].value, epoch, value_ws[key].data_type)
                if value is not None:
                    values[key] = value
                if cell.data_type == "f":
                    formulas[key] = str(cell.value)
                if cell.comment is not None:
                    comments[key] = cell.comment.text
        sheets.append({
            "name": ws.title,
            "values": values,
            "formulas": formulas,
            "merges": sorted(str(r) for r in ws.merged_cells.ranges),
            "comments": comments,
        })
    return {"schema_version": 1, "reader": "openpyxl-baseline", "sheets": sheets}


def normalize_formula(s: str) -> str:
    return s.lstrip("=").replace(" ", "").upper()


def compare_fixture(path: Path, rust_bin: str, wsl: bool, out_dir: Path, runs: int) -> dict:
    json_path = out_dir / f"{path.stem}_calamine.json"
    invoke_probe(rust_bin, path, json_path, wsl)  # warm-up
    calamine_ms, openpyxl_ms = [], []
    for _ in range(runs):
        invoke_probe(rust_bin, path, json_path, wsl)
        calamine_ms.append(json.loads(json_path.read_text(encoding="utf-8"))["elapsed_ms"])
        started = timer.perf_counter()
        openpyxl_dump(path)
        openpyxl_ms.append((timer.perf_counter() - started) * 1000)
    invoke_probe(rust_bin, path, json_path, wsl)
    actual = json.loads(json_path.read_text(encoding="utf-8"))
    expected = openpyxl_dump(path)
    expected_names = [sheet["name"] for sheet in expected["sheets"]]
    actual_names = [sheet["name"] for sheet in actual["sheets"]]
    results = {
        "sheet_names": {
            "matched": int(expected_names == actual_names),
            "total": 1,
            "rate": float(expected_names == actual_names),
            "differences": [] if expected_names == actual_names else [{"expected": expected_names, "actual": actual_names}],
        }
    }
    for field in ("values", "formulas", "merges", "comments"):
        exp_sheets = {s["name"]: s[field] for s in expected["sheets"]}
        got_sheets = {s["name"]: s[field] for s in actual["sheets"]}
        keys = sorted(set(exp_sheets) | set(got_sheets))
        matched = 0
        total = 0
        differences = []
        for name in keys:
            exp, got = exp_sheets.get(name, {}), got_sheets.get(name, {})
            if field == "merges":
                total += 1
                matched += int(exp == got)
                if exp != got:
                    differences.append({"sheet": name, "expected": exp, "actual": got})
            else:
                all_keys = sorted(set(exp) | set(got))
                for cell in all_keys:
                    total += 1
                    ev, gv = exp.get(cell), got.get(cell)
                    equal = ev == gv
                    if field == "formulas" and isinstance(ev, str) and isinstance(gv, str):
                        equal = normalize_formula(ev) == normalize_formula(gv)
                    if field == "values" and isinstance(ev, dict) and isinstance(gv, dict):
                        if ev.get("type") in ("integer", "number") and gv.get("type") in ("integer", "number"):
                            equal = math.isclose(float(ev["value"]), float(gv["value"]), rel_tol=0, abs_tol=1e-12)
                    matched += int(equal)
                    if not equal and len(differences) < 100:
                        differences.append({"sheet": name, "cell": cell, "expected": ev, "actual": gv})
        results[field] = {"matched": matched, "total": total, "rate": (matched / total if total else None), "differences": differences}
    return {
        "fixture": path.name,
        "calamine_json": str(json_path),
        "checks": results,
        "performance_ms": {
            "warmup_runs": 1,
            "measured_runs": runs,
            "calamine_internal": {"median": statistics.median(calamine_ms), "min": min(calamine_ms), "max": max(calamine_ms)},
            "openpyxl_read_baseline": {"median": statistics.median(openpyxl_ms), "min": min(openpyxl_ms), "max": max(openpyxl_ms)},
            "note": "Calamine value/formula/merge/comment export includes a namespace-aware comment supplement; OpenPyXL uses data_only values plus formula text.",
        },
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--rust-bin", default=os.getenv("RUST_FASTPATH_PROBE"))
    parser.add_argument("--wsl", action="store_true", help="invoke Linux binary through wsl.exe")
    parser.add_argument("--fixtures", type=Path, default=PROJECT / "tests" / "fixtures")
    parser.add_argument("--out-dir", type=Path, default=Path(tempfile.gettempdir()) / "open-data-audit-rust-read-compare")
    parser.add_argument("--report", type=Path)
    parser.add_argument("--runs", type=int, default=5)
    args = parser.parse_args()
    if not args.rust_bin:
        parser.error("--rust-bin or RUST_FASTPATH_PROBE is required")
    args.out_dir.mkdir(parents=True, exist_ok=True)
    fixtures = sorted(p for p in args.fixtures.iterdir() if p.suffix.lower() in {".xlsx", ".xlsm", ".xls", ".xlsb"} and not p.name.startswith("~$"))
    if args.runs < 1:
        parser.error("--runs must be positive")
    report = {"schema_version": 1, "fixtures": [compare_fixture(p, args.rust_bin, args.wsl, args.out_dir, args.runs) for p in fixtures]}
    totals = {}
    for field in ("sheet_names", "values", "formulas", "merges", "comments"):
        matched = sum(f["checks"][field]["matched"] for f in report["fixtures"])
        total = sum(f["checks"][field]["total"] for f in report["fixtures"])
        totals[field] = {"matched": matched, "total": total, "rate": matched / total if total else None}
    report["totals"] = totals
    report["status"] = "no_fixtures" if not fixtures else "completed"
    serialized = json.dumps(report, ensure_ascii=False, indent=2)
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(serialized + "\n", encoding="utf-8")
    print(serialized)


if __name__ == "__main__":
    main()
