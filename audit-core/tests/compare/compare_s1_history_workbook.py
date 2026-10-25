#!/usr/bin/env python3
"""Compare Rust history workbook ingestion with the read-only Python baseline."""

from __future__ import annotations

import argparse
from functools import lru_cache
import json
import subprocess
import sys
import tempfile
from dataclasses import asdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook

HISTORY_HEADERS = (
    "工作簿名",
    "工作表名",
    "定位单元格",
    "错误类型",
    "校验类型",
    "校验指标",
    "描述",
    "当前值",
    "对比值",
    "差值",
    "规则编号",
    "历史校验说明",
    "审核意见",
    "扩展列",
    "数据期",
)


def load_baseline(source_root: Path):
    sys.path.insert(0, str(source_root / "core" / "src"))
    from base_audit.history import classify_current_issues, merge_history
    from base_audit.models import Issue
    from base_audit.native.history_io import read_history_xlsx

    return Issue, classify_current_issues, merge_history, read_history_xlsx


def issue(issue_type, issue_id: str, period: str):
    return issue_type(
        issue_id=issue_id,
        period=period,
        batch_id="batch-old",
        audit_time="old-time",
        triggered=True,
        status="新增",
        first_seen_period=period,
        previous_seen_period="",
        consecutive_count=1,
        org_code="001",
        org_name="测试机构",
        report_code="R01",
        sheet_name="汇总表",
        rule_id=issue_id,
        severity="错误",
        formula_cell="C8",
        target_cell="D8",
        target_value=125,
        formula_result="错误",
        message="当前错误",
        source_file="机构一.xlsx",
        audit_file="机构一_审核版.xlsx",
    )


def normalize(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, dict):
        return {key: normalize(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [normalize(item) for item in value]
    return value


def first_difference(expected, actual, path="$",):
    if type(expected) is not type(actual):
        return path, expected, actual
    if isinstance(expected, dict):
        for key in sorted(set(expected) | set(actual)):
            if key not in expected or key not in actual:
                return f"{path}.{key}", expected.get(key), actual.get(key)
            difference = first_difference(expected[key], actual[key], f"{path}.{key}")
            if difference:
                return difference
    elif isinstance(expected, list):
        for index, (left, right) in enumerate(zip(expected, actual)):
            difference = first_difference(left, right, f"{path}[{index}]")
            if difference:
                return difference
        if len(expected) != len(actual):
            return f"{path}.length", len(expected), len(actual)
    elif expected != actual:
        return path, expected, actual
    return None


def write_history_sheet(workbook: Workbook, name: str, rules: tuple[str, ...]):
    sheet = workbook.create_sheet(name)
    sheet.append(HISTORY_HEADERS)
    sheet.append((
        "机构一.xlsx", "汇总表", "D8", "核实", "数值", "资产总额", "金额不一致",
        123.5, 100, 23.5, rules[0], "机构历史说明", "审核历史意见", "", "2026-02",
    ))
    if len(rules) > 1:
        sheet.append((
            "机构一.xlsx", "汇总表", "E9", "错误", "数值", "存款余额", "历史问题已整改",
            50, 40, 10, rules[1], "", "", "", "2026-02",
        ))
    # A row with content but no source file, sheet, target cell or rule ID is skipped.
    sheet.cell(sheet.max_row + 1, 6, "仅有指标的空身份行")
    return sheet


def synthetic_workbooks(directory: Path):
    standard_path = directory / "synthetic_history.xlsx"
    standard = Workbook()
    standard.remove(standard.active)
    write_history_sheet(standard, "核查表校验结果", (" RULE-CONT ", "RULE-RESOLVED"))
    write_history_sheet(standard, "自定义历史", ("RULE-CUSTOM",))
    write_history_sheet(standard, "历史审核结果", ("RULE-LEGACY",))
    standard.save(standard_path)

    legacy_path = directory / "legacy_history.xlsx"
    legacy = Workbook()
    legacy.remove(legacy.active)
    write_history_sheet(legacy, "历史审核结果", ("RULE-LEGACY",))
    legacy.save(legacy_path)
    return standard_path, legacy_path


def build_request(Issue, current, history):
    return {
        "period": "2026-03",
        "batch_id": "batch-new",
        "audit_time": "2026-10-24 12:00:00",
        "current_issues": [asdict(item) for item in current],
        # This sentinel proves workbook input replaces JSON history.
        "history": [asdict(issue(Issue, "JSON-DECOY", "2026-02"))] if history else [],
    }


@lru_cache(maxsize=None)
def to_wsl_path(path: str, distro: str) -> str:
    completed = subprocess.run(
        ["wsl.exe", "-d", distro, "--", "wslpath", "-a", path.replace("\\", "/")],
        check=True,
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
    )
    return completed.stdout.strip()


def run_rust(
    binary: Path,
    request: dict,
    workbook: Path,
    directory: Path,
    *,
    sheet=None,
    developer=False,
    wsl_distro=None,
):
    input_path = directory / f"input-{workbook.stem}-{'dev' if developer else 'user'}.json"
    input_path.write_text(json.dumps(normalize(request), ensure_ascii=False), encoding="utf-8")
    if wsl_distro:
        binary_arg = str(binary).replace("\\", "/")
        command = ["wsl.exe", "-d", wsl_distro, "--", binary_arg]
        input_arg = to_wsl_path(str(input_path), wsl_distro)
        workbook_arg = to_wsl_path(str(workbook), wsl_distro)
    else:
        command = [str(binary)]
        input_arg = str(input_path)
        workbook_arg = str(workbook)
    command.extend(("s1-history-classify", "--input", input_arg, "--history-workbook", workbook_arg))
    if sheet is not None:
        command.extend(("--history-sheet", sheet))
    if developer:
        command.append("--developer-view")
    completed = subprocess.run(
        command,
        check=True,
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
    )
    return json.loads(completed.stdout), completed.stderr.replace("\x00", "")


def expected_result(period, current, history, classify_current_issues, merge_history):
    merged = merge_history(history, current)
    classified, resolved = classify_current_issues(
        period,
        current,
        history,
        batch_id="batch-new",
        audit_time="2026-10-24 12:00:00",
    )
    return normalize({
        "current_issues": [asdict(item) for item in classified],
        "resolved_issues": [asdict(item) for item in resolved],
        "merged_history": [asdict(item) for item in merged],
    })


def compare_case(
    Issue, classify_current_issues, merge_history, read_history_xlsx,
    binary, workbook, directory, *, current_ids=(), sheet=None, developer=False,
    wsl_distro=None,
):
    history = read_history_xlsx(workbook, sheet_name=sheet or "")
    current = [issue(Issue, issue_id, "2026-03") for issue_id in current_ids]
    request = build_request(Issue, current, history)
    expected = expected_result("2026-03", current, history, classify_current_issues, merge_history)
    output, log = run_rust(
        binary,
        request,
        workbook,
        directory,
        sheet=sheet,
        developer=developer,
        wsl_distro=wsl_distro,
    )
    for field in ("current_issues", "resolved_issues", "merged_history"):
        actual = normalize(output[field])
        if actual != expected[field]:
            detail = first_difference(expected[field], actual, f"$.{field}")
            raise AssertionError(f"Python/Rust history workbook difference: {detail}")
    return output, log, history


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--audit-binary", required=True, type=Path)
    parser.add_argument("--wsl-distro", help="Run the Linux audit binary through this WSL distro.")
    parser.add_argument("--source-root", required=True, type=Path)
    parser.add_argument("--report", required=True, type=Path)
    parser.add_argument(
        "--real-history-workbook",
        type=Path,
        help="Optional read-only product workbook for an additional sanitized smoke comparison.",
    )
    args = parser.parse_args()
    Issue, classify_current_issues, merge_history, read_history_xlsx = load_baseline(args.source_root)
    args.report.parent.mkdir(parents=True, exist_ok=True)

    with tempfile.TemporaryDirectory(
        prefix="audit-s1-history-workbook-",
        dir=args.report.parent.resolve(),
    ) as temporary:
        directory = Path(temporary)
        standard_path, legacy_path = synthetic_workbooks(directory)
        standard_output, ordinary_log, standard_history = compare_case(
            Issue,
            classify_current_issues,
            merge_history,
            read_history_xlsx,
            args.audit_binary,
            standard_path,
            directory,
            current_ids=(" RULE-CONT ",),
            wsl_distro=args.wsl_distro,
        )
        developer_output, developer_log = run_rust(
            args.audit_binary,
            build_request(Issue, [issue(Issue, " RULE-CONT ", "2026-03")], standard_history),
            standard_path,
            directory,
            developer=True,
            wsl_distro=args.wsl_distro,
        )
        expected = expected_result(
            "2026-03",
            [issue(Issue, " RULE-CONT ", "2026-03")],
            standard_history,
            classify_current_issues,
            merge_history,
        )
        for field in ("current_issues", "resolved_issues", "merged_history"):
            if normalize(developer_output[field]) != expected[field]:
                detail = first_difference(expected[field], normalize(developer_output[field]), f"$.{field}")
                raise AssertionError(f"Developer-view parity difference: {detail}")
        requested_output, _, requested_history = compare_case(
            Issue,
            classify_current_issues,
            merge_history,
            read_history_xlsx,
            args.audit_binary,
            standard_path,
            directory,
            sheet="自定义历史",
            wsl_distro=args.wsl_distro,
        )
        legacy_output, _, legacy_history = compare_case(
            Issue,
            classify_current_issues,
            merge_history,
            read_history_xlsx,
            args.audit_binary,
            legacy_path,
            directory,
            wsl_distro=args.wsl_distro,
        )

        real_sample = None
        if args.real_history_workbook:
            real_history = read_history_xlsx(args.real_history_workbook)
            real_current = []
            real_request = {
                "period": "2099-12",
                "batch_id": "synthetic-smoke",
                "audit_time": "2099-12-31 00:00:00",
                "current_issues": [],
                "history": [],
            }
            real_output, _ = run_rust(
                args.audit_binary,
                real_request,
                args.real_history_workbook,
                directory,
                wsl_distro=args.wsl_distro,
            )
            real_expected = expected_result(
                "2099-12",
                real_current,
                real_history,
                classify_current_issues,
                merge_history,
            )
            for field in ("current_issues", "resolved_issues", "merged_history"):
                if normalize(real_output[field]) != real_expected[field]:
                    detail = first_difference(real_expected[field], normalize(real_output[field]), f"$.{field}")
                    raise AssertionError(f"Read-only real workbook parity difference: {detail}")
            real_sample = {
                "role": "read_only_real_history_workbook",
                "history_rows": len(real_history),
                "reader": real_output["history_workbook_read"]["reader"],
                "sheet_name": real_output["history_workbook_read"]["sheet_name"],
                "python_rust_field_differences": 0,
            }

    ordinary_events = ordinary_log
    developer_events = [
        json.loads(line)
        for line in developer_log.splitlines()
        if line.lstrip().startswith("{")
    ]
    if "引擎 calamine" not in ordinary_events or "耗时" not in ordinary_events:
        raise AssertionError("ordinary user log omits actual reader engine or elapsed time")
    if not any(
        event.get("step") == "history_workbook_read"
        and event.get("status") == "completed"
        and event.get("engine") == "calamine"
        and "elapsed_ms" in event
        for event in developer_events
    ):
        raise AssertionError("developer log omits completed Calamine timing event")

    summary = {
        "suite": "audit.s1_history_workbook",
        "baseline": "Open-Data-Audit base_audit.native.history_io.read_history_xlsx",
        "status": "passed",
        "synthetic_cases": {
            "standard_sheet": {
                "history_rows": len(standard_history),
                "current_issues": len(standard_output["current_issues"]),
                "resolved_issues": len(standard_output["resolved_issues"]),
                "python_rust_field_differences": 0,
            },
            "requested_sheet": {
                "sheet_name": requested_output["history_workbook_read"]["sheet_name"],
                "history_rows": len(requested_history),
                "python_rust_field_differences": 0,
            },
            "legacy_sheet": {
                "sheet_name": legacy_output["history_workbook_read"]["sheet_name"],
                "history_rows": len(legacy_history),
                "python_rust_field_differences": 0,
            },
        },
        "user_log_engine_and_timing": True,
        "developer_log_engine_and_timing": True,
        "source_workbooks_modified": False,
    }
    if real_sample is not None:
        summary["read_only_real_sample"] = real_sample
    args.report.write_text(json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
