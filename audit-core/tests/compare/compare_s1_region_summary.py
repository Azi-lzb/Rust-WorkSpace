#!/usr/bin/env python3
"""Compare the Rust S1-F02 named-region CLI with the read-only Python flow."""

from __future__ import annotations

import argparse
import hashlib
import json
import subprocess
import sys
import tempfile
import time
from datetime import date, datetime
from pathlib import Path
from typing import Any

sys.dont_write_bytecode = True
WORKSPACE = Path(__file__).resolve().parents[2]
SOURCE_CORE = Path(r"C:\Users\AZI\Desktop\VibeCode\3.open-data-audit\core")
sys.path.insert(0, str(SOURCE_CORE / "src"))

from base_audit.name_config import (  # noqa: E402
    FIXED_ROW_SUMMARY_FUNCTION,
    USED_RANGE_SUMMARY_FUNCTION,
    FeatureMapping,
)
from base_audit.region_summary import run_region_summaries  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402


def digest(path: Path) -> str:
    result = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            result.update(chunk)
    return result.hexdigest()


def write_source(path: Path, rows: dict[int, tuple[Any, Any]]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    for row_number, values in rows.items():
        sheet.cell(row_number, 2, values[0])
        sheet.cell(row_number, 3, values[1])
    workbook.save(path)


def equal_value(left: Any, right: Any) -> bool:
    if isinstance(left, (date, datetime)):
        left = left.isoformat()
    if isinstance(right, (date, datetime)):
        right = right.isoformat()
    if isinstance(left, bool) or isinstance(right, bool):
        return left is right
    if isinstance(left, (int, float)) and isinstance(right, (int, float)):
        return left == right or abs(left - right) <= max(abs(left), abs(right), 1.0) * 1e-15
    return left == right


def compare_workbooks(python_path: Path, rust_path: Path) -> dict[str, Any]:
    python = load_workbook(python_path, data_only=False, read_only=False)
    rust = load_workbook(rust_path, data_only=False, read_only=False)
    try:
        sheet_name_mismatches = python.sheetnames != rust.sheetnames
        value_mismatches = 0
        format_mismatches = 0
        header_style_mismatches = 0
        filter_mismatches = 0
        freeze_mismatches = 0
        dimension_mismatches = 0
        width_deltas: list[float] = []
        for name in python.sheetnames:
            if name not in rust.sheetnames:
                continue
            expected, actual = python[name], rust[name]
            if (expected.max_row, expected.max_column) != (actual.max_row, actual.max_column):
                dimension_mismatches += 1
            rows = max(expected.max_row, actual.max_row)
            columns = max(expected.max_column, actual.max_column)
            for row in range(1, rows + 1):
                for column in range(1, columns + 1):
                    left, right = expected.cell(row, column), actual.cell(row, column)
                    if not equal_value(left.value, right.value):
                        value_mismatches += 1
                    if left.number_format != right.number_format:
                        format_mismatches += 1
                    if row == 1:
                        left_color = str(left.fill.fgColor.rgb)[-6:]
                        right_color = str(right.fill.fgColor.rgb)[-6:]
                        if (
                            left.font.bold != right.font.bold
                            or left.fill.fill_type != right.fill.fill_type
                            or left_color != right_color
                        ):
                            header_style_mismatches += 1
            if expected.auto_filter.ref != actual.auto_filter.ref:
                filter_mismatches += 1
            if expected.freeze_panes != actual.freeze_panes:
                freeze_mismatches += 1
            for column in range(1, columns + 1):
                letter = expected.cell(1, column).column_letter
                left = expected.column_dimensions[letter].width
                right = actual.column_dimensions[letter].width
                width_deltas.append(abs(left - right))
        return {
            "status": "passed"
            if not any(
                (
                    sheet_name_mismatches,
                    value_mismatches,
                    format_mismatches,
                    header_style_mismatches,
                    filter_mismatches,
                    freeze_mismatches,
                    dimension_mismatches,
                )
            )
            else "failed",
            "sheet_name_mismatches": int(sheet_name_mismatches),
            "python_sheet_names": python.sheetnames,
            "rust_sheet_names": rust.sheetnames,
            "dimension_mismatches": dimension_mismatches,
            "value_mismatches": value_mismatches,
            "number_format_mismatches": format_mismatches,
            "header_style_mismatches": header_style_mismatches,
            "filter_mismatches": filter_mismatches,
            "freeze_mismatches": freeze_mismatches,
            "column_width_max_delta": max(width_deltas, default=0.0),
            "column_widths_exact": all(delta == 0 for delta in width_deltas),
            "data_rows_by_sheet": {
                name: max(0, python[name].max_row - 1) for name in python.sheetnames
            },
        }
    finally:
        python.close()
        rust.close()


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--audit-binary", type=Path, required=True)
    parser.add_argument("--report", type=Path)
    args = parser.parse_args()
    binary = args.audit_binary.resolve(strict=True)
    template_path = WORKSPACE / "tests" / "fixtures" / "s1_region_summary_template.xlsx"

    with tempfile.TemporaryDirectory(prefix="audit-s1-region-summary-") as temporary:
        root = Path(temporary)
        input_dir = root / "input"
        output_dir = root / "outputs"
        input_dir.mkdir()
        output_dir.mkdir()
        alpha = input_dir / "!Alpha_20260731.xlsx"
        beta = input_dir / "Beta_20260731.xlsx"
        write_source(alpha, {4: ("甲", 10), 5: (None, 0), 7: ("乙", 20)})
        write_source(beta, {4: ("丙", 30), 5: ("丁", 40)})
        input_files = [alpha, beta]
        before = {path: digest(path) for path in input_files}
        features = (
            FeatureMapping(
                "任意行汇总",
                USED_RANGE_SUMMARY_FUNCTION,
                ("任意行汇总区域",),
                False,
                "",
                "",
            ),
            FeatureMapping(
                "固定行汇总",
                FIXED_ROW_SUMMARY_FUNCTION,
                ("固定行汇总区域",),
                False,
                "",
                "",
            ),
        )

        progress: list[str] = []
        python_started = time.perf_counter()
        result = run_region_summaries(
            template_path=template_path,
            input_dir=input_dir,
            output_dir=output_dir / "python",
            config_path=template_path,
            summary_features=features,
            on_step=progress.append,
            engine_preference="自动",
        )
        python_elapsed_ms = round((time.perf_counter() - python_started) * 1000)

        rust_path = output_dir / "rust.xlsx"
        rust_started = time.perf_counter()
        command = [
            str(binary),
            "s1-region-summary",
            "--template-workbook",
            str(template_path),
            "--input-dir",
            str(input_dir),
            "--output-workbook",
            str(rust_path),
            "--feature-name",
            "任意行汇总",
            "--feature-name",
            "固定行汇总",
            "--developer-view",
        ]
        completed = subprocess.run(
            command,
            check=False,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True,
            encoding="utf-8",
        )
        rust_elapsed_ms = round((time.perf_counter() - rust_started) * 1000)
        if completed.returncode:
            raise RuntimeError(f"Rust CLI failed: {completed.stderr.strip()}")
        rust_payload = json.loads(completed.stdout)
        rust_events = [
            json.loads(line)
            for line in completed.stderr.splitlines()
            if line.strip().startswith("{")
        ]

        comparison = compare_workbooks(result.output_path, rust_path)
        width_exact = comparison["column_widths_exact"]
        status = comparison["status"]
        if status == "passed" and not width_exact:
            status = "passed_with_column_width_variance"
        report = {
            "feature": "S1-F02 named-region summary",
            "status": status,
            "fixture": "synthetic",
            "python_engine": next(
                (line.split("：", 1)[1] for line in progress if line.startswith("已连接表格引擎：")),
                "not_reported",
            ),
            "rust_reader": rust_payload["summary"]["reader"],
            "rust_writer": rust_payload["summary"]["writer"],
            "python_elapsed_ms": python_elapsed_ms,
            "rust_cli_elapsed_ms": rust_elapsed_ms,
            "rust_events": len(rust_events),
            "source_workbooks": rust_payload["summary"]["source_workbooks"],
            "source_sheets": rust_payload["summary"]["source_sheets"],
            "output_sheets": rust_payload["summary"]["output_sheets"],
            "data_rows": rust_payload["summary"]["data_rows"],
            "input_files_unchanged": all(digest(path) == value for path, value in before.items()),
            "comparison": comparison,
            "scope": "Windows synthetic formula-free XLSX; named arbitrary/fixed rows; merged multi-row header; values, types, sheet order, dimensions, number formats, header style, filter and freeze compared; no fallback template, extra metadata, history enrichment, real template or Linux runtime",
        }
        encoded = json.dumps(report, ensure_ascii=False, indent=2) + "\n"
        if args.report:
            args.report.parent.mkdir(parents=True, exist_ok=True)
            args.report.write_text(encoded, encoding="utf-8")
        print(encoded, end="")
        return 0 if status.startswith("passed") and report["input_files_unchanged"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
