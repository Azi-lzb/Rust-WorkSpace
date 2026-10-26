#!/usr/bin/env python3
"""Compare the Rust S5 run-outputs port with the Python write_run_outputs.

Both sides run the configured enabled business rule set against the same
inputs and write outputs under parallel temporary roots with a fixed run id.
Python uses `engine.run(output_dir=...)` with the openpyxl writer engine as the
styling baseline. The comparator diffs: audit copies (byte-level), the manifest
(path-normalized), the 审核结果.xlsx workbooks (values/types, hyperlinks,
widths, freeze, filter, row heights, header/body styles, numeric formats,
page setup, sheet inventory), and the ledger's stable fields. Aggregate
counts only.
The optional structure-only mode seeds Rust with Python-created snapshots and
compares migration continuity without making unrelated output deltas gate the
snapshot result.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl.utils import column_index_from_string, get_column_letter, range_boundaries

ROOT = Path(__file__).resolve().parents[3]
RUN_ID = "PARITYRUN"


def parse_current_source_spec(value: str) -> tuple[str, str, Path]:
    parts = value.split("=", 2)
    if len(parts) != 3 or not parts[0].strip() or not parts[2]:
        raise ValueError("current source must use SOURCE_ID=BUSINESS=FILE form")
    return parts[0].strip(), parts[1].strip(), Path(parts[2])


def parse_external_source_spec(value: str) -> tuple[str, str, Path]:
    parts = value.split("=", 2)
    if len(parts) != 3 or not parts[0].strip() or not parts[2]:
        raise ValueError("external source must use SOURCE_ID=BUSINESS=FILE form")
    return parts[0].strip(), parts[1].strip(), Path(parts[2])


def input_fingerprints(
    config: Path,
    source: Path,
    template_path: Path | None,
    template_dir: Path | None,
    additional_sources: list[tuple[str, str, Path]] | None = None,
    external_sources: list[tuple[str, str, Path, bool]] | None = None,
):
    inputs = {"config": config, "source": source}
    for index, (source_id, _business, path) in enumerate(additional_sources or []):
        inputs[f"current:{index}:{source_id}"] = path
    for index, (source_id, _business, path, _legacy_default) in enumerate(external_sources or []):
        inputs[f"external:{index}:{source_id}"] = path
    if template_path:
        inputs["template"] = template_path
    if template_dir:
        for path in sorted(template_dir.rglob("*.xlsx")):
            if path.is_file():
                inputs[f"template:{path.relative_to(template_dir).as_posix()}"] = path
    return {
        key: hashlib.sha256(path.read_bytes()).hexdigest()
        for key, path in inputs.items()
    }


def normalize_link(target: str, root: Path) -> str:
    text = str(target).replace("\\", "/")
    # Resolve relative copy links to (basename, fragment) pairs.
    if "#" in text:
        path_part, fragment = text.split("#", 1)
        basename = Path(path_part).name
        return f"{basename}#{fragment}"
    return Path(text).name


def normalize_value(value):
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        number = float(value)
        if number.is_integer() and abs(number) < 1e16:
            return f"{number:.1f}"
        return repr(number)
    if isinstance(value, str):
        try:
            number = float(value)
        except ValueError:
            return value
        if value.strip() != value:
            return value
        return value
    return value


def expand_widths(dimensions):
    """Expand grouped <col min max> spans (openpyxl records them under the
    first letter only) into per-letter widths."""
    widths = {}
    for dimension in dimensions.values():
        if dimension.width is None:
            continue
        min_col = dimension.min or column_index_from_string(dimension.index)
        max_col = dimension.max or min_col
        for index in range(min_col, max_col + 1):
            widths[get_column_letter(index)] = round(dimension.width, 6)
    return widths

def normalize_color(rgb):
    """Colors compare modulo the 8-digit alpha prefix: openpyxl writes 00/FF
    prefixes depending on the producing library while Excel semantics only
    cover the six-digit RGB."""
    if not rgb:
        return rgb
    text = str(rgb)
    return text[-6:]


def normalize_fit(value, fit_to_page):
    """fitToWidth=1 is the effective default under fitToPage; rust_xlsxwriter
    omits the attribute when it equals 1 while the openpyxl baseline writes it
    explicitly."""
    if fit_to_page and value is None:
        return 1
    return value

def read_workbook_facts(path: Path):
    from openpyxl import load_workbook

    workbook = load_workbook(path, data_only=False)
    facts = {}
    for sheet in workbook.worksheets:
        sheet_facts = {
            "freeze": sheet.freeze_panes,
            "filter": sheet.auto_filter.ref,
            "gridlines": sheet.sheet_view.showGridLines,
            "orientation": sheet.page_setup.orientation,
            "paper_size": sheet.page_setup.paperSize,
            "fit_to_width": normalize_fit(
                sheet.page_setup.fitToWidth,
                bool(
                    sheet.sheet_properties.pageSetUpPr
                    and sheet.sheet_properties.pageSetUpPr.fitToPage
                ),
            ),
            "fit_to_height": sheet.page_setup.fitToHeight,
            "print_titles": sheet.print_title_rows,
            "widths": expand_widths(sheet.column_dimensions),
            "row_heights": {
                key: value.height
                for key, value in sheet.row_dimensions.items()
                if value.height is not None
            },
            "header_cells": [],
            "cells": [],
            "hyperlinks": {},
        }
        header = next(sheet.iter_rows(min_row=1, max_row=1), ())
        for cell in header:
            font = cell.font
            fill = cell.fill
            sheet_facts["header_cells"].append({
                "value": cell.value,
                "font": (
                    font.name if font else None,
                    font.size if font else None,
                    bool(font.bold) if font else None,
                    normalize_color(font.color.rgb if font and font.color else None),
                ),
                "fill": (
                    normalize_color(fill.fgColor.rgb if fill and fill.fgColor else None),
                ),
                "align": (
                    cell.alignment.horizontal,
                    cell.alignment.vertical,
                    bool(cell.alignment.wrap_text),
                ),
            })
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if cell.value is None and cell.hyperlink is None:
                    continue
                entry = {
                    "ref": cell.coordinate,
                    "value": normalize_value(cell.value),
                    "type": cell.data_type,
                    "format": cell.number_format,
                }
                if cell.hyperlink is not None:
                    link = normalize_link(cell.hyperlink.target, path.parent)
                    # openpyxl accepts a URI fragment (`file.xlsx#Sheet!A1`)
                    # while rust_xlsxwriter stores the same destination in
                    # the OOXML hyperlink `location` attribute. Compare the
                    # resolved workbook destination, not those encodings.
                    if cell.hyperlink.location:
                        link = f"{link}#{cell.hyperlink.location}"
                    entry["link"] = link
                if cell.value is not None:
                    font = cell.font
                    entry["font"] = (
                        font.name if font else None,
                        font.size if font else None,
                    )
                    entry["align"] = (
                        cell.alignment.horizontal,
                        cell.alignment.vertical,
                        bool(cell.alignment.wrap_text),
                    )
                sheet_facts["cells"].append(entry)
        facts[sheet.title] = sheet_facts
    workbook.close()
    return facts


def diff(left, right, path: str, diffs: list[str]) -> None:
    if isinstance(left, dict) and isinstance(right, dict):
        for key in sorted(set(left) | set(right)):
            diff(left.get(key), right.get(key), f"{path}.{key}", diffs)
    elif isinstance(left, list) and isinstance(right, list):
        if len(left) != len(right):
            diffs.append(f"{path}: length {len(left)} vs {len(right)}")
        for index in range(min(len(left), len(right))):
            diff(left[index], right[index], f"{path}[{index}]", diffs)
    elif left != right:
        diffs.append(path)


def cell_value_differences(left: dict, right: dict) -> list[dict]:
    differences = []
    for sheet in sorted(set(left) | set(right)):
        left_cells = {cell["ref"]: cell for cell in left.get(sheet, {}).get("cells", [])}
        right_cells = {cell["ref"]: cell for cell in right.get(sheet, {}).get("cells", [])}
        for coordinate in sorted(set(left_cells) | set(right_cells)):
            left_cell = left_cells.get(coordinate, {})
            right_cell = right_cells.get(coordinate, {})
            if left_cell.get("value") != right_cell.get("value"):
                differences.append({
                    "sheet": sheet,
                    "cell": coordinate,
                    "python_type": left_cell.get("type", "missing"),
                    "rust_type": right_cell.get("type", "missing"),
                })
    return differences


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source-root", type=Path, required=True)
    parser.add_argument("--config-workbook", type=Path, required=True)
    parser.add_argument("--audit-copy", type=Path, required=True)
    parser.add_argument("--period", default="2026-07-31")
    parser.add_argument("--institution-code", default="PARITY")
    parser.add_argument("--institution-name", default="PARITY")
    parser.add_argument("--source-id", default="CURRENT")
    parser.add_argument("--business", action="append", default=[])
    parser.add_argument("--source-business", default="")
    parser.add_argument(
        "--current-source",
        action="append",
        default=[],
        metavar="SOURCE_ID=BUSINESS=FILE",
    )
    parser.add_argument("--external-source", type=Path)
    parser.add_argument("--external-source-id")
    parser.add_argument("--external-source-business", default="")
    parser.add_argument(
        "--external-source-binding",
        action="append",
        default=[],
        metavar="SOURCE_ID=BUSINESS=FILE",
    )
    parser.add_argument("--default-external-source")
    parser.add_argument("--template-path", type=Path)
    parser.add_argument("--template-dir", type=Path)
    parser.add_argument(
        "--rust-cli",
        action="store_true",
        help="compare against the product s5-audit-run CLI rather than its service harness",
    )
    parser.add_argument(
        "--rust-bin",
        type=Path,
        help="use this already-built open-data-audit binary for the CLI comparison",
    )
    parser.add_argument(
        "--structure-only",
        action="store_true",
        help="report parity for structure snapshot events without gating on other output deltas",
    )
    parser.add_argument("--workspace", type=Path, default=ROOT)
    parser.add_argument("--report", type=Path, required=True)
    args = parser.parse_args()

    source_root = args.source_root.resolve(strict=True)
    workspace = args.workspace.resolve(strict=True)
    report_path = args.report.resolve()
    template_path = args.template_path.resolve(strict=True) if args.template_path else None
    template_dir = args.template_dir.resolve(strict=True) if args.template_dir else None
    businesses = args.business or ["单位贷款"]
    primary_business = args.source_business or businesses[0]
    additional_sources = []
    known_source_ids = {args.source_id}
    try:
        for spec in args.current_source:
            source_id, business, path = parse_current_source_spec(spec)
            if source_id in known_source_ids:
                raise ValueError(f"duplicate current source id: {source_id}")
            known_source_ids.add(source_id)
            additional_sources.append((source_id, business, path.resolve(strict=True)))
    except (OSError, ValueError) as error:
        parser.error(str(error))
    external_sources: list[tuple[str, str, Path, bool]] = []
    known_external_ids: set[str] = set()
    if args.external_source:
        if not args.external_source_id or not args.external_source_id.strip():
            parser.error("--external-source requires --external-source-id")
        external_sources.append((
            args.external_source_id.strip(),
            args.external_source_business.strip(),
            args.external_source.resolve(strict=True),
            True,
        ))
        known_external_ids.add(args.external_source_id.strip())
    elif args.external_source_id:
        parser.error("--external-source-id requires --external-source")
    try:
        for spec in args.external_source_binding:
            source_id, business, path = parse_external_source_spec(spec)
            if source_id in known_external_ids:
                raise ValueError(f"duplicate external source id: {source_id}")
            known_external_ids.add(source_id)
            external_sources.append((source_id, business, path.resolve(strict=True), False))
    except (OSError, ValueError) as error:
        parser.error(str(error))
    default_external_source = args.default_external_source or (
        args.external_source_id.strip() if args.external_source_id else ""
    )
    if default_external_source and default_external_source not in known_external_ids:
        parser.error("--default-external-source must name a bound external source")
    if args.rust_bin:
        args.rust_bin = args.rust_bin.resolve(strict=True)
    use_rust_cli = (
        args.rust_cli or bool(additional_sources) or bool(external_sources) or args.rust_bin is not None
    )
    inputs_before = input_fingerprints(
        args.config_workbook.resolve(strict=True),
        args.audit_copy.resolve(strict=True),
        template_path,
        template_dir,
        additional_sources,
        external_sources,
    )
    python_snapshot_baseline_loaded_by_rust = False
    # The comparison imports the legacy baseline from its read-only source tree.
    sys.dont_write_bytecode = True
    sys.path.insert(0, str(source_root / "core" / "src"))
    os.environ["AUDIT3_WRITER_ENGINE"] = "openpyxl"

    with tempfile.TemporaryDirectory(prefix="audit-s5-outputs-") as temporary:
        directory = Path(temporary)
        py_schema_dir = directory / "py_schema_snapshots"
        rs_data_dir = directory / "rust_data"
        rs_schema_dir = rs_data_dir / "s5" / ".schema_snapshots"
        py_schema_dir.mkdir(parents=True)
        rs_schema_dir.mkdir(parents=True)
        shared_schema_key = os.urandom(32)
        (py_schema_dir / ".schema-key").write_bytes(shared_schema_key)
        (rs_schema_dir / ".schema-key").write_bytes(shared_schema_key)
        config_copy = directory / "config.xlsx"
        source_copy = directory / "source.xlsx"
        shutil.copy2(args.config_workbook.resolve(strict=True), config_copy)
        shutil.copy2(args.audit_copy.resolve(strict=True), source_copy)
        copied_current_sources = []
        for index, (source_id, business, source_path) in enumerate(additional_sources):
            suffix = source_path.suffix or ".xlsx"
            copied_path = directory / f"current-source-{index}{suffix}"
            shutil.copy2(source_path, copied_path)
            copied_current_sources.append((source_id, business, copied_path))
        copied_external_sources = []
        for index, (source_id, business, source_path, legacy_default) in enumerate(external_sources):
            suffix = source_path.suffix or ".xlsx"
            copied_path = directory / f"external-source-{index}{suffix}"
            shutil.copy2(source_path, copied_path)
            copied_external_sources.append((source_id, business, copied_path, legacy_default))
        py_root = directory / "py_out"
        rs_root = directory / "rs_out"

        # --- Python side: full run + write_run_outputs ---
        from s1audit_configured.engine import AuditEngine
        from s1audit_configured.model import RuntimeContext, SourceBinding
        from s1audit_configured.runtime import RuntimePaths

        engine = AuditEngine(config_copy)
        current_source_bindings = {
            args.source_id: SourceBinding(
                args.source_id,
                source_copy,
                business=primary_business,
                period=args.period,
                institution_code=args.institution_code,
                institution_name=args.institution_name,
            )
        }
        for source_id, business, copied_path in copied_current_sources:
            current_source_bindings[source_id] = SourceBinding(
                source_id,
                copied_path,
                business=business,
                period=args.period,
                institution_code=args.institution_code,
                institution_name=args.institution_name,
            )
        external_source_bindings = {
            source_id: SourceBinding(
                source_id,
                copied_path,
                business=business,
                period=args.period,
                institution_code=args.institution_code,
                institution_name=args.institution_name,
            )
            for source_id, business, copied_path, _legacy_default in copied_external_sources
        }
        context = RuntimeContext(
            period=args.period,
            institution_code=args.institution_code,
            institution_name=args.institution_name,
            current_sources=current_source_bindings,
            external_sources=external_source_bindings,
            default_external_source=default_external_source,
        )
        result = engine.run(
            context,
            output_dir=py_root,
            run_id=RUN_ID,
            businesses=businesses,
            schema_snapshot_dir=py_schema_dir,
            persist_schema_snapshots=False,
            runtime_paths=RuntimePaths(
                template_path=template_path,
                template_dir=template_dir,
            ),
        )
        py_run_dir = py_root / RUN_ID
        if args.structure_only:
            python_snapshot_files = list(py_schema_dir.glob("*.json"))
            for snapshot_path in python_snapshot_files:
                shutil.copy2(snapshot_path, rs_schema_dir / snapshot_path.name)
            python_snapshot_baseline_loaded_by_rust = bool(python_snapshot_files)

        # --- Rust side ---
        request = {
            "config": str(config_copy),
            "source": str(source_copy),
            "template_path": str(template_path) if template_path else None,
            "template_dir": str(template_dir) if template_dir else None,
            "period": args.period,
            "institution_code": args.institution_code,
            "institution_name": args.institution_name,
            "source_id": args.source_id,
            "run_id": RUN_ID,
            "output_root": str(rs_root),
        }
        input_path = directory / "request.json"
        input_path.write_text(json.dumps(request, ensure_ascii=False), encoding="utf-8")
        rust_report = directory / "rust_summary.json"
        environment = os.environ.copy()
        environment["BASE_AUDIT_DATA_DIR"] = str(rs_data_dir)
        if use_rust_cli:
            rust_command = (
                [str(args.rust_bin)]
                if args.rust_bin
                else [
                    "cargo", "run", "--offline", "--locked", "--quiet", "-p", "audit-core",
                    "--bin", "open-data-audit", "--",
                ]
            ) + [
                "s5-audit-run",
                "--config", str(config_copy), "--source", str(source_copy),
                "--source-id", args.source_id, "--source-business", primary_business,
                "--period", args.period,
                "--org-code", args.institution_code, "--org-name", args.institution_name,
                "--output-dir", str(rs_root), "--run-id", RUN_ID,
            ]
            for business in businesses:
                rust_command.extend(["--business", business])
            for source_id, business, copied_path in copied_current_sources:
                rust_command.extend([
                    "--current-source",
                    f"{source_id}={business}={copied_path}",
                ])
            for source_id, business, copied_path, legacy_default in copied_external_sources:
                if legacy_default:
                    rust_command.extend([
                        "--external-source", str(copied_path),
                        "--external-source-id", source_id,
                    ])
                else:
                    rust_command.extend([
                        "--external-source-binding",
                        f"{source_id}={business}={copied_path}",
                    ])
            if default_external_source and default_external_source != args.external_source_id:
                rust_command.extend([
                    "--default-external-source-id", default_external_source,
                ])
            if template_path:
                rust_command.extend(["--template-path", str(template_path)])
            if template_dir:
                rust_command.extend(["--template-dir", str(template_dir)])
            completed = subprocess.run(
                rust_command,
                cwd=workspace,
                env=environment,
                capture_output=True,
                text=True,
                encoding="utf-8",
                errors="replace",
                check=False,
            )
            try:
                summary = json.loads(completed.stdout) if completed.returncode == 0 else None
            except json.JSONDecodeError:
                summary = None
            rs_run_dir = rs_root / RUN_ID
        else:
            environment["AUDIT_S5_OUTPUTS_INPUT"] = str(input_path)
            environment["AUDIT_S5_OUTPUTS_REPORT"] = str(rust_report)
            completed = subprocess.run(
                [
                    "cargo", "test", "-p", "audit-core", "--test", "s5_outputs_parity",
                    "--locked", "--", "--ignored", "--nocapture",
                ],
                cwd=workspace,
                env=environment,
                capture_output=True,
                text=True,
                encoding="utf-8",
                errors="replace",
                check=False,
            )
            summary = (
                json.loads(rust_report.read_text(encoding="utf-8"))
                if completed.returncode == 0 and rust_report.is_file()
                else None
            )
            rs_run_dir = None
        if completed.returncode != 0 or summary is None:
            print(json.dumps({
                "status": "rust_runner_failed",
                "exit_code": completed.returncode,
                "stderr_tail": completed.stderr[-1500:].replace(str(workspace), "[WORKSPACE]"),
            }))
            return 1
        if rs_run_dir is None:
            rs_run_dir = Path(summary["run_dir"])

        # --- compare artifacts ---
        diffs: list[str] = []

        # 1. audit copies: same file set, byte-identical
        py_copies = sorted(p.name for p in (py_run_dir / "审核副本").iterdir())
        rs_copies = sorted(p.name for p in (rs_run_dir / "审核副本").iterdir())
        if py_copies != rs_copies:
            diffs.append(f"copies.set: {py_copies} vs {rs_copies}")
        else:
            for name in py_copies:
                if (py_run_dir / "审核副本" / name).read_bytes() != (
                    rs_run_dir / "审核副本" / name
                ).read_bytes():
                    diffs.append(f"copies.bytes:{name}")

        # 2. workbooks
        py_book = read_workbook_facts(py_run_dir / "审核结果.xlsx")
        rs_book = read_workbook_facts(rs_run_dir / "审核结果.xlsx")
        value_differences = cell_value_differences(py_book, rs_book)
        diff(py_book, rs_book, "workbook", diffs)

        # 3. ledger stable fields
        py_ledger = json.loads((py_run_dir / "运行总账.json").read_text(encoding="utf-8"))
        rs_ledger = json.loads((rs_run_dir / "运行总账.json").read_text(encoding="utf-8"))
        stable = {
            "run_id": (py_ledger.get("run_id"), rs_ledger.get("run_id")),
            "status_counts": (py_ledger.get("status_counts"), rs_ledger.get("status_counts")),
            "run_status": (py_ledger.get("run_status"), rs_ledger.get("run_status")),
        }
        for field, (left, right) in stable.items():
            if left != right:
                diffs.append(f"ledger.{field}")

        structure_fields = (
            "snapshot_id", "role", "source_id", "business", "sheet", "table_id",
            "row_header_width", "configured_row_path_count", "configured_column_path_count",
            "matched_row_path_count", "matched_column_path_count", "row_region", "column_region",
            "status", "detail",
        )

        def canonical_region_reference(reference):
            if not isinstance(reference, str) or not reference:
                return reference
            sheet, separator, coordinates = reference.rpartition("!")
            if not separator:
                return reference.replace("$", "").upper()
            if len(sheet) >= 2 and sheet[0] == "'" and sheet[-1] == "'":
                sheet = sheet[1:-1].replace("''", "'")
            return f"{sheet}!{coordinates.replace('$', '').upper()}"

        def structure_facts(ledger, normalize_python_baseline=False):
            events = ledger.get("structure_checks", [])
            return sorted(
                ({
                    field: canonical_region_reference(event.get(field))
                    if field in {"row_region", "column_region"}
                    else "UNCHANGED"
                    if normalize_python_baseline and field == "status"
                    and event.get(field) == "BASELINE_CREATED"
                    else event.get(field)
                    for field in structure_fields
                } for event in events),
                key=lambda event: event.get("snapshot_id") or "",
            )

        py_structure = structure_facts(py_ledger, normalize_python_baseline=args.structure_only)
        rs_structure = structure_facts(rs_ledger)
        python_events_by_id = {
            event.get("snapshot_id"): event for event in py_ledger.get("structure_checks", [])
        }
        rust_events_by_id = {
            event.get("snapshot_id"): event for event in rs_ledger.get("structure_checks", [])
        }
        region_reference_representation_difference_count = sum(
            python_event.get(field) != rust_events_by_id[snapshot_id].get(field)
            and canonical_region_reference(python_event.get(field))
            == canonical_region_reference(rust_events_by_id[snapshot_id].get(field))
            for snapshot_id, python_event in python_events_by_id.items()
            if snapshot_id in rust_events_by_id
            for field in ("row_region", "column_region")
        )
        region_geometry_differences = []
        for snapshot_id, python_event in python_events_by_id.items():
            rust_event = rust_events_by_id.get(snapshot_id)
            if rust_event is None:
                continue
            for field in ("row_region", "column_region"):
                python_reference = python_event.get(field)
                rust_reference = rust_event.get(field)
                if python_reference == rust_reference:
                    continue

                def bounds(reference):
                    if not isinstance(reference, str) or not reference:
                        return None
                    _, separator, coordinates = reference.rpartition("!")
                    if not separator:
                        coordinates = reference
                    try:
                        return range_boundaries(coordinates.replace("$", ""))
                    except ValueError:
                        return None

                if bounds(python_reference) != bounds(rust_reference):
                    region_geometry_differences.append({
                        "field": field,
                        "python_bounds": bounds(python_reference),
                        "rust_bounds": bounds(rust_reference),
                    })
        structure_diffs: list[str] = []
        diff(py_structure, rs_structure, "ledger.structure_checks", structure_diffs)
        if args.structure_only and (not py_structure or not rs_structure):
            structure_diffs.append("ledger.structure_checks.empty")
        structure_status_counts: dict[str, dict[str, int]] = {}
        for engine_name, events in (("python", py_ledger.get("structure_checks", [])),
                                    ("rust", rs_ledger.get("structure_checks", []))):
            counts: dict[str, int] = {}
            for event in events:
                status = str(event.get("status", "missing"))
                counts[status] = counts.get(status, 0) + 1
            structure_status_counts[engine_name] = counts

        # 4. manifest: roles/order/snapshot flags (paths normalized to names)
        def manifest_facts(ledger):
            facts = []
            for entry in ledger.get("audit_copy_manifest", []):
                facts.append({
                    "role": entry.get("role"),
                    "source_id": entry.get("source_id"),
                    "copy_name": Path(entry.get("copy_path", "")).name,
                    "source_name": Path(entry.get("source_path", "")).name,
                    "matches_snapshot": entry.get("matches_snapshot"),
                    "sha_before": entry.get("sha256_before"),
                    "sha_after": entry.get("sha256_after"),
                    "sha_copy": entry.get("sha256_copy"),
                })
            return facts

        diff(
            manifest_facts(py_ledger),
            manifest_facts(rs_ledger),
            "ledger.manifest",
            diffs,
        )

        inputs_unchanged = inputs_before == input_fingerprints(
            args.config_workbook.resolve(strict=True),
            args.audit_copy.resolve(strict=True),
            template_path,
            template_dir,
            additional_sources,
            external_sources,
        )
        if not inputs_unchanged:
            diffs.append("inputs.modified_during_run")

        selected_diffs = list(structure_diffs) if args.structure_only else list(diffs) + structure_diffs
        if not inputs_unchanged and args.structure_only:
            selected_diffs.append("inputs.modified_during_run")

        triggered = summary.get("triggered_rows")
        failures = summary.get("failure_rows")
        print(json.dumps({
            "rust_triggered_rows": triggered,
            "rust_failure_rows": failures,
            "py_status_counts": py_ledger.get("status_counts"),
            "structure_check_count": len(py_structure),
            "structure_diff_count": len(structure_diffs),
            "region_reference_representation_difference_count": region_reference_representation_difference_count,
            "region_geometry_differences": region_geometry_differences,
            "structure_diff_paths": sorted(set(structure_diffs)),
            "diff_count": len(selected_diffs),
            "diff_sample": [] if args.structure_only else sorted(set(selected_diffs))[:15],
        }, ensure_ascii=False))

    report = {
        "schema_version": 1,
        "suite": "audit.s5_run_outputs_parity",
        "comparison_scope": "structure_snapshots_only" if args.structure_only else "full_run_output_and_structure_snapshot",
        "status": "completed" if not selected_diffs else "mismatch",
        "sample_scope": {
            "business": businesses[0] if len(businesses) == 1 else "multiple",
            "business_count": len(businesses),
            "current_source_count": 1 + len(additional_sources),
            "additional_current_sources_count": len(additional_sources),
            "external_source_count": len(external_sources),
            "external_source_ids": [source_id for source_id, _business, _path, _legacy in external_sources],
            "default_external_source": default_external_source or None,
            "run_id": RUN_ID,
            "files_copied_to_temporary": True,
            "template_path_supplied": template_path is not None,
            "template_dir_supplied": template_dir is not None,
            "python_writer_engine": "openpyxl (styling baseline)",
            "rust_entry": "s5-audit-run CLI" if use_rust_cli else "service test harness",
        },
        "comparison": {
            "diff_count": len(selected_diffs),
            "diff_paths_sample": [] if args.structure_only else sorted(set(selected_diffs))[:20],
            "structure_check_count": len(py_structure),
            "structure_check_diff_count": len(structure_diffs),
            "region_reference_representation_difference_count": region_reference_representation_difference_count,
            "structure_check_status_counts": structure_status_counts,
            "non_structure_output_diff_count": len(diffs),
            "cell_value_difference_count": len(value_differences),
            "cell_value_differences": [] if args.structure_only else value_differences,
            "status_counts": {
                "python": py_ledger.get("status_counts"),
                "rust": rs_ledger.get("status_counts"),
            },
        },
        "verification": {
            "rust_runner": (
                "open-data-audit s5-audit-run CLI"
                if use_rust_cli
                else "audit-core/tests/s5_outputs_parity.rs"
            ),
            "python_baseline": (
                "s1audit_configured.engine.run(output_dir=...) with "
                "AUDIT3_WRITER_ENGINE=openpyxl"
            ),
            "python_baseline_version": sys.version.split()[0],
            "inputs_unchanged": inputs_unchanged,
            "raw_values_persisted": False,
            "python_snapshot_baseline_loaded_by_rust": python_snapshot_baseline_loaded_by_rust,
            "capability_status": "unverified",
            "production_route_eligible": False,
        },
        "limitations": [
            "Conditional-format and data-quality tables are compared empty; "
            "their populated paths belong to the S1 conditional-format bridge.",
            "Structure snapshots use one shared isolated HMAC key; structure-only mode seeds Rust with Python-created temporary snapshots and expects Rust UNCHANGED events.",
            "Hash-integrity downgrade remains outside this comparison.",
            "When comparison_scope is structure_snapshots_only, other output deltas are counted but do not affect status.",
            "Colors compare modulo the 8-digit alpha prefix and fitToWidth=1 "
            "compares equal to an omitted attribute under fitToPage: both are "
            "rust_xlsxwriter/openpyxl serializer representation deltas with "
            "identical Excel semantics.",
        ],
    }
    report_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.write_text(
        json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    print(json.dumps({"status": report["status"], "diffs": len(selected_diffs)}))
    return 0 if not selected_diffs else 1


if __name__ == "__main__":
    raise SystemExit(main())
