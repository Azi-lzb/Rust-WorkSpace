"""Compare the Rust S2 writer candidate with the read-only Python exporter."""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import sys
import tempfile
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def decode_value(value: Any) -> Any:
    if isinstance(value, dict) and value.get("type") == "decimal":
        parts = value["value"]
        return Decimal((parts["sign"], tuple(parts["digits"]), parts["exponent"]))
    if isinstance(value, dict) and value.get("type") in {
        "empty", "string", "number", "integer", "boolean", "datetime_serial",
        "datetime_iso", "duration_iso", "error",
    }:
        return decode_value(value.get("value"))
    return value


def findings_from_fixture(source_root: Path, rows: list[dict[str, Any]]):
    sys.path.insert(0, str(source_root / "core" / "src"))
    from base_audit.systems.s2_report_collection.models import (  # type: ignore[import-not-found]
        AuditFinding,
        RelatedRecord,
        SourceRef,
    )

    findings = []
    for row in rows:
        values = {key: decode_value(value) for key, value in row.items()}
        values["form_codes"] = tuple(values.get("form_codes", ()))
        values["source_refs"] = tuple(
            SourceRef(**source) for source in values.get("source_refs", ())
        )
        values["related_records"] = tuple(
            RelatedRecord(
                indicator_code=record.get("indicator_code", ""),
                indicator_name=record.get("indicator_name", ""),
                form_code=record.get("form_code", ""),
                period=record.get("period", ""),
                value=decode_value(record.get("value")),
                source=SourceRef(**record.get("source", {
                    "source_file": "", "source_sheet": "", "source_row": 0,
                })),
            )
            for record in values.get("related_records", ())
        )
        findings.append(AuditFinding(**values))
    return findings


def color_key(color):
    if color is None:
        return None
    rgb = color.rgb if color.type == "rgb" and isinstance(color.rgb, str) else None
    return (
        color.type,
        rgb[-6:].upper() if rgb else None,
        color.theme,
        color.indexed,
        round(float(color.tint or 0), 6),
    )


def style_key(cell):
    return {
        "bold": bool(cell.font.bold),
        "font_color": color_key(cell.font.color),
        "fill_type": cell.fill.fill_type,
        "fill_color": color_key(cell.fill.fgColor),
        "horizontal": cell.alignment.horizontal,
        "vertical": cell.alignment.vertical,
        "wrap_text": bool(cell.alignment.wrap_text),
        "number_format": cell.number_format,
    }


def value_matches(left, right, tolerance: float) -> bool:
    if left is None or left == "":
        return right is None or right == ""
    if right is None or right == "":
        return False
    if isinstance(left, bool) or isinstance(right, bool):
        return type(left) is type(right) and left == right
    if isinstance(left, (int, float)) and isinstance(right, (int, float)):
        left_number, right_number = float(left), float(right)
        if left_number == right_number:
            return True
        scale = max(abs(left_number), abs(right_number))
        return scale > 0 and abs(left_number - right_number) / scale <= tolerance
    return type(left) is type(right) and left == right


def column_width(sheet, column: int) -> float:
    """Resolve OOXML min/max column spans instead of assuming one key per column."""
    for dimension in sheet.column_dimensions.values():
        first = dimension.min
        last = dimension.max
        if first is not None and last is not None and first <= column <= last:
            return dimension.width
    return sheet.column_dimensions[get_column_letter(column)].width


def compare_workbooks(python_path: Path, rust_path: Path, tolerance: float):
    python_book = load_workbook(python_path, data_only=False)
    rust_book = load_workbook(rust_path, data_only=False)
    differences: list[dict[str, Any]] = []
    stats = {
        "sheets": 0,
        "cells": 0,
        "numeric_cells": 0,
        "formulas": 0,
        "cell_styles": 0,
        "column_widths": 0,
        "hidden_rows": 0,
        "worksheet_settings": 0,
    }

    def record(kind: str, location: str, expected: Any, actual: Any):
        if len(differences) < 100:
            differences.append({
                "kind": kind,
                "location": location,
                "python": expected,
                "rust": actual,
            })

    try:
        if python_book.sheetnames != rust_book.sheetnames:
            record("sheet_order", "workbook", python_book.sheetnames, rust_book.sheetnames)
        stats["sheets"] = min(len(python_book.worksheets), len(rust_book.worksheets))

        for python_sheet, rust_sheet in zip(python_book.worksheets, rust_book.worksheets):
            sheet_name = python_sheet.title
            max_row = max(python_sheet.max_row, rust_sheet.max_row)
            max_column = max(python_sheet.max_column, rust_sheet.max_column)
            if (python_sheet.max_row, python_sheet.max_column) != (rust_sheet.max_row, rust_sheet.max_column):
                record(
                    "dimensions", sheet_name,
                    [python_sheet.max_row, python_sheet.max_column],
                    [rust_sheet.max_row, rust_sheet.max_column],
                )

            for row in range(1, max_row + 1):
                for column in range(1, max_column + 1):
                    expected = python_sheet.cell(row, column)
                    actual = rust_sheet.cell(row, column)
                    address = f"{sheet_name}!{expected.coordinate}"
                    stats["cells"] += 1
                    if expected.data_type == "f" or actual.data_type == "f":
                        stats["formulas"] += 1
                        if expected.data_type != actual.data_type or expected.value != actual.value:
                            record("formula", address, [expected.data_type, expected.value], [actual.data_type, actual.value])
                    elif isinstance(expected.value, (int, float)) and not isinstance(expected.value, bool):
                        stats["numeric_cells"] += 1
                        if not value_matches(expected.value, actual.value, tolerance):
                            record("value", address, expected.value, actual.value)
                    elif not value_matches(expected.value, actual.value, tolerance):
                        record("value", address, expected.value, actual.value)

                    if style_key(expected) != style_key(actual):
                        stats["cell_styles"] += 1
                        record("style", address, style_key(expected), style_key(actual))

            expected_sheet_settings = {
                "freeze_panes": python_sheet.freeze_panes,
                "auto_filter": python_sheet.auto_filter.ref,
                "page_margins": tuple(
                    getattr(python_sheet.page_margins, name)
                    for name in ("left", "right", "top", "bottom", "header", "footer")
                ),
            }
            actual_sheet_settings = {
                "freeze_panes": rust_sheet.freeze_panes,
                "auto_filter": rust_sheet.auto_filter.ref,
                "page_margins": tuple(
                    getattr(rust_sheet.page_margins, name)
                    for name in ("left", "right", "top", "bottom", "header", "footer")
                ),
            }
            if expected_sheet_settings != actual_sheet_settings:
                stats["worksheet_settings"] += 1
                record("worksheet_settings", sheet_name, expected_sheet_settings, actual_sheet_settings)

            for row in range(1, max_row + 1):
                expected_hidden = bool(python_sheet.row_dimensions[row].hidden)
                actual_hidden = bool(rust_sheet.row_dimensions[row].hidden)
                if expected_hidden != actual_hidden:
                    stats["hidden_rows"] += 1
                    record("row_hidden", f"{sheet_name}!{row}", expected_hidden, actual_hidden)

            for column in range(1, max_column + 1):
                letter = get_column_letter(column)
                expected_width = column_width(python_sheet, column)
                actual_width = column_width(rust_sheet, column)
                if not math.isclose(expected_width, actual_width, rel_tol=0, abs_tol=0.1):
                    stats["column_widths"] += 1
                    record("column_width", f"{sheet_name}!{letter}", expected_width, actual_width)
    finally:
        python_book.close()
        rust_book.close()

    return stats, differences


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source-root", type=Path, required=True)
    parser.add_argument("--rust-workbook", type=Path, required=True)
    parser.add_argument("--report-dir", type=Path)
    args = parser.parse_args()

    source_root = args.source_root.resolve()
    rust_workbook = args.rust_workbook.resolve()
    workspace = Path(__file__).resolve().parents[2]
    fixture_path = workspace / "tests" / "fixtures" / "s2_exporter" / "cases.json"
    report_dir = args.report_dir or workspace / "probe" / "reports"
    report_dir.mkdir(parents=True, exist_ok=True)
    fixture = json.loads(fixture_path.read_text(encoding="utf-8"))
    findings = findings_from_fixture(source_root, fixture["findings"])

    with tempfile.TemporaryDirectory(prefix="s2-exporter-python-") as temp_name:
        python_output_dir = Path(temp_name)
        sys.path.insert(0, str(source_root / "core" / "src"))
        from base_audit.systems.s2_report_collection.exporter import ExcelExporter  # type: ignore[import-not-found]

        python_paths = ExcelExporter().export(
            output_dir=python_output_dir,
            current_records=[],
            previous_records=[],
            findings=findings,
            run_info={},
            hide_unchanged_period_rows=fixture["hide_unchanged_period_rows"],
            hide_matching_external_rows=fixture["hide_matching_external_rows"],
        )
        python_workbook = Path(python_paths.result)
        stats, differences = compare_workbooks(python_workbook, rust_workbook, 1e-15)

        formula_cache = load_workbook(rust_workbook, data_only=True)
        try:
            rust_formula_cache = formula_cache["校验规则"]["J2"].value
        finally:
            formula_cache.close()

        source_files = {
            "exporter.py": source_root / "core/src/base_audit/systems/s2_report_collection/exporter.py",
            "models.py": source_root / "core/src/base_audit/systems/s2_report_collection/models.py",
        }
        python_report = {
            "status": "python_baseline_generated",
            "baseline_python": sys.version.split()[0],
            "fixture": "tests/fixtures/s2_exporter/cases.json",
            "fixture_sha256": sha256(fixture_path),
            "case_findings": len(findings),
            "sheet_names": ["环比规则", "外部核对规则", "校验规则"],
            "source_sha256": {name: sha256(path) for name, path in source_files.items()},
        }
        (report_dir / "s2_exporter_python_baseline.json").write_text(
            json.dumps(python_report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
        )

        parity_report = {
            "status": "passed" if not differences and rust_formula_cache is None else "failed",
            "python_baseline": python_report,
            "rust_workbook": str(rust_workbook),
            "rust_workbook_sha256": sha256(rust_workbook),
            "numeric_relative_tolerance": 1e-15,
            "formula_cache_empty": rust_formula_cache is None,
            "stats": stats,
            "difference_count": len(differences),
            "differences": differences,
        }
        (report_dir / "s2_exporter_parity.json").write_text(
            json.dumps(parity_report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
        )

    print(json.dumps(parity_report, ensure_ascii=False, indent=2))
    return 0 if parity_report["status"] == "passed" else 1


if __name__ == "__main__":
    raise SystemExit(main())
