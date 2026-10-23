#!/usr/bin/env python3
"""Create sanitized XLSX inputs for the S2 external workbook importer parity probe."""

from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils.datetime import CALENDAR_MAC_1904


ROOT = Path(__file__).resolve().parents[2]
FIXTURES = ROOT / "tests" / "fixtures" / "s2_external_import"


def workbook(path: Path) -> Workbook:
    result = Workbook()
    result.remove(result.active)
    result.create_sheet("集中系统数据")
    return result


def save_plain() -> None:
    book = workbook(FIXTURES / "plain_layout.xlsx")
    sheet = book["集中系统数据"]
    sheet.append(["导出说明：普通格式"])
    sheet.append([])
    sheet.append(["报表项目", "余额", "指标X", "版本号"])
    sheet.append(["资产A", "1,234.50", 3, "v1"])
    sheet.append([])
    sheet.append([None, None, 2, None])
    sheet.append(["资产B", None, "无效", None])
    sheet.append(["资产A", "9.00", 5, None])

    preferred = book.create_sheet("机构参照表")
    preferred.append(["统一社会信用代码", "社会信用代码", "报表项目", "禁用"])
    preferred.append(["legacy-a", "code-a", "映射A", "否"])
    preferred.append(["legacy-b", "code-a", "映射A后行", "False"])
    preferred.append(["legacy-c", "code-b", "已禁用映射", "yes"])
    preferred.append(["legacy-d", "code-c", "映射C", None])

    legacy = book.create_sheet("参照表")
    legacy.append(["社会信用代码", "报表项目"])
    legacy.append(["code-a", "旧表映射不应优先"])
    book.save(FIXTURES / "plain_layout.xlsx")


def save_plain_fallback() -> None:
    book = workbook(FIXTURES / "plain_missing_label.xlsx")
    sheet = book["集中系统数据"]
    sheet.append(["项目编号", "余额"])
    sheet.append(["A01", 10])
    book.save(FIXTURES / "plain_missing_label.xlsx")


def save_plain_zero_headers() -> None:
    book = workbook(FIXTURES / "plain_zero_headers.xlsx")
    sheet = book["集中系统数据"]
    sheet.append(["行名称", 0, False, "余额"])
    sheet.append(["项目A", 123, 456, 99])
    book.save(FIXTURES / "plain_zero_headers.xlsx")


def save_plain_legacy_reference() -> None:
    book = workbook(FIXTURES / "plain_legacy_reference.xlsx")
    sheet = book["集中系统数据"]
    sheet.append(["行名称", "余额"])
    sheet.append(["项目A", 12])
    reference = book.create_sheet("参照表")
    reference.append(["社会信用代码", "报表项目", "禁用"])
    reference.append(["code-old", "旧版映射", "否"])
    book.save(FIXTURES / "plain_legacy_reference.xlsx")


def save_time_series() -> None:
    book = workbook(FIXTURES / "time_series.xlsx")
    sheet = book["集中系统数据"]
    sheet.append(["导出说明：时序格式"])
    sheet.append([])
    sheet.append(["数据日期", "行名称", "报表名称", "余额", "指标B", "行顺序号"])
    sheet.append([datetime(2026, 6, 30), "银行甲", "日报", "100.50", 3, 1])
    sheet.append(["2026-07-31T12:05:00", "银行乙", "日报", None, "无效", 2])
    sheet.append([20260630, "银行丙", "日报", 0, 8, 3])
    sheet.append([])
    sheet.append(["20260831", "银行丁", "日报", 7, 9, 4])
    sheet.append([None, "银行戊", "日报", 1, 0, 5])
    sheet.append(["2026/09/30 12:00:00", "银行己", "日报", 1, 0, 6])
    sheet.append(["2026-08-31", None, "日报", 1, 0, 6])
    book.save(FIXTURES / "time_series.xlsx")


def save_time_series_1904() -> None:
    book = workbook(FIXTURES / "time_series_1904.xlsx")
    book.epoch = CALENDAR_MAC_1904
    sheet = book["集中系统数据"]
    sheet.append(["数据日期", "行名称", "余额"])
    sheet.append([datetime(2024, 2, 29), "银行1904", 15])
    book.save(FIXTURES / "time_series_1904.xlsx")


def save_duplicate_time_series() -> None:
    book = workbook(FIXTURES / "time_series_duplicate.xlsx")
    sheet = book["集中系统数据"]
    sheet.append(["数据日期", "行名称", "余额"])
    sheet.append(["2026-06-30", "银行甲", 1])
    sheet.append(["20260630", "银行甲", 2])
    book.save(FIXTURES / "time_series_duplicate.xlsx")


def save_duplicate_title() -> None:
    book = workbook(FIXTURES / "time_series_duplicate_title.xlsx")
    sheet = book["集中系统数据"]
    sheet.append(["数据日期", "行名称", "余额", "余额"])
    sheet.append(["2026-06-30", "银行甲", 1, 2])
    book.save(FIXTURES / "time_series_duplicate_title.xlsx")


def save_partial_time_header() -> None:
    book = workbook(FIXTURES / "time_series_partial_header.xlsx")
    sheet = book["集中系统数据"]
    sheet.append(["数据日期", "余额"])
    sheet.append(["2026-06-30", 1])
    book.save(FIXTURES / "time_series_partial_header.xlsx")


def save_missing_main_sheet() -> None:
    book = Workbook()
    book.active.title = "其他数据"
    book.active.append(["报表项目", "余额"])
    book.active.append(["项目A", 1])
    book.save(FIXTURES / "missing_main_sheet.xlsx")


def save_bad_org_map() -> None:
    book = workbook(FIXTURES / "bad_org_map.xlsx")
    sheet = book["集中系统数据"]
    sheet.append(["行名称", "余额"])
    sheet.append(["项目A", 1])
    reference = book.create_sheet("机构参照表")
    reference.append(["机构代码", "项目名称"])
    reference.append(["code-a", "映射A"])
    book.save(FIXTURES / "bad_org_map.xlsx")


def main() -> None:
    FIXTURES.mkdir(parents=True, exist_ok=True)
    save_plain()
    save_plain_fallback()
    save_plain_zero_headers()
    save_plain_legacy_reference()
    save_time_series()
    save_time_series_1904()
    save_duplicate_time_series()
    save_duplicate_title()
    save_partial_time_header()
    save_missing_main_sheet()
    save_bad_org_map()
    (FIXTURES / "unsupported.csv").write_text("行名称,余额\n项目A,1\n", encoding="utf-8")

    cases = [
        {"id": "plain_layout", "file": "plain_layout.xlsx"},
        {"id": "plain_missing_label", "file": "plain_missing_label.xlsx"},
        {"id": "plain_zero_headers", "file": "plain_zero_headers.xlsx"},
        {"id": "plain_legacy_reference", "file": "plain_legacy_reference.xlsx"},
        {"id": "time_series", "file": "time_series.xlsx"},
        {"id": "time_series_1904", "file": "time_series_1904.xlsx"},
        {"id": "time_series_duplicate", "file": "time_series_duplicate.xlsx"},
        {"id": "time_series_duplicate_title", "file": "time_series_duplicate_title.xlsx"},
        {"id": "time_series_partial_header", "file": "time_series_partial_header.xlsx"},
        {"id": "missing_main_sheet", "file": "missing_main_sheet.xlsx"},
        {"id": "bad_org_map", "file": "bad_org_map.xlsx"},
        {"id": "unsupported_csv", "file": "unsupported.csv"},
    ]
    (FIXTURES / "cases.json").write_text(
        json.dumps({"cases": cases}, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    print(json.dumps({"fixture_dir": str(FIXTURES), "case_count": len(cases)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
