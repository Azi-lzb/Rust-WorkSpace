#!/usr/bin/env python3
"""Compare an Unverified Rust ImportRange candidate with Audit Python copy semantics.

The helper creates only temporary outputs. It never writes the supplied template
or audit workbook and emits aggregate counts only; no cell values or formulas are
persisted in the report.
"""

from __future__ import annotations

import argparse
import json
import math
import os
import shlex
import shutil
import subprocess
import sys
import tempfile
import time
import zipfile
from pathlib import Path
from types import SimpleNamespace
from xml.etree import ElementTree

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.utils.cell import range_boundaries


ROOT = Path(__file__).resolve().parents[3]
MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


class CandidateReadbackError(Exception):
    pass


def wsl_path(path: Path) -> str:
    drive = path.drive.rstrip(":").lower()
    if len(drive) != 1:
        raise ValueError("candidate path must be on a mounted Windows drive")
    relative = path.as_posix().split("/", 1)[1]
    return f"/mnt/{drive}/{relative}"


def make_synthetic_pair(directory: Path) -> tuple[Path, Path, list[dict[str, str]]]:
    template_path = directory / "template.xlsx"
    audit_path = directory / "audit.xlsx"
    ranges = ["A1:B2", "D1:E2", "G1:H2", "J1:K2"]

    template = Workbook()
    source = template.active
    source.title = "审核表"
    for index, address in enumerate(ranges, start=1):
        min_col, min_row, max_col, max_row = range_boundaries(address)
        for row in range(min_row, max_row + 1):
            for column in range(min_col, max_col + 1):
                cell = source.cell(row, column)
                cell.value = f"source-{index}-{row}-{column}"
                cell.font = Font(bold=(row == min_row), color="FF223344")
                cell.fill = PatternFill("solid", fgColor="FFEFEFEF")
                cell.alignment = Alignment(horizontal="center", wrap_text=True)
                cell.protection = Protection(locked=(column % 2 == 0), hidden=False)
                cell.number_format = "0.00"
        top_left = source.cell(min_row, min_col)
        top_left.value = f"=ROW()+COLUMN()+{index}"
        top_left.comment = Comment(f"synthetic-rule-{index}", "baseline")
    source["B2"].comment = Comment("synthetic-value-comment", "baseline")
    template.defined_names.add(
        DefinedName(
            "校验区域",
            attr_text=",".join(f"'审核表'!{address}" for address in ranges),
        )
    )
    template.save(template_path)
    template.close()

    audit = load_workbook(template_path, data_only=False, keep_links=True)
    destination = audit["审核表"]
    destination["M1"] = "target-only-sentinel"
    destination["M1"].comment = Comment("untouched-comment", "target")
    for index, address in enumerate(ranges, start=1):
        min_col, min_row, max_col, max_row = range_boundaries(address)
        for row in range(min_row, max_row + 1):
            for column in range(min_col, max_col + 1):
                destination.cell(row, column).value = f"old-{index}-{row}-{column}"
    untouched = audit.create_sheet("Untouched")
    untouched["A1"] = "preserve-sheet"
    untouched["B1"] = "=2+3"
    untouched["A1"].font = Font(italic=True)
    audit.save(audit_path)
    audit.close()

    return template_path, audit_path, [
        {"sheet_name": "审核表", "address": address} for address in ranges
    ]


def resolve_real_ranges(source_root: Path, template_path: Path, mapping_path: Path | None):
    sys.path.insert(0, str(source_root / "core" / "src"))
    from base_audit.name_config import (
        FORMULA_COPY_FUNCTION,
        ISSUE_EXTRACT_FUNCTION,
        STRUCTURE_COMPARE_FUNCTION,
        features_of_type,
        load_feature_mappings,
    )
    from base_audit.native.openpyxl_workbook import read_template

    mappings = load_feature_mappings(mapping_path, template_path)
    definition = read_template(
        template_path,
        formula_mappings=features_of_type(mappings, FORMULA_COPY_FUNCTION),
        structure_mappings=features_of_type(mappings, STRUCTURE_COMPARE_FUNCTION),
        extraction_mappings=features_of_type(mappings, ISSUE_EXTRACT_FUNCTION),
    )
    if definition.structured or len(definition.copy_ranges) != 4:
        raise ValueError("real candidate scope must be four named formula ranges")
    return definition


def cell_value_type_signature(cell):
    return cell.value, cell.data_type


def formula_text_signature(cell):
    if cell.data_type == "f" or (isinstance(cell.value, str) and cell.value.startswith("=")):
        return cell.value
    return None


def comment_signature(cell):
    return None if cell.comment is None else (cell.comment.text, cell.comment.author)


def hyperlink_signature(cell):
    if cell.hyperlink is None:
        return None
    return (
        cell.hyperlink.target,
        cell.hyperlink.location,
        cell.hyperlink.display,
        cell.hyperlink.tooltip,
    )


def cell_mismatch_counts(
    python_path: Path, rust_path: Path, template_path: Path, target_path: Path, ranges
):
    """Compare copied business content while excluding presentation styling.

    Values, types, formulas, comments, hyperlinks, rules, and workbook structure
    remain in scope. Ordinary cell styles, dimensions, and print appearance do
    not gate the S1 business-copy candidate.
    """
    python_book = load_workbook(python_path, data_only=False, keep_links=True)
    rust_book = load_workbook(rust_path, data_only=False, keep_links=True)
    template_book = load_workbook(template_path, data_only=False, keep_links=True)
    target_book = load_workbook(target_path, data_only=False, keep_links=True)
    target_cache_book = load_workbook(target_path, data_only=True, keep_links=True)
    rust_cache_book = load_workbook(rust_path, data_only=True, keep_links=True)
    counts = {
        "range_cells": 0,
        "range_value_type_mismatches": 0,
        "range_formula_text_mismatches": 0,
        "range_comment_mismatches": 0,
        "range_hyperlink_mismatches": 0,
        "range_mismatches": 0,
        "outside_cells": 0,
        "outside_value_type_mismatches": 0,
        "outside_formula_text_mismatches": 0,
        "outside_comment_mismatches": 0,
        "outside_hyperlink_mismatches": 0,
        "outside_formula_cache_cells": 0,
        "outside_formula_cache_mismatches": 0,
        "outside_formula_cache_nonempty_mismatches": 0,
        "outside_mismatches": 0,
        "outside_sheet_mismatches": 0,
        "outside_python_rust_value_type_mismatches": 0,
        "outside_python_rust_formula_text_mismatches": 0,
        "outside_python_rust_comment_mismatches": 0,
        "outside_python_rust_hyperlink_mismatches": 0,
        "outside_python_rust_mismatches": 0,
    }
    range_cells: dict[str, set[tuple[int, int]]] = {}

    def record_cell(prefix, expected_value, actual_value):
        counts[f"{prefix}_value_type_mismatches"] += int(
            cell_value_type_signature(expected_value) != cell_value_type_signature(actual_value)
        )
        counts[f"{prefix}_formula_text_mismatches"] += int(
            formula_text_signature(expected_value) != formula_text_signature(actual_value)
        )
        counts[f"{prefix}_comment_mismatches"] += int(
            comment_signature(expected_value) != comment_signature(actual_value)
        )
        counts[f"{prefix}_hyperlink_mismatches"] += int(
            hyperlink_signature(expected_value) != hyperlink_signature(actual_value)
        )

    for item in ranges:
        min_col, min_row, max_col, max_row = range_boundaries(item["address"].replace("$", ""))
        sheet_name = item["sheet_name"]
        if any(
            sheet_name not in book.sheetnames
            for book in (python_book, rust_book, template_book)
        ):
            raise CandidateReadbackError("range_sheet_missing")
        cells = {(row, col) for row in range(min_row, max_row + 1) for col in range(min_col, max_col + 1)}
        range_cells.setdefault(sheet_name, set()).update(cells)
        python_sheet = python_book[sheet_name]
        rust_sheet = rust_book[sheet_name]
        for row, col in cells:
            counts["range_cells"] += 1
            python_cell = python_sheet.cell(row, col)
            rust_cell = rust_sheet.cell(row, col)
            record_cell(
                "range",
                python_cell,
                rust_cell,
            )

    for sheet_name in sorted(
        set(target_book.sheetnames) | set(python_book.sheetnames) | set(rust_book.sheetnames)
    ):
        if any(
            sheet_name not in book.sheetnames
            for book in (target_book, python_book, rust_book)
        ):
            counts["outside_sheet_mismatches"] += 1
            continue
        target_sheet = target_book[sheet_name]
        python_sheet = python_book[sheet_name]
        rust_sheet = rust_book[sheet_name]
        target_cache_sheet = target_cache_book[sheet_name]
        rust_cache_sheet = rust_cache_book[sheet_name]
        coordinates = (
            set(target_sheet._cells) | set(python_sheet._cells) | set(rust_sheet._cells)
        )
        for row, col in coordinates:
            if (row, col) in range_cells.get(sheet_name, set()):
                continue
            counts["outside_cells"] += 1
            target_cell = target_sheet.cell(row, col)
            python_cell = python_sheet.cell(row, col)
            rust_cell = rust_sheet.cell(row, col)
            record_cell(
                "outside",
                python_cell,
                rust_cell,
            )
            counts["outside_python_rust_value_type_mismatches"] += int(
                cell_value_type_signature(python_cell)
                != cell_value_type_signature(rust_cell)
            )
            counts["outside_python_rust_formula_text_mismatches"] += int(
                formula_text_signature(python_cell)
                != formula_text_signature(rust_cell)
            )
            counts["outside_python_rust_comment_mismatches"] += int(
                comment_signature(python_cell) != comment_signature(rust_cell)
            )
            counts["outside_python_rust_hyperlink_mismatches"] += int(
                hyperlink_signature(python_cell) != hyperlink_signature(rust_cell)
            )
            target_formula = formula_text_signature(target_cell)
            if target_formula is not None and formula_text_signature(rust_cell) == target_formula:
                target_cache = target_cache_sheet.cell(row, col).value
                rust_cache = rust_cache_sheet.cell(row, col).value
                counts["outside_formula_cache_cells"] += 1
                counts["outside_formula_cache_mismatches"] += int(
                    not cached_value_equal(target_cache, rust_cache)
                )
                counts["outside_formula_cache_nonempty_mismatches"] += int(
                    (target_cache is None) != (rust_cache is None)
                )

    for prefix in ("range", "outside"):
        counts[f"{prefix}_mismatches"] = sum(
            counts[f"{prefix}_{kind}_mismatches"]
            for kind in ("value_type", "formula_text", "comment", "hyperlink")
        )
    counts["outside_mismatches"] += counts["outside_formula_cache_mismatches"]
    counts["outside_mismatches"] += counts["outside_sheet_mismatches"]
    counts["outside_python_rust_mismatches"] = sum(
        counts[f"outside_python_rust_{kind}_mismatches"]
        for kind in ("value_type", "formula_text", "comment", "hyperlink")
    )
    for book in (python_book, rust_book, template_book, target_book, target_cache_book, rust_cache_book):
        book.close()
    return counts


def workbook_metadata(path: Path):
    workbook = load_workbook(path, data_only=False, keep_links=True)
    result = {
        "sheet_names": tuple(workbook.sheetnames),
        "defined_names": tuple(sorted((item.name, item.attr_text, item.localSheetId, item.hidden) for item in workbook.defined_names.values())),
        "sheets": {},
    }
    for sheet in workbook.worksheets:
        validations = tuple(
            sorted(
                (
                    str(item.sqref), item.type, item.operator, item.formula1, item.formula2,
                    item.allow_blank, item.showErrorMessage, item.showInputMessage,
                )
                for item in sheet.data_validations.dataValidation
            )
        )
        conditional = tuple(
            sorted((str(group.sqref), repr(group.rules)) for group in sheet.conditional_formatting)
        )
        result["sheets"][sheet.title] = (
            tuple(sorted(str(item) for item in sheet.merged_cells.ranges)),
            validations,
            conditional,
            sheet.sheet_state,
        )
    workbook.close()
    return result


def calculation_flags(path: Path):
    with zipfile.ZipFile(path) as archive:
        root = ElementTree.fromstring(archive.read("xl/workbook.xml"))
    node = root.find(f"{{{MAIN_NS}}}calcPr")
    if node is None:
        return {"full_calc_on_load": None, "force_full_calc": None}
    return {
        "full_calc_on_load": node.get("fullCalcOnLoad"),
        "force_full_calc": node.get("forceFullCalc"),
    }


def formula_cache_count(path: Path) -> int:
    count = 0
    with zipfile.ZipFile(path) as archive:
        for name in archive.namelist():
            if name.startswith("xl/worksheets/") and name.endswith(".xml") and "/_rels/" not in name:
                root = ElementTree.fromstring(archive.read(name))
                for cell in root.findall(f".//{{{MAIN_NS}}}c"):
                    formula = cell.find(f"{{{MAIN_NS}}}f")
                    value = cell.find(f"{{{MAIN_NS}}}v")
                    if formula is not None and value is not None and value.text not in (None, ""):
                        count += 1
    return count


def formula_cache_snapshot(path: Path):
    formula_book = load_workbook(path, data_only=False, read_only=True, keep_links=False)
    cached_book = load_workbook(path, data_only=True, read_only=True, keep_links=False)
    values = {}
    try:
        for formula_sheet in formula_book.worksheets:
            cached_sheet = cached_book[formula_sheet.title]
            for row in formula_sheet.iter_rows():
                for cell in row:
                    if cell.data_type == "f":
                        values[(formula_sheet.title, cell.coordinate)] = (
                            cell.value,
                            cached_sheet[cell.coordinate].value,
                        )
    finally:
        formula_book.close()
        cached_book.close()
    return values


def cached_value_equal(left, right) -> bool:
    if isinstance(left, bool) or isinstance(right, bool):
        return type(left) is type(right) and left == right
    if isinstance(left, (int, float)) and isinstance(right, (int, float)):
        return math.isclose(float(left), float(right), rel_tol=1e-15, abs_tol=0.0)
    return type(left) is type(right) and left == right


def recalculate_and_compare(
    sidecar_path: Path,
    python_output: Path,
    rust_output: Path,
    directory: Path,
    expected_cache_values: dict[tuple[str, str], object] | None = None,
):
    started = time.perf_counter()
    python_recalculated = directory / "python_recalculated.xlsx"
    rust_recalculated = directory / "rust_recalculated.xlsx"
    if os.name == "nt":
        def recalculate_with_com(source: Path, destination: Path):
            try:
                result = subprocess.run(
                    [
                        str(sidecar_path),
                        "recalculate",
                        "--input",
                        str(source),
                        "--output",
                        str(destination),
                        "--engine",
                        "com",
                    ],
                    timeout=300,
                    capture_output=True,
                    text=True,
                    encoding="utf-8",
                    errors="replace",
                    check=False,
                )
            except subprocess.TimeoutExpired:
                return None, "excel_com_timeout"
            except OSError as error:
                return None, f"sidecar_start_{type(error).__name__}"

            if result.returncode != 0:
                diagnostic = (result.stderr + "\n" + result.stdout).casefold()
                if "recalculation changed formulas" in diagnostic:
                    return None, "formula_text_changed"
                if "workbooks.open" in diagnostic or "不能取得类 workbooks 的 open 属性" in diagnostic:
                    return None, "excel_com_workbook_open_failed"
                return None, f"sidecar_exit_{result.returncode}"
            try:
                report = json.loads(result.stdout)
            except (json.JSONDecodeError, TypeError):
                return None, "sidecar_invalid_json"
            if not isinstance(report, dict):
                return None, "sidecar_invalid_response"
            validation = report.get("validation", {})
            formula_count = validation.get("formula_count") if isinstance(validation, dict) else None
            if (
                report.get("status") != "recalculated"
                or report.get("engine") != "Excel.Application"
                or not isinstance(validation, dict)
                or validation.get("formula_text_preserved") is not True
                or validation.get("cache_nodes_complete") is not True
                or not isinstance(formula_count, int)
                or isinstance(formula_count, bool)
                or formula_count <= 0
                or validation.get("formula_cache_cells") != formula_count
                or not destination.is_file()
                or destination.stat().st_size == 0
            ):
                return None, "sidecar_validation_failed"
            return report, None

        python_report, python_error = recalculate_with_com(
            python_output, python_recalculated
        )
        rust_report, rust_error = recalculate_with_com(rust_output, rust_recalculated)
        if python_error or rust_error:
            return {
                "status": "failed",
                "engine_requested": "Excel.Application COM",
                "python_recalculation_status": "failed" if python_error else "completed",
                "python_error_category": python_error,
                "candidate_recalculation_status": "failed" if rust_error else "completed",
                "candidate_error_category": rust_error,
                "mismatch_count": 1,
                "elapsed_ms": round((time.perf_counter() - started) * 1000),
            }
        python_engine = python_report["engine"]
        sidecar_report = rust_report
    else:
        shutil.copy2(python_output, python_recalculated)

        from base_audit.engines.libreoffice import LibreOfficeCalculator
        calculator = LibreOfficeCalculator()

        try:
            python_result = calculator.recalculate_many(
                (python_recalculated,),
                batch_size=8,
            )
        except Exception as error:
            raise CandidateReadbackError(f"python_recalculation_{type(error).__name__}") from None

        try:
            sidecar_result = subprocess.run(
                [
                    str(sidecar_path),
                    "recalculate",
                    "--input",
                    str(rust_output),
                    "--output",
                    str(rust_recalculated),
                    "--engine",
                    "lo",
                ],
                timeout=calculator.timeout_seconds,
                capture_output=True,
                text=True,
                encoding="utf-8",
                errors="replace",
                check=False,
            )
        except subprocess.TimeoutExpired:
            raise CandidateReadbackError("sidecar_timeout") from None
        except OSError as error:
            raise CandidateReadbackError(f"sidecar_start_{type(error).__name__}") from None
        if sidecar_result.returncode != 0:
            raise CandidateReadbackError(f"sidecar_exit_{sidecar_result.returncode}")

        try:
            sidecar_report = json.loads(sidecar_result.stdout)
        except (json.JSONDecodeError, TypeError):
            raise CandidateReadbackError("sidecar_invalid_json") from None
        if not isinstance(sidecar_report, dict):
            raise CandidateReadbackError("sidecar_invalid_response")
        validation = sidecar_report.get("validation", {})
        if not isinstance(validation, dict):
            raise CandidateReadbackError("sidecar_invalid_validation")
        formula_count = validation.get("formula_count")
        formula_cache_cells = validation.get("formula_cache_cells")
        if (
            sidecar_report.get("status") != "recalculated"
            or not isinstance(sidecar_report.get("engine"), str)
            or "LibreOffice" not in sidecar_report["engine"]
            or validation.get("formula_text_preserved") is not True
            or validation.get("cache_nodes_complete") is not True
            or not isinstance(formula_count, int)
            or isinstance(formula_count, bool)
            or formula_count <= 0
            or formula_cache_cells != formula_count
            or not rust_recalculated.is_file()
            or rust_recalculated.stat().st_size == 0
        ):
            raise CandidateReadbackError("sidecar_validation_failed")
        python_engine = python_result.engine_display

    python_values = formula_cache_snapshot(python_recalculated)
    rust_values = formula_cache_snapshot(rust_recalculated)
    python_keys = set(python_values)
    rust_keys = set(rust_values)
    formula_set_difference_count = len(python_keys ^ rust_keys)
    formula_text_mismatch_count = 0
    cache_value_mismatch_count = 0
    for key in python_keys & rust_keys:
        python_formula, python_cache = python_values[key]
        rust_formula, rust_cache = rust_values[key]
        formula_text_mismatch_count += int(python_formula != rust_formula)
        cache_value_mismatch_count += int(not cached_value_equal(python_cache, rust_cache))

    python_expected_cache_mismatch_count = 0
    sidecar_expected_cache_mismatch_count = 0
    if expected_cache_values is not None:
        python_expected_cache_mismatch_count = sum(
            key not in python_values
            or not cached_value_equal(python_values[key][1], expected_value)
            for key, expected_value in expected_cache_values.items()
        )
        sidecar_expected_cache_mismatch_count = sum(
            key not in rust_values
            or not cached_value_equal(rust_values[key][1], expected_value)
            for key, expected_value in expected_cache_values.items()
        )
    expected_cache_mismatch_count = (
        python_expected_cache_mismatch_count + sidecar_expected_cache_mismatch_count
    )

    mismatch_count = (
        formula_set_difference_count
        + formula_text_mismatch_count
        + cache_value_mismatch_count
        + expected_cache_mismatch_count
    )
    return {
        "status": "matched" if mismatch_count == 0 else "mismatch",
        "python_engine": python_engine,
        "sidecar_engine": sidecar_report["engine"],
        "formula_count_python": len(python_keys),
        "formula_count_sidecar": len(rust_keys),
        "python_cached_values_nonempty": sum(value is not None for _, value in python_values.values()),
        "sidecar_cached_values_nonempty": sum(value is not None for _, value in rust_values.values()),
        "formula_set_difference_count": formula_set_difference_count,
        "formula_text_mismatch_count": formula_text_mismatch_count,
        "cache_value_mismatch_count": cache_value_mismatch_count,
        "expected_cache_cells": len(expected_cache_values or {}),
        "python_expected_cache_mismatch_count": python_expected_cache_mismatch_count,
        "sidecar_expected_cache_mismatch_count": sidecar_expected_cache_mismatch_count,
        "expected_cache_mismatch_count": expected_cache_mismatch_count,
        "numeric_relative_tolerance": 1e-15,
        "elapsed_ms": round((time.perf_counter() - started) * 1000),
        "mismatch_count": mismatch_count,
    }


def range_stats(template_path: Path, ranges):
    workbook = load_workbook(template_path, data_only=False, read_only=False)
    formulas = 0
    comments = 0
    for item in ranges:
        min_col, min_row, max_col, max_row = range_boundaries(item["address"].replace("$", ""))
        sheet = workbook[item["sheet_name"]]
        for row in sheet.iter_rows(
            min_row=min_row,
            max_row=max_row,
            min_col=min_col,
            max_col=max_col,
        ):
            for cell in row:
                formulas += int(cell.data_type == "f")
                comments += int(cell.comment is not None)
    workbook.close()
    return {"formula_cells": formulas, "range_comment_count_python": comments}


def invoke_candidate(
    workspace: Path,
    paths: dict[str, Path],
    distribution: str,
    runtime: str,
    engine: str,
    openpyxl_python: str | None,
    include_calculation_properties: bool,
):
    if runtime not in ("native", "wsl"):
        raise ValueError("candidate runtime must be native or wsl")
    environment_values = {key: str(value) for key, value in paths.items()}
    environment_values["AUDIT_S1_COPY_CANDIDATE_ENGINE"] = engine
    environment_values["AUDIT_S1_COPY_CANDIDATE_CALC_PROPERTIES"] = (
        "1" if include_calculation_properties else "0"
    )
    if engine == "openpyxl-worker":
        if not openpyxl_python:
            raise ValueError("openpyxl-worker requires --openpyxl-python")
        environment_values["RUST_WORKSPACE_OPENPYXL_PYTHON"] = openpyxl_python

    if runtime == "wsl":
        if os.name != "nt":
            raise ValueError("the explicit WSL candidate runtime is available from Windows only")
        wsl_environment_values = {
            key: wsl_path(value) if key in paths else value
            for key, value in environment_values.items()
        }
        assignments = " ".join(
            f"{key}={shlex.quote(value)}" for key, value in wsl_environment_values.items()
        )
        command = (
            f"cd {shlex.quote(wsl_path(workspace))} && {assignments} "
            "cargo test -p audit-core --test s1_formula_copy_candidate --locked -- --ignored --nocapture"
        )
        completed = subprocess.run(
            ["wsl.exe", "--distribution", distribution, "--exec", "bash", "-lc", command],
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            check=False,
        )
    else:
        environment = os.environ.copy()
        environment.update(environment_values)
        completed = subprocess.run(
            ["cargo", "test", "-p", "audit-core", "--test", "s1_formula_copy_candidate", "--locked", "--", "--ignored", "--nocapture"],
            cwd=workspace,
            env=environment,
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            check=False,
        )
    if completed.returncode:
        output = completed.stdout + "\n" + completed.stderr
        summary_lines = [
            line.strip()
            for line in output.splitlines()
            if "panicked at" in line or line.strip().startswith("error:")
        ]
        output_lines = output.splitlines()
        panic_index = next(
            (index for index, line in enumerate(output_lines) if "panicked at" in line),
            None,
        )
        if panic_index is not None:
            summary = " ".join(
                line.strip() for line in output_lines[panic_index : panic_index + 3] if line.strip()
            )
        else:
            summary = summary_lines[-1] if summary_lines else "candidate test failed without a diagnostic line"
        for path in paths.values():
            summary = summary.replace(str(path), "<workbook>")
        print(
            f"Rust S1 copy candidate exited {completed.returncode}: {summary}",
            file=sys.stderr,
        )
    return completed.returncode


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source-root", type=Path, required=True)
    parser.add_argument("--template-workbook", type=Path)
    parser.add_argument("--audit-workbook", type=Path)
    parser.add_argument("--mapping-config", type=Path)
    parser.add_argument("--workspace", type=Path, default=ROOT)
    parser.add_argument("--wsl-distribution", default="Ubuntu")
    parser.add_argument(
        "--runtime",
        choices=("native", "wsl"),
        default="native",
        help="Run the Rust candidate in the current Windows environment; WSL is explicit opt-in",
    )
    parser.add_argument(
        "--engine",
        choices=("umya", "openpyxl-worker"),
        default="umya",
        help="Direct Unverified candidate backend; this never changes production routing",
    )
    parser.add_argument(
        "--openpyxl-python",
        "--openpyxl-python-wsl",
        dest="openpyxl_python",
        help="Test-only worker Python path for the selected runtime; not a packaged product runtime",
    )
    parser.add_argument(
        "--scope",
        choices=("full-s1", "range-import-only"),
        default="full-s1",
        help="Use range-import-only solely to diagnose ImportRange independently of calc flags",
    )
    parser.add_argument("--report", type=Path, required=True)
    parser.add_argument(
        "--sidecar-path",
        type=Path,
        help="Explicit lo-convert binary; when supplied, compare Python and sidecar LibreOffice formula caches",
    )
    args = parser.parse_args()
    if bool(args.template_workbook) != bool(args.audit_workbook):
        parser.error("provide both real workbook inputs or neither")

    source_root = args.source_root.resolve(strict=True)
    sidecar_path = args.sidecar_path.resolve(strict=True) if args.sidecar_path else None
    if sidecar_path is not None and not sidecar_path.is_file():
        parser.error("--sidecar-path must point to the lo-convert executable")
    sys.path.insert(0, str(source_root / "core" / "src"))
    workspace = args.workspace.resolve(strict=True)
    report_path = args.report.resolve()

    with tempfile.TemporaryDirectory(prefix="audit-s1-formula-copy-candidate-") as temporary:
        directory = Path(temporary)
        if args.template_workbook:
            template_path = args.template_workbook.resolve(strict=True)
            audit_path = args.audit_workbook.resolve(strict=True)
            mapping_path = args.mapping_config.resolve(strict=True) if args.mapping_config else None
            before_template = template_path.read_bytes()
            before_audit = audit_path.read_bytes()
            definition = resolve_real_ranges(source_root, template_path, mapping_path)
            ranges = [
                {"sheet_name": item.sheet_name, "address": item.address}
                for item in definition.copy_ranges
            ]
            sample_kind = "real_named_formula_ranges"
        else:
            template_path, audit_path, ranges = make_synthetic_pair(directory)
            before_template = template_path.read_bytes()
            before_audit = audit_path.read_bytes()
            definition = SimpleNamespace(
                rules=[],
                copy_ranges=[SimpleNamespace(**item) for item in ranges],
                structured=False,
            )
            mapping_path = None
            sample_kind = "synthetic"

        if len(ranges) != 4:
            raise ValueError("candidate must contain exactly four ImportRange operations")
        python_output = directory / "python_output.xlsx"
        rust_output = directory / "rust_output.xlsx"
        rust_report_path = directory / "rust_candidate.json"
        shutil.copy2(audit_path, python_output)

        from base_audit.native.openpyxl_workbook import copy_formula_ranges_from_template

        template_book = load_workbook(template_path, data_only=False, keep_links=False)
        python_started = time.perf_counter()
        copy_formula_ranges_from_template(template_book, python_output, definition)
        python_elapsed_ms = round((time.perf_counter() - python_started) * 1000)
        template_book.close()

        rust_started = time.perf_counter()
        candidate_paths = {
            "AUDIT_S1_COPY_CANDIDATE_INPUT": audit_path,
            "AUDIT_S1_COPY_CANDIDATE_TEMPLATE": template_path,
            "AUDIT_S1_COPY_CANDIDATE_OUTPUT": rust_output,
            "AUDIT_S1_COPY_CANDIDATE_REPORT": rust_report_path,
        }
        if mapping_path is not None:
            candidate_paths["AUDIT_S1_COPY_CANDIDATE_MAPPING"] = mapping_path
        rust_exit = invoke_candidate(
            workspace,
            candidate_paths,
            args.wsl_distribution,
            args.runtime,
            args.engine,
            args.openpyxl_python,
            args.scope == "full-s1",
        )
        rust_elapsed_ms = round((time.perf_counter() - rust_started) * 1000)
        output_created = rust_output.is_file()
        candidate = json.loads(rust_report_path.read_text(encoding="utf-8")) if rust_report_path.exists() else {
            "status": "candidate_invocation_failed",
            "production_route_eligible": False,
        }

        recalculation = {"status": "not_requested", "mismatch_count": 0}
        if sidecar_path is not None:
            expected_cache_values = (
                {
                    ("审核表", "A1"): 3,
                    ("审核表", "D1"): 7,
                    ("审核表", "G1"): 11,
                    ("审核表", "J1"): 15,
                }
                if sample_kind == "synthetic"
                else None
            )
            if rust_output.exists():
                try:
                    recalculation = recalculate_and_compare(
                        sidecar_path,
                        python_output,
                        rust_output,
                        directory,
                        expected_cache_values,
                    )
                except CandidateReadbackError as error:
                    recalculation = {
                        "status": "failed",
                        "error_category": str(error),
                        "mismatch_count": 1,
                    }
            else:
                recalculation = {
                    "status": "failed",
                    "error_category": "rust_output_missing",
                    "mismatch_count": 1,
                }

        readback_failure_category = candidate.get("output_readback_failure_category")
        style_only_readback_diagnostic = (
            readback_failure_category == "imported_style_semantic_mismatch"
        )
        readback_failed = bool(candidate.get("output_readback_failed", False)) and not (
            style_only_readback_diagnostic
        )
        if style_only_readback_diagnostic:
            readback_failure_category = None
        if rust_output.exists():
            python_calc = calculation_flags(python_output)
            python_cache_count = formula_cache_count(python_output)
            try:
                rust_calc = calculation_flags(rust_output)
                candidate_cache_count = formula_cache_count(rust_output)
                calc_mismatch = int(python_calc != rust_calc)
            except Exception as error:
                readback_failed = True
                readback_failure_category = f"package_{type(error).__name__}"
                rust_calc = {"full_calc_on_load": None, "force_full_calc": None}
                candidate_cache_count = 0
                calc_mismatch = 1
            try:
                cell_comparison = cell_mismatch_counts(python_output, rust_output, template_path, audit_path, ranges)
            except CandidateReadbackError as error:
                readback_failed = True
                readback_failure_category = readback_failure_category or str(error)
                cell_comparison = {
                    "range_cells": 0, "range_mismatches": 0,
                    "outside_cells": 0, "outside_mismatches": 0,
                    "outside_formula_cache_cells": 0,
                    "outside_formula_cache_mismatches": 0,
                    "outside_formula_cache_nonempty_mismatches": 0,
                    "outside_python_rust_mismatches": 0,
                }
            try:
                python_metadata = workbook_metadata(python_output)
                target_metadata = workbook_metadata(audit_path)
                rust_metadata = workbook_metadata(rust_output)
                metadata_mismatch = int(python_metadata != rust_metadata)
                target_metadata_mismatch = int(target_metadata != rust_metadata)
            except Exception as error:
                readback_failed = True
                readback_failure_category = readback_failure_category or f"metadata_{type(error).__name__}"
                metadata_mismatch = 1
                target_metadata_mismatch = 1
        else:
            readback_failed = True
            readback_failure_category = "rust_output_missing"
            cell_comparison = {
                "range_cells": 0, "range_mismatches": 0,
                "outside_cells": 0, "outside_mismatches": 0,
                "outside_formula_cache_cells": 0,
                "outside_formula_cache_mismatches": 0,
                "outside_formula_cache_nonempty_mismatches": 0,
                "outside_python_rust_mismatches": 0,
            }
            metadata_mismatch = 1
            target_metadata_mismatch = 1
            calc_mismatch = 1
            python_calc = calculation_flags(python_output)
            rust_calc = {"full_calc_on_load": None, "force_full_calc": None}
            candidate_cache_count = 0
            python_cache_count = formula_cache_count(python_output)

        inputs_unchanged = (
            template_path.read_bytes() == before_template
            and audit_path.read_bytes() == before_audit
        )
        mismatches = (
            cell_comparison["range_mismatches"]
            + cell_comparison["outside_mismatches"]
            + metadata_mismatch
            + calc_mismatch
            + int(readback_failed)
        )
        mismatches += recalculation["mismatch_count"]
        scope_stats = range_stats(template_path, ranges)
        if rust_exit != 0 or not output_created:
            candidate_status = "candidate_unavailable"
        elif not inputs_unchanged:
            candidate_status = "input_mutated"
        elif readback_failed:
            candidate_status = "comparison_incomplete"
        elif mismatches:
            candidate_status = "candidate_mismatch"
        else:
            candidate_status = "candidate_match_unverified"
        report = {
            "schema_version": 3,
            "suite": "audit.s1_formula_copy_candidate",
            "status": candidate_status,
            "sample_kind": sample_kind,
            "candidate_engine": args.engine,
            "candidate_operation_scope": args.scope,
            "python": {"baseline": "base_audit.native.openpyxl_workbook.copy_formula_ranges_from_template", "elapsed_ms": python_elapsed_ms},
            "rust_candidate": {
                **candidate,
                "elapsed_ms_including_test_harness": rust_elapsed_ms,
                "candidate_test_exit_code": rust_exit,
                "formula_cache_nonempty_count": candidate_cache_count,
                "output_readback_failed": readback_failed,
                "output_readback_failure_category": readback_failure_category if readback_failed else None,
            },
            "copy_scope": {
                "range_count": len(ranges),
                **scope_stats,
            },
            "comparison": {
                **cell_comparison,
                "python_vs_rust_workbook_metadata_mismatch_count": metadata_mismatch,
                "target_vs_rust_workbook_metadata_mismatch_count": target_metadata_mismatch,
                "outside_formula_cache_cells": cell_comparison["outside_formula_cache_cells"],
                "outside_formula_cache_mismatches": cell_comparison["outside_formula_cache_mismatches"],
                "outside_formula_cache_nonempty_mismatches": cell_comparison["outside_formula_cache_nonempty_mismatches"],
                "calculation_flags_mismatch_count": calc_mismatch,
                "python_calculation_flags": python_calc,
                "rust_calculation_flags": rust_calc,
                "python_formula_cache_nonempty_count": python_cache_count,
                "presentation_style_comparison": "not performed; ordinary formatting is nonblocking",
                "style_only_readback_diagnostic_nonblocking": style_only_readback_diagnostic,
                "output_readback_failed": readback_failed,
                "mismatch_count": mismatches,
            },
            "recalculation": recalculation,
            "verification": {
                "input_workbooks_unchanged": inputs_unchanged,
                "formula_recalculation_performed": recalculation["status"] in ("matched", "mismatch"),
                "capability_status": "unverified",
                "production_route_eligible": False,
                "target_linux_validation": "not performed",
                "scope": (
                    "four cross-workbook ImportRange operations plus calculation flags; direct candidate only"
                    if args.scope == "full-s1"
                    else "four cross-workbook ImportRange operations only; calculation flags intentionally omitted for diagnostics"
                ),
                "presentation_style_comparison": "not performed; ordinary formatting is nonblocking",
                "outside_python_rust_mismatches": cell_comparison[
                    "outside_python_rust_mismatches"
                ],
                "outside_formula_caches_compared_to": "original target workbook where formula text matches",
            },
        }

    report_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(
        json.dumps(
            {
                "status": report["status"],
                "mismatch_count": mismatches,
                "sample_kind": sample_kind,
                "recalculation_status": recalculation["status"],
            }
        )
    )
    return 0 if inputs_unchanged else 1


if __name__ == "__main__":
    raise SystemExit(main())
