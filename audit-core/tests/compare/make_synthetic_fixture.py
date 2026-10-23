"""Create a clearly synthetic OOXML fixture for exercising the parity harness.

This helper is not a runtime backend. Never put business data in this fixture.
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "fixtures" / "synthetic_probe.xlsx"


def main() -> None:
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Synthetic"
    ws["A1"] = 1
    ws["A2"] = 2
    ws["A3"] = 3
    ws["B2"] = "original merged anchor"
    ws.merge_cells("B2:C2")
    ws["C4"] = 12.5
    ws["D4"] = "=SUM(A1:A3)"
    ws["E5"] = "original named target"
    ws["F2"] = 1
    ws["G2"] = "red"
    ws["H2"] = "comment cell"
    ws["H2"].comment = Comment("synthetic comment text", "fixture")
    for row in range(1, 8):
        for col in range(10, 18):
            ws.cell(row=row, column=col, value=f"synthetic-{row}-{col}")
    ws.conditional_formatting.add(
        "F2:F5",
        CellIsRule(
            operator="greaterThan",
            formula=["0"],
            fill=PatternFill(fill_type="solid", fgColor="FFFF00"),
        ),
    )
    validation = DataValidation(type="list", formula1='"red,green,blue"', allow_blank=True)
    validation.add("G2:G5")
    ws.add_data_validation(validation)
    wb.defined_names.add(DefinedName("SyntheticNamedCell", attr_text="'Synthetic'!$E$5"))
    ws.row_dimensions[2].height = 18
    ws.row_dimensions[3].height = 16
    ws.column_dimensions["I"].width = 12
    ws.freeze_panes = "C3"
    ws.auto_filter.ref = "A1:H10"
    ws.print_area = "A1:D10"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = 0.3
    ws.page_margins.right = 0.3
    ws.print_options.horizontalCentered = True
    ws.sheet_view.showGridLines = False
    wb.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
