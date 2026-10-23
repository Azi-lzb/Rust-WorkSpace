#!/usr/bin/env python3
"""Compare Rust S1 formula-result Issue extraction with the read-only Python baseline."""

from __future__ import annotations

import argparse
import dataclasses
import hashlib
import json
import os
import subprocess
import sys
import tempfile
import time
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from typing import Any

sys.dont_write_bytecode = True


def wsl_path(path: Path) -> str:
    drive = path.drive.rstrip(":").lower()
    if len(drive) != 1:
        raise ValueError(f"无法转换为 WSL 挂载路径：{path.drive}")
    relative = path.as_posix().split("/", 1)[1]
    return f"/mnt/{drive}/{relative}"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def normalise_path(value: str) -> str:
    text = value.replace("\\", "/")
    if text.startswith("/mnt/") and len(text) > 6 and text[5].isalpha() and text[6] == "/":
        text = f"{text[5].upper()}:/{text[7:]}"
    return text.casefold()


def json_value(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, Path):
        return str(value)
    if isinstance(value, dict):
        return {key: json_value(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [json_value(item) for item in value]
    return value


def normalise_issue(issue: dict[str, Any]) -> dict[str, Any]:
    result = json_value(issue)
    for field in ("source_file", "audit_file"):
        result[field] = normalise_path(result[field])
    return result


def compare_issue_lists(python_issues: list[Any], rust_issues: list[dict[str, Any]]) -> dict[str, Any]:
    python_by_id = {
        issue.issue_id: normalise_issue(dataclasses.asdict(issue)) for issue in python_issues
    }
    rust_by_id = {issue["issue_id"]: normalise_issue(issue) for issue in rust_issues}
    if len(python_by_id) != len(python_issues) or len(rust_by_id) != len(rust_issues):
        raise ValueError("Issue identity is not unique; refusing to hide duplicate formula results")
    python_ids = set(python_by_id)
    rust_ids = set(rust_by_id)
    field_differences: Counter[str] = Counter()
    for issue_id in python_ids & rust_ids:
        python_issue = python_by_id[issue_id]
        rust_issue = rust_by_id[issue_id]
        for field in sorted(set(python_issue) | set(rust_issue)):
            if python_issue.get(field) != rust_issue.get(field):
                field_differences[field] += 1
    missing_in_rust = len(python_ids - rust_ids)
    missing_in_python = len(rust_ids - python_ids)
    return {
        "python_issue_count": len(python_ids),
        "rust_issue_count": len(rust_ids),
        "missing_in_rust": missing_in_rust,
        "missing_in_python": missing_in_python,
        "field_difference_counts": dict(sorted(field_differences.items())),
        "mismatch_count": missing_in_rust + missing_in_python + sum(field_differences.values()),
    }


def invoke_rust(args: argparse.Namespace, template: Path, audit: Path) -> tuple[dict[str, Any], int]:
    executable = args.rust_executable.resolve(strict=True)
    if os.name == "nt" and executable.suffix.lower() != ".exe":
        command = [
            "wsl.exe",
            "--distribution",
            args.wsl_distribution,
            "--exec",
            wsl_path(executable),
        ]
        template_arg = wsl_path(template)
        audit_arg = wsl_path(audit)
        source_arg = wsl_path(args.source_file.resolve()) if args.source_file else None
        mapping_arg = wsl_path(args.mapping_config.resolve(strict=True)) if args.mapping_config else None
    else:
        command = [str(executable)]
        template_arg = str(template)
        audit_arg = str(audit)
        source_arg = str(args.source_file.resolve()) if args.source_file else None
        mapping_arg = str(args.mapping_config.resolve(strict=True)) if args.mapping_config else None
    command.extend(
        [
            "s1-formula-audit",
            "--template-workbook",
            template_arg,
            "--audit-workbook",
            audit_arg,
            "--period",
            args.period,
            "--batch-id",
            args.batch_id,
            "--audit-time",
            args.audit_time,
            "--org-code",
            args.org_code,
            "--org-name",
            args.org_name,
        ]
    )
    if source_arg:
        command.extend(["--source-file", source_arg])
    if mapping_arg:
        command.extend(["--mapping-config", mapping_arg])
    started = time.perf_counter()
    result = subprocess.run(
        command,
        check=False,
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
    )
    elapsed_ms = round((time.perf_counter() - started) * 1000)
    if result.returncode:
        raise RuntimeError(
            f"Rust S1 公式审核 CLI 返回 {result.returncode}：{result.stderr.strip()}"
        )
    return json.loads(result.stdout), elapsed_ms


def run_synthetic_error_regression(args: argparse.Namespace, source_root: Path) -> dict[str, Any]:
    from openpyxl import Workbook, load_workbook
    from openpyxl.workbook.defined_name import DefinedName

    sys.path.insert(0, str(source_root / "core" / "src"))
    from base_audit.name_config import (
        FORMULA_COPY_FUNCTION,
        ISSUE_EXTRACT_FUNCTION,
        STRUCTURE_COMPARE_FUNCTION,
        features_of_type,
        load_feature_mappings,
    )
    from base_audit.native.openpyxl_workbook import (
        _read_issue_value_blocks,
        extract_issues,
        read_template,
    )

    with tempfile.TemporaryDirectory(prefix="audit-s1-formula-error-") as directory:
        root = Path(directory)
        template_path = root / "synthetic-template.xlsx"
        audit_path = root / "synthetic-results.xlsx"

        template = Workbook()
        sheet = template.active
        sheet.title = "检查表"
        formulas = (
            '=IF(TRUE,"错误|复核|指标|说明|10|20|10","")',
            '=IF(TRUE,"通过","")',
            '=IF(TRUE,"错误","")',
            "=1+1",
        )
        for row, formula in enumerate(formulas, start=1):
            sheet.cell(row=row, column=1, value=formula)
        template.defined_names.add(
            DefinedName("校验区域", attr_text="'检查表'!$A$1:$A$4")
        )
        template.save(template_path)
        template.close()

        audit = Workbook()
        sheet = audit.active
        sheet.title = "检查表"
        for row, value in enumerate(
            (
                "错误|复核|指标|说明|10|20|10",
                "#REF!",
                0,
                "#VALUE!",
            ),
            start=1,
        ):
            sheet.cell(row=row, column=1, value=value)
        audit.save(audit_path)
        audit.close()

        error_count = 0
        check = load_workbook(audit_path, read_only=True, data_only=True, keep_links=False)
        try:
            error_count = sum(check["检查表"][address].data_type == "e" for address in ("A2", "A4"))
        finally:
            check.close()
        if error_count != 2:
            raise AssertionError("Synthetic Excel error values were not stored as OOXML error cells")

        mappings = load_feature_mappings(None, template_path)
        definition = read_template(
            template_path,
            formula_mappings=features_of_type(mappings, FORMULA_COPY_FUNCTION),
            structure_mappings=features_of_type(mappings, STRUCTURE_COMPARE_FUNCTION),
            extraction_mappings=features_of_type(mappings, ISSUE_EXTRACT_FUNCTION),
        )
        cached_values = _read_issue_value_blocks(audit_path, definition.rules)
        issue_source = args.source_file.resolve() if args.source_file else audit_path
        python_issues = extract_issues(
            audit_path,
            definition.rules,
            period=args.period,
            batch_id=args.batch_id,
            org_code=args.org_code,
            org_name=args.org_name,
            source_file=issue_source,
            extraction_ranges=definition.extraction_ranges,
            cached_values=cached_values,
        )
        for issue in python_issues:
            issue.audit_time = args.audit_time

        rust_payload, _elapsed_ms = invoke_rust(args, template_path, audit_path)
        rust_issues = rust_payload.get("issues")
        if not isinstance(rust_issues, list):
            raise ValueError("Rust S1 公式审核 CLI 没有返回合成 Issue 数组")
        comparison = compare_issue_lists(python_issues, rust_issues)
        formula_error_issues = sum(
            isinstance(issue.formula_result, str) and issue.formula_result.startswith("#")
            for issue in python_issues
        )
        return {
            "input": "synthetic stored results with OOXML error cells; no recalculation",
            "stored_error_cells": error_count,
            "formula_error_issues": formula_error_issues,
            **comparison,
        }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source-root", type=Path, required=True)
    parser.add_argument("--template-workbook", type=Path, required=True)
    parser.add_argument("--audit-workbook", type=Path, required=True)
    parser.add_argument("--source-file", type=Path)
    parser.add_argument("--rust-executable", type=Path, required=True)
    parser.add_argument("--mapping-config", type=Path)
    parser.add_argument("--wsl-distribution", default="Ubuntu")
    parser.add_argument("--period", default="2026-07")
    parser.add_argument("--batch-id", default="formula-parity")
    parser.add_argument("--audit-time", default="2026-10-24 12:00:00")
    parser.add_argument("--org-code", default="")
    parser.add_argument("--org-name", default="")
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()

    source_root = args.source_root.resolve(strict=True)
    template_path = args.template_workbook.resolve(strict=True)
    audit_path = args.audit_workbook.resolve(strict=True)
    source_path = args.source_file.resolve() if args.source_file else audit_path
    mapping_config = args.mapping_config.resolve(strict=True) if args.mapping_config else None
    workbook_hashes_before = (sha256(template_path), sha256(audit_path))
    sys.path.insert(0, str(source_root / "core" / "src"))
    from base_audit.name_config import (
        FORMULA_COPY_FUNCTION,
        ISSUE_EXTRACT_FUNCTION,
        STRUCTURE_COMPARE_FUNCTION,
        features_of_type,
        load_feature_mappings,
    )
    from base_audit.native.openpyxl_workbook import (
        _read_issue_value_blocks,
        extract_issues,
        read_template,
    )

    python_started = time.perf_counter()
    mappings = load_feature_mappings(mapping_config, template_path)
    definition = read_template(
        template_path,
        formula_mappings=features_of_type(mappings, FORMULA_COPY_FUNCTION),
        structure_mappings=features_of_type(mappings, STRUCTURE_COMPARE_FUNCTION),
        extraction_mappings=features_of_type(mappings, ISSUE_EXTRACT_FUNCTION),
    )
    selected_rules = [
        rule
        for rule in definition.rules
        if rule.enabled
        and (
            not definition.extraction_ranges
            or any(
                item.sheet_name == rule.sheet_name and _in_range(rule.formula_cell, item.address)
                for item in definition.extraction_ranges
            )
        )
    ]
    cached_values = _read_issue_value_blocks(audit_path, selected_rules)
    python_issues = extract_issues(
        audit_path,
        definition.rules,
        period=args.period,
        batch_id=args.batch_id,
        org_code=args.org_code,
        org_name=args.org_name,
        source_file=source_path,
        extraction_ranges=definition.extraction_ranges,
        cached_values=cached_values,
    )
    for issue in python_issues:
        issue.audit_time = args.audit_time
    formula_result_values_nonempty = sum(
        1
        for rule in selected_rules
        if (value := cached_values.get((rule.sheet_name, rule.formula_cell)))
        not in (None, False, 0)
        and str(value).strip()
    )
    python_elapsed_ms = round((time.perf_counter() - python_started) * 1000)

    rust_payload, rust_cli_elapsed_ms = invoke_rust(args, template_path, audit_path)
    rust_issues = rust_payload.get("issues")
    if not isinstance(rust_issues, list):
        raise ValueError("Rust S1 公式审核 CLI 没有返回 Issue 数组")
    comparison = compare_issue_lists(python_issues, rust_issues)
    mismatch_count = comparison["mismatch_count"]
    workbooks_unchanged = workbook_hashes_before == (
        sha256(template_path),
        sha256(audit_path),
    )
    mismatch_count += int(not workbooks_unchanged)
    rust_summary = rust_payload.get("summary", {})
    python_summary = {
        "rules_total": len(definition.rules),
        "rules_enabled": sum(rule.enabled for rule in definition.rules),
        "rules_selected": len(selected_rules),
        "formula_result_values_nonempty": formula_result_values_nonempty,
        "issue_count": len(python_issues),
        "severity_counts": dict(Counter(issue.severity for issue in python_issues)),
        "elapsed_ms": python_elapsed_ms,
    }
    report = {
        "schema_version": 1,
        "suite": "audit.s1_formula_issue_extraction",
        "baseline": "Open-Data-Audit base_audit.native.openpyxl_workbook.extract_issues",
        "status": "passed" if mismatch_count == 0 else "mismatch",
        "sample": {
            "template_workbook": "read-only real S1 template",
            "audit_workbook": "read-only saved formula-result workbook",
            "issue_identity_input": "provided source file" if args.source_file else "audit workbook default",
            "mapping_config": "provided" if mapping_config else "source defaults",
        },
        "python": python_summary,
        "rust": {
            **{
                field: rust_summary.get(field)
                for field in (
                    "reader",
                    "rules_total",
                    "rules_enabled",
                    "rules_selected",
                    "formula_result_values_nonempty",
                    "issue_count",
                    "read_time_ms",
                    "issue_build_time_ms",
                    "total_time_ms",
                )
            },
            "cli_elapsed_ms": rust_cli_elapsed_ms,
        },
        "comparison": {
            **comparison,
        },
        "verification": {
            "input_workbooks_unchanged": workbooks_unchanged,
            "raw_issue_values_or_source_paths_persisted": False,
            "formula_recalculation": "not performed",
            "scope": "one real template/result workbook and one synthetic OOXML error-cell case; full Issue objects compared in memory",
        },
    }
    synthetic_regression = None
    if mapping_config is None:
        synthetic_regression = run_synthetic_error_regression(args, source_root)
        report["synthetic_regression"] = synthetic_regression
        mismatch_count += synthetic_regression["mismatch_count"]
        report["status"] = "passed" if mismatch_count == 0 else "mismatch"
        report["comparison"]["mismatch_count"] = mismatch_count
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if mismatch_count == 0 else 1


def _in_range(cell: str, address: str) -> bool:
    try:
        from openpyxl.utils.cell import coordinate_to_tuple, range_boundaries

        row, column = coordinate_to_tuple(cell.replace("$", ""))
        min_col, min_row, max_col, max_row = range_boundaries(address.replace("$", ""))
        return min_row <= row <= max_row and min_col <= column <= max_col
    except ValueError:
        return False


if __name__ == "__main__":
    raise SystemExit(main())
