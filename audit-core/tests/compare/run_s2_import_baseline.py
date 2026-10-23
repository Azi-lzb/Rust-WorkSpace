#!/usr/bin/env python3
"""Create a synthetic S2 workbook and capture the read-only Python importer baseline."""

from __future__ import annotations

import argparse
from dataclasses import asdict, replace
import hashlib
import importlib
import json
import sys
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook


WORKSPACE = Path(__file__).resolve().parents[2]
FIXTURE_DIR = WORKSPACE / "tests/fixtures/s2_import"
CONFIG_PATH = FIXTURE_DIR / "config.json"
PRIMARY_FILE = "9144#2026-06-30#x#20202#甲银行.xlsx"
NAME_MAPPED_FILE = "#2026-06-30#x#20202#乙银行.xlsx"


CONFIG = {
    "unit_settings": {"source_unit": "元", "output_file_unit": "万元"},
    "institutions_by_id": {
        "9144": {
            "institution_id": "9144",
            "institution_name": "配置名称不覆盖文件名",
            "social_credit_code": "CREDIT-9144",
            "institution_type": "农商行",
            "handling_branch": "一部",
            "region": "示例地区",
        }
    },
    "institutions_by_name": {
        "乙银行": {
            "institution_id": "9200",
            "institution_name": "乙银行配置名称",
            "social_credit_code": "CREDIT-9200",
            "institution_type": "城商行",
            "handling_branch": "二部",
            "region": "另一个地区",
        }
    },
    "indicators": {
        "20202001": {"indicator_name": "存款余额配置名", "form_code": "20202", "data_type": "余额"},
        "20202002": {"indicator_name": "累计金额", "form_code": "20202", "data_type": "累发"},
        "20201028": {"indicator_name": "持股比例", "form_code": "20202", "data_type": "百分数"},
        "20202003": {"indicator_name": "机构数量", "form_code": "20202", "data_type": "个数"},
        "20202004": {"indicator_name": "真实零值", "form_code": "20202", "data_type": "余额"},
        "20202005": {"indicator_name": "逗号金额", "form_code": "20202", "data_type": "余额"},
        "20202006": {"indicator_name": "无效金额", "form_code": "20202", "data_type": "余额"},
        "20202007": {"indicator_name": "文字示例", "form_code": "20202", "data_type": "文字"},
        "20202008": {"indicator_name": "占位文字", "form_code": "20202", "data_type": "文字"},
    },
}


def write_sheet(book: Workbook, title: str, records: list[list[object]]) -> None:
    sheet = book.create_sheet(title)
    sheet.append(["报表", None, None])
    sheet.append([None, None, None])
    sheet.append(["指标代码", "指标名称", "本期情况"])
    for record in records:
        sheet.append(record)


def create_fixtures() -> None:
    FIXTURE_DIR.mkdir(parents=True, exist_ok=True)
    CONFIG_PATH.write_text(json.dumps(CONFIG, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    book = Workbook()
    book.remove(book.active)
    write_sheet(
        book,
        "20202",
        [
            ["20202001", "存款余额", 1_234_567.89],
            ["20202002", "累计金额", 12_345_678.9],
            ["20201028", "持股比例", 9.92],
            ["20202003", "机构数量", 42],
            ["20202004", "真实零值", 0],
            ["20202005", "逗号金额", "1,234.50"],
            ["20202006", "无效金额", "待核实"],
            ["20202007", "文字示例", " 保留 文本 "],
            ["20202008", "占位文字", 0],
            ["20209999", "未配置指标", 8.5],
            [None, "没有代码的行", 123],
            [2**63, "超过有符号整数范围的代码", 1],
        ],
    )
    write_sheet(
        book,
        "20203",
        [
            ["20202001", "重复余额", 9_999_999],
            ["20202004", "重复零值", 0],
        ],
    )
    book.save(FIXTURE_DIR / PRIMARY_FILE)
    book.close()

    name_mapped = Workbook()
    sheet = name_mapped.active
    sheet.title = "20202"
    sheet.append(["报表", None, None])
    sheet.append([None, None, None])
    sheet.append(["指标代码", "指标名称", "本期情况"])
    sheet.append(["20202001", "乙银行存款余额", 10_000])
    name_mapped.save(FIXTURE_DIR / NAME_MAPPED_FILE)
    name_mapped.close()

    # These are filtered before opening; the unrecognized file and lock file
    # intentionally contain invalid workbook bytes.
    (FIXTURE_DIR / "unrecognized.xlsx").write_bytes(b"not an xlsx")
    (FIXTURE_DIR / "~$ignored.xlsx").write_bytes(b"not an xlsx")


def decimal_tuple(value: Decimal) -> dict[str, object]:
    parts = value.as_tuple()
    return {"sign": parts.sign, "digits": list(parts.digits), "exponent": parts.exponent}


def serial_value(value: datetime | date) -> float:
    base = datetime(1899, 12, 30)
    current = datetime(value.year, value.month, value.day)
    if isinstance(value, datetime):
        current = value.replace(tzinfo=None)
    delta = current - base
    return delta.days + delta.seconds / 86_400 + delta.microseconds / 86_400_000_000


def tagged(value: object) -> dict[str, object]:
    if value is None:
        return {"type": "empty"}
    if isinstance(value, bool):
        return {"type": "boolean", "value": value}
    if isinstance(value, Decimal):
        return {"type": "decimal", "value": decimal_tuple(value)}
    if isinstance(value, (datetime, date)):
        return {"type": "datetime_serial", "value": serial_value(value)}
    if isinstance(value, int):
        return {"type": "integer", "value": value}
    if isinstance(value, float):
        return {"type": "number", "value": value if value == value and abs(value) != float("inf") else None}
    if isinstance(value, str):
        return {"type": "string", "value": value}
    raise TypeError(f"unsupported fixture value: {type(value).__name__}")


def record_payload(record) -> dict[str, object]:
    return {
        "institution_id": record.institution_id,
        "institution_name": record.institution_name,
        "social_credit_code": record.social_credit_code,
        "institution_type": record.institution_type,
        "handling_branch": record.handling_branch,
        "region": record.region,
        "form_code": record.form_code,
        "indicator_code": record.indicator_code,
        "indicator_name": record.indicator_name,
        "period": record.period,
        "source_date": record.source_date,
        "value": tagged(record.value),
        "value_type": record.value_type,
        "source_unit": record.source_unit,
        "normalized_value": tagged(record.normalized_value),
        "normalized_unit": record.normalized_unit,
        "source": {
            "source_file": Path(record.source.source_file).name,
            "source_sheet": record.source.source_sheet,
            "source_row": record.source.source_row,
        },
    }


def finding_payload(finding) -> dict[str, object]:
    payload = asdict(finding)
    for source in payload["source_refs"]:
        source["source_file"] = Path(source["source_file"]).name
    trace = payload["calculation_trace"]
    if trace:
        normalized = []
        for segment in trace.split("；"):
            label, detail = segment.split("：", 1)
            location, row = detail.split("!第", 1)
            source_path, sheet = location.rsplit("/", 1)
            normalized.append(f"{label}：{Path(source_path).name}/{sheet}!第{row}")
        payload["calculation_trace"] = "；".join(normalized)
    return payload


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source-root", type=Path, default=Path(r"C:\Users\AZI\Desktop\VibeCode\3.open-data-audit"))
    parser.add_argument("--output", type=Path, default=WORKSPACE / "probe/reports/s2_import_python_baseline.json")
    parser.add_argument("--validation-output", type=Path, default=WORKSPACE / "probe/reports/s2_validation_python_baseline.json")
    args = parser.parse_args()
    create_fixtures()
    sys.dont_write_bytecode = True
    source_root = args.source_root.resolve()
    source_package_root = str(source_root / "core" / "src")
    sys.path.insert(0, source_package_root)
    importer = importlib.import_module("base_audit.systems.s2_report_collection.importer")
    models = importlib.import_module("base_audit.systems.s2_report_collection.models")
    validation = importlib.import_module("base_audit.systems.s2_report_collection.validation")

    config = models.ReportAuditConfig(
        unit_settings=models.UnitSettings("元", "元", "万元"),
        institutions_by_id={
            "9144": models.InstitutionProfile(
                "9144", "配置名称不覆盖文件名", "CREDIT-9144", "农商行", "一部", "示例地区"
            )
        },
        institutions_by_name={
            "乙银行": models.InstitutionProfile(
                "9200", "乙银行配置名称", "CREDIT-9200", "城商行", "二部", "另一个地区"
            )
        },
        indicators={
            code: models.IndicatorConfig(
                code,
                indicator_name=item["indicator_name"],
                form_code=item["form_code"],
                data_type=item["data_type"],
            )
            for code, item in CONFIG["indicators"].items()
        },
    )
    result = importer.import_period_directory(FIXTURE_DIR, config=config, label="本期")
    wire_result = {
        "records": [record_payload(record) for record in result.records],
        "duplicates": [
            [record_payload(first), record_payload(second)]
            for first, second in result.duplicates
        ],
        "warnings": result.warnings,
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    report = {
        "status": "passed_on_synthetic_fixture",
        "source": {
            "project": str(source_root),
            "module": "core/src/base_audit/systems/s2_report_collection/importer.py",
            "sha256": hashlib.sha256(
                (source_root / "core/src/base_audit/systems/s2_report_collection/importer.py").read_bytes()
            ).hexdigest().upper(),
            "baseline_python": sys.version.split()[0],
            "reader": "source importer selected python-calamine when available; legacy fallback unchanged",
        },
        "fixture": {
            "directory": "tests/fixtures/s2_import",
            "workbooks": [PRIMARY_FILE, NAME_MAPPED_FILE],
            "scenario_count": len(wire_result["records"])
            + len(wire_result["duplicates"])
            + len(wire_result["warnings"]),
        },
        "result": wire_result,
        "scope": "Synthetic S2 .xlsx directory import: filename identity, first-three-column rows, institution/indicator enrichment, amount normalization, duplicate retention and warnings. Not a real workbook or end-to-end S2 business flow.",
    }
    args.output.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    # Exercise the validation-only branches that normal filename-derived rows
    # cannot produce, while keeping the importer baseline itself unchanged.
    missing_identity_record = replace(
        result.records[-1],
        institution_id="",
        form_code="",
        indicator_code="MISSING-IDENTITY",
        source=models.SourceRef("synthetic-missing.xlsx", "Synthetic", 99),
    )
    validation_input = importer.ImportResult(
        records=[*result.records, missing_identity_record],
        duplicates=list(result.duplicates),
        warnings=list(result.warnings),
    )
    config_warnings = ["合成配置缺少字段", "合成配置存在警告"]
    external_warnings = ["外部数据缺失", "外部值不是数值"]
    validation_result = {
        "import_findings": [finding_payload(item) for item in validation.validate_import(validation_input)],
        "config_findings": [finding_payload(item) for item in validation.validate_config_warnings(config_warnings)],
        "external_findings": [finding_payload(item) for item in validation.validate_external_warnings(external_warnings)],
    }
    validation_report = {
        "status": "python_baseline_generated",
        "source": {
            "project": str(source_root),
            "module": "core/src/base_audit/systems/s2_report_collection/validation.py",
            "sha256": hashlib.sha256(
                (source_root / "core/src/base_audit/systems/s2_report_collection/validation.py").read_bytes()
            ).hexdigest().upper(),
            "baseline_python": sys.version.split()[0],
        },
        "inputs": {
            "import_records": len(result.records),
            "duplicates": len(result.duplicates),
            "import_warnings": len(result.warnings),
            "validation_only_missing_identity_records": 1,
            "config_warnings": config_warnings,
            "external_warnings": external_warnings,
        },
        "result": validation_result,
        "scope": "S2 import-quality, config-warning and external-warning finding conversion only; no rule engine or report-writing workflow.",
    }
    args.validation_output.parent.mkdir(parents=True, exist_ok=True)
    args.validation_output.write_text(
        json.dumps(validation_report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    print(json.dumps({"baseline": str(args.output), "records": len(wire_result["records"]), "duplicates": len(wire_result["duplicates"]), "warnings": len(wire_result["warnings"]), "source_sha256": report["source"]["sha256"], "validation_baseline": str(args.validation_output), "validation_findings": sum(len(values) for values in validation_result.values()), "validation_source_sha256": validation_report["source"]["sha256"]}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
