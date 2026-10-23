#!/usr/bin/env python3
"""Generate isolated synthetic workbooks for S2 configuration checker parity."""

from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook

WORKSPACE = Path(__file__).resolve().parents[2]
FIXTURE_DIR = WORKSPACE / "tests" / "fixtures" / "s2_config"

GENERAL_HEADERS = ["设置项", "设置值", "说明"]
INDICATOR_HEADERS = [
    "指标代码", "指标名称", "表单代码", "数据属性", "环比启用", "环比策略组", "最小变动值", "禁用", "备注",
]
STRATEGY_HEADERS = ["策略组", "规则编号", "数据属性", "下限", "上限", "审核级别", "分类说明", "禁用"]
RULE_HEADERS = ["规则编号", "类型", "描述", "规则内容", "Thd值(万元)", "取反", "级别", "禁用", "备注"]
EXTERNAL_HEADERS = ["规则编号", "指标代码", "比较方式", "外部指标名称", "容差", "级别", "启用"]


def make_valid_config(
    path: Path,
    *,
    legacy_severity: bool = False,
    gap_duplicate: bool = False,
    datetime_indicator_flag: bool = False,
) -> None:
    book = Workbook()
    general = book.active
    general.title = "通用设置"
    general.append(GENERAL_HEADERS)
    for row in [
        ["源数据单位", "元", ""],
        ["外部文件单位", "元", ""],
        ["输出文件单位", "元", ""],
        ["换算范围", "", ""],
        ["是否隐藏无变动指标", "否", ""],
        ["是否隐藏与大集中一致指标", "否", ""],
        ["文字变动默认级别", "硬性", ""],
        ["本期缺失默认级别", "硬性", ""],
        ["上期缺失默认级别", "硬性", ""],
    ]:
        general.append(row)

    indicators = book.create_sheet("指标参照")
    indicators.append(INDICATOR_HEADERS)
    enabled: str | datetime = datetime(2024, 1, 2) if datetime_indicator_flag else "是"
    indicators.append(["1001", "示例指标", "F01", "余额", enabled, "G1", 0, "否", "合成夹具"])

    institutions = book.create_sheet("机构参照")
    institutions.append(["机构名称", "机构代码", "机构类别", "承接行", "归属行", "报表项目", "禁用"])

    strategies = book.create_sheet("环比策略")
    strategies.append(STRATEGY_HEADERS)
    level = "严重" if legacy_severity else "硬性"
    if gap_duplicate:
        strategies.append(["G1", "B1", "余额", None, -10, level, "负值区间", "否"])
        strategies.append(["G1", "B1", "余额", -9, None, level, "正值区间", "否"])
    else:
        strategies.append(["G1", "B1", "余额", None, None, level, "全区间", "否"])

    rules = book.create_sheet("校验规则")
    rules.append(RULE_HEADERS)

    external = book.create_sheet("外部核对规则")
    external.append(EXTERNAL_HEADERS)

    book.save(path)
    book.close()


def main() -> int:
    FIXTURE_DIR.mkdir(parents=True, exist_ok=True)
    make_valid_config(FIXTURE_DIR / "check_valid.xlsx")
    make_valid_config(FIXTURE_DIR / "check_legacy_severity.xlsx", legacy_severity=True)
    make_valid_config(FIXTURE_DIR / "check_gap_duplicate.xlsx", gap_duplicate=True)
    make_valid_config(
        FIXTURE_DIR / "check_datetime_indicator_flag.xlsx",
        datetime_indicator_flag=True,
    )

    original_cases = json.loads((FIXTURE_DIR / "cases.json").read_text(encoding="utf-8"))["cases"]
    checker_cases = [
        *original_cases,
        {"id": "check_valid", "file": "check_valid.xlsx"},
        {"id": "check_legacy_severity", "file": "check_legacy_severity.xlsx"},
        {"id": "check_gap_duplicate", "file": "check_gap_duplicate.xlsx"},
        {"id": "check_datetime_indicator_flag", "file": "check_datetime_indicator_flag.xlsx"},
    ]
    (FIXTURE_DIR / "check_cases.json").write_text(
        json.dumps({"cases": checker_cases}, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    print(json.dumps({"fixture_dir": str(FIXTURE_DIR), "checker_case_count": len(checker_cases)}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
