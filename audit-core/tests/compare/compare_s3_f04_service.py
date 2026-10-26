#!/usr/bin/env python3
"""Compare the Rust S3-F04 CSV service with the read-only Python baseline."""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import subprocess
import sys
import tempfile
from pathlib import Path

sys.dont_write_bytecode = True


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def excel_value(row: dict, header: str, target_unit: str):
    base_name = header.replace("(%)", "").replace("(单位)", "")
    if base_name in {"左值", "右值", "差异", "差异绝对值"} and header != "差异幅度(%)":
        label = f"{base_name}({target_unit})" if base_name == "差异" else base_name
        return row.get(label, row.get(base_name))
    return row.get(header)


def same_value(expected, actual) -> bool:
    if expected in (None, ""):
        return actual in (None, "")
    if isinstance(expected, (int, float)) and not isinstance(expected, bool):
        return isinstance(actual, (int, float)) and math.isclose(
            float(expected), float(actual), rel_tol=1e-12, abs_tol=1e-12
        )
    return expected == actual


def mismatch_kind(expected, actual) -> str:
    if expected in (None, ""):
        return "expected_blank"
    if actual in (None, ""):
        return "actual_blank"
    if type(expected) is not type(actual):
        return "type"
    if isinstance(expected, str) and isinstance(actual, str):
        return "text_same_length" if len(expected) == len(actual) else "text_length"
    return "different_value"


def plain_cell(value):
    return "" if value is None else value


def column_width_for_index(sheet, column_index: int):
    """Read grouped OOXML <col min/max> ranges as the width of each column."""
    for dimension in sheet.column_dimensions.values():
        if (
            dimension.min is not None
            and dimension.max is not None
            and dimension.min <= column_index <= dimension.max
        ):
            return dimension.width
    default_width = sheet.sheet_format.defaultColWidth
    return 13.0 if default_width is None else default_width


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--python-project", type=Path, required=True)
    parser.add_argument("--rust-cli", type=Path, required=True)
    parser.add_argument("--current-csv", type=Path, action="append", required=True)
    parser.add_argument("--previous-csv", type=Path, action="append", default=[])
    parser.add_argument("--config-workbook", type=Path, action="append", required=True)
    args = parser.parse_args()

    project = args.python_project.resolve()
    sys.path.insert(0, str(project / "core" / "src"))
    from base_audit.systems.s3_central_statistics.config import load_central_config
    from base_audit.systems.s3_central_statistics.comparison_engine import get_unit_factor
    from base_audit.systems.s3_central_statistics.cross_period_engine import (
        DEFAULT_TITLE,
        run_cross_period,
    )
    from base_audit.systems.s3_central_statistics.service import _load_dataset

    inputs = [*args.current_csv, *args.previous_csv, *args.config_workbook]
    before = {path.resolve(): sha256(path) for path in inputs}
    config = load_central_config(args.config_workbook)
    source_unit = config.param("默认源数据单位", "元")
    target_unit = config.param("默认目标单位", "亿元")
    factor = get_unit_factor(source_unit) / get_unit_factor(target_unit)
    exemptions = config.exempt_indicators()
    current = _load_dataset(args.current_csv, exempt=exemptions, factor=factor)
    previous = (
        _load_dataset(args.previous_csv, exempt=exemptions, factor=factor)
        if args.previous_csv
        else None
    )
    expected_rows, rule_logs = run_cross_period(
        current=current,
        previous=previous,
        config=config,
        target_unit=target_unit,
    )
    disabled = sum(
        str(row.get("禁用") or "").strip() in {"是", "1", "true", "True", "Y", "y"}
        for row in config.cross_rules
    )
    log_rows = [
        (
            "配置载入",
            f"本期数值核对规则：读取 {len(config.cross_rules)} 条，启用 {len(config.cross_rules) - disabled} 条，禁用 {disabled} 条",
            "",
            "",
            "",
        ),
        (
            "配置载入",
            "数据日期：本期 "
            + (current.record_date or "未知")
            + (
                f"；上期 {previous.record_date}"
                if previous is not None
                else "（未选上期，{...} 规则跳过）"
            ),
            "",
            "",
            "",
        ),
        ("规则命中", f"输出 {len(expected_rows)} 行", "", "", ""),
        *rule_logs,
    ]

    with tempfile.TemporaryDirectory(prefix="audit-s3-f04-parity-") as temporary:
        temporary_dir = Path(temporary)
        output_path = temporary_dir / "result.xlsx"
        log_path = temporary_dir / "run-log.xlsx"
        command = [
            str(args.rust_cli.resolve()),
            "s3-check-current-values",
        ]
        for path in args.current_csv:
            command.extend(["--current-csv", str(path.resolve())])
        for path in args.previous_csv:
            command.extend(["--previous-csv", str(path.resolve())])
        for path in args.config_workbook:
            command.extend(["--config-workbook", str(path.resolve())])
        command.extend(
            [
                "--output-workbook",
                str(output_path),
                "--log-workbook",
                str(log_path),
                "--developer-view",
            ]
        )
        completed = subprocess.run(command, text=True, capture_output=True, check=False)
        if completed.returncode != 0:
            print(
                json.dumps(
                    {
                        "status": "failed",
                        "phase": "rust_cli",
                        "exit_code": completed.returncode,
                    },
                    ensure_ascii=False,
                )
            )
            return 2

        try:
            report = json.loads(completed.stdout)
        except json.JSONDecodeError:
            print(json.dumps({"status": "failed", "phase": "rust_report_json"}))
            return 2

        from openpyxl import load_workbook

        rust_book = load_workbook(output_path, data_only=False, read_only=False)
        rust_sheet = rust_book["对比结果"]
        mismatches = 0
        structure_mismatches = 0
        header_mismatches = 0
        value_mismatches = 0
        value_mismatches_by_column = {}
        value_mismatch_kinds = {}
        workbook_feature_mismatches = 0
        width_mismatches = 0
        style_mismatches = 0
        number_format_mismatches = 0
        if rust_book.sheetnames != ["对比结果"]:
            structure_mismatches += 1
        if rust_sheet.max_row != len(expected_rows) + 1 or rust_sheet.max_column != len(DEFAULT_TITLE):
            structure_mismatches += 1
        if rust_sheet.freeze_panes != "A2":
            structure_mismatches += 1

        for column, header in enumerate(DEFAULT_TITLE, start=1):
            if rust_sheet.cell(1, column).value != header:
                header_mismatches += 1
        for row_index, expected in enumerate(expected_rows, start=2):
            for column, header in enumerate(DEFAULT_TITLE, start=1):
                actual = rust_sheet.cell(row_index, column).value
                wanted = excel_value(expected, header, target_unit)
                if not same_value(wanted, actual):
                    value_mismatches += 1
                    value_mismatches_by_column[header] = value_mismatches_by_column.get(header, 0) + 1
                    kinds = value_mismatch_kinds.setdefault(header, {})
                    kind = mismatch_kind(wanted, actual)
                    kinds[kind] = kinds.get(kind, 0) + 1

        if rust_sheet.merged_cells.ranges:
            workbook_feature_mismatches += 1
        if any(cell.data_type == "f" for row in rust_sheet.iter_rows() for cell in row):
            workbook_feature_mismatches += 1
        expected_widths = (12, 10, 22, 10, 10, 10, 22, 10, 14, 12, 12, 12, 12, 12, 24, 24, 44)
        for column, expected_width in enumerate(expected_widths, start=1):
            width = column_width_for_index(rust_sheet, column)
            if width is None or not math.isclose(width, expected_width, abs_tol=0.1):
                width_mismatches += 1
        header_fill = rust_sheet.cell(1, 1).fill.fgColor.rgb or ""
        if not header_fill.endswith("1F4E78") or not rust_sheet.cell(1, 1).font.bold:
            style_mismatches += 1
        for row_index in range(2, len(expected_rows) + 2):
            for column in range(10, 15):
                value = rust_sheet.cell(row_index, column).value
                if value is not None and rust_sheet.cell(row_index, column).number_format != "0.00;-0.00;":
                    number_format_mismatches += 1
        mismatches += sum(
            (
                structure_mismatches,
                header_mismatches,
                value_mismatches,
                workbook_feature_mismatches,
                width_mismatches,
                style_mismatches,
                number_format_mismatches,
            )
        )
        rust_book.close()

        log_book = load_workbook(log_path, data_only=True, read_only=True)
        log_sheet = log_book["运行日志"]
        actual_logs = [
            tuple(plain_cell(log_sheet.cell(row, column).value) for column in range(2, 7))
            for row in range(2, log_sheet.max_row + 1)
        ]
        expected_logs = [tuple(plain_cell(value) for value in row) for row in log_rows]
        if actual_logs != expected_logs:
            mismatches += 1
        log_mismatches = int(actual_logs != expected_logs)
        log_book.close()

        after = {path.resolve(): sha256(path) for path in inputs}
        inputs_unchanged = before == after
        result = {
            "status": "passed" if mismatches == 0 and inputs_unchanged else "failed",
            "current_file_count": len(args.current_csv),
            "previous_file_count": len(args.previous_csv),
            "config_file_count": len(args.config_workbook),
            "current_record_count": len(current.records),
            "previous_record_count": len(previous.records) if previous else 0,
            "configured_rule_count": len(config.cross_rules),
            "enabled_rule_count": len(config.cross_rules) - disabled,
            "python_result_rows": len(expected_rows),
            "rust_result_rows": report.get("rows"),
            "python_log_rows": len(expected_logs),
            "rust_log_rows": len(actual_logs),
            "workbook_mismatches": mismatches,
            "mismatch_counts": {
                "structure": structure_mismatches,
                "headers": header_mismatches,
                "row_values": value_mismatches,
                "row_values_by_column": value_mismatches_by_column,
                "row_value_kinds_by_column": value_mismatch_kinds,
                "workbook_features": workbook_feature_mismatches,
                "column_widths": width_mismatches,
                "header_style": style_mismatches,
                "number_formats": number_format_mismatches,
                "run_log": log_mismatches,
            },
            "inputs_unchanged": inputs_unchanged,
            "rust_engines": {
                "config_reader": report.get("config_reader"),
                "input_reader": report.get("input_reader"),
                "writer": report.get("writer"),
            },
            "rust_timings_ms": {
                "config_read": report.get("config_read_time_ms"),
                "input_read": report.get("input_read_time_ms"),
                "calculation": report.get("calculation_time_ms"),
                "workbook_write": report.get("workbook_write_time_ms"),
                "total": report.get("total_time_ms"),
            },
        }
        print(json.dumps(result, ensure_ascii=False, indent=2))
        return 0 if result["status"] == "passed" else 1


if __name__ == "__main__":
    raise SystemExit(main())
