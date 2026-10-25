#!/usr/bin/env python3
"""Compare the Rust S3-F02 CLI with the read-only Python exporter on synthetic XLSX files."""

from __future__ import annotations

import argparse
import json
import math
import platform
import subprocess
import sys
import tempfile
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")


HEADERS = (
    "机构类代码", "机构类名称", "地区代码", "地区名称", "指标代码", "指标名称",
    "数据值", "上期值", "输出说明", "说明内容", "增减额（先出现）", "增减额（后出现）",
)
ROWS = (
    ("7020001", "示例村镇银行", "4401", "惠城区", "I-01", "指标一", 0.418462, 0.25, "是", "补充说明一", -0.168462, 999),
    ("12345", "普通银行", "510101", "省级", "I-02", "指标二", 10.5, 10, "否", "不应导出", 0.5, 998),
    ("12345", "普通银行", "510102", "城北区", "I-03", "指标三", 25, 20, " YES ", "补充说明三", 5, 997),
    ("99999", "地方机构", "6101", "新区", "I-04", "指标四", 40, 35, "✓", "补充说明四", 5, 996),
    ("4567", "未映射机构", "3201", "南区", "I-05", "指标五", 12, 11, "1", "补充说明五", 1, 995),
)
CONFIG_ROWS = (
    ("7020001", "4401", "精确映射机构"),
    ("12345", "510102", "精确市级机构"),
    ("12345", "*", "显式通配机构"),
    ("99999", "*", "全国通配机构"),
)


def load_baseline(source_root: Path):
    sys.path.insert(0, str(source_root / "core" / "src"))
    from base_audit.systems.s3_central_statistics.explanation_exporter import (
        export_selected_explanations,
    )

    return export_selected_explanations


def create_inputs(comparison_path: Path, config_path: Path) -> None:
    comparison = Workbook()
    cover = comparison.active
    cover.title = "封面"
    cover["A1"] = "不是比较结果"
    result = comparison.create_sheet("比较结果")
    result.append(HEADERS)
    for row in ROWS:
        result.append(row)
    comparison.save(comparison_path)

    config = Workbook()
    sheet = config.active
    sheet.title = "机构地区参照"
    sheet.append(("机构类代码", "地区代码", "机构名称"))
    for row in CONFIG_ROWS:
        sheet.append(row)
    config.save(config_path)


def color_value(color):
    if color is None:
        return None
    if color.type == "rgb":
        return ("rgb", color.rgb[-6:].upper())
    if color.type == "theme" and color.theme == 1:
        return None
    return (color.type, color.indexed if color.type == "indexed" else color.value)


def style_value(cell):
    fill = cell.fill
    return {
        "font": (
            cell.font.name or "Calibri",
            cell.font.bold,
            cell.font.sz or 11,
            cell.font.italic,
            cell.font.underline,
            cell.font.strike,
            color_value(cell.font.color),
        ),
        "fill": (
            fill.fill_type,
            color_value(fill.fgColor) if fill.fill_type else None,
        ),
        "alignment": (
            cell.alignment.horizontal,
            cell.alignment.vertical,
            cell.alignment.wrap_text,
            cell.alignment.text_rotation,
            cell.alignment.shrink_to_fit,
            cell.alignment.indent,
            cell.alignment.readingOrder,
        ),
        "number_format": cell.number_format,
        "border": tuple(
            (
                getattr(cell.border, side).style,
                color_value(getattr(cell.border, side).color),
            )
            for side in ("left", "right", "top", "bottom", "diagonal")
        ),
        "protection": (cell.protection.locked, cell.protection.hidden),
    }


def column_width(sheet, column: str) -> float:
    index = ord(column) - ord("A") + 1
    for dimension in sheet.column_dimensions.values():
        if dimension.min is not None and dimension.max is not None:
            if dimension.min <= index <= dimension.max:
                return dimension.width
    dimension = sheet.column_dimensions.get(column)
    return dimension.width if dimension is not None else sheet.sheet_format.defaultColWidth


def column_hidden(sheet, column: str) -> bool:
    index = ord(column) - ord("A") + 1
    return any(
        dimension.hidden
        for dimension in sheet.column_dimensions.values()
        if dimension.min is not None
        and dimension.max is not None
        and dimension.min <= index <= dimension.max
    )


def row_value(sheet, row: int) -> tuple[float | None, bool, int, bool]:
    dimension = sheet.row_dimensions.get(row)
    if dimension is None:
        return sheet.sheet_format.defaultRowHeight, False, 0, False
    height = dimension.height
    if height is None:
        height = sheet.sheet_format.defaultRowHeight
    return height, dimension.hidden, dimension.outlineLevel, dimension.collapsed


def page_setup_value(sheet) -> tuple[object, ...]:
    page_setup = sheet.page_setup
    page_setup_properties = sheet.sheet_properties.pageSetUpPr
    margins = sheet.page_margins
    return (
        None if page_setup_properties is None else page_setup_properties.fitToPage,
        page_setup.orientation or "portrait",
        page_setup.fitToWidth,
        page_setup.fitToHeight,
        page_setup.scale,
        page_setup.paperSize,
        str(sheet.print_area),
        margins.left,
        margins.right,
        margins.top,
        margins.bottom,
        margins.header,
        margins.footer,
        sheet.print_options.horizontalCentered,
        sheet.print_options.verticalCentered,
    )


def defined_names_value(workbook) -> list[tuple[object, ...]]:
    return sorted(
        (
            name,
            definition.attr_text,
            definition.localSheetId,
            definition.hidden,
        )
        for name, definition in workbook.defined_names.items()
    )


def equal_value(left, right) -> bool:
    if isinstance(left, (int, float)) and not isinstance(left, bool):
        return isinstance(right, (int, float)) and not isinstance(right, bool) and math.isclose(
            left, right, rel_tol=0.0, abs_tol=1e-12
        )
    return left == right


def compare_workbooks(expected_path: Path, actual_path: Path) -> tuple[list[str], int]:
    expected = load_workbook(expected_path, data_only=False)
    actual = load_workbook(actual_path, data_only=False)
    differences: list[str] = []

    def add_difference(difference: str) -> None:
        if difference not in differences:
            differences.append(difference)

    try:
        if expected.sheetnames != actual.sheetnames:
            add_difference("sheetnames")
        if defined_names_value(expected) != defined_names_value(actual):
            add_difference("defined_names")
        left = expected["说明文件"]
        right = actual["说明文件"]
        expected_info = expected["生成信息"]
        actual_info = actual["生成信息"]
        for sheet in (expected_info, actual_info):
            timestamp = sheet["B3"].value
            try:
                datetime.strptime(timestamp, "%Y-%m-%d %H:%M:%S")
            except (TypeError, ValueError):
                add_difference("生成信息.B3.timestamp_format")
                break

        for sheet_name in expected.sheetnames:
            if sheet_name not in actual.sheetnames:
                continue
            source = expected[sheet_name]
            output = actual[sheet_name]
            if source.sheet_state != output.sheet_state:
                add_difference(f"{sheet_name}.sheet_state")
            for row in range(1, max(source.max_row, output.max_row) + 1):
                for column in range(1, max(source.max_column, output.max_column) + 1):
                    expected_cell = source.cell(row, column)
                    actual_cell = output.cell(row, column)
                    if not (
                        sheet_name == "生成信息"
                        and expected_cell.coordinate == "B3"
                    ) and not equal_value(expected_cell.value, actual_cell.value):
                        add_difference(
                            f"{sheet_name}.{expected_cell.coordinate}.value_or_formula"
                        )
                    if style_value(expected_cell) != style_value(actual_cell):
                        add_difference(f"{sheet_name}.{expected_cell.coordinate}.style")
                    expected_comment = (
                        None
                        if expected_cell.comment is None
                        else (expected_cell.comment.text, expected_cell.comment.author)
                    )
                    actual_comment = (
                        None
                        if actual_cell.comment is None
                        else (actual_cell.comment.text, actual_cell.comment.author)
                    )
                    if expected_comment != actual_comment:
                        add_difference(f"{sheet_name}.{expected_cell.coordinate}.comment")
                    expected_link = (
                        None if expected_cell.hyperlink is None else expected_cell.hyperlink.target
                    )
                    actual_link = (
                        None if actual_cell.hyperlink is None else actual_cell.hyperlink.target
                    )
                    if expected_link != actual_link:
                        add_difference(f"{sheet_name}.{expected_cell.coordinate}.hyperlink")

            if set(map(str, source.merged_cells.ranges)) != set(
                map(str, output.merged_cells.ranges)
            ):
                add_difference(f"{sheet_name}.merged_ranges")
            if source.freeze_panes != output.freeze_panes:
                add_difference(f"{sheet_name}.freeze_panes")
            if source.auto_filter.ref != output.auto_filter.ref:
                add_difference(f"{sheet_name}.auto_filter")
            if page_setup_value(source) != page_setup_value(output):
                add_difference(f"{sheet_name}.page_setup")
            for column_index in range(1, max(source.max_column, output.max_column) + 1):
                column = get_column_letter(column_index)
                left_width = column_width(source, column)
                right_width = column_width(output, column)
                if (
                    left_width is None
                    or right_width is None
                    or not math.isclose(left_width, right_width, rel_tol=0.0, abs_tol=0.05)
                ):
                    add_difference(f"{sheet_name}.column_width.{column}")
                if column_hidden(source, column) != column_hidden(output, column):
                    add_difference(f"{sheet_name}.column_hidden.{column}")
            for row in range(1, max(source.max_row, output.max_row) + 1):
                if row_value(source, row) != row_value(output, row):
                    add_difference(f"{sheet_name}.row_dimensions.{row}")
            if len(source.data_validations.dataValidation) != len(
                output.data_validations.dataValidation
            ):
                add_difference(f"{sheet_name}.data_validations")
            if len(source.conditional_formatting) != len(output.conditional_formatting):
                add_difference(f"{sheet_name}.conditional_formatting")
            if len(source.tables) != len(output.tables):
                add_difference(f"{sheet_name}.tables")
            if len(source._images) != len(output._images):
                add_difference(f"{sheet_name}.images")
            if len(source._charts) != len(output._charts):
                add_difference(f"{sheet_name}.charts")
        return differences, left.max_row - 2
    finally:
        expected.close()
        actual.close()


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--audit-binary", required=True, type=Path)
    parser.add_argument("--source-root", required=True, type=Path)
    parser.add_argument("--report", required=True, type=Path)
    parser.add_argument(
        "--comparison-file",
        type=Path,
        help="Use a real or sanitized S3 comparison workbook instead of synthetic input.",
    )
    parser.add_argument(
        "--config-workbook",
        type=Path,
        help="Optional organization/region mapping workbook for real-input comparison.",
    )
    args = parser.parse_args()
    if args.config_workbook is not None and args.comparison_file is None:
        parser.error("--config-workbook requires --comparison-file")
    export_selected_explanations = load_baseline(args.source_root)

    result = {
        "feature": "S3-F02",
        "fixture": "real_s3_workbook" if args.comparison_file else "synthetic",
        "baseline": "read-only Python export_selected_explanations",
        "cli_entry": "s3-export-explanations",
        "python_version": platform.python_version(),
        "status": "failed",
    }
    try:
        with tempfile.TemporaryDirectory(prefix="audit-s3-f02-parity-") as temporary:
            root = Path(temporary)
            expected = root / "python-explanations.xlsx"
            comparison = args.comparison_file or root / "comparison.xlsx"
            config = args.config_workbook or root / "mapping.xlsx"
            if args.comparison_file is None:
                create_inputs(comparison, config)
            comparison_before = comparison.read_bytes()
            config_before = config.read_bytes() if config.is_file() else None
            selected = export_selected_explanations(
                comparison,
                expected,
                config_path=config if config_before is not None else None,
            )
            differences = []
            view_reports = {}
            output_rows = 0
            for view in ("user", "developer"):
                actual = root / f"rust-explanations-{view}.xlsx"
                command = [
                    str(args.audit_binary),
                    "s3-export-explanations",
                    "--comparison-file",
                    str(comparison),
                    "--output-workbook",
                    str(actual),
                ]
                if config_before is not None:
                    command.extend(("--config-workbook", str(config)))
                if view == "developer":
                    command.append("--developer-view")
                completed = subprocess.run(
                    command,
                    check=True,
                    capture_output=True,
                    text=True,
                    encoding="utf-8",
                )
                report = json.loads(completed.stdout)
                workbook_differences, output_rows = compare_workbooks(expected, actual)
                differences.extend(f"{view}_view:{item}" for item in workbook_differences)
                if report.get("status") != "completed":
                    differences.append(f"{view}_view:status")
                if report.get("selected_rows") != selected:
                    differences.append(f"{view}_view:selected_rows")
                view_reports[view] = {
                    "status": report.get("status"),
                    "selected_rows": report.get("selected_rows"),
                }
                if view == "developer":
                    events = [
                        json.loads(line)
                        for line in completed.stderr.splitlines()
                        if line.strip()
                    ]
                    event_statuses = [event.get("status") for event in events]
                    if event_statuses != ["started", "completed"]:
                        differences.append("developer_view:progress_events")
                    if report.get("reader") != "calamine":
                        differences.append("developer_view:reader")
                    if report.get("writer") != "rust_xlsxwriter":
                        differences.append("developer_view:writer")
                    completed_event = events[-1] if events else {}
                    timing_fields = (
                        "comparison_read_time_ms",
                        "config_read_time_ms",
                        "workbook_write_time_ms",
                        "elapsed_ms",
                    )
                    timings = {
                        field: completed_event.get(field) for field in timing_fields
                    }
                    if not all(isinstance(value, int) for value in timings.values()):
                        differences.append("developer_view:timings")
                    view_reports[view].update(
                        {
                            "reader": report.get("reader"),
                            "writer": report.get("writer"),
                            "progress_statuses": event_statuses,
                            "timings_ms": timings,
                        }
                    )
            if comparison.read_bytes() != comparison_before:
                differences.append("comparison_input_changed")
            if config_before is not None and config.read_bytes() != config_before:
                differences.append("configuration_input_changed")
            comparison_unchanged = comparison.read_bytes() == comparison_before
            config_unchanged = (
                config_before is None or config.read_bytes() == config_before
            )
            result.update(
                {
                    "status": "passed" if not differences else "failed",
                    "selected_rows_python": selected,
                    "selected_rows_rust": view_reports["user"]["selected_rows"],
                    "selected_rows_rust_developer": view_reports["developer"][
                        "selected_rows"
                    ],
                    "output_rows_compared": output_rows,
                    "rust_reader": view_reports["developer"]["reader"],
                    "rust_writer": view_reports["developer"]["writer"],
                    "views": view_reports,
                    "differences": differences,
                    "comparison_input_unchanged": comparison_unchanged,
                    "configuration_input_unchanged": config_unchanged,
                    "checks": [
                        "selected rows and cell values",
                        "region code and exact/wildcard institution names",
                        "header order and comparison data preserved without rounding",
                        "all output values, formula text, effective styles and names",
                        "sheet order, merges, freeze panes, filters, row/column sizes",
                        "page setup including margins",
                        "metadata content and timestamp format",
                        "user and developer CLI views with measured engine timings",
                        "comparison and configuration workbooks remain inputs only",
                    ],
                }
            )
    except Exception as error:  # keep the report small and avoid recording local paths
        result["error"] = type(error).__name__

    args.report.parent.mkdir(parents=True, exist_ok=True)
    args.report.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0 if result["status"] == "passed" else 1


if __name__ == "__main__":
    raise SystemExit(main())
