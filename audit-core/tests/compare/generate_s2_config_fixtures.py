#!/usr/bin/env python3
"""Build synthetic S2 configuration workbooks for the Rust/Python parity probe."""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook

WORKSPACE = Path(__file__).resolve().parents[2]
FIXTURE_DIR = WORKSPACE / "tests" / "fixtures" / "s2_config"

GENERAL_HEADERS = ["设置项", "设置值", "说明"]
INDICATOR_HEADERS = [
    "指标代码", "指标名称", "表单代码", "数据属性", "环比启用", "环比策略组", "最小变动值", "禁用", "备注",
]
EXTERNAL_HEADERS = ["规则编号", "指标代码", "比较方式", "外部指标名称", "容差", "级别", "启用"]


def workbook_with_general() -> Workbook:
    book = Workbook()
    general = book.active
    general.title = "通用设置"
    general.append(GENERAL_HEADERS)
    general.append(["源数据单位", "亿元", ""])
    general.append(["外部文件单位", "万元", ""])
    general.append(["输出文件单位", "元", ""])
    general.append(["换算范围", "", "仅金额"])
    general.append(["是否隐藏无变动指标", "是", ""])
    general.append(["是否隐藏与大集中一致指标", "否", ""])
    general.append(["文字变动默认级别", "软性", ""])
    general.append(["本期缺失默认级别", "软性", ""])
    general.append(["上期缺失默认级别", "硬性", ""])
    return book


def add_required_sheets(book: Workbook, *, indicator_headers=INDICATOR_HEADERS, external_headers=EXTERNAL_HEADERS) -> None:
    indicators = book.create_sheet("指标参照")
    indicators.append(indicator_headers)
    external = book.create_sheet("外部核对规则")
    external.append(external_headers)


def save(book: Workbook, name: str) -> None:
    book.save(FIXTURE_DIR / name)
    book.close()


def main() -> int:
    FIXTURE_DIR.mkdir(parents=True, exist_ok=True)
    cases: list[dict[str, str]] = []

    book = workbook_with_general()
    indicators = book.create_sheet("指标参照")
    indicators.append(INDICATOR_HEADERS)
    indicators.append([20202001, "存款余额", "20202", "余额", "是", "BAL01", 0.5, "否", "主指标"])
    indicators.append([20202001, "重复指标", "20202", "余额", "是", "BAL01", 0, "否", "重复行跳过"])
    indicators.append([20202002, "缺表单", None, "个数", None, None, None, "TRUE", ""]) 
    indicators.append(["20201028.0", "持股比例", "20201", "百分数", "否", "PCT01", None, "否", "字符串数值代码"])

    institutions = book.create_sheet("机构参照")
    institutions.append(["机构名称", "机构代码", "机构类别", "承接行", "归属行", "报表项目", "禁用"])
    institutions.append(["甲银行", 9144, "农商行", "一部", "华东", "存款项目", "否"])
    institutions.append(["停用银行", 9999, "银行", "二部", "华北", "忽略项目", "是"])
    institutions.append(["仅名称机构", None, "", "", "华南", "仅名称项目", "否"])
    institutions.append(["甲银行", 9144, "更新类别", "三部", "华中", "更新项目", "否"])

    strategies = book.create_sheet("环比策略")
    strategies.append(["策略组", "规则编号", "数据属性", "下限", "上限", "审核级别", "分类说明", "禁用"])
    strategies.append(["BAL01", None, "余额", None, -96, "错误", "降幅低于-96%", "否"])
    strategies.append(["BAL01", "B002", "余额", -96, -90, "提示", "降幅(-96%,-90%]", "否"])
    strategies.append(["BAL01", "B003", "余额", -80, -50, "异常级别", "降幅(-80%,50%)", "是"])
    strategies.append(["BAL01", "B004", "余额", 4, 3, "软性", "非法区间", "否"])
    strategies.append(["PCT01", "P001", "百分数", 0, 30, "", "增长区间", "否"])

    rules = book.create_sheet("校验规则")
    rules.append(["规则编号", "类型", "描述", "规则内容", "Thd值(万元)", "取反", "级别", "禁用", "备注"])
    rules.append(["R001", "表达式", "未知指标引用", "[999999] > 0", 0.1, "否", "严重", "否", "基线规则"])
    rules.append(["R002", "表达式", "有效规则", "[20202001] > Thd", 0.05, "yes", "提示", "否", "取反规则"])
    rules.append(["R003", "未知类型", "不支持", "[20202001] > 0", None, "否", "硬性", "否", "显式跳过"])
    rules.append([None, None, "使用默认编号", "{20201028} = 0", None, None, "BAD", "否", None])

    external = book.create_sheet("外部核对规则")
    external.append(EXTERNAL_HEADERS)
    external.append(["D001", 20202001, "等于", "存款余额", 0, "严重", None])
    external.append([None, 20202001, "等于", "缺少规则号", 0, "硬性", "是"])
    external.append(["D003", 999999, "等于", "未知指标", 0, "硬性", "是"])
    external.append(["D004", 20202001, "等于", None, 0, "硬性", "是"])
    external.append(["D005", 20202001, "大于", "非等值容差", 1, "硬性", "是"])
    external.append(["D006", 20202001, "等于", "无效级别转硬性", 0, "BAD", "是"])
    save(book, "comprehensive.xlsx")
    cases.append({"id": "comprehensive", "file": "comprehensive.xlsx"})

    book = workbook_with_general()
    add_required_sheets(book)
    book["指标参照"].append(["A", "指标A", "20201", "余额", "是", "", 0, "否", ""])
    old_bands = book.create_sheet("环比规则")
    old_bands.append(["规则编号", "下限", "上限", "分类说明", "级别", "禁用"])
    old_bands.append(["OLD1", None, None, "空行跳过", "", "否"])
    old_bands.append([None, -10, 0, "旧格式降幅", "核实", "否"])
    cases.append({"id": "legacy_band_headers", "file": "legacy_band_headers.xlsx"})
    save(book, "legacy_band_headers.xlsx")

    book = Workbook()
    general = book.active
    general.title = "通用设置"
    general.append(GENERAL_HEADERS)
    add_required_sheets(book)
    cases.append({"id": "missing_optional_setting_rows", "file": "missing_optional_setting_rows.xlsx"})
    save(book, "missing_optional_setting_rows.xlsx")

    book = workbook_with_general()
    book["通用设置"]["A1"] = "配置项"
    add_required_sheets(book)
    cases.append({"id": "invalid_general_header", "file": "invalid_general_header.xlsx"})
    save(book, "invalid_general_header.xlsx")

    book = workbook_with_general()
    book["通用设置"].append(["源数据单位", "万元", "重复"])
    add_required_sheets(book)
    cases.append({"id": "duplicate_general_setting", "file": "duplicate_general_setting.xlsx"})
    save(book, "duplicate_general_setting.xlsx")

    book = workbook_with_general()
    book["通用设置"]["B2"] = "美元"
    add_required_sheets(book)
    cases.append({"id": "invalid_unit", "file": "invalid_unit.xlsx"})
    save(book, "invalid_unit.xlsx")

    book = workbook_with_general()
    indicators = book.create_sheet("指标参照")
    indicators.append([*INDICATOR_HEADERS, "源数据单位", "大集中数据单位"])
    external = book.create_sheet("外部核对规则")
    external.append(EXTERNAL_HEADERS)
    cases.append({"id": "inline_unit_columns", "file": "inline_unit_columns.xlsx"})
    save(book, "inline_unit_columns.xlsx")

    book = workbook_with_general()
    book.create_sheet("其他")
    external = book.create_sheet("外部核对规则")
    external.append(EXTERNAL_HEADERS)
    cases.append({"id": "missing_indicator_sheet", "file": "missing_indicator_sheet.xlsx"})
    save(book, "missing_indicator_sheet.xlsx")

    book = workbook_with_general()
    book.create_sheet("指标参照").append(INDICATOR_HEADERS)
    book.create_sheet("外部核对规则").append(["规则编号", "指标代码", "外部系统", "比较方式", "外部指标名称", "容差", "级别", "启用"])
    cases.append({"id": "external_header_order", "file": "external_header_order.xlsx"})
    save(book, "external_header_order.xlsx")

    book = workbook_with_general()
    book.create_sheet("指标参照").append(INDICATOR_HEADERS)
    book.create_sheet("可选表")
    cases.append({"id": "missing_external_sheet", "file": "missing_external_sheet.xlsx"})
    save(book, "missing_external_sheet.xlsx")

    book = Workbook()
    book.active.title = "其他"
    add_required_sheets(book)
    cases.append({"id": "missing_general_sheet", "file": "missing_general_sheet.xlsx"})
    save(book, "missing_general_sheet.xlsx")

    book = workbook_with_general()
    book["通用设置"]["B6"] = "true"
    add_required_sheets(book)
    cases.append({"id": "invalid_boolean_setting", "file": "invalid_boolean_setting.xlsx"})
    save(book, "invalid_boolean_setting.xlsx")

    book = workbook_with_general()
    book["通用设置"]["B8"] = "严重"
    add_required_sheets(book)
    cases.append({"id": "invalid_optional_severity", "file": "invalid_optional_severity.xlsx"})
    save(book, "invalid_optional_severity.xlsx")

    book = workbook_with_general()
    book["通用设置"].append(["未知设置", "启用", ""])
    add_required_sheets(book)
    cases.append({"id": "unknown_general_setting", "file": "unknown_general_setting.xlsx"})
    save(book, "unknown_general_setting.xlsx")

    cases.append({"id": "missing_file", "file": "does_not_exist.xlsx"})
    (FIXTURE_DIR / "cases.json").write_text(
        json.dumps({"cases": cases}, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    print(json.dumps({"fixture_dir": str(FIXTURE_DIR), "case_count": len(cases)}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
