"""Compare a source workbook against the isolated umya probe output."""
from __future__ import annotations

import argparse
import json
import random
import re
from pathlib import Path

from openpyxl import load_workbook


def norm_ref(value):
    return re.sub(r"'([^']+)'!", r"\1!", str(value or "")).replace("$", "").upper()


def cf_snapshot(ws):
    out = []
    for conditional_formatting in ws.conditional_formatting:
        sqref = str(conditional_formatting.sqref)
        for rule in conditional_formatting.rules:
            dxf = None
            if rule.dxf is not None:
                dxf = {
                    "font_bold": rule.dxf.font.bold if rule.dxf.font else None,
                    "font_italic": rule.dxf.font.italic if rule.dxf.font else None,
                    "fill_type": rule.dxf.fill.fill_type if rule.dxf.fill else None,
                    "fill_fg": rule.dxf.fill.fgColor.type + ":" + str(rule.dxf.fill.fgColor.rgb or rule.dxf.fill.fgColor.indexed or rule.dxf.fill.fgColor.theme) if rule.dxf.fill else None,
                    "number_format": rule.dxf.numFmt.numFmt if rule.dxf.numFmt else None,
                }
            out.append({
                "sqref": sqref,
                "type": rule.type,
                "operator": rule.operator,
                "formula": list(rule.formula or []),
                "stop_if_true": rule.stopIfTrue,
                "priority": rule.priority,
                "dxf": dxf,
            })
    return out


def dv_snapshot(ws):
    return sorted(({
        "type": dv.type,
        "operator": dv.operator,
        "formula1": dv.formula1,
        "formula2": dv.formula2,
        "allow_blank": dv.allow_blank,
        "sqref": str(dv.sqref),
        "show_error_message": dv.showErrorMessage,
        "show_input_message": dv.showInputMessage,
    } for dv in ws.data_validations.dataValidation), key=lambda x: x["sqref"])


def names_snapshot(wb):
    result = {}
    for name, definition in wb.defined_names.items():
        result[f"global:{name}"] = norm_ref(definition.attr_text)
    for ws in wb.worksheets:
        for name, definition in ws.defined_names.items():
            result[f"{ws.title}:{name}"] = norm_ref(definition.attr_text)
    return result


def dims_snapshot(ws):
    rows = {str(i): {"height": d.height, "hidden": d.hidden, "style": d.style_id, "outline": d.outlineLevel}
            for i, d in ws.row_dimensions.items() if d.height is not None or d.hidden or d.style_id or d.outlineLevel}
    cols = {k: {"width": d.width, "hidden": d.hidden, "style": d.style, "min": d.min, "max": d.max}
            for k, d in ws.column_dimensions.items() if d.width is not None or d.hidden or d.style or d.min or d.max}
    return rows, cols


def print_settings_snapshot(ws):
    return {
        "orientation": ws.page_setup.orientation,
        "paperSize": ws.page_setup.paperSize,
        "fitToWidth": ws.page_setup.fitToWidth,
        "fitToHeight": ws.page_setup.fitToHeight,
        "scale": ws.page_setup.scale,
        "firstPageNumber": ws.page_setup.firstPageNumber,
        "useFirstPageNumber": ws.page_setup.useFirstPageNumber,
        "fitToPage": ws.sheet_properties.pageSetUpPr.fitToPage if ws.sheet_properties.pageSetUpPr else None,
        "margins": {"left": ws.page_margins.left, "right": ws.page_margins.right, "top": ws.page_margins.top, "bottom": ws.page_margins.bottom, "header": ws.page_margins.header, "footer": ws.page_margins.footer},
        "print_options": {"horizontalCentered": ws.print_options.horizontalCentered, "verticalCentered": ws.print_options.verticalCentered, "gridLines": ws.print_options.gridLines, "headings": ws.print_options.headings},
    }


def cell_value_snapshot(ws):
    return {c.coordinate: (c.data_type, c.value) for row in ws.iter_rows() for c in row if c.value is not None}


def run(source: Path, output: Path, max_random: int) -> dict:
    checks = []
    try:
        src_wb = load_workbook(source, data_only=False)
        out_wb = load_workbook(output, data_only=False)
        checks.append({"name": "openpyxl_openability", "passed": True, "detail": "source and output both open"})
    except Exception as exc:
        return {"source": str(source), "output": str(output), "checks": [{"name": "openpyxl_openability", "passed": False, "detail": repr(exc)}], "passed": 0, "total": 1, "rate": 0.0}

    def add(name, passed, detail=""):
        checks.append({"name": name, "passed": bool(passed), "detail": detail})

    add("sheet_names_and_order", src_wb.sheetnames == out_wb.sheetnames, f"source={src_wb.sheetnames}; output={out_wb.sheetnames}")
    for sheet_name in src_wb.sheetnames:
        sw, ow = src_wb[sheet_name], out_wb[sheet_name]
        prefix = sheet_name + ":"
        add(prefix + "merged_ranges", sorted(map(str, sw.merged_cells.ranges)) == sorted(map(str, ow.merged_cells.ranges)))
        src_cf, out_cf = cf_snapshot(sw), cf_snapshot(ow)
        add(prefix + "conditional_formatting_and_dxf", src_cf == out_cf, f"source_count={len(src_cf)}; output_count={len(out_cf)}")
        src_dv, out_dv = dv_snapshot(sw), dv_snapshot(ow)
        add(prefix + "data_validations", src_dv == out_dv, f"source_count={len(src_dv)}; output_count={len(out_dv)}")
        src_comments = {c.coordinate: c.comment.text for row in sw.iter_rows() for c in row if c.comment is not None}
        out_comments = {c.coordinate: c.comment.text for row in ow.iter_rows() for c in row if c.comment is not None}
        add(prefix + "comment_text", src_comments == out_comments, f"source={src_comments}; output={out_comments}")
        src_names, out_names = names_snapshot(src_wb), names_snapshot(out_wb)
        expected_names = dict(src_names)
        parea_key = f"{sheet_name}:_xlnm.Print_Area"
        if parea_key in expected_names:
            expected_names[parea_key] = f"{sheet_name}!A1:D10"
        add(prefix + "defined_names", expected_names == out_names, f"source={src_names}; expected={expected_names}; output={out_names}")

        src_rows, src_cols = dims_snapshot(sw)
        out_rows, out_cols = dims_snapshot(ow)
        expected_rows = {k: dict(v) for k, v in src_rows.items()}
        expected_rows.setdefault("2", {"height": None, "hidden": False, "style": 0, "outline": 0})["hidden"] = True
        expected_rows.setdefault("3", {"height": None, "hidden": False, "style": 0, "outline": 0})["height"] = 30.0
        add(prefix + "requested_row_dimensions", expected_rows == out_rows,
            f"expected={expected_rows}; output={out_rows}")
        expected_cols = {k: dict(v) for k, v in src_cols.items()}
        expected_cols.setdefault("I", {"width": None, "hidden": False, "style": 0, "min": 0, "max": 0})["width"] = 25.0
        add(prefix + "requested_column_width", out_cols.get("I", {}).get("width") == 25.0,
            f"requested I=25; output={out_cols.get('I')}")
        add(prefix + "other_column_dimensions_preserved", expected_cols == out_cols,
            f"expected={expected_cols}; output={out_cols}")
        add(prefix + "freeze_panes", sw.freeze_panes == ow.freeze_panes, f"{sw.freeze_panes} -> {ow.freeze_panes}")
        add(prefix + "auto_filter", sw.auto_filter.ref == ow.auto_filter.ref, f"{sw.auto_filter.ref} -> {ow.auto_filter.ref}")
        add(prefix + "print_area", norm_ref(sw.print_area) == norm_ref(ow.print_area), f"{sw.print_area} -> {ow.print_area}")
        add(prefix + "print_settings", print_settings_snapshot(sw) == print_settings_snapshot(ow),
            f"source={print_settings_snapshot(sw)}; output={print_settings_snapshot(ow)}")

        expected_formulas = {c.coordinate: c.value for row in sw.iter_rows() for c in row if c.data_type == "f"}
        expected_formulas["D4"] = "=SUM(A1:A3)"
        out_formulas = {c.coordinate: c.value for row in ow.iter_rows() for c in row if c.data_type == "f"}
        add(prefix + "formula_text", expected_formulas == out_formulas, f"expected={expected_formulas}; output={out_formulas}")

        expected_cell_values = {
            "A1": "umya-probe ordinary value",
            "B2": "umya-probe merged anchor",
            "E5": "umya-probe named cell",
            "F2": "umya-probe conditional formatting cell",
            "G2": "umya-probe validation cell",
            "H2": "umya-probe comment cell",
        }
        expected_cells = cell_value_snapshot(sw)
        for ref, val in expected_cell_values.items():
            expected_cells[ref] = ("s", val)
        expected_cells["D4"] = ("f", "=SUM(A1:A3)")
        out_cells = cell_value_snapshot(ow)
        add(prefix + "requested_cell_changes", all(out_cells.get(k) == v for k, v in expected_cells.items() if k in expected_cell_values or k == "D4"),
            "checked A1, B2, E5, F2, G2, H2 and D4")
        c4 = ow["C4"]
        fill_rgb = c4.fill.fgColor.rgb if c4.fill.fgColor.type == "rgb" else None
        add(prefix + "font_fill_border_number_format",
            c4.font.bold is True and fill_rgb in ("FFFFFF00", "FFFF00") and c4.border.left.style == "thin" and c4.number_format == "0.00",
            f"bold={c4.font.bold}; fill={fill_rgb}; left={c4.border.left.style}; number_format={c4.number_format}")

        original_values = cell_value_snapshot(sw)
        modified = {"A1", "B2", "D4", "E5", "F2", "G2", "H2"}
        candidates = [ref for ref in original_values if ref not in modified]
        rng = random.Random(20261022)
        selected = rng.sample(candidates, min(max_random, len(candidates)))
        random_ok = all(out_cells.get(ref) == original_values[ref] for ref in selected)
        add(prefix + "random_unmodified_cells", random_ok, f"seed=20261022; checked={len(selected)}; cells={selected}")

    passed = sum(c["passed"] for c in checks)
    return {"source": str(source), "output": str(output), "checks": checks, "passed": passed, "total": len(checks), "rate": passed / len(checks) if checks else None}


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--report", type=Path)
    parser.add_argument("--random-cells", type=int, default=50)
    args = parser.parse_args()
    result = run(args.input, args.output, args.random_cells)
    serialized = json.dumps(result, ensure_ascii=False, indent=2)
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(serialized + "\n", encoding="utf-8")
    print(serialized)


if __name__ == "__main__":
    main()
