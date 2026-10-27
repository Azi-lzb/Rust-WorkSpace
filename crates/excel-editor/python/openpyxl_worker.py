#!/usr/bin/env python3
"""Isolated, batch-oriented XLSX modification worker for ExcelEditPlan.

This module contains spreadsheet operations only. It is invoked once for a
whole workbook plan and never imports either product's business code.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import posixpath
import sys
import time
import traceback
import zipfile
from copy import copy
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple
from xml.etree import ElementTree
from xml.parsers import expat

SCHEMA_VERSION = 1
SUPPORTED_OPERATIONS = {
    "set_cell_value",
    "set_formula_text",
    "set_comment",
    "set_column_width",
    "set_row_height",
    "rename_sheet",
    "create_sheet",
    "create_sheet_at",
    "set_freeze_panes",
    "set_auto_filter",
    "delete_sheet",
    "set_cell_fill_color",
    "copy_range",
    "import_range",
    "set_workbook_calculation_properties",
    "set_page_settings",
    "replace_header_text",
}


class WorkerError(Exception):
    def __init__(self, kind: str, message: str, *, unsupported: bool = False) -> None:
        super().__init__(message)
        self.kind = kind
        self.unsupported = unsupported


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for block in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _formula_cache_cells(path: Path) -> List[Tuple[str, str]]:
    """Return coordinates with non-empty formula caches, resolving sheet rels.

    openpyxl drops cached formula results when it saves a workbook. The guard
    therefore needs sheet names as well as cell addresses so it can distinguish
    planned destination writes from caches that must remain untouched.
    """
    main_ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    document_rel_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    package_rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    with zipfile.ZipFile(str(path), "r") as package:
        workbook = ElementTree.fromstring(package.read("xl/workbook.xml"))
        relationships = ElementTree.fromstring(package.read("xl/_rels/workbook.xml.rels"))
        targets = {
            relationship.attrib["Id"]: relationship.attrib["Target"]
            for relationship in relationships.findall("{%s}Relationship" % package_rel_ns)
            if relationship.attrib.get("TargetMode") != "External"
        }
        cached_cells = []
        sheets = workbook.find("{%s}sheets" % main_ns)
        if sheets is None:
            raise WorkerError("invalid_workbook", "workbook sheet list is missing", unsupported=True)
        for sheet in sheets.findall("{%s}sheet" % main_ns):
            sheet_name = sheet.attrib.get("name")
            relationship_id = sheet.attrib.get("{%s}id" % document_rel_ns)
            target = targets.get(relationship_id or "")
            if not sheet_name or not target:
                raise WorkerError(
                    "unsupported_workbook_structure",
                    "cannot resolve worksheet relationship for formula-cache inspection",
                    unsupported=True,
                )
            part = posixpath.normpath(
                target.lstrip("/") if target.startswith("/") else posixpath.join("xl", target)
            )
            if part.startswith("../") or not part.startswith("xl/") or part not in package.namelist():
                raise WorkerError(
                    "unsupported_workbook_structure",
                    "worksheet relationship points outside the supported XLSX package layout",
                    unsupported=True,
                )
            root = ElementTree.fromstring(package.read(part))
            for cell in root.iter("{%s}c" % main_ns):
                formula = cell.find("{%s}f" % main_ns)
                cached = cell.find("{%s}v" % main_ns)
                address = cell.attrib.get("r")
                if formula is not None and cached is not None and cached.text not in (None, ""):
                    if not address:
                        raise WorkerError(
                            "unsupported_workbook_structure",
                            "formula cell with cached value has no coordinate",
                            unsupported=True,
                        )
                    cached_cells.append((sheet_name, address.replace("$", "").upper()))
        return cached_cells


def _range_write_covers_cell(sheet: str, address: str, operation: Dict[str, Any]) -> bool:
    """Check only operations that deterministically write every cell in a range."""
    if operation.get("operation") not in ("import_range", "copy_range"):
        return False
    if operation.get("target_sheet") != sheet:
        return False
    # `_copy_cells` skips blank source cells unless materialize_blank_cells is
    # set. Requiring both flags means even blank source cells explicitly clear
    # the destination formula/cache; style/comment copying alone is insufficient.
    if not operation.get("copy_values") or not operation.get("materialize_blank_cells"):
        return False
    try:
        from openpyxl.utils.cell import coordinate_to_tuple, range_boundaries
        source_min_col, source_min_row, source_max_col, source_max_row = range_boundaries(
            operation["source_range"]
        )
        target_row, target_col = coordinate_to_tuple(operation["target_start"].replace("$", ""))
        row, column = coordinate_to_tuple(address)
    except (KeyError, TypeError, ValueError):
        return False
    if not all((source_min_col, source_min_row, source_max_col, source_max_row)):
        return False
    target_max_row = target_row + source_max_row - source_min_row
    target_max_col = target_col + source_max_col - source_min_col
    return target_row <= row <= target_max_row and target_col <= column <= target_max_col


def _formula_cache_is_explicitly_replaced(
    sheet: str, address: str, operations: List[Dict[str, Any]]
) -> bool:
    current_sheet = sheet
    for operation in operations:
        kind = operation.get("operation")
        if kind == "delete_sheet" and operation.get("name") == current_sheet:
            return True
        if kind == "rename_sheet" and operation.get("old_name") == current_sheet:
            current_sheet = operation.get("new_name", current_sheet)
            continue
        if kind in ("set_cell_value", "set_formula_text"):
            target = operation.get("cell", "").replace("$", "").upper()
            if operation.get("sheet") == current_sheet and target == address:
                return True
        elif _range_write_covers_cell(current_sheet, address, operation):
            return True
    return False


def _uncovered_formula_cache_cells(path: Path, operations: List[Dict[str, Any]]) -> List[Tuple[str, str]]:
    return [
        (sheet, address)
        for sheet, address in _formula_cache_cells(path)
        if not _formula_cache_is_explicitly_replaced(sheet, address, operations)
    ]


def _preflight(request: Dict[str, Any]) -> Tuple[Path, Path, List[Dict[str, Any]]]:
    if request.get("schema_version") != SCHEMA_VERSION:
        raise WorkerError("protocol_version", "unsupported request schema", unsupported=True)
    request_id = request.get("request_id")
    if not isinstance(request_id, str) or not request_id.strip():
        raise WorkerError("invalid_request", "request_id must be a non-empty string", unsupported=True)
    if request.get("format") != "xlsx":
        raise WorkerError("unsupported_format", "openpyxl worker accepts XLSX only", unsupported=True)
    input_path = Path(request.get("input", "")).resolve(strict=True)
    output_path = Path(request.get("output", "")).resolve(strict=False)
    if input_path.suffix.lower() != ".xlsx" or output_path.suffix.lower() != ".xlsx":
        raise WorkerError("unsupported_format", "input and output paths must end in .xlsx", unsupported=True)
    if input_path == output_path:
        raise WorkerError("invalid_paths", "input and output paths must be different", unsupported=True)
    if not input_path.is_file():
        raise WorkerError("input_missing", "input workbook does not exist", unsupported=True)
    if output_path.exists() and output_path.stat().st_size != 0:
        raise WorkerError("output_exists", "worker output must be a new or empty candidate", unsupported=True)
    if not zipfile.is_zipfile(str(input_path)):
        raise WorkerError("unsupported_format", "input is not a valid XLSX ZIP package", unsupported=True)
    plan = request.get("plan")
    if not isinstance(plan, dict) or plan.get("request_id") != request_id:
        raise WorkerError("invalid_request", "embedded ExcelEditPlan does not match request_id", unsupported=True)
    operations = plan.get("operations")
    if not isinstance(operations, list) or not operations:
        raise WorkerError("invalid_request", "ExcelEditPlan has no operations", unsupported=True)
    unsupported = sorted({
        str(operation.get("operation", "<missing>")) if isinstance(operation, dict) else "<invalid>"
        for operation in operations
        if not isinstance(operation, dict) or operation.get("operation") not in SUPPORTED_OPERATIONS
    })
    if unsupported:
        raise WorkerError(
            "unsupported_operation",
            "openpyxl worker does not support the complete batch: " + ", ".join(unsupported),
            unsupported=True,
        )
    workbook = plan.get("workbook") or {}
    if not workbook.get("inventory_complete", False):
        raise WorkerError("unknown_package_parts", "OOXML inventory is incomplete", unsupported=True)
    uncovered_formula_caches = _uncovered_formula_cache_cells(input_path, operations)
    if uncovered_formula_caches:
        raise WorkerError(
            "formula_cache_not_preserved",
            "openpyxl would discard {} non-empty formula cache(s) outside explicit value-write operations; use an engine with matching cache evidence".format(
                len(uncovered_formula_caches)
            ),
            unsupported=True,
        )
    for operation in operations:
        if operation["operation"] == "import_range":
            value = operation.get("source_workbook", "")
            source = Path(value).resolve(strict=True)
            if source.suffix.lower() != ".xlsx" or not source.is_file() or not zipfile.is_zipfile(str(source)):
                raise WorkerError("unsupported_import_source", "ImportRange currently accepts valid XLSX sources only", unsupported=True)
    return input_path, output_path, operations


def _cell_value(value: Dict[str, Any]) -> Any:
    kind = value.get("type")
    content = value.get("value")
    if kind == "text":
        return content
    if kind == "decimal":
        from decimal import Decimal
        return Decimal(str(content))
    if kind == "date_time":
        return Decimal(str(content["serial"]))
    if kind == "boolean":
        return bool(content)
    if kind == "error":
        return content
    if kind == "blank":
        return None
    raise WorkerError("invalid_operation", "unknown CellValue variant: " + str(kind))


def _preserve_significant_string_whitespace(path: Path) -> None:
    """Restore OOXML's xml:space marker after openpyxl serializes string cells.

    openpyxl preserves leading/trailing spaces when it reads an inline string,
    but its writer can omit xml:space="preserve". That changes OOXML's declared
    whitespace semantics and can make Rust readers observe a different value.
    Re-serialize only worksheet/shared-string XML parts that contain such text.
    """
    xml_space = "http://www.w3.org/XML/1998/namespace}space"
    main_text = "http://schemas.openxmlformats.org/spreadsheetml/2006/main}t"
    replacement = path.with_name(path.name + ".whitespace.tmp")
    if replacement.exists():
        raise WorkerError("save_failed", "whitespace-preservation temporary output already exists")

    def with_namespaces(data: bytes) -> bytes:
        parser = expat.ParserCreate(namespace_separator="}")
        patches = []
        active_text = []

        def start_element(name: str, attributes: Dict[str, str]) -> None:
            if name == main_text:
                active_text.append((parser.CurrentByteIndex, attributes, []))

        def character_data(value: str) -> None:
            if active_text:
                active_text[-1][2].append(value)

        def end_element(name: str) -> None:
            if name != main_text:
                return
            start, attributes, chunks = active_text.pop()
            value = "".join(chunks)
            if not value or (value[0] not in " \t\r\n" and value[-1] not in " \t\r\n"):
                return
            existing = attributes.get(xml_space)
            if existing == "preserve":
                return
            if existing is not None:
                raise WorkerError(
                    "unsupported_whitespace_semantics",
                    "cannot safely change a non-preserve xml:space attribute",
                    unsupported=True,
                )
            quote = 0
            end = None
            for index in range(start, len(data)):
                byte = data[index]
                if quote:
                    if byte == quote:
                        quote = 0
                elif byte in (ord('"'), ord("'")):
                    quote = byte
                elif byte == ord(">"):
                    end = index
                    break
            if end is None:
                raise WorkerError("invalid_xml", "could not locate OOXML text element end")
            patches.append((end, b' xml:space="preserve"'))

        parser.StartElementHandler = start_element
        parser.CharacterDataHandler = character_data
        parser.EndElementHandler = end_element
        parser.Parse(data, True)
        if not patches:
            return data
        result = data
        for position, insertion in reversed(patches):
            result = result[:position] + insertion + result[position:]
        return result

    try:
        changed_parts = set()
        with zipfile.ZipFile(str(path), "r") as source:
            for item in source.infolist():
                if not (
                    item.filename.startswith("xl/worksheets/")
                    or item.filename == "xl/sharedStrings.xml"
                ):
                    continue
                data = source.read(item.filename)
                if with_namespaces(data) != data:
                    changed_parts.add(item.filename)
        if not changed_parts:
            return
        with zipfile.ZipFile(str(path), "r") as source, zipfile.ZipFile(
            str(replacement), "w"
        ) as target:
            for item in source.infolist():
                data = source.read(item.filename)
                if item.filename in changed_parts:
                    data = with_namespaces(data)
                target.writestr(item, data)
        os.replace(str(replacement), str(path))
    except Exception as error:
        try:
            replacement.unlink()
        except OSError:
            pass
        if isinstance(error, WorkerError):
            raise
        raise WorkerError("save_failed", "could not preserve significant string whitespace: " + str(error))


def _translate_formula(value: Any, source: Any, target: Any) -> Any:
    if not isinstance(value, str) or not value.startswith("="):
        return value
    from openpyxl.formula.translate import Translator
    origin = source if isinstance(source, str) else source.coordinate
    return Translator(value, origin=origin).translate_formula(target.coordinate)


def _copy_style(source: Any, target: Any, style_mode: str) -> None:
    if style_mode == "none":
        return
    if style_mode == "full":
        target.font = copy(source.font)
        target.fill = copy(source.fill)
        target.border = copy(source.border)
        target.alignment = copy(source.alignment)
        target.number_format = source.number_format
        target.protection = copy(source.protection)
        return
    if style_mode == "basic":
        target.font = copy(source.font)
        target.fill = copy(source.fill)
        target.alignment = copy(source.alignment)
        target.number_format = source.number_format
        return
    raise WorkerError("invalid_operation", "unknown range style mode: " + str(style_mode))


def _copy_comment(source: Any, target: Any) -> None:
    if source.comment is None:
        target.comment = None
        return
    from openpyxl.comments import Comment
    target.comment = Comment(source.comment.text, source.comment.author)


def _range_snapshot(source: Any, data_only_source: Any, address: str) -> Tuple[Any, Any, Any, Any]:
    from openpyxl.utils.cell import range_boundaries
    min_col, min_row, max_col, max_row = range_boundaries(address)
    if not min_col or not min_row or not max_col or not max_row:
        raise WorkerError("invalid_range", "range must have finite cell boundaries")
    if max_row < min_row or max_col < min_col:
        raise WorkerError("invalid_range", "range end precedes range start")
    return min_row, min_col, max_row, max_col


def _copy_cells(source: Any, cached: Any, target: Any, bounds: Tuple[int, int, int, int],
                target_start: str, copy_values: bool, copy_formulas: bool,
                style_mode: str, copy_comments: bool, materialize_blanks: bool,
                hide_zero: bool) -> None:
    from openpyxl.utils.cell import coordinate_to_tuple
    from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula
    min_row, min_col, max_row, max_col = bounds
    target_row, target_col = coordinate_to_tuple(target_start.replace("$", ""))
    # Snapshot first so overlapping ranges on the same worksheet behave like
    # copy/paste rather than reading values already written by this operation.
    snapshots = []
    for row in source.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for src in row:
            value = src.value
            if isinstance(value, (ArrayFormula, DataTableFormula)):
                raise WorkerError("unsupported_formula_kind", "array/data-table formula copying is not supported", unsupported=True)
            if value is None and style_mode == "none" and not copy_comments and not materialize_blanks:
                continue
            cached_value = cached.cell(src.row, src.column).value
            snapshots.append((
                src.coordinate,
                value,
                cached_value,
                # Match the Python baseline's `if source.has_style` contract.
                # Missing cells can inherit row/column defaults in Excel, but
                # they do not carry a per-cell style for ImportRange to copy.
                (copy(src.font), copy(src.fill), copy(src.border), copy(src.alignment),
                 src.number_format, copy(src.protection))
                if style_mode != "none" and src.has_style else None,
                None if src.comment is None else (src.comment.text, src.comment.author),
            ))
    for source_coordinate, value, cached_value, style, comment in snapshots:
        source_row, source_col = coordinate_to_tuple(source_coordinate)
        dst = target.cell(target_row + source_row - min_row, target_col + source_col - min_col)
        if style is not None:
            dst.font, dst.fill, dst.border, dst.alignment, dst.number_format, dst.protection = style
        if copy_formulas and isinstance(value, str) and value.startswith("="):
            dst.value = _translate_formula(value, source_coordinate, dst)
        elif copy_values:
            copied = cached_value if isinstance(value, str) and value.startswith("=") else value
            from decimal import Decimal
            if hide_zero and isinstance(copied, (int, float, Decimal)) and copied == 0:
                copied = None
            dst.value = copied
            if isinstance(copied, str) and copied.startswith("#"):
                dst.data_type = "e"
        if copy_comments:
            if comment is None:
                dst.comment = None
            else:
                from openpyxl.comments import Comment
                dst.comment = Comment(comment[0], comment[1])


def _copy_merges(source: Any, target: Any, bounds: Tuple[int, int, int, int],
                 target_start: str) -> None:
    from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter
    min_row, min_col, max_row, max_col = bounds
    target_row, target_col = coordinate_to_tuple(target_start.replace("$", ""))
    for merged in source.merged_cells.ranges:
        if not (merged.min_row <= max_row and merged.max_row >= min_row
                and merged.min_col <= max_col and merged.max_col >= min_col):
            continue
        if not (min_row <= merged.min_row and merged.max_row <= max_row
                and min_col <= merged.min_col and merged.max_col <= max_col):
            raise WorkerError("partial_merge_copy", "source range cuts through a merged range", unsupported=True)
        offset_row = target_row - min_row
        offset_col = target_col - min_col
        new_range = "{}{}:{}{}".format(
            get_column_letter(merged.min_col + offset_col), merged.min_row + offset_row,
            get_column_letter(merged.max_col + offset_col), merged.max_row + offset_row,
        )
        if any(existing.coord == new_range for existing in target.merged_cells.ranges):
            continue
        target.merge_cells(new_range)


def _copy_dimensions(source: Any, target: Any, bounds: Tuple[int, int, int, int],
                     target_start: str, operation: Dict[str, Any]) -> None:
    from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter, column_index_from_string
    min_row, min_col, max_row, max_col = bounds
    target_row, target_col = coordinate_to_tuple(target_start.replace("$", ""))
    row_offset = target_row - min_row
    col_offset = target_col - min_col
    all_rows = operation.get("copy_all_row_dimensions", False)
    all_columns = operation.get("copy_all_column_dimensions", False)
    if operation.get("copy_row_heights") or operation.get("copy_row_hidden"):
        for row, dimension in source.row_dimensions.items():
            if not all_rows and not min_row <= row <= max_row:
                continue
            destination = target.row_dimensions[row + row_offset]
            if operation.get("copy_row_heights"):
                destination.height = dimension.height
            if operation.get("copy_row_hidden"):
                destination.hidden = dimension.hidden
            if operation.get("copy_outline_levels"):
                destination.outlineLevel = dimension.outlineLevel
    if operation.get("copy_column_widths") or operation.get("copy_column_hidden"):
        for key, dimension in source.column_dimensions.items():
            first = dimension.min or column_index_from_string(key)
            last = dimension.max or first
            if not all_columns and (last < min_col or first > max_col):
                continue
            first = max(first, min_col) if not all_columns else first
            last = min(last, max_col) if not all_columns else last
            for column in range(first, last + 1):
                destination_col = column + col_offset
                destination = target.column_dimensions[get_column_letter(destination_col)]
                if operation.get("copy_column_widths"):
                    destination.width = dimension.width
                if operation.get("copy_column_hidden"):
                    destination.hidden = dimension.hidden
                if operation.get("copy_outline_levels"):
                    destination.outlineLevel = dimension.outlineLevel


def _apply_operation(workbook: Any, cached_workbook: Any, operation: Dict[str, Any]) -> None:
    from openpyxl.comments import Comment
    from openpyxl.styles import PatternFill
    from openpyxl.utils.cell import get_column_letter, coordinate_to_tuple, range_boundaries

    kind = operation["operation"]
    payload = operation
    if kind == "set_cell_value":
        cell = workbook[payload["sheet"]][payload["cell"]]
        value = _cell_value(payload["value"])
        cell.value = value
        if payload["value"].get("type") == "date_time":
            cell.number_format = payload["value"]["value"]["number_format"]
        if payload["value"].get("type") == "error":
            cell.data_type = "e"
    elif kind == "set_formula_text":
        formula = payload["formula"]
        workbook[payload["sheet"]][payload["cell"]].value = formula if formula.startswith("=") else "=" + formula
    elif kind == "set_comment":
        cell = workbook[payload["sheet"]][payload["cell"]]
        text = payload.get("text")
        cell.comment = Comment(text, payload.get("author") or "pytools") if text is not None else None
    elif kind == "set_column_width":
        workbook[payload["sheet"]].column_dimensions[get_column_letter(payload["column_index"])].width = payload["width"]
    elif kind == "set_row_height":
        workbook[payload["sheet"]].row_dimensions[payload["row_index"]].height = payload["height"]
    elif kind == "rename_sheet":
        workbook[payload["old_name"]].title = payload["new_name"]
    elif kind == "create_sheet":
        workbook.create_sheet(payload["name"])
    elif kind == "create_sheet_at":
        workbook.create_sheet(payload["name"], payload["index"])
    elif kind == "set_freeze_panes":
        workbook[payload["sheet"]].freeze_panes = payload.get("cell")
    elif kind == "set_auto_filter":
        workbook[payload["sheet"]].auto_filter.ref = payload.get("range")
    elif kind == "delete_sheet":
        if len(workbook.worksheets) <= 1:
            raise WorkerError("invalid_operation", "cannot delete the only worksheet")
        del workbook[payload["name"]]
    elif kind == "set_cell_fill_color":
        color = payload["argb"].strip().lstrip("#")
        if len(color) == 6:
            color = "FF" + color
        if len(color) != 8 or any(char not in "0123456789abcdefABCDEF" for char in color):
            raise WorkerError("invalid_operation", "fill color must be RGB or ARGB hexadecimal")
        workbook[payload["sheet"]][payload["cell"]].fill = PatternFill(fill_type="solid", fgColor=color.upper())
    elif kind in ("copy_range", "import_range"):
        if kind == "copy_range":
            source_workbook = workbook
            cached_source = cached_workbook
            source_name = payload["source_sheet"]
            source_range = payload["source_range"]
            style_mode = "full" if payload.get("copy_styles") else "none"
            copy_comments = payload.get("copy_comments", False)
            copy_merges = False
            dimension_options: Dict[str, Any] = {}
            materialize = False
            hide_zero = False
        else:
            source_workbook = __import__("openpyxl").load_workbook(payload["source_workbook"], data_only=False, keep_links=True)
            cached_source = __import__("openpyxl").load_workbook(payload["source_workbook"], data_only=True, keep_links=True)
            source_name = payload["source_sheet"]
            source_range = payload["source_range"]
            style_mode = payload.get("style_mode", "none")
            copy_comments = payload.get("copy_comments", False)
            copy_merges = payload.get("copy_merges", False)
            dimension_options = payload
            materialize = payload.get("materialize_blank_cells", False)
            hide_zero = payload.get("hide_zero", False)
        source = source_workbook[source_name]
        cached_source_sheet = cached_source[source_name]
        target = workbook[payload["target_sheet"]]
        bounds = _range_snapshot(source, cached_source_sheet, source_range)
        _copy_cells(
            source, cached_source_sheet, target, bounds, payload["target_start"],
            payload.get("copy_values", False), payload.get("copy_formulas", False),
            style_mode, copy_comments, materialize, hide_zero,
        )
        if copy_merges:
            _copy_merges(source, target, bounds, payload["target_start"])
        if dimension_options:
            _copy_dimensions(source, target, bounds, payload["target_start"], dimension_options)
        if kind == "import_range":
            source_workbook.close()
            cached_source.close()
    elif kind == "set_page_settings":
        sheet = workbook[payload["sheet"]]
        settings = payload["settings"]
        if settings.get("orientation") is not None:
            sheet.page_setup.orientation = settings["orientation"]
        if settings.get("fit_to_page") is not None:
            if sheet.sheet_properties.pageSetUpPr is None:
                from openpyxl.worksheet.properties import PageSetupProperties
                sheet.sheet_properties.pageSetUpPr = PageSetupProperties()
            sheet.sheet_properties.pageSetUpPr.fitToPage = settings["fit_to_page"]
        if settings.get("fit_to_width") is not None:
            sheet.page_setup.fitToWidth = settings["fit_to_width"]
        if settings.get("fit_to_height") is not None:
            sheet.page_setup.fitToHeight = settings["fit_to_height"]
        if settings.get("clear_scale"):
            sheet.page_setup.scale = None
        if settings.get("paper_size") is not None:
            sheet.page_setup.paperSize = str(settings["paper_size"])
        if settings.get("print_area") is not None:
            sheet.print_area = settings["print_area"]
        if settings.get("horizontal_centered") is not None:
            sheet.print_options.horizontalCentered = settings["horizontal_centered"]
        if settings.get("vertical_centered") is not None:
            sheet.print_options.verticalCentered = settings["vertical_centered"]
        if settings.get("footer") is not None:
            sheet.oddFooter.center.text = settings["footer"]
    elif kind == "set_workbook_calculation_properties":
        calculation = workbook.calculation
        if calculation is None:
            from openpyxl.workbook.properties import CalcProperties
            calculation = CalcProperties()
            workbook.calculation = calculation
        if payload.get("full_calc_on_load") is not None:
            calculation.fullCalcOnLoad = payload["full_calc_on_load"]
        if payload.get("force_full_calc") is not None:
            calculation.forceFullCalc = payload["force_full_calc"]
    elif kind == "replace_header_text":
        sheet = workbook[payload["sheet"]]
        search, replacement = payload["search"], payload["replacement"]
        for collection in (sheet.oddHeader, sheet.evenHeader, sheet.firstHeader,
                           sheet.oddFooter, sheet.evenFooter, sheet.firstFooter):
            for section in (collection.left, collection.center, collection.right):
                if section.text:
                    section.text = section.text.replace(search, replacement)
    else:
        raise WorkerError("unsupported_operation", "operation is not implemented: " + kind, unsupported=True)


def _verify_style_table_integrity(output_path: Path) -> None:
    """Reject a saved workbook whose cell style references exceed its style table.

    A styles table with fewer xf entries than a worksheet cell, row, or column
    style reference makes Excel/umya fall back to the default style and
    silently drop formatting semantics. Component-style assignment has been
    observed to produce such output on a real template; the worker must fail
    loudly instead of publishing a corrupted candidate.

    Parse the OOXML parts with namespace-aware parsers. Column dimensions refer
    to a cell style through their style attribute, not s.
    """
    import zipfile as _zipfile

    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    et_style_sheet_tag = "{{{}}}styleSheet".format(namespace)
    et_cell_xfs_tag = "{{{}}}cellXfs".format(namespace)
    et_xf_tag = "{{{}}}xf".format(namespace)
    expat_namespace = namespace + "}"
    expat_worksheet_tag = expat_namespace + "worksheet"
    expat_cell_tag = expat_namespace + "c"
    expat_row_tag = expat_namespace + "row"
    expat_column_tag = expat_namespace + "col"

    def checked_style_index(value: str, part_name: str, element_name: str) -> int:
        try:
            index = int(value)
        except ValueError:
            raise WorkerError(
                "style_table_corrupted",
                "saved workbook has an invalid {} style reference in {}".format(
                    element_name, part_name
                ),
            )
        if index < 0:
            raise WorkerError(
                "style_table_corrupted",
                "saved workbook has a negative {} style reference in {}".format(
                    element_name, part_name
                ),
            )
        return index

    try:
        with _zipfile.ZipFile(str(output_path)) as archive:
            styles_root = ElementTree.fromstring(archive.read("xl/styles.xml"))
            if styles_root.tag != et_style_sheet_tag:
                raise WorkerError(
                    "style_table_corrupted",
                    "saved workbook styles part has an unexpected root element",
                )
            cell_xfs_parts = [
                child for child in styles_root if child.tag == et_cell_xfs_tag
            ]
            if len(cell_xfs_parts) != 1:
                raise WorkerError(
                    "style_table_corrupted",
                    "saved workbook has no unique cellXfs style table",
                )
            cell_xfs = cell_xfs_parts[0]
            try:
                declared_count = int(cell_xfs.attrib["count"])
            except (KeyError, ValueError):
                raise WorkerError(
                    "style_table_corrupted",
                    "saved workbook cellXfs has no valid count attribute",
                )
            if declared_count < 0:
                raise WorkerError(
                    "style_table_corrupted",
                    "saved workbook cellXfs has a negative count attribute",
                )
            xf_elements = sum(child.tag == et_xf_tag for child in cell_xfs)
            worksheet_parts = [
                part_name
                for part_name in archive.namelist()
                if posixpath.dirname(part_name) == "xl/worksheets"
                and part_name.endswith(".xml")
            ]
            if not worksheet_parts:
                raise WorkerError(
                    "style_table_corrupted",
                    "saved workbook has no worksheet XML parts to validate",
                )
            for part_name in worksheet_parts:
                parser = expat.ParserCreate(namespace_separator="}")
                root_seen = False

                def check_worksheet_element(
                    element_name: str, attributes: Dict[str, str]
                ) -> None:
                    nonlocal root_seen
                    if not root_seen:
                        root_seen = True
                        if element_name != expat_worksheet_tag:
                            raise WorkerError(
                                "style_table_corrupted",
                                "saved workbook worksheet part has an unexpected root element: "
                                + part_name,
                            )
                    if element_name in (expat_cell_tag, expat_row_tag):
                        attribute_name = "s"
                    elif element_name == expat_column_tag:
                        attribute_name = "style"
                    else:
                        return
                    value = attributes.get(attribute_name)
                    if value is None:
                        return
                    style_index = checked_style_index(value, part_name, element_name)
                    if style_index >= xf_elements:
                        raise WorkerError(
                            "style_table_corrupted",
                            "saved workbook references cell style {} beyond its {} cellXfs entries in {}".format(
                                style_index, xf_elements, part_name
                            ),
                        )

                parser.StartElementHandler = check_worksheet_element
                with archive.open(part_name) as worksheet_xml:
                    parser.ParseFile(worksheet_xml)
                if not root_seen:
                    raise WorkerError(
                        "style_table_corrupted",
                        "saved workbook worksheet part is empty: " + part_name,
                    )
    except WorkerError:
        raise
    except Exception as error:
        raise WorkerError(
            "style_table_corrupted", "style-table integrity check failed: " + str(error)
        )
    if declared_count != xf_elements:
        raise WorkerError(
            "style_table_corrupted",
            "saved workbook cellXfs declares count={} but carries {} xf entries".format(
                declared_count, xf_elements
            ),
        )


def _apply(request: Dict[str, Any]) -> Dict[str, Any]:
    started = time.monotonic()
    input_path, output_path, operations = _preflight(request)
    before_hash = _sha256(input_path)
    try:
        import openpyxl
    except Exception as error:
        raise WorkerError("worker_unavailable", "openpyxl import failed: " + str(error), unsupported=True)
    try:
        workbook = openpyxl.load_workbook(str(input_path), data_only=False, keep_links=True, read_only=False)
        cached_workbook = openpyxl.load_workbook(str(input_path), data_only=True, keep_links=True, read_only=False)
    except Exception as error:
        raise WorkerError("workbook_open_failed", str(error))
    try:
        initial_sheet_names = list(workbook.sheetnames)
        for operation in operations:
            try:
                _apply_operation(workbook, cached_workbook, operation)
            except WorkerError:
                raise
            except Exception as error:
                raise WorkerError("operation_failed", "{}: {}".format(operation.get("operation"), error))
        if not workbook.worksheets:
            raise WorkerError("invalid_output", "workbook must retain at least one worksheet")
        try:
            workbook.save(str(output_path))
        except Exception as error:
            raise WorkerError("save_failed", str(error))
    finally:
        workbook.close()
        cached_workbook.close()
    _preserve_significant_string_whitespace(output_path)
    after_hash = _sha256(input_path)
    if after_hash != before_hash:
        raise WorkerError("source_changed", "worker input changed while applying the edit plan")
    if not output_path.is_file() or output_path.stat().st_size == 0:
        raise WorkerError("save_failed", "openpyxl did not create a non-empty output workbook")
    _verify_style_table_integrity(output_path)
    try:
        verification = openpyxl.load_workbook(str(output_path), read_only=False, data_only=False, keep_links=True)
        if not verification.worksheets:
            raise WorkerError("output_reopen_failed", "saved workbook has no worksheets")
        _verify_plan(verification, operations, initial_sheet_names)
        verification.close()
    except WorkerError:
        raise
    except Exception as error:
        raise WorkerError("output_reopen_failed", str(error))
    return {
        "schema_version": SCHEMA_VERSION,
        "request_id": request["request_id"],
        "status": "applied",
        "operation_count": len(operations),
        "verified_operation_count": len(operations),
        "input_sha256_before": before_hash,
        "input_sha256_after": after_hash,
        "output_sha256": _sha256(output_path),
        "diagnostics": {
            "python_version": sys.version.split()[0],
            "openpyxl_version": openpyxl.__version__,
            "elapsed_ms": int((time.monotonic() - started) * 1000),
        },
    }


def _final_sheet_names(initial_names: List[str], operations: List[Dict[str, Any]]) -> List[str]:
    """Apply sheet-name operations to the original order for final-state checks."""
    expected = list(initial_names)
    for operation in operations:
        kind = operation["operation"]
        if kind == "rename_sheet":
            old_name, new_name = operation["old_name"], operation["new_name"]
            try:
                index = expected.index(old_name)
            except ValueError:
                raise WorkerError(
                    "output_verification_failed",
                    "worksheet rename source missing from planned state",
                )
            expected[index] = new_name
        elif kind == "create_sheet":
            expected.append(operation["name"])
        elif kind == "create_sheet_at":
            index = operation["index"]
            if index < 0 or index > len(expected):
                raise WorkerError(
                    "output_verification_failed",
                    "worksheet insertion index is invalid in planned state",
                )
            expected.insert(index, operation["name"])
        elif kind == "delete_sheet":
            name = operation["name"]
            try:
                expected.remove(name)
            except ValueError:
                raise WorkerError(
                    "output_verification_failed",
                    "worksheet deletion target missing from planned state",
                )
    return expected


def _verify_plan(
    workbook: Any,
    operations: List[Dict[str, Any]],
    initial_sheet_names: List[str],
) -> None:
    """Read back simple edits with openpyxl; Rust separately reopens the file."""
    from decimal import Decimal
    from openpyxl.utils.cell import get_column_letter

    expected_sheet_names = _final_sheet_names(initial_sheet_names, operations)
    if workbook.sheetnames != expected_sheet_names:
        raise WorkerError(
            "output_verification_failed",
            "worksheet names/order readback mismatch",
        )

    for operation in operations:
        kind = operation["operation"]
        payload = operation
        if kind == "set_cell_value":
            cell = workbook[payload["sheet"]][payload["cell"]]
            expected = _cell_value(payload["value"])
            if isinstance(expected, Decimal):
                if cell.value is None or Decimal(str(cell.value)) != expected:
                    raise WorkerError("output_verification_failed", "numeric value readback mismatch at {}!{}".format(payload["sheet"], payload["cell"]))
            elif cell.value != expected:
                raise WorkerError("output_verification_failed", "value readback mismatch at {}!{}".format(payload["sheet"], payload["cell"]))
            if payload["value"].get("type") == "date_time" and cell.number_format != payload["value"]["value"]["number_format"]:
                raise WorkerError("output_verification_failed", "date number format readback mismatch")
        elif kind == "set_formula_text":
            formula = payload["formula"]
            formula = formula if formula.startswith("=") else "=" + formula
            if workbook[payload["sheet"]][payload["cell"]].value != formula:
                raise WorkerError("output_verification_failed", "formula text readback mismatch")
        elif kind == "set_comment":
            actual = workbook[payload["sheet"]][payload["cell"]].comment
            expected = payload.get("text")
            if (actual.text if actual else None) != expected:
                raise WorkerError("output_verification_failed", "comment text readback mismatch")
            if actual and actual.author != (payload.get("author") or "pytools"):
                raise WorkerError("output_verification_failed", "comment author readback mismatch")
        elif kind == "set_column_width":
            actual = workbook[payload["sheet"]].column_dimensions[get_column_letter(payload["column_index"])].width
            if actual is None or abs(float(actual) - float(payload["width"])) > 1e-10:
                raise WorkerError("output_verification_failed", "column width readback mismatch")
        elif kind == "set_row_height":
            actual = workbook[payload["sheet"]].row_dimensions[payload["row_index"]].height
            if actual is None or abs(float(actual) - float(payload["height"])) > 1e-10:
                raise WorkerError("output_verification_failed", "row height readback mismatch")
        elif kind == "set_freeze_panes":
            actual = workbook[payload["sheet"]].freeze_panes
            actual = None if actual is None else str(actual)
            if actual != payload.get("cell"):
                raise WorkerError("output_verification_failed", "freeze panes readback mismatch")
        elif kind == "set_auto_filter":
            actual = workbook[payload["sheet"]].auto_filter.ref
            if actual != payload.get("range"):
                raise WorkerError("output_verification_failed", "AutoFilter readback mismatch")
        elif kind == "set_cell_fill_color":
            expected = payload["argb"].strip().lstrip("#").upper()
            if len(expected) == 6:
                expected = "FF" + expected
            actual = workbook[payload["sheet"]][payload["cell"]].fill.fgColor.rgb
            if not actual or str(actual).upper() != expected:
                raise WorkerError("output_verification_failed", "cell fill readback mismatch")
        elif kind == "set_page_settings":
            sheet = workbook[payload["sheet"]]
            settings = payload["settings"]
            checks = (
                ("orientation", sheet.page_setup.orientation),
                ("fit_to_width", sheet.page_setup.fitToWidth),
                ("fit_to_height", sheet.page_setup.fitToHeight),
                ("paper_size", sheet.page_setup.paperSize),
                ("horizontal_centered", sheet.print_options.horizontalCentered),
                ("vertical_centered", sheet.print_options.verticalCentered),
            )
            for key, actual in checks:
                expected = settings.get(key)
                if expected is not None and str(actual).lower() != str(expected).lower():
                    raise WorkerError("output_verification_failed", "page setting readback mismatch: " + key)
            if settings.get("fit_to_page") is not None and sheet.sheet_properties.pageSetUpPr.fitToPage != settings["fit_to_page"]:
                raise WorkerError("output_verification_failed", "fit-to-page readback mismatch")
            if settings.get("clear_scale") and sheet.page_setup.scale is not None:
                raise WorkerError("output_verification_failed", "page scale was not cleared")
            if settings.get("print_area") is not None:
                actual = str(sheet.print_area)
                if settings["print_area"].replace("$", "") not in actual.replace("$", ""):
                    raise WorkerError("output_verification_failed", "print area readback mismatch")
            if settings.get("footer") is not None and sheet.oddFooter.center.text != settings["footer"]:
                raise WorkerError("output_verification_failed", "footer readback mismatch")
        elif kind == "set_workbook_calculation_properties":
            calculation = workbook.calculation
            if payload.get("full_calc_on_load") is not None and calculation.fullCalcOnLoad != payload["full_calc_on_load"]:
                raise WorkerError("output_verification_failed", "fullCalcOnLoad readback mismatch")
            if payload.get("force_full_calc") is not None and calculation.forceFullCalc != payload["force_full_calc"]:
                raise WorkerError("output_verification_failed", "forceFullCalc readback mismatch")
        elif kind in ("copy_range", "import_range", "replace_header_text"):
            # These operations receive independent range/metadata checks in the
            # Rust validator and dedicated fixture tests; opening the saved
            # package here still verifies it can be deserialized.
            continue


def _response(request_id: Optional[str], status: str, kind: str, message: str,
              diagnostics: Optional[Dict[str, Any]] = None) -> Dict[str, Any]:
    return {
        "schema_version": SCHEMA_VERSION,
        "request_id": request_id,
        "status": status,
        "error": {"kind": kind, "message": message},
        "diagnostics": diagnostics or {},
    }


def _run(request_path: Path, response_path: Path) -> int:
    request_id = None
    try:
        request = json.loads(request_path.read_text(encoding="utf-8"))
        if isinstance(request, dict):
            request_id = request.get("request_id")
        result = _apply(request)
        response_path.write_text(json.dumps(result, ensure_ascii=False), encoding="utf-8")
        return 0
    except WorkerError as error:
        status = "unsupported_before_write" if error.unsupported else "failed"
        result = _response(request_id, status, error.kind, str(error), {"traceback": traceback.format_exc(limit=8)})
        response_path.write_text(json.dumps(result, ensure_ascii=False), encoding="utf-8")
        return 0 if error.unsupported else 1
    except Exception as error:
        result = _response(request_id, "failed", "internal_error", str(error), {"traceback": traceback.format_exc(limit=12)})
        response_path.write_text(json.dumps(result, ensure_ascii=False), encoding="utf-8")
        return 1


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--request", required=True)
    parser.add_argument("--response", required=True)
    args = parser.parse_args()
    return _run(Path(args.request), Path(args.response))


if __name__ == "__main__":
    raise SystemExit(main())
