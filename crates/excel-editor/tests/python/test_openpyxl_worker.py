import hashlib
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
import zipfile
from xml.etree import ElementTree as ET

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side

WORKER = Path(__file__).resolve().parents[2] / "python" / "openpyxl_worker.py"


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def run_worker(plan, output):
    request = output.with_suffix(".request.json")
    response = output.with_suffix(".response.json")
    request.write_text(json.dumps({
        "schema_version": 1,
        "request_id": plan["request_id"],
        "format": "xlsx",
        "input": plan["input"],
        "output": str(output),
        "plan": plan,
    }), encoding="utf-8")
    process = subprocess.run(
        [sys.executable, "-I", str(WORKER), "--request", str(request), "--response", str(response)],
        capture_output=True,
        text=True,
        timeout=30,
    )
    return process, json.loads(response.read_text(encoding="utf-8"))


def operation(kind, **values):
    return {"operation": kind, **values}


class OpenpyxlWorkerTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="openpyxl-worker-test-")
        self.root = Path(self.temp.name)
        self.source = self.root / "source.xlsx"
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = "Data"
        sheet["A1"] = "before"
        sheet["B1"] = 3
        sheet["C2"] = '=A1&"!"'
        sheet.merge_cells("D1:E1")
        sheet["D1"] = "merged"
        sheet["A1"].comment = openpyxl.comments.Comment("old note", "tester")
        sheet["B2"].comment = openpyxl.comments.Comment("retain note", "tester")
        book.create_sheet("Notes")["A1"] = "keep"
        book.create_sheet("DeleteMe")
        book.save(self.source)
        self.import_source = self.root / "import.xlsx"
        imported = openpyxl.Workbook()
        src = imported.active
        src.title = "Source"
        src["A1"] = "imported"
        src["B1"] = 9
        src["A1"].font = Font(bold=True)
        src["A1"].comment = openpyxl.comments.Comment("copied note", "source-author")
        src["C1"] = "=1+1"
        src.merge_cells("A2:B2")
        src["A2"] = "merged import"
        src.row_dimensions[1].height = 25
        src.column_dimensions["A"].width = 22
        imported.save(self.import_source)
        self._add_formula_cache(self.import_source, "xl/worksheets/sheet1.xml", "C1", "2")

    @staticmethod
    def _add_formula_cache(path, part_name, address, value):
        replacement = path.with_name(path.stem + ".cached.xlsx")
        with zipfile.ZipFile(path, "r") as original, zipfile.ZipFile(replacement, "w", zipfile.ZIP_DEFLATED) as patched:
            for item in original.infolist():
                data = original.read(item.filename)
                if item.filename == part_name:
                    root = ET.fromstring(data)
                    main_ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
                    cell = root.find(f".//{{{main_ns}}}c[@r='{address}']")
                    formula = cell.find(f"{{{main_ns}}}f")
                    if formula is None:
                        raise AssertionError("expected formula cell")
                    cached_value = cell.find(f"{{{main_ns}}}v")
                    if cached_value is None:
                        cached_value = ET.SubElement(cell, f"{{{main_ns}}}v")
                    cached_value.text = value
                    ET.register_namespace("", main_ns)
                    data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
                patched.writestr(item, data)
        replacement.replace(path)

    def tearDown(self):
        self.temp.cleanup()

    def base_plan(self, output, operations):
        return {
            "request_id": "worker-test-1",
            "input": str(self.source),
            "output": str(output),
            "format": "xlsx",
            "workbook": {"format": "xlsx", "inventory_complete": True, "features_present": []},
            "operations": operations,
            "preserve": [],
        }

    def test_import_range_blank_unstyled_source_preserves_target_style(self):
        target_path = self.root / "blank-style-target.xlsx"
        target_book = openpyxl.Workbook()
        target_sheet = target_book.active
        target_sheet.title = "Target"
        target_sheet["A1"] = "preserve target style"
        target_sheet["A1"].font = Font(name="Target Font", size=14, bold=True)
        target_sheet["A1"].fill = PatternFill(fill_type="solid", fgColor="FF112233")
        target_sheet["A2"] = "preserve second target style"
        target_sheet["A2"].font = Font(name="Target Font 2", size=13, bold=True)
        target_sheet["A2"].fill = PatternFill(fill_type="solid", fgColor="FF445566")
        target_book.save(target_path)
        target_book.close()

        source_path = self.root / "blank-style-source.xlsx"
        source_book = openpyxl.Workbook()
        source_sheet = source_book.active
        source_sheet.title = "Source"
        source_sheet.row_dimensions[1].font = Font(name="Source Row Font", size=12)
        source_sheet.row_dimensions[1].alignment = Alignment(vertical="center")
        source_sheet.row_dimensions[2].font = Font(name="Source Row Font 2", size=11)
        source_sheet.column_dimensions["A"].font = Font(name="Source Column Font", size=10)
        source_sheet["A2"] = ""
        source_book.save(source_path)
        source_book.close()

        source_check = openpyxl.load_workbook(source_path, data_only=False)
        self.assertFalse(source_check["Source"]["A1"].has_style)
        self.assertFalse(source_check["Source"]["A2"].has_style)
        self.assertTrue(source_check["Source"].row_dimensions[1].has_style)
        self.assertTrue(source_check["Source"].row_dimensions[2].has_style)
        self.assertTrue(source_check["Source"].column_dimensions["A"].has_style)
        source_check.close()
        with zipfile.ZipFile(source_path) as source_zip:
            source_sheet_xml = ET.fromstring(source_zip.read("xl/worksheets/sheet1.xml"))
        namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
        self.assertIsNone(source_sheet_xml.find(f".//{{{namespace}}}c[@r='A1']"))
        self.assertIsNotNone(source_sheet_xml.find(f".//{{{namespace}}}c[@r='A2']"))

        output = self.root / "blank-style-output.xlsx"
        plan = self.base_plan(output, [
            operation(
                "import_range", source_workbook=str(source_path), source_sheet="Source",
                source_range="A1:A2", target_sheet="Target", target_start="A1",
                copy_values=True, copy_formulas=True, style_mode="full", copy_comments=False,
                copy_merges=False, copy_row_heights=False, copy_column_widths=False,
                copy_row_hidden=False, copy_column_hidden=False, copy_outline_levels=False,
                copy_all_row_dimensions=False, copy_all_column_dimensions=False,
                materialize_blank_cells=True, hide_zero=False,
            ),
        ])
        plan["input"] = str(target_path)
        process, response = run_worker(plan, output)

        self.assertEqual(process.returncode, 0, process.stderr)
        self.assertEqual(response["status"], "applied", response)
        saved = openpyxl.load_workbook(output, data_only=False)
        sheet = saved["Target"]
        self.assertIsNone(sheet["A1"].value)
        self.assertEqual(sheet["A1"].font.name, "Target Font")
        self.assertEqual(sheet["A1"].font.sz, 14)
        self.assertTrue(sheet["A1"].font.bold)
        self.assertEqual(sheet["A1"].fill.fgColor.rgb, "FF112233")
        self.assertIsNone(sheet["A2"].value)
        self.assertEqual(sheet["A2"].font.name, "Target Font 2")
        self.assertEqual(sheet["A2"].font.sz, 13)
        self.assertTrue(sheet["A2"].font.bold)
        self.assertEqual(sheet["A2"].fill.fgColor.rgb, "FF445566")
        saved.close()

    def test_import_range_unstyled_source_matches_legacy_style_component_overlays(self):
        target_path = self.root / "unstyled-source-target.xlsx"
        target_book = openpyxl.Workbook()
        target_sheet = target_book.active
        target_sheet.title = "Target"
        target_sheet["A1"] = "replace with blank"
        target_sheet["A1"].font = Font(name="Target Font", size=14, bold=True)
        target_sheet["A1"].fill = PatternFill(fill_type="solid", fgColor="FF112233")
        target_sheet["A1"].border = Border(left=Side(style="thin", color="FF445566"))
        target_sheet["A1"].alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
        target_sheet["A1"].number_format = "#,##0.00"
        target_sheet["A1"].protection = Protection(locked=False, hidden=True)
        target_book.save(target_path)
        target_book.close()

        source_path = self.root / "unstyled-source.xlsx"
        source_book = openpyxl.Workbook()
        source_sheet = source_book.active
        source_sheet.title = "Source"
        source_sheet.row_dimensions[1].font = Font(name="Inherited Row Font", size=12)
        source_sheet.row_dimensions[1].alignment = Alignment(horizontal="right")
        source_sheet.column_dimensions["A"].font = Font(name="Inherited Column Font", size=10)
        source_book.save(source_path)
        source_book.close()

        source_check = openpyxl.load_workbook(source_path, data_only=False)
        source_cell = source_check["Source"]["A1"]
        self.assertFalse(source_cell.has_style)
        self.assertTrue(source_check["Source"].row_dimensions[1].has_style)
        self.assertTrue(source_check["Source"].column_dimensions["A"].has_style)
        self.assertEqual(source_cell.number_format, "General")
        self.assertIsNone(source_cell.alignment.horizontal)
        self.assertTrue(source_cell.protection.locked)
        self.assertFalse(source_cell.protection.hidden)
        source_check.close()

        output = self.root / "unstyled-source-output.xlsx"
        plan = self.base_plan(output, [
            operation(
                "import_range", source_workbook=str(source_path), source_sheet="Source",
                source_range="A1:A1", target_sheet="Target", target_start="A1",
                copy_values=True, copy_formulas=True, style_mode="full", copy_comments=False,
                copy_merges=False, copy_row_heights=False, copy_column_widths=False,
                copy_row_hidden=False, copy_column_hidden=False, copy_outline_levels=False,
                copy_all_row_dimensions=False, copy_all_column_dimensions=False,
                materialize_blank_cells=True, hide_zero=False,
            ),
        ])
        plan["input"] = str(target_path)
        process, response = run_worker(plan, output)

        self.assertEqual(process.returncode, 0, process.stderr)
        self.assertEqual(response["status"], "applied", response)
        saved = openpyxl.load_workbook(output, data_only=False)
        result = saved["Target"]["A1"]
        self.assertIsNone(result.value)
        self.assertEqual(result.font.name, "Target Font")
        self.assertEqual(result.font.sz, 14)
        self.assertTrue(result.font.bold)
        self.assertEqual(result.fill.fgColor.rgb, "FF112233")
        self.assertEqual(result.border.left.style, "thin")
        self.assertEqual(result.border.left.color.rgb, "FF445566")
        self.assertEqual(result.number_format, "General")
        self.assertIsNone(result.alignment.horizontal)
        self.assertIsNone(result.alignment.vertical)
        self.assertIsNone(result.alignment.wrap_text)
        self.assertTrue(result.protection.locked)
        self.assertFalse(result.protection.hidden)
        saved.close()

    def test_openpyxl_round_trip_normalizes_stored_empty_inline_string(self):
        source = self.root / "empty-inline-string.xlsx"
        book = openpyxl.Workbook()
        book.save(source)
        book.close()

        namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
        replacement = source.with_name("empty-inline-string.patched.xlsx")
        ET.register_namespace("", namespace)
        with zipfile.ZipFile(source, "r") as original, zipfile.ZipFile(
            replacement, "w", zipfile.ZIP_DEFLATED
        ) as patched:
            for item in original.infolist():
                data = original.read(item.filename)
                if item.filename == "xl/worksheets/sheet1.xml":
                    root = ET.fromstring(data)
                    sheet_data = root.find(f"{{{namespace}}}sheetData")
                    row = ET.SubElement(sheet_data, f"{{{namespace}}}row", {"r": "10"})
                    cell = ET.SubElement(
                        row, f"{{{namespace}}}c", {"r": "E10", "t": "inlineStr"}
                    )
                    inline_string = ET.SubElement(cell, f"{{{namespace}}}is")
                    ET.SubElement(inline_string, f"{{{namespace}}}t").text = ""
                    data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
                patched.writestr(item, data)
        replacement.replace(source)

        loaded = openpyxl.load_workbook(source, data_only=False)
        self.assertEqual(loaded.active["E10"].value, "")
        output = self.root / "empty-inline-string-round-trip.xlsx"
        loaded.save(output)
        loaded.close()

        round_trip = openpyxl.load_workbook(output, data_only=False)
        self.assertIsNone(round_trip.active["E10"].value)
        round_trip.close()

    def test_import_range_maps_style_components_when_workbook_style_ids_collide(self):
        target_path = self.root / "style-index-collision-target.xlsx"
        target_book = openpyxl.Workbook()
        target_sheet = target_book.active
        target_sheet.title = "Target"
        target_sheet["A1"] = "target"
        target_sheet["A1"].font = Font(name="Target Font", size=14, bold=True, color="FFAA0000")
        target_sheet["A1"].fill = PatternFill(fill_type="solid", fgColor="FFFFCC00")
        target_book.save(target_path)
        target_book.close()

        source_path = self.root / "style-index-collision-source.xlsx"
        source_book = openpyxl.Workbook()
        source_sheet = source_book.active
        source_sheet.title = "Source"
        source_sheet["A1"] = "source"
        source_sheet["A1"].font = Font(name="Source Font", size=11, italic=True, color="FF0033AA")
        source_sheet["A1"].fill = PatternFill(fill_type="solid", fgColor="FF33CC99")
        source_book.save(source_path)
        source_book.close()

        target_check = openpyxl.load_workbook(target_path)
        source_check = openpyxl.load_workbook(source_path)
        self.assertEqual(target_check["Target"]["A1"].style_id, 1)
        self.assertEqual(source_check["Source"]["A1"].style_id, 1)
        self.assertNotEqual(
            target_check["Target"]["A1"].font.name,
            source_check["Source"]["A1"].font.name,
        )
        target_check.close()
        source_check.close()

        output = self.root / "style-index-collision-output.xlsx"
        plan = self.base_plan(output, [
            operation(
                "import_range", source_workbook=str(source_path), source_sheet="Source",
                source_range="A1:A1", target_sheet="Target", target_start="A1",
                copy_values=True, copy_formulas=True, style_mode="full", copy_comments=False,
                copy_merges=False, copy_row_heights=False, copy_column_widths=False,
                copy_row_hidden=False, copy_column_hidden=False, copy_outline_levels=False,
                copy_all_row_dimensions=False, copy_all_column_dimensions=False,
                materialize_blank_cells=True, hide_zero=False,
            ),
        ])
        plan["input"] = str(target_path)
        process, response = run_worker(plan, output)

        self.assertEqual(process.returncode, 0, process.stderr)
        self.assertEqual(response["status"], "applied", response)
        saved = openpyxl.load_workbook(output, data_only=False)
        cell = saved["Target"]["A1"]
        self.assertEqual(cell.value, "source")
        self.assertEqual(cell.font.name, "Source Font")
        self.assertEqual(cell.font.sz, 11)
        self.assertTrue(cell.font.italic)
        self.assertEqual(cell.font.color.rgb, "FF0033AA")
        self.assertEqual(cell.fill.fgColor.rgb, "FF33CC99")
        saved.close()

    def test_import_range_copies_nondefault_xf_zero_on_existing_cell_without_s_attribute(self):
        source_path = self.root / "implicit-xf0-source.xlsx"
        source_book = openpyxl.Workbook()
        source_book.active.title = "Source"
        source_book.active["A1"] = "source uses XF0"
        source_book.save(source_path)
        source_book.close()

        patched_path = self.root / "implicit-xf0-source-patched.xlsx"
        namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
        with zipfile.ZipFile(source_path, "r") as original, zipfile.ZipFile(
            patched_path, "w", zipfile.ZIP_DEFLATED
        ) as patched:
            for item in original.infolist():
                data = original.read(item.filename)
                if item.filename == "xl/styles.xml":
                    styles = ET.fromstring(data)
                    xf = styles.find(f"{{{namespace}}}cellXfs/{{{namespace}}}xf")
                    xf.set("applyAlignment", "1")
                    ET.SubElement(
                        xf, f"{{{namespace}}}alignment", {"vertical": "center"}
                    )
                    ET.register_namespace("", namespace)
                    data = ET.tostring(styles, encoding="utf-8", xml_declaration=True)
                patched.writestr(item, data)
        patched_path.replace(source_path)

        source_check = openpyxl.load_workbook(source_path, data_only=False)
        source_cell = source_check["Source"]["A1"]
        self.assertTrue(source_cell.has_style)
        self.assertEqual(source_cell.alignment.vertical, "center")
        source_check.close()
        with zipfile.ZipFile(source_path) as source_zip:
            source_sheet_xml = ET.fromstring(source_zip.read("xl/worksheets/sheet1.xml"))
        self.assertIsNone(source_sheet_xml.find(f".//{{{namespace}}}c[@r='A1']").get("s"))

        target_path = self.root / "implicit-xf0-target.xlsx"
        target_book = openpyxl.Workbook()
        target_sheet = target_book.active
        target_sheet.title = "Target"
        target_sheet["A1"] = "old target"
        target_sheet["A1"].font = Font(name="Target Font", size=14, bold=True)
        target_book.save(target_path)
        target_book.close()

        output = self.root / "implicit-xf0-output.xlsx"
        plan = self.base_plan(output, [operation(
            "import_range", source_workbook=str(source_path), source_sheet="Source",
            source_range="A1:A1", target_sheet="Target", target_start="A1",
            copy_values=True, copy_formulas=True, style_mode="full", copy_comments=False,
            copy_merges=False, copy_row_heights=False, copy_column_widths=False,
            copy_row_hidden=False, copy_column_hidden=False, copy_outline_levels=False,
            copy_all_row_dimensions=False, copy_all_column_dimensions=False,
            materialize_blank_cells=True, hide_zero=False,
        )])
        plan["input"] = str(target_path)
        process, response = run_worker(plan, output)

        self.assertEqual(process.returncode, 0, process.stderr)
        self.assertEqual(response["status"], "applied", response)
        saved = openpyxl.load_workbook(output, data_only=False)
        cell = saved["Target"]["A1"]
        self.assertEqual(cell.value, "source uses XF0")
        self.assertEqual(cell.alignment.vertical, "center")
        self.assertNotEqual(cell.font.name, "Target Font")
        saved.close()

    def test_batch_modifies_one_workbook_and_preserves_supported_structures(self):
        output = self.root / "candidate.xlsx"
        source_hash = digest(self.source)
        import_hash = digest(self.import_source)
        operations = [
            operation("set_cell_value", sheet="Data", cell="A1", value={"type": "decimal", "value": "12.75"}),
            operation("set_formula_text", sheet="Data", cell="B3", formula="=A1*2"),
            operation("set_comment", sheet="Data", cell="A1", text="new note", author="worker"),
            operation("set_column_width", sheet="Data", column_index=2, width=18.5),
            operation("set_row_height", sheet="Data", row_index=3, height=27.0),
            operation("set_cell_fill_color", sheet="Data", cell="A1", argb="FF112233"),
            operation("set_page_settings", sheet="Data", settings={
                "orientation": "landscape", "fit_to_page": True, "fit_to_width": 1,
                "fit_to_height": 0, "clear_scale": True, "paper_size": 9,
                "print_area": "A1:E3", "horizontal_centered": True,
                "vertical_centered": False, "footer": "&P",
            }),
            operation("copy_range", source_sheet="Data", source_range="A1:B1",
                      target_sheet="Data", target_start="H5", copy_values=True,
                      copy_formulas=False, copy_styles=True, copy_comments=True),
            operation("import_range", source_workbook=str(self.import_source), source_sheet="Source",
                      source_range="A1:C2", target_sheet="Data", target_start="J1",
                      copy_values=True, copy_formulas=True, style_mode="full", copy_comments=True,
                      copy_merges=True, copy_row_heights=True, copy_column_widths=True,
                      copy_row_hidden=False, copy_column_hidden=False, copy_outline_levels=False,
                      copy_all_row_dimensions=False, copy_all_column_dimensions=False,
                      materialize_blank_cells=True, hide_zero=False),
            operation("set_workbook_calculation_properties", full_calc_on_load=True, force_full_calc=True),
            operation("rename_sheet", old_name="Notes", new_name="Information"),
            operation("create_sheet", name="Generated"),
            operation("create_sheet_at", name="History", index=0),
            operation("set_freeze_panes", sheet="History", cell="A2"),
            operation("set_auto_filter", sheet="History", range="A1:M2"),
            operation("delete_sheet", name="DeleteMe"),
        ]
        plan = self.base_plan(output, operations)
        process, response = run_worker(plan, output)
        self.assertEqual(process.returncode, 0, process.stderr)
        self.assertEqual(response["status"], "applied", response)
        self.assertEqual(response["operation_count"], len(operations))
        self.assertEqual(response["input_sha256_before"], response["input_sha256_after"])
        self.assertEqual(digest(self.source), source_hash)
        self.assertEqual(digest(self.import_source), import_hash)

        saved = openpyxl.load_workbook(output, data_only=False)
        self.assertEqual(saved.sheetnames, ["History", "Data", "Information", "Generated"])
        self.assertEqual(saved["History"].freeze_panes, "A2")
        self.assertEqual(saved["History"].auto_filter.ref, "A1:M2")
        sheet = saved["Data"]
        self.assertEqual(sheet["A1"].value, 12.75)
        self.assertEqual(sheet["B3"].value, "=A1*2")
        self.assertEqual(sheet["C2"].value, '=A1&"!"')
        self.assertEqual(sheet["A1"].comment.text, "new note")
        self.assertEqual(sheet["B2"].comment.text, "retain note")
        self.assertEqual(sheet["A1"].fill.fgColor.rgb, "FF112233")
        self.assertAlmostEqual(sheet.column_dimensions["B"].width, 18.5)
        self.assertAlmostEqual(sheet.row_dimensions[3].height, 27.0)
        self.assertEqual(sheet.print_area, "'Data'!$A$1:$E$3")
        self.assertEqual(sheet.page_setup.orientation, "landscape")
        self.assertTrue(sheet.sheet_properties.pageSetUpPr.fitToPage)
        self.assertEqual(sheet.oddFooter.center.text, "&P")
        self.assertEqual(sheet["H5"].value, 12.75)
        self.assertEqual(sheet["I5"].value, 3)
        self.assertEqual(sheet["H5"].comment.text, "new note")
        self.assertEqual(sheet["J1"].value, "imported")
        self.assertTrue(sheet["J1"].font.bold)
        self.assertEqual(sheet["J1"].comment.text, "copied note")
        self.assertEqual(sheet["L1"].value, "=1+1")
        self.assertTrue(saved.calculation.fullCalcOnLoad)
        self.assertTrue(saved.calculation.forceFullCalc)
        self.assertIn("J2:K2", {item.coord for item in sheet.merged_cells.ranges})
        self.assertAlmostEqual(sheet.row_dimensions[1].height, 25)
        self.assertAlmostEqual(sheet.column_dimensions["J"].width, 22)
        saved.close()

    def test_delete_and_recreate_same_sheet_validates_final_names_and_order(self):
        output = self.root / "recreated-sheet.xlsx"
        plan = self.base_plan(output, [
            operation("delete_sheet", name="DeleteMe"),
            operation("create_sheet_at", name="DeleteMe", index=1),
        ])

        process, response = run_worker(plan, output)

        self.assertEqual(process.returncode, 0, process.stderr)
        self.assertEqual(response["status"], "applied", response)
        self.assertEqual(response["operation_count"], 2)
        saved = openpyxl.load_workbook(output, data_only=False)
        self.assertEqual(saved.sheetnames, ["Data", "DeleteMe", "Notes"])
        saved.close()

    def test_unsupported_row_shift_is_rejected_before_output_write(self):
        output = self.root / "unsupported.xlsx"
        plan = self.base_plan(output, [operation("insert_rows", sheet="Data", first_row=2, count=1)])
        process, response = run_worker(plan, output)
        self.assertEqual(process.returncode, 0)
        self.assertEqual(response["status"], "unsupported_before_write")
        self.assertEqual(response["error"]["kind"], "unsupported_operation")
        self.assertFalse(output.exists())

    def test_formula_cache_is_rejected_before_write(self):
        cached = self.root / "cached.xlsx"
        with zipfile.ZipFile(self.source, "r") as original, zipfile.ZipFile(cached, "w", zipfile.ZIP_DEFLATED) as patched:
            for item in original.infolist():
                data = original.read(item.filename)
                if item.filename == "xl/worksheets/sheet1.xml":
                    root = ET.fromstring(data)
                    main_ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
                    cell = root.find(f".//{{{main_ns}}}c[@r='C2']")
                    self.assertIsNotNone(cell)
                    cached_value = cell.find(f"{{{main_ns}}}v")
                    self.assertIsNotNone(cached_value)
                    cached_value.text = "1"
                    ET.register_namespace("", main_ns)
                    data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
                patched.writestr(item, data)
        self.assertTrue(zipfile.is_zipfile(cached))
        self.source = cached
        output = self.root / "cached-output.xlsx"
        plan = self.base_plan(output, [operation("set_cell_value", sheet="Data", cell="A1", value={"type": "text", "value": "changed"})])
        process, response = run_worker(plan, output)
        self.assertEqual(process.returncode, 0)
        self.assertEqual(response["status"], "unsupported_before_write")
        self.assertEqual(response["error"]["kind"], "formula_cache_not_preserved")
        self.assertFalse(output.exists())

    def test_formula_cache_inside_materialized_import_range_can_be_replaced(self):
        cached = self.root / "cached-target.xlsx"
        openpyxl.Workbook().save(cached)
        workbook = openpyxl.load_workbook(cached)
        workbook.active.title = "Data"
        workbook["Data"]["C2"] = "=1+1"
        workbook.save(cached)
        self._add_formula_cache(cached, "xl/worksheets/sheet1.xml", "C2", "2")

        output = self.root / "cached-target-output.xlsx"
        plan = self.base_plan(output, [operation(
            "import_range",
            source_workbook=str(self.import_source), source_sheet="Source",
            source_range="A1:C2", target_sheet="Data", target_start="A1",
            copy_values=True, copy_formulas=True, style_mode="full", copy_comments=True,
            copy_merges=True, copy_row_heights=False, copy_column_widths=False,
            copy_row_hidden=False, copy_column_hidden=False, copy_outline_levels=False,
            copy_all_row_dimensions=False, copy_all_column_dimensions=False,
            materialize_blank_cells=True, hide_zero=False,
        )])
        plan["input"] = str(cached)
        process, response = run_worker(plan, output)
        self.assertEqual(process.returncode, 0, process.stderr)
        self.assertEqual(response["status"], "applied", response)
        saved_values = openpyxl.load_workbook(output, data_only=True)
        self.assertIsNone(saved_values["Data"]["C2"].value)
        saved_values.close()
        saved_formulas = openpyxl.load_workbook(output, data_only=False)
        self.assertIsNone(saved_formulas["Data"]["C2"].value)
        saved_formulas.close()

    def test_materialized_import_range_does_not_cover_formula_cache_outside_target(self):
        cached = self.root / "cached-target-outside.xlsx"
        openpyxl.Workbook().save(cached)
        workbook = openpyxl.load_workbook(cached)
        workbook.active.title = "Data"
        workbook["Data"]["C2"] = "=1+1"
        workbook.save(cached)
        self._add_formula_cache(cached, "xl/worksheets/sheet1.xml", "C2", "2")

        output = self.root / "cached-target-outside-output.xlsx"
        plan = self.base_plan(output, [operation(
            "import_range",
            source_workbook=str(self.import_source), source_sheet="Source",
            source_range="A1:B2", target_sheet="Data", target_start="A1",
            copy_values=True, copy_formulas=True, style_mode="full", copy_comments=True,
            copy_merges=True, copy_row_heights=False, copy_column_widths=False,
            copy_row_hidden=False, copy_column_hidden=False, copy_outline_levels=False,
            copy_all_row_dimensions=False, copy_all_column_dimensions=False,
            materialize_blank_cells=True, hide_zero=False,
        )])
        plan["input"] = str(cached)
        process, response = run_worker(plan, output)
        self.assertEqual(process.returncode, 0, process.stderr)
        self.assertEqual(response["status"], "unsupported_before_write")
        self.assertEqual(response["error"]["kind"], "formula_cache_not_preserved")
        self.assertFalse(output.exists())

    def test_operation_error_is_structured_and_does_not_publish_output(self):
        output = self.root / "failed.xlsx"
        plan = self.base_plan(output, [operation("set_cell_value", sheet="Missing", cell="A1", value={"type": "text", "value": "x"})])
        process, response = run_worker(plan, output)
        self.assertNotEqual(process.returncode, 0)
        self.assertEqual(response["status"], "failed")
        self.assertEqual(response["error"]["kind"], "operation_failed")
        self.assertFalse(output.exists())


class StyleTableIntegrityTest(unittest.TestCase):
    """Direct coverage for the saved-output style-table integrity assertion."""

    def setUp(self):
        import importlib.util
        spec = importlib.util.spec_from_file_location("openpyxl_worker_module", WORKER)
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        self.integrity = module._verify_style_table_integrity
        self.root = Path(tempfile.mkdtemp(prefix="worker-style-integrity-"))
        self.addCleanup(self._cleanup)

    def _cleanup(self):
        import shutil
        shutil.rmtree(self.root, ignore_errors=True)

    def _workbook_with_styles(self, path, cellxfs_body, declared_count, sheet_xml):
        styles = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
            '<fonts count="1"><font><sz val="11"/><name val="Calibri"/></font></fonts>'
            '<fills count="2"><fill><patternFill patternType="none"/></fill>'
            '<fill><patternFill patternType="gray125"/></fill></fills>'
            '<borders count="1"><border><left/><right/><top/><bottom/><diagonal/></border></borders>'
            '<cellStyleXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/></cellStyleXfs>'
            f'<cellXfs count="{declared_count}">{cellxfs_body}</cellXfs>'
            '<cellStyles count="1"><cellStyle name="Normal" xfId="0" builtinId="0"/></cellStyles>'
            '</styleSheet>'
        )
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("[Content_Types].xml",
                '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
                '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
                '<Default Extension="xml" ContentType="application/xml"/>'
                '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
                '<Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
                '<Override PartName="/xl/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/></Types>')
            archive.writestr("_rels/.rels",
                '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/></Relationships>')
            archive.writestr("xl/workbook.xml",
                '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
                '<sheets><sheet name="Data" sheetId="1" r:id="rId1"/></sheets></workbook>')
            archive.writestr("xl/_rels/workbook.xml.rels",
                '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/>'
                '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/></Relationships>')
            archive.writestr("xl/styles.xml", styles)
            archive.writestr("xl/worksheets/sheet1.xml", sheet_xml)

    def test_out_of_range_style_reference_is_rejected(self):
        target = self.root / "oob.xlsx"
        self._workbook_with_styles(
            target,
            cellxfs_body='<xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/>',
            declared_count=1,
            sheet_xml='<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                      '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
                      '<sheetData><row r="1"><c r="A1" t="inlineStr" s="5"><is><t>x</t></is></c></row></sheetData></worksheet>',
        )
        with self.assertRaises(Exception) as caught:
            self.integrity(target)
        self.assertIn("beyond its", str(caught.exception))

    def test_out_of_range_column_dimension_style_is_rejected(self):
        target = self.root / "column-style-oob.xlsx"
        self._workbook_with_styles(
            target,
            cellxfs_body='<xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/>',
            declared_count=1,
            sheet_xml=(
                '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                '<x:worksheet xmlns:x="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
                '<x:cols><x:col min="1" max="1" style="1"/></x:cols>'
                '<x:sheetData><x:row r="1"><x:c r="A1" t="inlineStr"><x:is><x:t>x</x:t>'
                '</x:is></x:c></x:row></x:sheetData></x:worksheet>'
            ),
        )
        with self.assertRaises(Exception) as caught:
            self.integrity(target)
        self.assertIn("cell style 1 beyond its 1 cellXfs entries", str(caught.exception))

    def test_count_mismatch_is_rejected(self):
        target = self.root / "mismatch.xlsx"
        self._workbook_with_styles(
            target,
            cellxfs_body='<xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/>',
            declared_count=9,
            sheet_xml='<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                      '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
                      '<sheetData><row r="1"><c r="A1" t="inlineStr" s="0"><is><t>x</t></is></c></row></sheetData></worksheet>',
        )
        with self.assertRaises(Exception) as caught:
            self.integrity(target)
        self.assertIn("declares count=", str(caught.exception))

    def test_consistent_table_passes(self):
        target = self.root / "consistent.xlsx"
        self._workbook_with_styles(
            target,
            cellxfs_body='<xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/>'
                         '<xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0" applyAlignment="1"><alignment vertical="center"/></xf>',
            declared_count=2,
            sheet_xml='<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                      '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
                      '<sheetData><row r="1"><c r="A1" t="inlineStr" s="1"><is><t>x</t></is></c></row></sheetData></worksheet>',
        )
        self.integrity(target)


if __name__ == "__main__":
    unittest.main()
